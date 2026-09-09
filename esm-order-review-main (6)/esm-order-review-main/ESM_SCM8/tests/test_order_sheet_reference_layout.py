from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from backend.analysis import order_template_review_basis
from core import common, export_excel, export_excel_order_sheet, export_excel_report, export_excel_util, order_review, order_review_report, validation


PRODUCT_NAME = "\uc0c1\ud488\uba85"
PRODUCT_CODE = "\uc0c1\ud488\ucf54\ub4dc"
BRAND = "\ube0c\ub79c\ub4dc"
EU_LOCAL_STOCK = "EU \ud604\uc9c0 \uc7ac\uace0"
SAFETY_STOCK = "\uc548\uc804\uc7ac\uace0"
MONTHLY_SALES = "\uc6d4\ud3c9\uade0"
SHIPPING_STOCK = "\uc6b4\uc1a1\uc911"
OPEN_PO_QTY = "\ubbf8\uc785\uace0 \uc218\ub7c9"
HQ_EU_STOCK = "\ubcf8\uc0ac EU\ucc3d\uace0"
ORDER_NEEDED_QTY = "\ubc1c\uc8fc\ud544\uc694\uc218\ub7c9"
ORDER_AMOUNT_EUR = "\ubd80\uc871\uae08\uc561_EUR"
EU_COVER_DAYS = "EU \ud604\uc9c0 \ucee4\ubc84\uc77c\uc218"
DEPLETION_DATE = "\uc608\uc0c1 \uc18c\uc9c4\uc77c"
SALES_QTY = "PA+CA \ud310\ub9e4\uc218\ub7c9"


def _settings():
    return {
        "safety_months": 3,
        "base_date": date(2026, 5, 20),
        "lead_times": {
            "\ud56d\uacf5": 15,
            "\ud574\uc6b4": 80,
            "\ucca0\uc1a1": 35,
            "\ud2b8\ub7ed": 28,
        },
        "eur_krw_rate": 1500,
    }


def test_order_template_review_basis_uses_stock_file_rows_including_zero_price_exclusions():
    combined = pd.DataFrame(
        {
            PRODUCT_CODE: ["SKU-ZERO", "SKU-NORMAL", "SKU-SHIPPING-ONLY"],
            BRAND: ["Brand A", "Brand A", "Brand A"],
            "재고파일 존재여부": ["Y", "Y", ""],
            "제외 SKU": [True, False, False],
            "제외유형": ["0단가/무상", "", ""],
        }
    )
    filtered = combined[~combined["제외 SKU"]].copy()

    basis = order_template_review_basis(combined, filtered, {**_settings(), "brand": "Brand A"})

    assert basis[PRODUCT_CODE].tolist() == ["SKU-ZERO", "SKU-NORMAL"]


def _report_df():
    return pd.DataFrame(
        {
            PRODUCT_CODE: ["SKU-AIR", "SKU-NOSALES"],
            PRODUCT_NAME: ["Need Air", "No Sales"],
            BRAND: ["Brand A", "Brand B"],
            EU_LOCAL_STOCK: [80, 0],
            SAFETY_STOCK: [120, 0],
            MONTHLY_SALES: [40, 0],
            SHIPPING_STOCK: [30, 0],
            OPEN_PO_QTY: [7, 0],
            HQ_EU_STOCK: [50, 0],
            ORDER_NEEDED_QTY: [90, 0],
            ORDER_AMOUNT_EUR: [450.25, 0],
            EU_COVER_DAYS: [2, 999],
            DEPLETION_DATE: [date(2026, 5, 22), "\ud310\ub9e4\uc5c6\uc74c"],
            SALES_QTY: [1000, 0],
        }
    )


def test_report_summary_splits_open_po_included_and_excluded_order_totals():
    report = _report_df().copy()
    report["1차 부족수량"] = [999, 0]

    summary = export_excel_report.build_order_review_report_summary(report, {**_settings(), "eur_krw_rate": 1500})

    expected_unit_price = 450.25 / 90
    assert summary["total_order_qty"] == 10
    assert round(summary["total_order_eur"], 6) == round(expected_unit_price * 10, 6)
    assert round(summary["total_order_krw"], 6) == round(expected_unit_price * 10 * 1500, 6)
    assert summary["total_with_open_qty"] == 3
    assert round(summary["total_with_open_eur"], 6) == round(expected_unit_price * 3, 6)
    assert round(summary["total_with_open_krw"], 6) == round(expected_unit_price * 3 * 1500, 6)


def test_report_df_applies_export_settings_to_canonical_order_metrics():
    report = order_review_report.build_order_review_report_df(
        _report_df(),
        {**_settings(), "safety_months": 5},
    )

    first = report.iloc[0]
    expected_unit_price = 450.25 / 90
    assert first["발주필요수량"] == 90
    assert first["부족수량"] == 90
    assert first["1차 부족수량"] == 90
    assert first["최종 부족수량"] == 90
    assert first["안전재고 목표수량"] == 200
    assert first["조치 요약"] == "신규 발주 90개"
    assert first["추가 발주 필요 수량"] == 90
    assert round(float(first["부족금액_EUR"]), 6) == round(expected_unit_price * 90, 6)
    assert round(float(first["부족금액_KRW"]), 6) == round(expected_unit_price * 90 * 1500, 6)


def test_stock_eta_order_and_report_totals_share_canonical_metrics():
    settings = {**_settings(), "safety_months": 5}
    review = pd.DataFrame(
        {
            PRODUCT_CODE: ["SKU-CONSISTENT", "SKU-ZERO"],
            PRODUCT_NAME: ["Consistent", "Zero"],
            BRAND: ["Brand A", "Brand A"],
            "최근 3개월 판매수량": [120, 0],
            "월평균 판매수량": [40, 0],
            "EU 현지 가용수량": [80, 0],
            "운송중 수량": [30, 0],
            "미입고수량": [7, 0],
            "본사 EU창고 가용수량": [0, 0],
            "EU 입고단가": [5, 0],
            "EU 현지 커버일수": [30, 999],
            "고갈 예상일": [date(2026, 6, 19), "판매없음"],
        }
    )

    report = order_review_report.build_order_review_report_df(review, settings)
    stock_eta = export_excel_order_sheet.build_stock_eta_sheet_df(review, settings=settings)
    order_sheet = export_excel_order_sheet.build_order_sheet_df(report, settings)
    summary = export_excel_report.build_order_review_report_summary(report, settings)

    expected_qty = 90
    expected_eur = expected_qty * 5
    expected_krw = expected_eur * settings["eur_krw_rate"]
    assert float(stock_eta["발주필요수량"].sum()) == expected_qty
    assert float(order_sheet["발주총\n필요수량"].sum()) == expected_qty
    assert float(report["발주필요수량"].sum()) == expected_qty
    assert summary["total_order_qty"] == expected_qty
    assert float(order_sheet["총발주금액"].sum()) == expected_krw
    assert summary["total_order_eur"] == expected_eur
    assert summary["total_order_krw"] == expected_krw


def test_order_sheet_maps_report_values_to_reference_columns():
    df = export_excel_order_sheet.build_order_sheet_df(_report_df(), _settings())
    cols = common.ORDER_SHEET_COLUMNS

    assert list(df.columns) == cols
    assert cols[cols.index("③ 합산\n(M)") + 1 : cols.index("미입고\n현황")] == [
        "안전재고 기준(M)",
        "확보재고 수량",
        "안전재고 목표수량",
        "⚠ 알람",
    ]
    assert "MOQ" not in cols
    first = df.iloc[0]
    assert first["SKU"] == "SKU-AIR"
    assert first["상품명"] == "Need Air"
    assert first["현지재고"] == 80
    assert first["월평균판매량"] == 40
    assert first["운송재고"] == 30
    assert first["① 현지\n(M)"] == 2.0
    assert first["② 운송중\n(M)"] == 0.8
    assert first["③ 합산\n(M)"] == 2.8
    assert first["안전재고 기준(M)"] == 3
    assert first["확보재고 수량"] == 110
    assert first["안전재고 목표수량"] == 120
    assert first["⚠ 알람"] == "-"
    assert first["미입고\n현황"] == 7
    assert first["현지창고 재고"] == 50
    assert first["발주총\n필요수량"] == 10
    assert round(float(first["총발주금액"]), 6) == round((450.25 / 90) * 10 * 1500, 6)
    assert first["운송수단\n추천"] == "\u2460 \ud56d\uacf5 \uae34\uae09"
    assert first["항공\n필요량"] == 10
    assert first["철송\n필요량"] == 0
    assert first["해운\n필요량"] == 0

    second = df.iloc[1]
    assert second["① 현지\n(M)"] == "\u2014"
    assert second["발주총\n필요수량"] == 0
    assert second["운송수단\n추천"] == "-"
    assert second["항공\n필요량"] == 0


def test_order_sheet_sku_column_uses_product_code_not_product_name():
    report = _report_df().copy()
    report.loc[1, PRODUCT_NAME] = report.loc[0, PRODUCT_NAME]

    df = export_excel_order_sheet.build_order_sheet_df(report, _settings())

    assert df["SKU"].tolist() == ["SKU-AIR", "SKU-NOSALES"]


def test_order_sheet_transport_recommendation_matches_reference_boundaries():
    recommendations = export_excel_order_sheet._order_sheet_transport_recommendation(
        pd.Series([1, 1, 1, 1, 1, 1, 1, 0]),
        pd.Series([1, 15, 16, 35, 36, 80, 81, 10]),
        _settings()["lead_times"],
    )

    assert recommendations.tolist() == [
        "\u2460 \ud56d\uacf5 \uae34\uae09",
        "\u2460 \ud56d\uacf5 \uae34\uae09",
        "\u2461 \ud56d\uacf5",
        "\u2461 \ud56d\uacf5",
        "\u2462 \ucca0\uc1a1",
        "\u2462 \ucca0\uc1a1",
        "\u2463 \ud574\uc6b4 \uac00\ub2a5",
        "-",
    ]


def test_order_sheet_format_matches_reference_header_blocks_and_formulas():
    order_df = export_excel_order_sheet.build_order_sheet_df(_report_df(), _settings())
    top_sales_df = export_excel_order_sheet.build_order_sheet_top_sales_df(_report_df())
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, order_df, start_row=10)

    export_excel_order_sheet.format_order_sheet(ws, _settings(), top_sales_df)

    headers = {ws.cell(10, col).value: col for col in range(1, ws.max_column + 1)}
    assert "MOQ" not in headers
    assert "=에 따른 발주금액" not in headers
    assert "총발주금액" in headers
    signal_col = headers["발주\n신호등"]
    recommendation_col = headers["운송수단\n추천"]
    required_col = headers["발주총\n필요수량"]
    amount_col = headers["총발주금액"]
    depletion_col = headers["고갈\n예정일"]
    alert_col = headers["⚠ 알람"]
    open_po_col = headers["미입고\n현황"]

    assert ws["A3"].value.startswith("\u25bc \ud30c\ub77c\ubbf8\ud130 \uc124\uc815")
    merged_ranges = {str(merged_range) for merged_range in ws.merged_cells.ranges}
    assert "A3:F3" in merged_ranges
    assert "A3:H3" not in merged_ranges
    assert ws["G3"].fill.fill_type is None
    assert ws["H3"].fill.fill_type is None
    parameter_labels = [ws.cell(4, col).value for col in range(1, 9)]
    assert parameter_labels == [
        "\uc548\uc804\uc7ac\uace0(M)",
        "\ud56d\uacf5 L/T(\uc77c)",
        "\ud574\uc6b4 L/T(\uc77c)",
        "\ucca0\uc1a1\nL/T(\uc77c)",
        "\ud2b8\ub7ed\nL/T(\uc77c)",
        "\uae30\uc900\uc77c",
        None,
        None,
    ]
    assert ws["B5"].value == 15
    assert ws["C5"].value == 80
    assert ws["F5"].value == date(2026, 5, 20)
    assert ws.column_dimensions["F"].width >= 12
    assert ws["A6"].value is None
    assert ws.row_dimensions[6].hidden is False
    assert ws["A4"].font.size <= 9
    assert ws["A3"].font.size <= 11
    assert ws["K3"].value == "\uc6b4\uc1a1\uc218\ub2e8 \ucd94\ucc9c \uae30\uc900"
    assert "L3:M3" in {str(merged_range) for merged_range in ws.merged_cells.ranges}
    assert ws.column_dimensions["K"].width >= 19
    assert ws.column_dimensions["L"].width >= 24
    assert ws.column_dimensions["M"].width >= 18
    assert ws.row_dimensions[4].height >= 34
    assert ws.row_dimensions[5].height >= 22
    assert ws["K6"].value == "\u2462 \ucca0\uc1a1"
    assert ws["K7"].value == "\u2463 \ud574\uc6b4 \uac00\ub2a5"
    assert ws["S2"].value == "SKU"
    assert ws["U2"].value == "\ucd5c\uadfc 3\uac1c\uc6d4 \ud310\ub9e4\ub7c9"
    assert ws.column_dimensions["U"].width >= 20
    assert ws["S8"].value == "\ud569\uacc4 (\uc0c1\uc704 5\uac1c)"
    assert ws.cell(9, required_col).value == f"=SUM({get_column_letter(required_col)}11:{get_column_letter(required_col)}{ws.max_row})"
    assert ws.cell(9, amount_col).value == f"=SUM({get_column_letter(amount_col)}11:{get_column_letter(amount_col)}{ws.max_row})"
    assert ws.cell(9, amount_col).number_format == "#,##0"
    assert ws.cell(9, required_col).fill.fgColor.rgb == "00FCE4D6"
    assert ws.cell(9, amount_col).font.color.rgb == "00C00000"
    assert ws.cell(10, 1).value == "No"
    assert alert_col == open_po_col - 1
    assert ws.cell(10, alert_col).value == "⚠ 알람"
    assert ws.cell(10, alert_col).fill.fgColor.rgb == "00FFFF00"
    assert ws.cell(11, alert_col).value == '=IF(E11<=0,"OOS",IF(Q11>0,"발주필요","-"))'
    assert ws.protection.sheet is True
    assert ws.cell(11, alert_col).protection.hidden is True
    assert ws.cell(11, alert_col).protection.locked is True
    assert ws.cell(11, required_col).protection.locked is False
    assert ws.cell(10, signal_col).value == "발주\n신호등"
    assert ws.cell(11, 1).border.left.style == "thin"
    assert ws.cell(11, 2).border.right.style == "thin"
    assert ws.cell(12, signal_col).border.bottom.style == "thin"
    assert str(ws.cell(11, signal_col).value).startswith("=IF(")
    assert "● 안전" in str(ws.cell(11, signal_col).value)
    assert "발주완료" not in str(ws.cell(11, signal_col).value)
    assert ws.cell(11, recommendation_col).font.color.rgb == "00C00000"
    assert ws.cell(12, recommendation_col).value == "-"
    assert ws.cell(11, depletion_col).number_format == "yyyy-mm-dd"
    for cover_month_header in ["\u2460 \ud604\uc9c0\n(M)", "\u2461 \uc6b4\uc1a1\uc911\n(M)", "\u2462 \ud569\uc0b0\n(M)"]:
        assert ws.cell(11, headers[cover_month_header]).number_format == "#,##0.##"
    assert ws.freeze_panes == "D11"


def test_order_sheet_fast_mode_keeps_data_cell_borders():
    order_df = export_excel_order_sheet.build_order_sheet_df(_report_df(), _settings())
    large_order_df = pd.concat([order_df] * 151, ignore_index=True)
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, large_order_df, start_row=10)

    export_excel_order_sheet.format_order_sheet(ws, _settings(), pd.DataFrame())

    headers = {ws.cell(10, col).value: col for col in range(1, ws.max_column + 1)}
    signal_col = headers["발주\n신호등"]
    deep_row = 310

    assert ws.cell(deep_row, 1).border.left.style == "thin"
    assert ws.cell(deep_row, 2).border.right.style == "thin"
    assert ws.cell(deep_row, signal_col).border.bottom.style == "thin"


def test_generated_excel_places_stock_eta_sheet_first_and_removes_eta_timeline():
    settings = validation.default_settings_for_validation()
    review = validation.apply_order_review_filters(order_review.order_review_df(settings), settings)

    xlsx = export_excel.generate_order_review_excel_fast(settings, review)
    wb = load_workbook(BytesIO(xlsx), read_only=True)

    assert wb.sheetnames[:5] == ["재고 ETA", "발주", "보고서", "입고예정", "확인필요"]
    assert "OUT.발주" not in wb.sheetnames
    assert "ETA 타임라인_검토" not in wb.sheetnames


def test_generated_excel_includes_report_sheet_layout():
    settings = validation.default_settings_for_validation()
    review = validation.apply_order_review_filters(order_review.order_review_df(settings), settings)
    xlsx = export_excel.generate_order_review_excel_fast(settings, review)
    wb = load_workbook(BytesIO(xlsx))
    ws = wb["보고서"]

    assert ws.freeze_panes == "A13"
    assert "U1:Y1" in {str(merged_range) for merged_range in ws.merged_cells.ranges}
    assert "P11:S11" in {str(merged_range) for merged_range in ws.merged_cells.ranges}
    assert "O11:R11" not in {str(merged_range) for merged_range in ws.merged_cells.ranges}
    assert ws.row_dimensions[1].height == 18
    assert ws.row_dimensions[6].height == 18
    assert ws.row_dimensions[8].height == 18
    assert ws["U1"].value == "▼ 파라미터 설정  | 솔팅 가능"
    assert ws["U1"].fill.fgColor.rgb == "004F6228"
    assert ws["U1"].font.color.rgb == "00FFFFFF"
    assert ws["U2"].value == "안전재고(M)"
    assert ws["V2"].value == "AirL/T(일)"
    assert ws["W2"].value == "해운L/T(일)"
    assert ws["X2"].value == "철송L/T(일)"
    assert ws["Y2"].value == "트럭L/T(일)"
    assert ws["V4"].value == "환율(KRW/EUR)"
    assert ws["V5"].value == settings["eur_krw_rate"]
    assert ws["V5"].number_format == "#,##0.00"
    assert ws["U3"].fill.fgColor.rgb == "00FFF2CC"
    assert ws["V3"].fill.fgColor.rgb == "00FFF2CC"
    assert ws["W3"].fill.fgColor.rgb == "00FFF2CC"
    assert ws["X3"].fill.fgColor.rgb == "00FFF2CC"
    assert ws["Y3"].fill.fgColor.rgb == "00FFF2CC"
    assert ws["V6"].fill.fgColor.rgb == "0092D050"
    assert ws["V7"].fill.fgColor.rgb == "00E2F0D9"
    assert ws["U6"].value is None
    assert ws["U6"].border.left is None or ws["U6"].border.left.style is None
    assert ws["U9"].border.bottom is None or ws["U9"].border.bottom.style is None
    assert ws["V6"].value == "총 발주 필요수량"
    assert ws["W6"].value == "총 발주금액(EUR)"
    assert ws["X6"].value == "총 발주 금액(KRW)"
    assert ws["V8"].value == "미입고 해소 후 총 발주 필요수량"
    assert ws["W8"].value == "미입고 해소 후 총 발주금액(EUR)"
    assert ws["X8"].value == "(미입고 해소 후)총 발주 금액(KRW)"
    assert str(ws["V7"].value).startswith("=SUMIF(")
    assert "'_보고서원본'!$N$2:$N$" in str(ws["V7"].value)
    assert "'_보고서원본'!$H$2:$H$" in str(ws["V7"].value)
    assert "'_보고서원본'!$I$2:$I$" in str(ws["W7"].value)
    assert "'_보고서원본'!$J$2:$J$" in str(ws["X7"].value)
    assert ws["X7"].number_format == "#,##0"
    assert "'_보고서원본'!$K$2:$K$" in str(ws["V9"].value)
    assert "'_보고서원본'!$L$2:$L$" in str(ws["W9"].value)
    assert "'_보고서원본'!$M$2:$M$" in str(ws["X9"].value)
    assert ws["X9"].number_format == "#,##0"
    for coord in ("V7", "W7", "X7", "V9", "W9", "X9"):
        assert ws[coord].protection.hidden is True
        assert ws[coord].protection.locked is True
    assert ws["A12"].value == "코드"
    assert ws["C12"].value == "SKU"
    assert ws["D12"].value == "브랜드"
    assert ws["E12"].value == "미입고현황"
    assert ws["F12"].value == "현지 창고재고"
    assert ws["G12"].value == "⚠ 알람"
    assert ws["G12"].fill.fgColor.rgb == "00FFFF00"
    assert ws["H12"].value == "발주 필요수량"
    assert ws["I12"].value == "=발주 금액(EUR)"
    assert ws["J12"].value == "=발주 금액(KRW)"
    assert ws["K12"].value == "(미입고 해소 후)발주 필요수량"
    assert ws["L12"].value == "=(미입고 해소 후)발주 금액(EUR)"
    assert ws["M12"].value == "=(미입고 해소 후)발주 금액(KRW)"
    assert str(ws.auto_filter.ref).startswith("A12:M")
    assert "ReportBrandFilter" in ws.tables
    assert ws["D13"].border.left.style == "thin"
    assert ws["D13"].border.left.color.rgb == "00000000"
    assert ws["M13"].border.bottom.style == "thin"
    assert ws.column_dimensions["A"].hidden is False
    assert ws.column_dimensions["D"].hidden is False
    assert ws.column_dimensions["E"].hidden is False
    assert ws.column_dimensions["M"].hidden is False
    assert ws.column_dimensions["M"].width >= 18
    assert ws.column_dimensions["N"].hidden is True
    assert ws.column_dimensions["U"].hidden is False
    assert ws["B2"].value == "브랜드 조회"
    assert ws["B3"].value == "브랜드 검색"
    assert ws["C3"].value == "전체"
    assert ws["B4"].value == "검색 후보"
    assert str(ws["B5"].value).startswith("=IFERROR(INDEX('_브랜드_선택값'!$A$2:$A")
    assert "MATCH(ROWS($B$5:B5),'_브랜드_선택값'!$B$2:$B" in str(ws["B5"].value)
    assert ws["H2"].border.left is None or ws["H2"].border.left.style is None
    assert "_보고서원본" in wb.sheetnames
    assert "_브랜드_선택값" in wb.sheetnames
    assert wb["_보고서원본"].sheet_state == "hidden"
    assert wb["_브랜드_선택값"].sheet_state == "hidden"
    assert str(ws["A13"].value).startswith("=IFERROR(INDEX('_보고서원본'!$A$2:$A$")
    assert "MATCH(ROWS($A$13:A13),'_보고서원본'!$N$2:$N$" in str(ws["A13"].value)
    assert str(ws["D13"].value).startswith("=IFERROR(INDEX('_보고서원본'!$D$2:$D$")
    assert "MATCH(ROWS($A$13:A13),'_보고서원본'!$N$2:$N$" in str(ws["D13"].value)
    source_ws = wb["_보고서원본"]
    assert source_ws["B2"].value == 1
    assert source_ws["G1"].value == "⚠ 알람"
    assert source_ws["G2"].value == common.ALERT_URGENT_LABEL
    assert str(ws["G13"].value).startswith("=IFERROR(INDEX('_보고서원본'!$G$2:$G$")
    assert ws.protection.sheet is True
    assert ws["G13"].protection.hidden is True
    assert ws["G13"].protection.locked is True
    assert ws["H13"].protection.locked is False
    assert source_ws["J2"].value == source_ws["I2"].value * settings["eur_krw_rate"]
    assert source_ws["M2"].value == source_ws["L2"].value * settings["eur_krw_rate"]
    assert ws["J13"].number_format == "#,##0"
    assert ws["M13"].number_format == "#,##0"
    helper_formula = str(source_ws["N2"].value)
    assert 'SEARCH(SUBSTITUTE(TRIM(\'보고서\'!$C$3)," ",""),SUBSTITUTE(TRIM($D2)," ",""))' in helper_formula
    brand_ws = wb["_브랜드_선택값"]
    assert brand_ws["B1"].value == "_검색순번"
    assert 'SEARCH(SUBSTITUTE(TRIM(\'보고서\'!$C$3)," ",""),SUBSTITUTE(TRIM($A2)," ",""))' in str(brand_ws["B2"].value)
    brand_validations = [
        validation
        for validation in ws.data_validations.dataValidation
        if "C3" in str(validation.sqref)
    ]
    assert not brand_validations
    return
    report = order_review_report.build_order_review_report_df(review)
    assert ws["C13"].value == report.iloc[0]["상품코드"]
    assert ws["D13"].value == report.iloc[0]["상품명"]
