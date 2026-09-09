from __future__ import annotations

from datetime import date
from pathlib import Path
import re

import pandas as pd
import pytest
from openpyxl import Workbook

from core import common, export_excel_order_sheet, export_excel_report, export_excel_util, order_review
from core.loaders import order_review_settings_signature
from core.session import SessionContext
from core.transport import (
    normalize_transport_code,
    prepare_shipping,
    transport_recommendation_by_depletion_days,
)
from core.validation import default_settings_for_validation


def test_uk_korean_air_labels_require_and_preserve_explicit_service_type():
    assert normalize_transport_code("항공(DIR)", "UK") == "AIR_DIR"
    assert normalize_transport_code("항공(DIR·직항)", "UK") == "AIR_DIR"
    assert normalize_transport_code("항공(T/S)", "UK") == "AIR_TS"
    assert normalize_transport_code("항공(T/S·환적)", "UK") == "AIR_TS"
    assert normalize_transport_code("항공", "UK") == "AIR"


def test_prepare_shipping_uses_exact_uk_air_method_and_preserves_source_eta():
    source = pd.DataFrame(
        [
            {"itemcode": "DIR", "qty": 1, "shipdate": "2026-07-23", "mode": "항공(DIR)"},
            {"itemcode": "TS", "qty": 1, "shipdate": "2026-07-23", "mode": "항공(T/S)"},
            {
                "itemcode": "GENERIC",
                "qty": 1,
                "shipdate": "2026-07-23",
                "ETA": "2026-08-01",
                "mode": "항공",
            },
        ]
    )

    prepared = prepare_shipping(source, default_settings_for_validation("UK"))

    assert prepared["운송수단 코드"].tolist() == ["AIR_DIR", "AIR_TS", "AIR"]
    assert prepared["운송 L/T"].tolist()[:2] == [5, 12]
    assert pd.isna(prepared.loc[2, "운송 L/T"])
    assert prepared["ETA"].tolist() == ["2026-07-28", "2026-08-04", "2026-08-01"]
    assert prepared.loc[2, "ETA 구분"] == "원본ETA"


def test_prepare_shipping_prefers_explicit_actual_arrival_over_eta_and_default():
    source = pd.DataFrame(
        [
            {
                "itemcode": "ACTUAL",
                "qty": 1,
                "shipdate": "2026-07-01",
                "ETA": "2026-07-20",
                "실제 입고일": "2026-07-18",
                "mode": "OCEAN",
            }
        ]
    )

    prepared = prepare_shipping(source, default_settings_for_validation("PL"))

    assert prepared.loc[0, "ETA"] == "2026-07-18"
    assert prepared.loc[0, "실제 도착일"] == "2026-07-18"
    assert prepared.loc[0, "ETA 구분"] == "실제입고일"
    assert prepared.loc[0, "운송 L/T"] == 70


def test_dynamic_recommendation_never_falls_back_to_pl_only_methods():
    usa = default_settings_for_validation("USA")
    label, days, _ = transport_recommendation_by_depletion_days(
        20,
        usa["applied_lead_times"],
        lead_times_by_code=usa["lead_times_by_code"],
        entity_code="USA",
    )
    assert label == "② 항공"
    assert days == 5
    assert "철송" not in label

    hq = default_settings_for_validation("HQ")
    label, days, reason = transport_recommendation_by_depletion_days(
        100,
        hq["applied_lead_times"],
        lead_times_by_code=hq["lead_times_by_code"],
        entity_code="HQ",
    )
    assert (label, days) == ("판단불가", 0)
    assert "지원되는 운송수단 없음" in reason


def test_pl_recommendation_keeps_existing_rail_boundary_and_ignores_trucking_tie():
    pl = default_settings_for_validation("PL")
    label, days, _ = transport_recommendation_by_depletion_days(
        31,
        pl["applied_lead_times"],
        lead_times_by_code=pl["lead_times_by_code"],
        entity_code="PL",
    )
    assert (label, days) == ("③ 철송", 30)


def test_uk_rough_distribution_requires_air_service_selection_instead_of_using_five_days():
    uk_policy = order_review._rough_transport_lead_time_policy(
        default_settings_for_validation("UK")
    )
    assert uk_policy == (None, None, True, True)

    pl_policy = order_review._rough_transport_lead_time_policy(
        default_settings_for_validation("PL")
    )
    assert pl_policy == (15, 30, True, False)


def test_order_sheet_recommendation_uses_entity_method_set():
    usa = default_settings_for_validation("USA")
    recommendations = export_excel_order_sheet._order_sheet_transport_recommendation(
        pd.Series([1, 1]),
        pd.Series([20, 31]),
        usa["lead_times"],
        lead_times_by_code=usa["lead_times_by_code"],
        entity_code="USA",
    )
    assert recommendations.tolist() == ["② 항공", "③ 해운 가능"]


def test_usa_order_sheet_criteria_only_lists_air_and_ocean():
    usa = default_settings_for_validation("USA")
    criteria = export_excel_order_sheet._order_sheet_transport_criteria(usa)

    assert criteria == [
        ("운송수단 추천 기준", "조건"),
        ("① 항공 긴급", "고갈까지 ≤ 항공 L/T일"),
        ("② 항공", "항공 L/T < 고갈 ≤ 해운 L/T"),
        ("③ 해운 가능", "해운 L/T < 고갈"),
        ("● 발주불필요", "안전재고 충족"),
    ]

    wb = Workbook()
    ws = wb.active
    blank_order_row = pd.DataFrame([{column: "" for column in common.ORDER_SHEET_COLUMNS}])
    export_excel_util.append_df(ws, blank_order_row, start_row=10)
    export_excel_order_sheet.format_order_sheet(ws, usa, pd.DataFrame())

    visible_criteria = [ws.cell(row, column).value for row in range(3, 9) for column in range(11, 14)]
    assert "철송" not in " ".join(str(value or "") for value in visible_criteria)
    assert ws["K4"].value == "① 항공 긴급"
    assert ws["K5"].value == "② 항공"
    assert ws["K6"].value == "③ 해운 가능"
    assert ws["K7"].value == "● 발주불필요"
    assert ws["K8"].value is None


def test_usa_order_sheet_omits_unsupported_transport_need_columns():
    usa = default_settings_for_validation("USA")
    report = pd.DataFrame(
        {
            "\uc0c1\ud488\ucf54\ub4dc": ["SKU-USA"],
            "\uc0c1\ud488\uba85": ["USA product"],
            "\ube0c\ub79c\ub4dc": ["Brand"],
        }
    )

    df = export_excel_order_sheet.build_order_sheet_df(report, usa)

    air_column = "\ud56d\uacf5\n\ud544\uc694\ub7c9"
    rail_column = "\ucca0\uc1a1\n\ud544\uc694\ub7c9"
    sea_column = "\ud574\uc6b4\n\ud544\uc694\ub7c9"
    assert air_column in df.columns
    assert sea_column in df.columns
    assert rail_column not in df.columns
    # The rest of the fixed layout must survive untouched.
    assert list(df.columns) == [
        column for column in common.ORDER_SHEET_COLUMNS if column != rail_column
    ]


def test_pl_order_sheet_keeps_all_transport_need_columns():
    pl = default_settings_for_validation("PL")
    report = pd.DataFrame(
        {
            "\uc0c1\ud488\ucf54\ub4dc": ["SKU-PL"],
            "\uc0c1\ud488\uba85": ["PL product"],
            "\ube0c\ub79c\ub4dc": ["Brand"],
        }
    )

    df = export_excel_order_sheet.build_order_sheet_df(report, pl)

    assert list(df.columns) == list(common.ORDER_SHEET_COLUMNS)


def test_order_review_cache_signature_includes_entity_and_method_overrides():
    context = SessionContext()
    pl = default_settings_for_validation("PL")
    usa = default_settings_for_validation("USA")
    overridden = default_settings_for_validation("PL")
    overridden["lead_time_overrides"] = {"OCEAN": 71}
    overridden["lead_times_by_code"] = {**overridden["lead_times_by_code"], "OCEAN": 71}

    assert order_review_settings_signature(pl, context=context) != order_review_settings_signature(
        usa,
        context=context,
    )
    assert order_review_settings_signature(pl, context=context) != order_review_settings_signature(
        overridden,
        context=context,
    )


def test_reference_dates_remain_calendar_day_additions():
    settings = default_settings_for_validation("UK")
    source = pd.DataFrame(
        [{"itemcode": "DIR", "qty": 1, "shipdate": date(2026, 7, 23), "mode": "AIR_DIR"}]
    )
    prepared = prepare_shipping(source, settings)
    assert prepared.loc[0, "ETA"] == "2026-07-28"


@pytest.mark.parametrize(
    ("entity_code", "transport_code", "expected_eta"),
    [
        ("PL", "OCEAN", "2026-10-01"),
        ("USA", "OCEAN", "2026-08-22"),
        ("VN", "AIR", "2026-07-29"),
        ("UK", "AIR_DIR", "2026-07-28"),
        ("UK", "AIR_TS", "2026-08-04"),
    ],
)
def test_entity_method_lead_time_is_connected_to_shipping_eta(
    entity_code: str,
    transport_code: str,
    expected_eta: str,
):
    prepared = prepare_shipping(
        pd.DataFrame(
            [
                {
                    "itemcode": f"{entity_code}-{transport_code}",
                    "qty": 1,
                    "shipdate": "2026-07-23",
                    "mode": transport_code,
                }
            ]
        ),
        default_settings_for_validation(entity_code),
    )
    assert prepared.loc[0, "ETA"] == expected_eta


def test_uk_direct_and_transshipment_shipping_eta_are_seven_days_apart():
    prepared = prepare_shipping(
        pd.DataFrame(
            [
                {"itemcode": "DIR", "qty": 1, "shipdate": "2026-07-23", "mode": "AIR_DIR"},
                {"itemcode": "TS", "qty": 1, "shipdate": "2026-07-23", "mode": "AIR_TS"},
            ]
        ),
        default_settings_for_validation("UK"),
    )
    eta = pd.to_datetime(prepared["ETA"])
    assert eta.iloc[1] - eta.iloc[0] == pd.Timedelta(days=7)


def test_export_summary_keeps_pl_official_values_and_never_fills_hq_from_pl():
    pl_summary = export_excel_report.build_order_review_report_summary(
        pd.DataFrame(),
        default_settings_for_validation("PL"),
    )
    assert (
        pl_summary["air_lt"],
        pl_summary["sea_lt"],
        pl_summary["rail_lt"],
        pl_summary["truck_lt"],
    ) == (15, 70, 30, 30)

    hq_summary = export_excel_report.build_order_review_report_summary(
        pd.DataFrame(),
        default_settings_for_validation("HQ"),
    )
    assert (
        hq_summary["air_lt"],
        hq_summary["sea_lt"],
        hq_summary["rail_lt"],
        hq_summary["truck_lt"],
    ) == ("-", "-", "-", "-")


def test_order_logic_v2_screen_transport_codes_match_entity_lead_time_config():
    """The V2 settings card must not offer a transport an entity does not run.

    ``core.lead_times.LEAD_TIME_CONFIG`` is the reviewed source for which
    transports a legal entity operates. The V2 screen keeps its own AIR/RAIL/SEA
    list, so this guards the two from drifting apart (for example USA rail).
    """

    from core.lead_times import lead_time_methods

    source = (
        Path(__file__).resolve().parents[1]
        / "frontend"
        / "components"
        / "redesign"
        / "screens"
        / "order-v2"
        / "useNewOrderLogicScreenModel.tsx"
    ).read_text(encoding="utf-8")

    block = re.search(
        r"const ENTITY_TRANSPORT_CODES[^=]*=\s*\{(.*?)\};",
        source,
        re.DOTALL,
    )
    assert block is not None, "V2 화면에 법인별 운송수단 목록이 없습니다."

    screen_codes = {
        entity: set(re.findall(r'"([A-Z_]+)"', codes))
        for entity, codes in re.findall(r"(\w+):\s*\[([^\]]*)\]", block.group(1))
    }
    assert set(screen_codes) == {"PL", "USA", "HQ"}

    # The V2 engine names ocean "SEA" while the entity reference uses "OCEAN".
    v2_name = {"OCEAN": "SEA", "AIR": "AIR", "RAIL": "RAIL"}
    for entity, codes in screen_codes.items():
        supported = {
            v2_name[method.transport_code]
            for method in lead_time_methods(entity)
            if method.transport_code in v2_name
        }
        assert codes == supported, (
            f"{entity} 화면 운송수단 {sorted(codes)}이 "
            f"법인 리드타임 구성 {sorted(supported)}과 다릅니다."
        )

    assert "RAIL" not in screen_codes["USA"]
    assert "RAIL" in screen_codes["PL"]
