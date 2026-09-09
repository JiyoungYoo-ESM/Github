from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.services.order_logic_v2_excel import (
    CHECK_REQUIRED_HEADER_ROW,
    LOGISTICS_HEADER_ROW,
    ORDER_HEADER_ROW,
    generate_order_logic_v2_excel,
)


AMOUNT_HEADERS = {
    "단가(EUR)",
    "제안금액(EUR)",
    "단가(KRW)",
    "제안금액(KRW)",
}
EXCLUDED_LEGACY_HEADERS = {
    "MOQ입력",
    "상한금액",
}
OPERATIONAL_HEADERS = {
    "발주까지 여유\n(주, 평균)",
    "예상 발주일\n(평균)",
    "보수적 발주 여유\n(주)",
    "조기경보일\n(CV·서비스수준 반영)",
}


def _result_payload() -> dict:
    return {
        "job_id": "job-v2-001",
        "logic_version": "2.0.0-beta.2",
        "as_of": "2026-07-27",
        "policy_mode": "shortage",
        "calculated_at": "2026-07-27T17:45:00+09:00",
        "period_start": "2026-04-27",
        "period_end": "2026-07-26",
        "source_fetched_at": "2026-07-27T17:30:00+09:00",
        "exchange_rate_krw": 1666.16,
        "exchange_rate_date": "2026-07-28",
        "exchange_rate_checked_date": "2026-07-29",
        "exchange_rate_source": "api",
        "settings": {
            "service_level": 0.95,
            "review_cycle_weeks": 2,
            "lead_time_sigma_weeks": 2.18,
            "lt_air_days": 16.4,
            "lt_rail_days": 36.6,
            "lt_sea_days": 72.9,
        },
        "rows": [
            {
                "sku_code": "SKU-001",
                "product_name": "테스트 선크림",
                "brand": "Silicon2",
                "barcode": "8800000000001",
                "sales_13w_qty": 1567,
                "data_status": "정상",
                "grade": "MAJOR",
                "demand_avg": 120.5,
                "demand_sigma": 86.76,
                "cv": 0.72,
                "z_applied": 1.68,
                "safety_stock": 410,
                "reorder_point": 1391,
                "target_stock": 1632,
                "open_qty": 50,
                "incoming_qty": 100,
                "pnfm_qty": 30,
                "inbound_progress_qty": 20,
                "inbound_completed_qty": 10,
                "eu_available_qty": 50,
                "transit_qty": 200,
                "next_eta": "2026-08-10",
                "shipping_schedule": [
                    {"eta": "2026-07-20", "qty": 25},
                    {"eta": "2026-08-10", "qty": 175},
                ],
                "shipping_eta_details": [
                    {
                        "eta": "2026-07-20",
                        "qty": 25,
                        "eta_status": "원천 ETA",
                        "ship_date": "2026-06-01",
                        "transport_label": "해운",
                        "lead_time_days": None,
                    },
                    {
                        "eta": "2026-08-10",
                        "qty": 175,
                        "eta_status": "출고일 추정 ETA",
                        "ship_date": "2026-07-04",
                        "transport_label": "철송",
                        "lead_time_days": 36.6,
                    },
                ],
                "eta_actual_qty": 25,
                "eta_estimated_qty": 175,
                "eta_missing_qty": 0,
                "local_available_qty": 300,
                "inventory_position": 650,
                "depletion_weeks": 4.1,
                "suggested_qty": 1000,
                "upper_suggested_qty": 1100,
                "confirmed_qty": 900,
                "memo": "담당자 검토 완료",
                "order_signal": "발주",
                "unit_price_eur": 4.25,
                "suggested_amount_eur": 4250.0,
                "confirmed_amount_eur": 3825.0,
                "unit_price_krw": 6_960.0,
                "suggested_amount_krw": 6_960_000.0,
                "confirmed_amount_krw": 6_264_000.0,
                "warnings": ["ETA 확인", "간헐수요 점검"],
            },
            {
                "sku_code": "SKU-002",
                "product_name": "테스트 세럼",
                "brand": "Silicon2",
                "barcode": "8800000000002",
                "sales_13w_qty": 260,
                "data_status": "⚠확인(간헐)",
                "grade": "MINOR",
                "demand_avg": 20,
                "demand_sigma": 28,
                "cv": 1.4,
                "z_applied": 1.28,
                "safety_stock": 130,
                "reorder_point": 293,
                "target_stock": 333,
                "open_qty": 0,
                "incoming_qty": 0,
                "pnfm_qty": 0,
                "inbound_progress_qty": 0,
                "inbound_completed_qty": 0,
                "eu_available_qty": 10,
                "transit_qty": 0,
                "next_eta": None,
                "shipping_schedule": [],
                "shipping_eta_details": [
                    {
                        "eta": None,
                        "qty": 12,
                        "eta_status": "ETA 미확인",
                        "ship_date": "2026-07-27",
                        "transport_label": None,
                        "lead_time_days": None,
                    }
                ],
                "eta_actual_qty": 0,
                "eta_estimated_qty": 0,
                "eta_missing_qty": 12,
                "local_available_qty": 25,
                "inventory_position": 35,
                "depletion_weeks": 1.25,
                "suggested_qty": 300,
                "upper_suggested_qty": 300,
                "confirmed_qty": 0,
                "memo": None,
                "order_signal": "확인후발주",
                "unit_price_eur": 8.5,
                "suggested_amount_eur": 2550.0,
                "confirmed_amount_eur": 0.0,
                "unit_price_krw": 13_920.0,
                "suggested_amount_krw": 4_176_000.0,
                "confirmed_amount_krw": 0.0,
                "warnings": [],
            },
        ],
    }


def _headers(sheet) -> list[str]:
    return [
        str(cell.value)
        for cell in sheet[ORDER_HEADER_ROW]
        if cell.value not in (None, "")
    ]


def test_export_is_valid_xlsx_with_expected_sheets_and_layout() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=True)

    assert content.startswith(b"PK")
    with ZipFile(BytesIO(content)) as archive:
        members = set(archive.namelist())
        assert "[Content_Types].xml" in members
        assert "xl/workbook.xml" in members
        assert "xl/worksheets/sheet1.xml" in members
        assert "xl/worksheets/sheet2.xml" in members

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == [
        "발주제안",
        "물류전망",
        "확인필요",
        "_계산기준",
        "_운송원천",
    ]
    assert workbook["_계산기준"].sheet_state == "hidden"
    assert workbook["_운송원천"].sheet_state == "hidden"

    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    assert "적용 시나리오 : 쇼티지 방어" in sheet["A2"].value
    assert headers[:6] == ["상품코드", "상품명", "브랜드", "상태", "등급", "판매량\n(13주)"]
    assert "월평균\n판매량" in headers
    assert "주평균\n판매량" in headers
    assert "재고보유\n개월수(MOI)" in headers
    assert "제안수량\n(미입고 인정)" in headers
    assert "상한수량\n(미입고 제외)" not in headers
    assert "확정수량\n(담당자 발주)" not in headers
    assert not any(header.startswith("확정금액") for header in headers)
    assert headers[-4:] == ["단가(EUR)", "제안금액(EUR)", "단가(KRW)", "제안금액(KRW)"]
    assert "본사 창고\n가용재고" in headers
    assert "현지 가용재고" in headers
    assert "본사 EU창고\n가용재고" not in headers
    assert "현지가용(SKO)" not in headers
    inbound_index = headers.index("① 미입고 수량")
    assert headers[inbound_index : inbound_index + 4] == [
        "① 미입고 수량",
        "② PNFM확정 수량",
        "③ 입고진행중 수량",
        "④ 입고완료 수량",
    ]
    assert AMOUNT_HEADERS <= set(headers)
    assert OPERATIONAL_HEADERS <= set(headers)
    assert "운송수단\n추천" not in headers
    assert not any(
        excluded in header
        for excluded in EXCLUDED_LEGACY_HEADERS
        for header in headers
    )
    next_eta_column = headers.index("다음 입고예정\n(운송중 최소 ETA)") + 1
    early_warning_column = headers.index("조기경보일\n(CV·서비스수준 반영)") + 1
    assert sheet.column_dimensions[get_column_letter(next_eta_column)].width == 24
    assert sheet.column_dimensions[get_column_letter(early_warning_column)].width == 28
    assert sheet.freeze_panes == f"A{ORDER_HEADER_ROW + 1}"
    expected_last_column = get_column_letter(len(headers))
    assert sheet.auto_filter.ref == (
        f"A{ORDER_HEADER_ROW}:{expected_last_column}{ORDER_HEADER_ROW + 2}"
    )
    assert "BETA" in sheet["A1"].value
    assert "참고값" in sheet["A4"].value
    assert "unit_cost_krw" in sheet["A4"].value
    assert "경고" not in headers
    workbook.close()


def test_export_clears_stale_template_group_bar_above_headers() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=False)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)

    assert not any(
        merged_range.min_row <= 6 <= merged_range.max_row
        for merged_range in sheet.merged_cells.ranges
    )
    for column_index in range(1, len(headers) + 1):
        cell = sheet.cell(6, column_index)
        assert cell.value is None
        assert cell.fill.fill_type is None
        assert cell.style_id == sheet.cell(5, column_index).style_id
    workbook.close()


def test_export_does_not_add_amount_group_label_above_headers() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]

    assert not any(
        merged_range.min_row <= 6 <= merged_range.max_row
        for merged_range in sheet.merged_cells.ranges
    )
    assert all(cell.value is None for cell in sheet[6])
    workbook.close()


def test_export_writes_cms_inbound_status_values_next_to_incoming() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=False)
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}
    first_row = ORDER_HEADER_ROW + 1

    assert sheet.cell(first_row, index["① 미입고 수량"]).value == 50
    assert sheet.cell(first_row, index["② PNFM확정 수량"]).value == 30
    assert sheet.cell(first_row, index["③ 입고진행중 수량"]).value == 20
    assert sheet.cell(first_row, index["④ 입고완료 수량"]).value == 10
    assert sheet.cell(first_row, index["② PNFM확정 수량"]).number_format == "#,##0"
    workbook.close()


def test_export_marks_skus_without_inbound_status_source_as_dash() -> None:
    payload = _result_payload()
    payload["rows"] = [
        {**payload["rows"][0], "inbound_status_source_present": True},
        {**payload["rows"][1], "inbound_status_source_present": False},
    ]
    content = generate_order_logic_v2_excel(payload, include_amounts=False)
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}

    status_headers = [
        "① 미입고 수량",
        "② PNFM확정 수량",
        "③ 입고진행중 수량",
        "④ 입고완료 수량",
    ]
    first_row = ORDER_HEADER_ROW + 1
    second_row = first_row + 1
    assert [sheet.cell(first_row, index[header]).value for header in status_headers] == [
        50,
        30,
        20,
        10,
    ]
    assert [sheet.cell(second_row, index[header]).value for header in status_headers] == [
        "-",
        "-",
        "-",
        "-",
    ]
    workbook.close()


def test_export_writes_v41_operational_formulas_without_changing_order_values() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}
    first_row = ORDER_HEADER_ROW + 1

    assert (
        sheet.cell(first_row, index["제안수량\n(미입고 인정)"]).value
        == 1000
    )
    assert "상한수량\n(미입고 제외)" not in headers
    assert sheet.cell(
        first_row,
        index["예상 발주일\n(평균)"],
    ).value.startswith("=IF(")
    assert "SQRT" in sheet.cell(
        first_row,
        index["조기경보일\n(CV·서비스수준 반영)"],
    ).value
    assert sheet.cell(first_row, index["단가(KRW)"]).value == 6_960.0
    assert sheet.cell(first_row, index["제안금액(KRW)"]).value == 6_960_000.0
    assert "확정금액(KRW)" not in headers
    workbook.close()


def test_export_adds_logistics_sheet_with_eta_bucket_formulas() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["물류전망"]
    headers = [
        str(cell.value)
        for cell in sheet[LOGISTICS_HEADER_ROW]
        if cell.value not in (None, "")
    ]
    first_row = LOGISTICS_HEADER_ROW + 1

    assert headers[:10] == [
        "상품코드",
        "상품명",
        "바코드",
        "판매량\n(13주)",
        "월평균\n판매량",
        "주평균\n판매량",
        "현재고\n(현지가용)",
        "운송중\n합계",
        "재고보유\n개월수(MOI)",
        "안전재고\n(상하한 적용)",
    ]
    assert headers[10] == "지연\n(07/27 이전 ETA)"
    assert headers[-1] == "W44+\n10/26~ ETA"
    assert sheet.cell(first_row, 1).value == "SKU-001"
    assert sheet.cell(first_row, 3).value == "8800000000001"
    assert sheet.cell(first_row, 5).value.startswith("=IF(")
    assert "SUMIFS" in sheet.cell(first_row, 11).value
    assert '"<>"' in sheet.cell(first_row, 11).value
    assert "_운송원천" in sheet.cell(first_row, 12).value
    assert sheet.cell(first_row, 26).value == 1391
    assert sheet.cell(first_row + 1, 27).value == "SKU-000002"
    assert sheet.freeze_panes == f"A{LOGISTICS_HEADER_ROW + 1}"
    workbook.close()


def test_logistics_conditional_formats_use_one_text_color() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["물류전망"]
    font_colors = []
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            font = rule.dxf.font if rule.dxf else None
            if font and font.color and font.color.type == "rgb":
                font_colors.append(font.color.rgb)

    assert font_colors
    assert set(font_colors) == {"FF20232A"}
    workbook.close()


def test_logistics_hides_eta_missing_column_when_no_missing_qty() -> None:
    payload = _result_payload()
    for row in payload["rows"]:
        row["eta_missing_qty"] = 0
        row["shipping_eta_details"] = [
            item
            for item in row.get("shipping_eta_details", [])
            if item.get("eta_status") != "ETA 미확인"
        ]

    content = generate_order_logic_v2_excel(payload, include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["물류전망"]
    headers = [
        str(cell.value)
        for cell in sheet[LOGISTICS_HEADER_ROW]
        if cell.value not in (None, "")
    ]

    assert "ETA 미확인\n물량" not in headers
    assert headers[-1] == "W44+\n10/26~ ETA"
    # 마지막 표시 컬럼이 하나 줄어들면 자동필터 범위도 함께 좁혀져야 한다.
    assert sheet.auto_filter.ref.startswith(f"A{LOGISTICS_HEADER_ROW}:")
    assert sheet.auto_filter.ref.split(":")[1].startswith(
        get_column_letter(len(headers))
    )
    assert "ETA 미확인 0개" in sheet["A3"].value
    workbook.close()


def test_logistics_eta_buckets_distinguish_skus_that_differ_only_by_case() -> None:
    payload = _result_payload()
    case_variant = {
        **payload["rows"][0],
        "sku_code": "sku-001",
        "shipping_eta_details": [
            {
                "eta": "2026-07-20",
                "qty": 800,
                "eta_status": "원천 ETA",
                "ship_date": "2026-06-01",
                "transport_label": "해운",
                "lead_time_days": None,
            }
        ],
        "shipping_schedule": [{"eta": "2026-07-20", "qty": 800}],
        "eta_actual_qty": 800,
        "eta_estimated_qty": 0,
        "eta_missing_qty": 0,
        "transit_qty": 800,
    }
    payload["rows"].append(case_variant)

    content = generate_order_logic_v2_excel(payload, include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    logistics = workbook["물류전망"]
    shipping_source = workbook["_운송원천"]
    first_row = LOGISTICS_HEADER_ROW + 1
    case_variant_row = LOGISTICS_HEADER_ROW + 3

    assert shipping_source["H1"].value == "SKU 대소문자 구분키"
    assert shipping_source["H2"].value == "SKU-000001"
    assert shipping_source["H5"].value == "SKU-000003"
    assert logistics["AA7"].value == "SKU-000001"
    assert logistics["AA9"].value == "SKU-000003"
    assert logistics.column_dimensions["AA"].hidden
    assert "$H$2:$H$5,$AA7" in logistics.cell(first_row, 11).value
    assert "$H$2:$H$5,$AA9" in logistics.cell(case_variant_row, 11).value
    assert "$A$2:$A$5,$A7" not in logistics.cell(first_row, 11).value
    workbook.close()


def test_export_writes_rows_and_numeric_formats() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}

    assert sheet.cell(ORDER_HEADER_ROW + 1, index["상품코드"]).value == "SKU-001"
    assert sheet.cell(ORDER_HEADER_ROW + 1, index["등급"]).value == "주력"
    assert sheet.cell(ORDER_HEADER_ROW + 1, index["신호등"]).value == "🔴발주"
    assert sheet.cell(ORDER_HEADER_ROW + 2, index["등급"]).value == "일반"
    assert (
        sheet.cell(ORDER_HEADER_ROW + 2, index["신호등"]).value
        == "⚠확인후발주"
    )
    assert (
        sheet.cell(ORDER_HEADER_ROW + 1, index["제안수량\n(미입고 인정)"]).value
        == 1000
    )
    assert (
        sheet.cell(ORDER_HEADER_ROW + 1, index["제안수량\n(미입고 인정)"]).number_format
        == "#,##0"
    )
    assert "상한수량\n(미입고 제외)" not in headers
    assert (
        sheet.cell(ORDER_HEADER_ROW + 1, index["단가(EUR)"]).number_format
        == "€#,##0.00"
    )
    assert (
        sheet.cell(ORDER_HEADER_ROW + 1, index["제안금액(KRW)"]).number_format
        == "₩#,##0"
    )
    workbook.close()


def test_export_removes_amount_columns_without_permission() -> None:
    content = generate_order_logic_v2_excel(_result_payload(), include_amounts=False)
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)

    assert AMOUNT_HEADERS.isdisjoint(headers)
    assert "제안수량\n(미입고 인정)" in headers
    assert "상한수량\n(미입고 제외)" not in headers
    assert "확정수량\n(담당자 발주)" not in headers
    assert "금액 조회 권한이 없어" in sheet["A4"].value
    assert "참고금액" in sheet["A4"].value
    workbook.close()


def test_usa_export_uses_usd_and_actual_krw_amount_headers() -> None:
    payload = _result_payload()
    payload.update(
        {
            "entity_code": "USA",
            "currency_code": "USD",
            "exchange_rate_krw": 1375.5,
            "exchange_rate_date": "2026-08-04",
        }
    )
    content = generate_order_logic_v2_excel(
        payload,
        include_amounts=True,
        entity_code="USA",
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}

    assert "단가(USD)" in headers
    assert "제안금액(USD)" in headers
    assert "제안금액(KRW)" in headers
    assert not any(header.startswith("확정금액") for header in headers)
    assert headers[-4:] == ["단가(USD)", "제안금액(USD)", "단가(KRW)", "제안금액(KRW)"]
    assert "단가(EUR)" not in headers
    assert "unit_cost_krw" in sheet["A4"].value
    assert "본사 창고\n가용재고" in headers
    assert "현지 가용재고" in headers
    assert "본사 미주창고\n가용재고" not in headers
    assert "현지가용(USA)" not in headers
    assert "현지가용(SKO)" not in headers
    assert (
        sheet.cell(ORDER_HEADER_ROW + 1, index["단가(USD)"]).number_format
        == "$#,##0.00"
    )
    workbook.close()


def test_export_signal_colors_follow_signal_text_for_all_rows() -> None:
    payload = _result_payload()
    payload["rows"] = [
        {**payload["rows"][0], "order_signal": "발주"},
        {**payload["rows"][1], "order_signal": "발주"},
        {**payload["rows"][0], "order_signal": "확인후발주"},
        {**payload["rows"][1], "order_signal": "확인후발주"},
        {**payload["rows"][0], "order_signal": "충분"},
        {**payload["rows"][1], "order_signal": "충분"},
    ]
    content = generate_order_logic_v2_excel(payload, include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    signal_column = headers.index("신호등") + 1
    expected_styles = {
        "🔴발주": ("FFFFE4E8", "FFB4002D"),
        "⚠확인후발주": ("FFFFF3CD", "FF7A4B00"),
        "🟢충분": ("FFE8F7EE", "FF087A3D"),
    }

    for row_index in range(ORDER_HEADER_ROW + 1, ORDER_HEADER_ROW + 1 + len(payload["rows"])):
        cell = sheet.cell(row_index, signal_column)
        expected_fill, expected_font = expected_styles[cell.value]
        assert cell.fill.fgColor.rgb == expected_fill
        assert cell.font.color.rgb == expected_font
    workbook.close()


def test_export_sold_out_status_uses_gray_style() -> None:
    payload = _result_payload()
    payload["rows"] = [
        {
            **payload["rows"][0],
            "data_status": "판매없음",
            "order_signal": "충분",
        }
    ]
    content = generate_order_logic_v2_excel(payload, include_amounts=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    status_column = _headers(sheet).index("상태") + 1
    status_cell = sheet.cell(ORDER_HEADER_ROW + 1, status_column)

    assert status_cell.fill.fgColor.rgb.endswith("F6F7F9")
    assert status_cell.font.color.rgb.endswith("6B7280")
    workbook.close()


def test_check_required_sheet_uses_entity_specific_stock_master_copy() -> None:
    expected_copy = {
        "PL": (
            "SKO 최초 입고 대상 SKU · 재고 마스터 미등록",
            "현지·본사 EU창고 재고 마스터에 없는 SKU입니다.",
            "13주 판매수량",
        ),
        "USA": (
            "USA 최초 입고 대상 SKU · 재고 마스터 미등록",
            "USA 현지·본사 미주창고 재고 마스터에 없는 SKU입니다.",
            "13주 판매수량",
        ),
        "HQ": (
            "본사 오포창고 최초 입고 대상 SKU · 재고 마스터 미등록",
            "본사 오포창고 재고 마스터에 없는 SKU입니다.",
            "91일 판매수량",
        ),
    }

    for entity_code, (
        expected_title,
        expected_description,
        expected_sales_label,
    ) in expected_copy.items():
        payload = _result_payload()
        payload.update(
            {
                "entity_code": entity_code,
                "period_unit": "day" if entity_code == "HQ" else "week",
                "demand_period_count": 91 if entity_code == "HQ" else 13,
            }
        )
        content = generate_order_logic_v2_excel(
            payload,
            include_amounts=False,
            entity_code=entity_code,
        )
        workbook = load_workbook(BytesIO(content), data_only=True)
        sheet = workbook["확인필요"]
        assert sheet["A1"].value == expected_title
        assert expected_description in sheet["A2"].value
        assert sheet.cell(CHECK_REQUIRED_HEADER_ROW, 4).value == expected_sales_label
        assert (
            sheet.cell(CHECK_REQUIRED_HEADER_ROW, 6).value
            == "미입고/입고예정 수량"
        )
        workbook.close()


def test_export_keeps_metadata_in_document_properties_and_hidden_calc_basis() -> None:
    payload = _result_payload()
    content = generate_order_logic_v2_excel(payload, include_amounts=False)
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert "실행정보" not in workbook.sheetnames
    assert payload["logic_version"] in workbook.properties.title
    calc_sheet = workbook["_계산기준"]
    assert calc_sheet["B1"].value.date().isoformat() == payload["as_of"]
    assert calc_sheet["B2"].value == payload["policy_mode"].upper()
    assert calc_sheet["B12"].value == "CMS stock/local unit_cost_krw"
    assert "재고 조회 기준일" in calc_sheet["B13"].value
    assert calc_sheet["B14"].value == "제안수량 × unit_price_krw"
    assert calc_sheet["B15"].value == payload["exchange_rate_source"]
    workbook.close()


def test_hq_export_uses_krw_source_amounts_without_conversion_columns() -> None:
    payload = _result_payload()
    payload.update(
        {
            "entity_code": "HQ",
            "currency_code": "KRW",
            "period_unit": "day",
            "demand_period_count": 91,
            # HQ amounts are already KRW, so a run carries no exchange rate.
            "exchange_rate_krw": None,
            "exchange_rate_date": "",
            "exchange_rate_source": "",
        }
    )
    content = generate_order_logic_v2_excel(
        payload,
        include_amounts=True,
        entity_code="HQ",
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}

    assert "단가(KRW)" in headers
    assert "제안금액(KRW)" in headers
    assert "단가(EUR)" not in headers
    assert "제안금액(EUR)" not in headers
    # HQ has one local KRW proposal amount and no conversion duplicate.
    assert headers.count("제안금액(KRW)") == 1
    assert "확정금액(KRW)" not in headers

    first_row = ORDER_HEADER_ROW + 1
    suggested_amount = sheet.cell(first_row, index["제안금액(KRW)"]).value
    assert suggested_amount == payload["rows"][0]["suggested_amount_eur"]
    assert (
        sheet.cell(first_row, index["단가(KRW)"]).number_format == "₩#,##0.00"
    )

    # The notice must not read as a failed lookup for a KRW-native result.
    assert "원화 원천값이라 환율을 적용하지 않습니다" in sheet["A4"].value
    assert "환율을 조회하지 못해" not in sheet["A4"].value
    workbook.close()


def test_hq_export_uses_opo_available_stock_and_hides_local_sko_column() -> None:
    payload = _result_payload()
    payload.update(
        {
            "entity_code": "HQ",
            "currency_code": "KRW",
            "period_unit": "day",
            "demand_period_count": 91,
        }
    )

    content = generate_order_logic_v2_excel(
        payload,
        include_amounts=False,
        entity_code="HQ",
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["발주제안"]
    headers = _headers(sheet)
    index = {header: column for column, header in enumerate(headers, start=1)}

    assert "본사 창고\n가용재고" in headers
    assert "본사창고\n가용재고" not in headers
    assert "본사 본사창고\n가용재고" not in headers
    assert "현지가용(SKO)" not in headers
    first_row = ORDER_HEADER_ROW + 1
    assert sheet.cell(first_row, index["본사 창고\n가용재고"]).value == 300
    assert "현지가용(SKO)" not in sheet.cell(first_row, 1).value
    workbook.close()


def test_metadata_and_calc_basis_describe_cms_actual_krw_for_currency_entities() -> None:
    payload = _result_payload()
    content = generate_order_logic_v2_excel(payload, include_amounts=False)
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert "실행정보" not in workbook.sheetnames
    assert payload["logic_version"] in workbook.properties.title
    calc_sheet = workbook["_계산기준"]
    assert calc_sheet["B1"].value.date().isoformat() == payload["as_of"]
    assert calc_sheet["B2"].value == payload["policy_mode"].upper()
    assert calc_sheet["B12"].value == "CMS stock/local unit_cost_krw"
    assert "재고 조회 기준일" in calc_sheet["B13"].value
    assert calc_sheet["B14"].value == "제안수량 × unit_price_krw"
    assert calc_sheet["B15"].value == payload["exchange_rate_source"]
    workbook.close()
