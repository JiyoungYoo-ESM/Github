from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from core import common, export_excel_order_sheet, export_excel_sheets, export_excel_util


BRAND = "\ube0c\ub79c\ub4dc"
SALES_3M = "3\uac1c\uc6d4 \ud310\ub9e4\ub7c9"
MONTHLY = "\uc6d4\ud3c9\uade0 \ud310\ub9e4\ub7c9"
SAFETY_STOCK = "\uc548\uc804\uc7ac\uace0"
EU_STOCK_1 = "1. \ud604\uc9c0 \uc7ac\uace0"
SHIPPING_STOCK_2 = "2. \uc6b4\uc1a1 \uc7ac\uace0"
COMBINED_STOCK_3 = "3. \ud604\uc9c0+\uc6b4\uc1a1"
EU_MONTHS_1 = "1. \ud604\uc9c0 \uc7ac\uace0(M)"
SHIPPING_MONTHS_2 = "2. \uc6b4\uc1a1\uc911(M)"
COMBINED_MONTHS_3 = "3. \ud569\uc0b0(M)"
STATUS = "\uc0c1\ud0dc"
STOCK_ETA_STATUS = "\uc7ac\uace0 ETA \uc0c1\ud0dc"
ALERT = "\uc54c\ub9bc"
EU_STOCK = "\uc720\ub7fd \uc7ac\uace0"
OUTBOUND = "(-) \uc608\uc0c1 \uc18c\ube44\ub7c9\n(\uc548\uc804\uc7ac\uace0 \uae30\uac04)"
FUTURE_OUTBOUND = "(-) \ucd9c\uace0 \ubb3c\ub7c9\n(\ubbf8\ub798 \uc548\uc804\uc7ac\uace0 \uae30\uac04\n\ub3d9\uc548 \uc18c\uc9c4\ub420 \uc608\uc0c1\ub7c9)"
PIPELINE = "(+) \uc785\uace0\ubb3c\ub7c9\n(\ud30c\uc774\ud504\ub77c\uc778 \ud569\uacc4)"
AVAILABLE = "\uac00\uc6a9 \uc7ac\uace0"
OPEN_PO_STATUS = "\ubbf8\uc785\uace0 \ud604\ud669"
HQ_EU_STOCK_HEADER = "\ud604\uc9c0 \ucc3d\uace0 \uc7ac\uace0"
ORDER_REQUIRED = "\ubc1c\uc8fc\ud544\uc694\uc218\ub7c9"
REPLENISHMENT = "\ud544\uc694 \ubcf4\ucda9\ub7c9"

PRODUCT_CODE = "\uc0c1\ud488\ucf54\ub4dc"
PRODUCT_NAME = "\uc0c1\ud488\uba85"
RECENT_3M_SALES = "\ucd5c\uadfc 3\uac1c\uc6d4 \ud310\ub9e4\uc218\ub7c9"
MONTHLY_SALES = "\uc6d4\ud3c9\uade0 \ud310\ub9e4\uc218\ub7c9"
SAFETY_TARGET = "\uc548\uc804\uc7ac\uace0 \ubaa9\ud45c\uc218\ub7c9"
EU_AVAILABLE_QTY = "EU \ud604\uc9c0 \uac00\uc6a9\uc218\ub7c9"
SHIPPING_QTY = "\uc6b4\uc1a1\uc911 \uc218\ub7c9"
OPEN_PO_QTY = "\ubbf8\uc785\uace0\uc218\ub7c9"
HQ_EU_AVAILABLE_QTY = "\ubcf8\uc0ac EU\ucc3d\uace0 \uac00\uc6a9\uc218\ub7c9"
ORDER_NEEDED_QTY = "\ubc1c\uc8fc\ud544\uc694\uc218\ub7c9"
ARRIVAL_DATE = "\ub3c4\ucc29\uc77c"
QTY = "\uc218\ub7c9"


def _review_df(shipping_qty=30, order_needed_qty=90):
    return pd.DataFrame(
        {
            PRODUCT_CODE: ["SKU-1"],
            PRODUCT_NAME: ["Sample Product"],
            BRAND: ["Brand"],
            RECENT_3M_SALES: [120],
            MONTHLY_SALES: [40],
            SAFETY_TARGET: [120],
            EU_AVAILABLE_QTY: [80],
            SHIPPING_QTY: [shipping_qty],
            OPEN_PO_QTY: [7],
            HQ_EU_AVAILABLE_QTY: [50],
            ORDER_NEEDED_QTY: [order_needed_qty],
            STATUS: ["원본상태"],
            EU_STOCK: [81],
            AVAILABLE: [777],
            OPEN_PO_STATUS: [8],
            ORDER_REQUIRED: [order_needed_qty],
            REPLENISHMENT: [888],
        }
    )


def test_stock_eta_columns_use_reference_order_without_legacy_reference_columns():
    assert common.UNIFIED_ESM_ORDER_REVIEW_COLUMNS == [
        "\ucf54\ub4dc",
        "No",
        "SKU",
        BRAND,
        SALES_3M,
        MONTHLY,
        SAFETY_STOCK,
        EU_STOCK_1,
        SHIPPING_STOCK_2,
        COMBINED_STOCK_3,
        EU_MONTHS_1,
        SHIPPING_MONTHS_2,
        COMBINED_MONTHS_3,
        ALERT,
        OPEN_PO_STATUS,
        HQ_EU_STOCK_HEADER,
        ORDER_REQUIRED,
    ]
    assert REPLENISHMENT not in common.UNIFIED_ESM_ORDER_REVIEW_COLUMNS


def test_stock_eta_sheet_uses_review_values_and_pipeline_calendar():
    arrival_df = pd.DataFrame(
        {
            "SKU": ["SKU-1"],
            ARRIVAL_DATE: [date(2026, 5, 21)],
            QTY: [30],
        }
    )

    review = pd.concat([_review_df(), _review_df()], ignore_index=True)
    review.loc[1, PRODUCT_CODE] = "SKU-2"
    review.loc[1, PRODUCT_NAME] = "No Pipeline Product"

    result = export_excel_order_sheet.build_stock_eta_sheet_df(review, arrival_df)
    row = result.loc[result["코드"].eq("SKU-1")].iloc[0]
    no_pipeline_row = result.loc[result["코드"].eq("SKU-2")].iloc[0]

    assert list(result.columns[: len(common.UNIFIED_ESM_ORDER_REVIEW_COLUMNS)]) == common.UNIFIED_ESM_ORDER_REVIEW_COLUMNS
    assert row["SKU"] == "Sample Product"
    assert row[SALES_3M] == 120
    assert row[EU_STOCK_1] == 80
    assert row[SHIPPING_STOCK_2] == 30
    assert row[COMBINED_STOCK_3] == 110
    assert row[EU_MONTHS_1] == 2.0
    assert row[SHIPPING_MONTHS_2] == 0.8
    assert row[COMBINED_MONTHS_3] == 2.8
    assert "재고 ETA 상태" not in result.columns
    assert "상태" not in result.columns
    assert row[ALERT] == "발주필요"
    assert row[OPEN_PO_STATUS] == 8
    assert row[HQ_EU_STOCK_HEADER] == 50
    assert row[ORDER_REQUIRED] == 10
    assert row["2026-5-21"] == 30
    assert no_pipeline_row["2026-5-21"] == "-"

    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, result)
    export_excel_sheets.format_unified_esm_order_review_sheet(ws, {"base_date": date(2026, 5, 20)})

    assert [ws.cell(row_idx, 5).value for row_idx in range(1, 6)] == [
        "SKU",
        "현지 재고 수량",
        "현지+운송 수량",
        "OOS SKU",
        "발주필요 SKU",
    ]
    assert [ws.cell(row_idx, 6).value for row_idx in range(1, 6)] == [
        '=COUNTIFS(C10:C1048576,"<>",C10:C1048576,"<>0")',
        "=SUM(H10:H100000)",
        "=SUM(J10:J100000)",
        '=COUNTIFS(H10:H100000,"<=0",C10:C100000,"<>",C10:C100000,"<>0")',
        '=IF($B$2="미입고 해소 후",SUMPRODUCT(--(Q10:Q100000>O10:O100000)),COUNTIFS(Q10:Q100000,">0",C10:C100000,"<>",C10:C100000,"<>0"))',
    ]
    assert ws["A2"].value == "미입고 반영 기준"
    assert ws["B2"].value == "미입고 해소 전"
    assert len(ws.data_validations.dataValidation) == 1
    assert [ws.cell(row_idx, 5).value for row_idx in range(6, 8)] == [None, None]
    assert [ws.cell(row_idx, 6).value for row_idx in range(6, 8)] == [None, None]


def test_stock_eta_summary_uses_after_inbound_scenario_when_selected():
    result = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df())
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, result)

    export_excel_sheets.format_unified_esm_order_review_sheet(
        ws,
        {"base_date": date(2026, 5, 20), "inbound_scenario": "after"},
    )

    assert ws["B2"].value == "미입고 해소 후"
    assert ws["F5"].value == '=IF($B$2="미입고 해소 후",SUMPRODUCT(--(Q10:Q100000>O10:O100000)),COUNTIFS(Q10:Q100000,">0",C10:C100000,"<>",C10:C100000,"<>0"))'


def test_scenario_renderer_starts_below_the_two_row_header(tmp_path):
    result = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df())
    wb = Workbook()
    ws = wb.active
    ws.title = "재고 ETA"
    export_excel_util.append_df(ws, result)
    export_excel_sheets.format_unified_esm_order_review_sheet(
        ws,
        {"base_date": date(2026, 5, 20)},
    )
    source = tmp_path / "order-review.xlsx"
    wb.save(source)

    rendered = export_excel_sheets.render_order_review_workbook_for_scenario(source, "after")
    rendered_wb = load_workbook(BytesIO(rendered), data_only=False)
    rendered_ws = rendered_wb["재고 ETA"]

    assert rendered_ws["B2"].value == "미입고 해소 후"
    assert rendered_ws["F5"].value == '=IF($B$2="미입고 해소 후",SUMPRODUCT(--(Q10:Q100000>O10:O100000)),COUNTIFS(Q10:Q100000,">0",C10:C100000,"<>",C10:C100000,"<>0"))'


def test_stock_eta_formula_keeps_excluded_review_rows_at_zero():
    review = pd.concat([_review_df(), _review_df()], ignore_index=True)
    review.loc[0, PRODUCT_CODE] = "SKU-ELIGIBLE"
    review.loc[1, PRODUCT_CODE] = "SKU-EXCLUDED"
    review.loc[1, STATUS] = "발주제외"

    result = export_excel_order_sheet.build_stock_eta_sheet_df(review)
    assert result[common.ORDER_REVIEW_ELIGIBILITY_HELPER].tolist() == [1, 0]
    assert result[ORDER_REQUIRED].tolist() == [10, 0]

    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, result)
    export_excel_sheets.format_unified_esm_order_review_sheet(
        ws,
        {"safety_months": 3, "base_date": date(2026, 5, 20)},
    )

    helper_col = next(
        col
        for col in range(1, ws.max_column + 1)
        if ws.cell(9, col).value == common.ORDER_REVIEW_ELIGIBILITY_HELPER
    )
    helper_letter = get_column_letter(helper_col)
    assert ws.column_dimensions[helper_letter].hidden is True
    assert ws["Q10"].value == f"=IF(${helper_letter}10=0,0,IF(F10=0,0,MAX(0,G10-J10)))"
    assert ws["Q11"].value == f"=IF(${helper_letter}11=0,0,IF(F11=0,0,MAX(0,G11-J11)))"
    assert ws.auto_filter.ref == "A9:Q11"


def test_stock_eta_sheet_displays_empty_pipeline_as_dash():
    result = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df(shipping_qty=0, order_needed_qty=0))
    row = result.iloc[0]

    assert PIPELINE not in result.columns
    assert row[ALERT] == "-"
    assert row[ORDER_REQUIRED] == 40


def test_stock_eta_alert_uses_local_oos_and_order_qty():
    review = _review_df(order_needed_qty=10)
    review[SHIPPING_QTY] = [200]

    result = export_excel_order_sheet.build_stock_eta_sheet_df(review)

    assert result.iloc[0][ALERT] == "발주필요"

    review[EU_AVAILABLE_QTY] = [0]
    result = export_excel_order_sheet.build_stock_eta_sheet_df(review)

    assert result.iloc[0][ALERT] == "OOS"


def test_stock_eta_status_rows_use_reference_colors():
    review = pd.concat([_review_df(shipping_qty=60, order_needed_qty=10), _review_df(order_needed_qty=0)], ignore_index=True)
    review.loc[1, PRODUCT_CODE] = "SKU-OOS"
    review.loc[1, EU_AVAILABLE_QTY] = 0
    review.loc[1, SHIPPING_QTY] = 0
    df = export_excel_order_sheet.build_stock_eta_sheet_df(review)
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, df)

    export_excel_sheets.format_unified_esm_order_review_sheet(ws, {"base_date": date(2026, 5, 20)})

    assert ws.cell(10, 1).fill.fgColor.rgb == "00FCE4D6"
    assert ws.cell(11, 1).fill.fgColor.rgb == "00FFC7CE"


def test_stock_eta_month_columns_hide_trailing_decimal_zeroes_for_large_exports():
    base = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df())
    df = pd.concat([base] * 301, ignore_index=True)
    df["코드"] = [f"SKU-{idx}" for idx in range(len(df))]

    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, df)

    export_excel_sheets.format_unified_esm_order_review_sheet(ws, {"base_date": date(2026, 5, 20)})

    for row_idx in (10, 310):
        assert ws[f"K{row_idx}"].number_format == "#,##0.##"
        assert ws[f"L{row_idx}"].number_format == "#,##0.##"
        assert ws[f"M{row_idx}"].number_format == "#,##0.##"
        assert ws[f"Q{row_idx}"].number_format == "#,##0"


def test_stock_eta_excel_header_uses_runtime_lead_times_and_renamed_column():
    arrival_df = pd.DataFrame(
        {
            "SKU": ["SKU-1"],
            ARRIVAL_DATE: [date(2026, 5, 21)],
            QTY: [30],
        }
    )
    df = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df(), arrival_df)
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, df)

    export_excel_sheets.format_unified_esm_order_review_sheet(
        ws,
            {
                "safety_months": 3,
                "base_date": date(2026, 5, 20),
                "lead_times": {
                    "\ud56d\uacf5": 15,
                "\ud574\uc6b4": 80,
                "\ucca0\uc1a1": 30,
                "\ud2b8\ub7ed": 30,
            },
        },
    )

    headers = [ws.cell(9, col).value or ws.cell(8, col).value for col in range(1, ws.max_column + 1)]
    date_col = headers.index("2026-5-21") + 1
    merged_ranges = {str(merged_range) for merged_range in ws.merged_cells.ranges}

    assert ws["H2"].value is None
    assert "H2:O2" not in merged_ranges
    assert ws["H2"].fill.fill_type is None
    parameter_labels = [ws.cell(3, col).value for col in range(8, 16)]
    assert parameter_labels == [
        "\uc548\uc804\uc7ac\uace0(M)",
        "Air\nL/T(\uc77c)",
        "\ud574\uc6b4\nL/T(\uc77c)",
        "\ucca0\uc1a1\nL/T(\uc77c)",
        "\ud2b8\ub7ed\ud0b9\nL/T(\uc77c)",
        "\uae30\uc900\uc77c",
        None,
        None,
    ]
    assert ws["H4"].value == 3
    assert ws["I4"].value == 15
    assert ws["J4"].value == 80
    assert ws["K4"].value == 30
    assert ws["L4"].value == 30
    assert ws["M4"].value == date(2026, 5, 20)
    assert ws.column_dimensions["M"].width >= 8
    assert ws.column_dimensions["C"].width >= 40
    assert ws.column_dimensions["H"].width >= 10
    assert ws.column_dimensions["I"].width >= 13
    assert ws.column_dimensions["J"].width >= 11
    assert ws.column_dimensions["K"].width >= 8
    assert ws.cell(7, 1).value is None
    assert ws.cell(7, 1).fill.fill_type is None
    assert not any(str(merged_range).startswith("A7:") for merged_range in merged_ranges)
    assert ws.row_dimensions[9].height >= 40.5
    assert ws.cell(5, 8).value is None
    assert "H5:O5" not in merged_ranges
    assert ws.cell(5, 8).fill.fill_type is None
    assert {"A8:A9", "B8:B9", "C8:C9", "D8:D9", "E8:E9", "F8:F9", "G8:G9"}.issubset(merged_ranges)
    assert {"H8:J8", "K8:M8"}.issubset(merged_ranges)
    assert ws["G8"].value == SAFETY_STOCK
    assert ws["H8"].value == "재고 수량"
    assert ws["H9"].value == "1. 현지\n재고\n(가용재고)"
    assert ws["I9"].value == "2. 운송\n재고"
    assert ws["J9"].value == "3. 현지\n+운송"
    assert ws["K8"].value == "재고 보유개월수"
    assert ws["K9"].value == "1. 현지\n재고(M)"
    assert ws["L9"].value == "2. 운송중\n(M)"
    assert ws["M9"].value == "3. 합산\n(M)"
    assert ws["K10"].number_format == "#,##0.##"
    assert ws["L10"].number_format == "#,##0.##"
    assert ws["M10"].number_format == "#,##0.##"
    assert ws["N8"].value == "⚠ 알람"
    assert ws["O8"].value == OPEN_PO_STATUS
    assert ws["P8"].value == HQ_EU_STOCK_HEADER
    assert ws["Q8"].value == "발주\n필요수량"
    assert ws["N8"].fill.fgColor.rgb == "00FFFF00"
    assert ws["O8"].fill.fgColor.rgb == "00FCE4D6"
    assert ws["P8"].fill.fgColor.rgb == "00FCE4D6"
    assert ws["Q8"].fill.fgColor.rgb == "00F4B183"
    assert ws["N10"].value == '=IF(H10<=0,"OOS",IF(Q10>0,"발주필요","-"))'
    assert ws["G10"].value == "=ROUND(F10*$H$4,0)"
    helper_col = next(
        col
        for col in range(1, ws.max_column + 1)
        if ws.cell(9, col).value == common.ORDER_REVIEW_ELIGIBILITY_HELPER
    )
    helper_letter = get_column_letter(helper_col)
    assert ws["Q10"].value == f"=IF(${helper_letter}10=0,0,IF(F10=0,0,MAX(0,G10-J10)))"
    assert ws.column_dimensions[helper_letter].hidden is True
    assert ws["Q10"].number_format == "#,##0"
    assert ws.protection.sheet is True
    assert ws["N10"].protection.hidden is True
    assert ws["N10"].protection.locked is True
    assert ws["H10"].protection.locked is False
    assert ws["Q8"].comment is not None
    assert "MAX(0" in ws["Q8"].comment.text
    assert f"{get_column_letter(date_col)}8:{get_column_letter(date_col)}9" in merged_ranges
    assert HQ_EU_STOCK_HEADER in headers
    assert OPEN_PO_STATUS in headers
    date_col_letter = get_column_letter(date_col)
    assert ws.cell(5, date_col).value == "ETA 날짜별 총 도착 수량"
    assert ws.cell(6, date_col).value == f"=SUM({date_col_letter}10:{date_col_letter}100000)"
    assert ws.cell(6, date_col).number_format == "#,##0"
    assert ws.cell(8, date_col).value == "2026-5-21"


def test_stock_eta_pipeline_lookup_is_case_sensitive_for_renewal_skus():
    review = pd.DataFrame(
        {
            PRODUCT_CODE: ["DRAS01-CASreu", "DRAS01-CASReu"],
            PRODUCT_NAME: ["Renew current", "Renew inbound"],
            BRAND: ["닥터엘시아", "닥터엘시아"],
            RECENT_3M_SALES: [56724, 0],
            MONTHLY_SALES: [18908, 0],
            SAFETY_TARGET: [56724, 0],
            EU_AVAILABLE_QTY: [57499, 0],
            SHIPPING_QTY: [0, 182040],
            OPEN_PO_QTY: [0, 0],
            HQ_EU_AVAILABLE_QTY: [0, 0],
            ORDER_NEEDED_QTY: [0, 0],
        }
    )
    arrival_df = pd.DataFrame(
        {
            "SKU": ["DRAS01-CASReu"],
            ARRIVAL_DATE: [date(2026, 6, 13)],
            QTY: [182040],
        }
    )

    result = export_excel_order_sheet.build_stock_eta_sheet_df(review, arrival_df)
    lower_r_row = result.loc[result["코드"].eq("DRAS01-CASreu")].iloc[0]
    upper_r_row = result.loc[result["코드"].eq("DRAS01-CASReu")].iloc[0]

    assert lower_r_row["2026-6-13"] == "-"
    assert upper_r_row["2026-6-13"] == 182040


def test_stock_eta_fast_mode_keeps_data_cell_borders():
    base = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df())
    df = pd.concat([base] * 301, ignore_index=True)
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, df)

    export_excel_sheets.format_unified_esm_order_review_sheet(ws, {"base_date": date(2026, 5, 20)})

    assert ws.cell(10, 1).border.left.style == "thin"
    assert ws.cell(100, 2).border.right.style == "thin"


def test_stock_eta_fast_mode_applies_status_row_colors():
    order_needed = export_excel_order_sheet.build_stock_eta_sheet_df(_review_df(shipping_qty=60, order_needed_qty=10))
    oos_review = _review_df(order_needed_qty=10, shipping_qty=0)
    oos_review.loc[0, PRODUCT_CODE] = "SKU-OOS"
    oos_review.loc[0, EU_AVAILABLE_QTY] = 0
    oos = export_excel_order_sheet.build_stock_eta_sheet_df(oos_review)
    df = pd.concat([order_needed, oos, order_needed] * 101, ignore_index=True)
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, df)

    export_excel_sheets.format_unified_esm_order_review_sheet(ws, {"base_date": date(2026, 5, 20)})

    assert ws.cell(10, 1).fill.fgColor.rgb == "00FCE4D6"
    assert ws.cell(11, 1).fill.fgColor.rgb == "00FFC7CE"
    assert ws.cell(310, 15).border.bottom.style == "thin"
