from datetime import date

import pandas as pd
from openpyxl import Workbook

from core import common, eta, export_excel_util
from core.lead_times import analysis_lead_time_settings


def _settings():
    return {
        "base_date": date(2026, 5, 20),
        "eur_krw_rate": 1748,
    }


def _calendar_df():
    return pd.DataFrame(
        {
            "\ub3c4\ucc29\uc77c": [date(2026, 2, 24), date(2026, 5, 9), date(2026, 5, 10)],
            "\ub3c4\ucc29 \uc608\uc815\uc6d4": ["2026-02", "2026-05", "2026-05"],
            "\uc6b4\uc1a1\uc218\ub2e8": ["\ud56d\uacf5", "\ud574\uc6b4", "\ud2b8\ub7ed"],
            "\ube0c\ub79c\ub4dc": ["Brand A", "Brand B", "Brand C"],
            "SKU": ["SKU-1", "SKU-2", "SKU-3"],
            "\uc0c1\ud488\uba85": ["Sample 1", "Sample 2", "Sample 3"],
            "\uc218\ub7c9": [100, 200, 50],
            "\ub3c4\ucc29 \uc608\uc815 \uae08\uc561_EUR": [10, 20, 5],
            "\ub3c4\ucc29 \uc608\uc815 \uae08\uc561_KRW": [17480, 34960, 8740],
            "\ucd9c\uace0\uc77c": [date(2026, 2, 9), date(2026, 4, 19), date(2026, 4, 25)],
            "\uc6b4\uc1a1 L/T": [15, 20, 15],
        }
    )


def test_arrival_reference_summary_maps_calendar_values_by_month_and_mode():
    df = eta.build_arrival_reference_summary_df(_settings(), _calendar_df())

    assert list(df.columns) == common.ARRIVAL_REFERENCE_COLUMNS

    feb = df[df["\uc785\uace0 \uc608\uc815"].eq("2026-02")].iloc[0]
    assert feb["\ud56d\uacf5"] == 100
    assert feb["\ud569\uacc4"] == 100
    assert feb["\uae08\uc561(EUR)"] == 10
    assert feb["\uae08\uc561(KRW)"] == 17480
    assert feb["금액(억원)"] == "0억"

    may = df[df["\uc785\uace0 \uc608\uc815"].eq("2026-05")].iloc[0]
    assert may["\ud574\uc6b4"] == 200
    assert may["\ud2b8\ub7ed"] == 50
    assert may["\ud569\uacc4"] == 250
    assert may["\uae08\uc561(EUR)"] == 25
    assert may["\uae08\uc561(KRW)"] == 43700
    assert may["금액(억원)"] == "0억"

    total = df[df["\uc785\uace0 \uc608\uc815"].eq("\ud569\uacc4")].iloc[0]
    assert total["\ud574\uc6b4"] == 200
    assert total["\ud56d\uacf5"] == 100
    assert total["\ud2b8\ub7ed"] == 50
    assert total["\ud569\uacc4"] == 350

    eur = df[df["\uc785\uace0 \uc608\uc815"].eq("\uae08\uc561(EUR)")].iloc[0]
    assert eur["\ud574\uc6b4"] == 20
    assert eur["\ud56d\uacf5"] == 10
    assert eur["\ud2b8\ub7ed"] == 5


def test_arrival_reference_summary_format_matches_reference_cells():
    df = eta.build_arrival_reference_summary_df(_settings(), _calendar_df())
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, df)

    eta.format_arrival_reference_summary_sheet(ws, _settings())

    assert ws["A1"].value == "\uc785\uace0 \uc608\uc815"
    assert ws["B1"].value == "\ud574\uc6b4"
    assert ws["E1"].value == "\ud2b8\ub7ed"
    assert ws["F1"].value == "\ud569\uacc4"
    assert ws["I1"].value == "금액(억원)"
    assert ws["J1"].value is None
    assert ws["A1"].fill.fgColor.rgb == "004F6228"
    assert ws["A2"].fill.fgColor.rgb == "00E2F0D9"
    assert ws["H2"].fill.fgColor.rgb == "00FCE4D6"
    assert ws["I2"].fill.fgColor.rgb == "00FCE4D6"
    assert ws["J1"].fill.fill_type is None
    assert ws.column_dimensions["H"].width >= 20
    assert ws.column_dimensions["I"].width >= 12
    assert ws.column_dimensions["J"].width >= 10
    assert ws["A14"].value == "\ud569\uacc4"
    assert ws["A15"].value == "\uae08\uc561(EUR)"
    assert ws["A16"].value == "\uae08\uc561(KRW)"
    assert ws["B14"].value == 200
    assert ws["C14"].value == 100
    assert ws["E14"].value == 50
    assert ws["F14"].value == 350
    assert ws.freeze_panes is None


def test_us_arrival_reference_excludes_review_rows_from_the_arrival_sheet():
    settings = {
        **analysis_lead_time_settings("USA"),
        "base_date": date(2026, 8, 5),
        "currency_code": "USD",
        "currency_krw_rate": 1300,
    }
    calendar = pd.DataFrame(
        {
            "도착 예정월": ["2026-08", "2026-08", "2026-08"],
            "운송수단 코드": ["OCEAN", "AIR", ""],
            "운송수단": ["해운", "항공", "확인필요"],
            "수량": [100, 20, 30],
            "도착 예정 금액_USD": [1000, 200, 300],
            "도착 예정 금액_KRW": [1_300_000, 260_000, 390_000],
        }
    )

    result = eta.build_arrival_reference_summary_df(settings, calendar)

    assert list(result.columns) == [
        "입고 예정", "해운", "항공", "합계",
        "금액(USD)", "금액(KRW)", "금액(억원)",
    ]
    august = result[result["입고 예정"].eq("2026-08")].iloc[0]
    assert august["해운"] == 100
    assert august["항공"] == 20
    assert "확인필요" not in result.columns
    assert august["합계"] == 120
    assert august["해운"] + august["항공"] == august["합계"]
    assert august["금액(USD)"] == 1200


def test_arrival_reference_detail_excludes_review_rows():
    settings = {
        **analysis_lead_time_settings("USA"),
        "base_date": date(2026, 8, 5),
        "currency_code": "USD",
        "currency_krw_rate": 1300,
    }
    calendar = pd.DataFrame(
        {
            "도착일": [date(2026, 8, 20), date(2026, 8, 20), date(2026, 8, 21)],
            "운송수단 코드": ["OCEAN", "AIR", ""],
            "운송수단": ["해운", "항공", "확인필요"],
            "SKU": ["SKU-OCEAN", "SKU-AIR", "SKU-REVIEW"],
            "수량": [100, 20, 30],
            "도착 예정 금액_KRW": [1_300_000, 260_000, 390_000],
        }
    )

    detail = eta.build_arrival_reference_detail_df(settings, calendar)

    assert set(detail["SKU"]) == {"SKU-OCEAN", "SKU-AIR"}
    assert detail.set_index("SKU")["운송 L/T"].to_dict() == {
        "SKU-OCEAN": 30,
        "SKU-AIR": 5,
    }
    assert "ETA 구분" not in detail.columns
    assert "ETA 출처" not in detail.columns
    assert "확인필요" not in set(detail["운송수단"])


def test_uk_arrival_reference_splits_direct_and_transshipment_air():
    settings = {
        **analysis_lead_time_settings("UK"),
        "base_date": date(2026, 8, 5),
        "currency_code": "EUR",
        "currency_krw_rate": 1700,
    }
    calendar = pd.DataFrame(
        {
            "도착 예정월": ["2026-08", "2026-08", "2026-08"],
            "운송수단 코드": ["OCEAN", "AIR_DIR", "AIR_TS"],
            "운송수단": ["해운", "항공(DIR·직항)", "항공(T/S·환적)"],
            "수량": [100, 20, 30],
            "도착 예정 금액_EUR": [1000, 200, 300],
            "도착 예정 금액_KRW": [1_700_000, 340_000, 510_000],
        }
    )

    result = eta.build_arrival_reference_summary_df(settings, calendar)
    assert list(result.columns[:5]) == [
        "입고 예정", "해운", "항공(DIR·직항)", "항공(T/S·환적)", "합계",
    ]
    august = result[result["입고 예정"].eq("2026-08")].iloc[0]
    assert august["항공(DIR·직항)"] == 20
    assert august["항공(T/S·환적)"] == 30


def test_arrival_reference_sheet_includes_sku_arrival_schedule_section():
    wb = Workbook()
    ws = wb.active

    eta.write_arrival_reference_sheet(ws, _settings(), _calendar_df())

    assert ws["A19"].value == "SKU\ubcc4 \ub3c4\ucc29 \uc77c\uc815"
    assert ws["A20"].value == "\ub3c4\ucc29\uc77c"
    assert ws["B20"].value == "\uc6b4\uc1a1\uc218\ub2e8"
    assert ws["G20"].value == "\ub3c4\ucc29\uc608\uc815\uae08\uc561(KRW)"
    assert ws["I20"].value == "\uc6b4\uc1a1 L/T"
    assert ws["J20"].value == "\ub3c4\ucc29\uc608\uc815\uc77c"
    assert ws["K20"].value == "D-Day"
    assert ws["A21"].value == date(2026, 2, 24)
    assert ws["B21"].value == "\ud56d\uacf5"
    assert ws["D21"].value == "SKU-1"
    assert ws["E21"].value == "Sample 1"
    assert ws["F21"].value == 100
    assert ws["G21"].value == 17480
    assert ws["I21"].value == 15
    assert ws["J21"].value == date(2026, 2, 24)
    assert ws["K21"].value == -85
    assert ws["J21"].number_format == "yyyy-mm-dd"
    assert ws["K21"].number_format == '"D-"0;"D+"0;"D-Day"'
    assert ws.auto_filter.ref == "A20:K23"
    assert ws.column_dimensions["C"].width >= 18
    assert ws.column_dimensions["E"].width >= 58
    assert ws.column_dimensions["G"].width >= 22
    assert ws.row_dimensions[20].height >= 30
    assert ws.row_dimensions[21].height >= 36
