"""HQ/OPO order-logic v2 contract tests.

These cover the approved HQ decisions: day-unit periodic review, the 91-day
demand window, the OPO sales population, the exclusive inventory-position
stages and entity isolation from PL/USA.
"""

from __future__ import annotations

from datetime import date
import io
import statistics

import pytest

from fastapi.testclient import TestClient

from backend.main import app as main_app
from backend.routers import order_logic_v2 as router
from backend.services import order_logic_v2_service as service
from backend.services.order_logic_v2_hq_source import (
    COMPLETED_DAY_COUNT,
    build_hq_order_logic_source,
    completed_day_window,
    is_hq_demand_row,
    resolve_hq_lead_time,
)
from core.order_logic_v2 import (
    GRADE_MAJOR,
    POLICY_CASH,
    POLICY_SHORTAGE,
    HQOrderLogicConfig,
    OrderLogicValidationError,
    SALES_STATUS_BULK_INCLUDED,
    SALES_STATUS_NORMAL,
    SALES_STATUS_NO_SALES,
    SKUOrderInput,
    calculate_hq_sku_order,
)


def _daily_vector(mean: float, stdev_target: float) -> list[float]:
    """Build 91 daily values with an exact mean and exact ``STDEV.S``."""

    pairs = COMPLETED_DAY_COUNT // 2
    delta = (stdev_target**2 * (COMPLETED_DAY_COUNT - 1) / (2 * pairs)) ** 0.5
    values: list[float] = []
    for _ in range(pairs):
        values.append(mean + delta)
        values.append(mean - delta)
    if COMPLETED_DAY_COUNT % 2:
        values.append(mean)
    return values


def test_daily_vector_helper_matches_requested_moments() -> None:
    values = _daily_vector(15.0, 9.0)
    assert len(values) == COMPLETED_DAY_COUNT
    assert statistics.fmean(values) == pytest.approx(15.0)
    assert statistics.stdev(values) == pytest.approx(9.0)


def test_hq_approved_hand_check_vector_returns_660() -> None:
    # docs/ORDER_LOGIC_V2_R_S_SPEC.md section 9.1.
    config = HQOrderLogicConfig(lt_days=12.0, sigma_l_days=6.0)
    item = SKUOrderInput(
        sku_code="SKU-HQ",
        weekly_sales=_daily_vector(15.0, 9.0),
        revenue_amt=1000.0,
        qty_local_available=150.0,
    )

    result = calculate_hq_sku_order(
        item,
        grade=GRADE_MAJOR,
        policy_mode=POLICY_SHORTAGE,
        config=config,
    )

    assert result.p_weeks == pytest.approx(40.0)
    assert result.layer1 == pytest.approx(180.0)
    assert result.layer2_ss_raw == pytest.approx(178.90, abs=0.01)
    assert result.ss_floor == pytest.approx(210.0)
    assert result.ss_cap == pytest.approx(525.0)
    assert result.layer2_ss == pytest.approx(210.0)
    assert result.layer3 == pytest.approx(420.0)
    assert result.target_stock_s == pytest.approx(810.0)
    assert result.suggested_qty == 660.0
    # The approved model has no reorder point and no IP < s gate.
    assert result.reorder_point_s is None


def test_hq_scenarios_share_lead_time_and_differ_only_by_z() -> None:
    config = HQOrderLogicConfig(lt_days=12.0, sigma_l_days=6.0)
    item = SKUOrderInput(
        sku_code="SKU-HQ",
        weekly_sales=_daily_vector(15.0, 9.0),
        revenue_amt=1000.0,
        qty_local_available=150.0,
    )

    cash = calculate_hq_sku_order(
        item, grade=GRADE_MAJOR, policy_mode=POLICY_CASH, config=config
    )
    shortage = calculate_hq_sku_order(
        item, grade=GRADE_MAJOR, policy_mode=POLICY_SHORTAGE, config=config
    )

    assert cash.lt_days == shortage.lt_days
    assert cash.sigma_l_weeks == shortage.sigma_l_weeks
    assert cash.z_applied == 1.28
    assert shortage.z_applied == 1.68


def test_hq_order_qty_is_zero_when_position_covers_target() -> None:
    config = HQOrderLogicConfig(lt_days=12.0, sigma_l_days=6.0)
    item = SKUOrderInput(
        sku_code="SKU-HQ",
        weekly_sales=_daily_vector(15.0, 9.0),
        revenue_amt=1000.0,
        qty_local_available=5_000.0,
    )

    result = calculate_hq_sku_order(
        item, grade=GRADE_MAJOR, policy_mode=POLICY_SHORTAGE, config=config
    )

    assert result.suggested_qty == 0.0
    assert result.need_order is False


def test_hq_config_requires_a_measured_lead_time() -> None:
    with pytest.raises(OrderLogicValidationError, match="lt_days"):
        HQOrderLogicConfig(lt_days=0.0, sigma_l_days=6.0)


def _hq_bulk_case(values: list[float]):
    config = HQOrderLogicConfig(lt_days=12.0, sigma_l_days=6.0)
    item = SKUOrderInput(
        sku_code="SKU-HQ",
        weekly_sales=tuple(values),
        revenue_amt=1000.0,
    )
    return calculate_hq_sku_order(
        item,
        grade=GRADE_MAJOR,
        policy_mode=POLICY_SHORTAGE,
        config=config,
    )


def test_hq_bulk_screen_compares_against_the_selling_day_mean() -> None:
    """대량포함은 실제 판매일 평균과 비교한다.

    무판매일을 포함한 91일 평균과 비교하면, 판매일이 31일 미만인 SKU는
    수치와 무관하게 항상 3배를 넘는다(최대값 >= 총량/k, 평균 = 총량/91).
    그래서 본사 국내조달 계산 대상 전건이 대량포함으로 몰렸다.
    """

    # 이틀에 28개씩 균등: 판매된 날들 사이 편차가 없으므로 대량포함이 아니다.
    even = _hq_bulk_case([28.0] * 2 + [0.0] * (COMPLETED_DAY_COUNT - 2))
    assert even.weeks_with_sales == 2
    assert even.has_bulk_week is False

    # 매일 팔리는데 하루만 폭증: 진짜 대량 건은 계속 잡아낸다.
    spike = _hq_bulk_case([10.0] * (COMPLETED_DAY_COUNT - 1) + [500.0])
    assert spike.has_bulk_week is True
    assert spike.sales_status == SALES_STATUS_BULK_INCLUDED

    # 매일 균등하면 정상이며, 이는 91일 평균 기준에서도 같은 결과다.
    flat = _hq_bulk_case([10.0] * COMPLETED_DAY_COUNT)
    assert flat.has_bulk_week is False
    assert flat.sales_status == SALES_STATUS_NORMAL


def test_hq_order_quantity_still_uses_the_full_window_mean() -> None:
    """판정 기준만 바뀌고 발주량 산식의 d_bar는 91일 평균을 유지한다."""

    result = _hq_bulk_case([28.0] * 2 + [0.0] * (COMPLETED_DAY_COUNT - 2))
    assert result.d_bar == pytest.approx(56.0 / COMPLETED_DAY_COUNT)


def test_hq_does_not_apply_the_intermittency_screen() -> None:
    """HQ는 간헐 판정을 적용하지 않는다(RS-018).

    PL/USA는 13주 표본에서도 7주 판매 기준을 유지한다. 이를 일 단위로 옮긴
    `91일 중 49일`은 주말·공휴일이 0으로 고정돼 HQ 정상 조달 패턴을 과도하게
    간헐로 분류하므로 본사에는 적용하지 않는다.
    """

    # 판매일이 하루뿐이어도 간헐로 분류하지 않는다.
    single_day = _hq_bulk_case([56.0] + [0.0] * (COMPLETED_DAY_COUNT - 1))
    assert single_day.weeks_with_sales == 1
    assert single_day.is_intermittent is False
    assert single_day.sales_status == SALES_STATUS_NORMAL

    # 과거 49일 기준에 걸렸던 구간도 정상으로 계산된다.
    thirty_days = _hq_bulk_case(
        [10.0] * 30 + [0.0] * (COMPLETED_DAY_COUNT - 30)
    )
    assert thirty_days.is_intermittent is False
    assert thirty_days.sales_status == SALES_STATUS_NORMAL

    # 판매가 전혀 없으면 여전히 판매없음이며, 대량포함 판정도 그대로 살아 있다.
    assert _hq_bulk_case([0.0] * COMPLETED_DAY_COUNT).sales_status == (
        SALES_STATUS_NO_SALES
    )
    assert _hq_bulk_case(
        [10.0] * (COMPLETED_DAY_COUNT - 1) + [500.0]
    ).sales_status == (
        SALES_STATUS_BULK_INCLUDED
    )


def test_hq_wrong_period_count_is_data_insufficient() -> None:
    config = HQOrderLogicConfig(lt_days=12.0, sigma_l_days=6.0)
    item = SKUOrderInput(
        sku_code="SKU-HQ",
        weekly_sales=[10.0] * 13,  # weekly grain must not be accepted for HQ
        revenue_amt=100.0,
    )

    result = calculate_hq_sku_order(
        item, grade=GRADE_MAJOR, policy_mode=POLICY_SHORTAGE, config=config
    )

    assert result.is_data_insufficient is True
    assert result.suggested_qty is None


def test_completed_day_window_excludes_as_of_and_spans_91_days() -> None:
    window = completed_day_window("2026-08-14")

    assert window.period_end == date(2026, 8, 13)
    assert window.period_start == date(2026, 5, 15)
    assert len(window.days) == COMPLETED_DAY_COUNT
    assert window.days[0] == window.period_start
    assert window.days[-1] == window.period_end


@pytest.mark.parametrize(
    ("warehouse", "biz_type", "expected"),
    [
        ("OPO", "KR-OVERSEAS", True),
        ("OPO", "KR-DOMESTIC", True),
        ("OPO", "KR-DOMESTIC 0%", True),
        ("OPO", "STAFF-SALES", False),
        ("OPO", "FREE SAMPLE", False),
        ("OPO", "자사간거래-커미션", False),
        ("OPO", "ETC", False),
        ("KR-W", "KR-OVERSEAS", False),
        ("BR-EU", "KR-OVERSEAS", False),
    ],
)
def test_hq_demand_population_filter(warehouse, biz_type, expected) -> None:
    row = {"whouse_nm": warehouse, "biz_type": biz_type}
    assert is_hq_demand_row(row) is expected


def test_hq_lead_time_blocks_below_two_samples() -> None:
    with pytest.raises(ValueError, match="2건 미만"):
        resolve_hq_lead_time(
            {"sample_count": 1, "avg_days": "12.03", "stddev_days": "16.45"}
        )


def test_hq_lead_time_reads_measured_statistic() -> None:
    resolved = resolve_hq_lead_time(
        {
            "sample_count": 5464,
            "avg_days": "12.03",
            "stddev_days": "16.45",
            "day_type": "calendar",
        }
    )

    assert resolved["sample_size"] == 5464
    assert resolved["mean_days"] == pytest.approx(12.03)
    assert resolved["stdev_days"] == pytest.approx(16.45)
    assert resolved["observation_months"] == 12
    assert resolved["warehouse_code"] == "OPO"


def _raw_hq_source(ship_date: str = "2026-08-13") -> dict[str, object]:
    return {
        "products": [
            {
                "prod_cd": "SKU-A",
                "prod_nm": "마스터 상품 A",
                "brand_nm": "마스터 브랜드",
            },
            {
                "prod_cd": "SKU-MASTER-ONLY",
                "prod_nm": "계산대상 아님",
                "brand_nm": "마스터 브랜드",
            },
        ],
        "sales_history": [
            {
                "prod_cd": "SKU-A",
                "prod_nm": "상품 A",
                "brand_nm": "브랜드",
                "whouse_nm": "OPO",
                "biz_type": "KR-OVERSEAS",
                "qty": 30,
                "amount_krw": "300000",
                "ship_dt": ship_date,
            },
            {
                "prod_cd": "SKU-A",
                "prod_nm": "상품 A",
                "brand_nm": "브랜드",
                "whouse_nm": "OPO",
                "biz_type": "FREE SAMPLE",
                "qty": 999,
                "amount_krw": "0",
                "ship_dt": ship_date,
            },
            {
                "prod_cd": "SKU-A",
                "prod_nm": "상품 A",
                "brand_nm": "브랜드",
                "whouse_nm": "KR-W",
                "biz_type": "KR-OVERSEAS",
                "qty": 777,
                "amount_krw": "10000",
                "ship_dt": ship_date,
            },
        ],
        "inventory": [
            {
                "sku": "SKU-A",
                "available_qty": 40,
                "avg_unit_cost": "5500",
                "unit_cost": "5500",
                "stock_status": "normal",
            },
            {
                "sku": "SKU-TROUBLE",
                "available_qty": 900,
                "avg_unit_cost": "1000",
                "unit_cost": "1000",
                "stock_status": "trouble",
            },
        ],
        "open_po": [{"sku": "SKU-A", "remaining_qty": 100}],
        "inbound_confirmed": [
            {
                "sku": "SKU-A",
                "remaining_qty": 25,
                "pnfm_confirmed_qty": 20,
                "inbound_in_progress_qty": 5,
                "completed_qty": 10,
            }
        ],
        "leadtime_stats": {
            "sample_count": 5464,
            "avg_days": "12.03",
            "stddev_days": "16.45",
        },
    }


def test_hq_source_keeps_only_opo_demand_and_zero_fills_days() -> None:
    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")

    row = next(row for row in prepared.rows if row["sku_code"] == "SKU-A")
    assert row["product_name"] == "마스터 상품 A"
    assert row["brand"] == "마스터 브랜드"
    assert all(row["sku_code"] != "SKU-MASTER-ONLY" for row in prepared.rows)
    demand = row["weekly_sales"]
    assert len(demand) == COMPLETED_DAY_COUNT
    # Only the accepted OPO row counts; the sample and other warehouse are out.
    assert sum(demand) == 30.0
    assert demand[-1] == 30.0
    assert demand[0] == 0.0
    audit = prepared.source_audit
    assert audit["sales_excluded_other_warehouse"] == 1
    assert audit["sales_excluded_biz_type"] == 1
    assert audit["warehouse_code"] == "OPO"
    assert audit["currency_code"] == "KRW"


def test_hq_source_excludes_negative_amount_positive_qty_from_demand() -> None:
    raw = _raw_hq_source()
    raw["sales_history"].append(
        {
            "prod_cd": "SKU-A",
            "prod_nm": "HQ item",
            "brand_nm": "HQ brand",
            "whouse_nm": "OPO",
            "biz_type": "KR-OVERSEAS",
            "qty": 5,
            "amount_krw": -1000,
            "ship_dt": "2026-08-13",
        }
    )

    prepared = build_hq_order_logic_source(raw, as_of="2026-08-14")
    row = next(row for row in prepared.rows if row["sku_code"] == "SKU-A")
    audit = prepared.source_audit

    assert sum(row["weekly_sales"]) == 30.0
    assert audit["sales_demand_excluded_rows"] == 1
    assert audit["sales_demand_excluded_qty"] == 5.0
    assert audit["sales_demand_excluded_amount"] == -1000.0
    assert "NEGATIVE_AMOUNT_POSITIVE_QTY_EXCLUDED" in row["warnings"]


def test_hq_inventory_position_uses_open_po_123_and_excludes_shipping() -> None:
    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")

    row = next(row for row in prepared.rows if row["sku_code"] == "SKU-A")
    assert row["open_qty"] == 100.0
    assert row["incoming_qty"] == 125.0
    assert row["pnfm_qty"] == 20.0
    assert row["transit_qty"] == 0.0
    assert row["inbound_progress_qty"] == 5.0
    assert row["inbound_completed_qty"] == 10.0
    assert row["incoming_qty"] == (
        row["open_qty"] + row["pnfm_qty"] + row["inbound_progress_qty"]
    )
    assert row["local_available_qty"] == 40.0
    # HQ has no forwarding warehouse stage.
    assert row["eu_available_qty"] == 0.0
    assert row["unit_price_local"] == 5500.0


@pytest.mark.parametrize(
    ("raw_price", "expected_price", "expected_amount"),
    [
        (0, 0.0, 0.0),
        (None, None, None),
        ("", None, None),
    ],
)
def test_hq_cms_unit_price_is_valid_but_null_or_blank_is_missing(
    raw_price: object,
    expected_price: float | None,
    expected_amount: float | None,
) -> None:
    raw = _raw_hq_source()
    raw["inventory"][0]["unit_cost"] = raw_price
    raw["inventory"][0]["avg_unit_cost"] = "999999"
    raw["sales_history"][0]["qty"] = 3_000
    raw["sales_history"][0]["amount_krw"] = "30000000"

    prepared = build_hq_order_logic_source(raw, as_of="2026-08-14")
    config = HQOrderLogicConfig(lt_days=12.03, sigma_l_days=16.45)
    result = service.build_hq_order_logic_v2_result(
        prepared,
        cache_info={"created_at": "2026-08-14T00:00:00+00:00", "age_seconds": 10},
        applied_mode="SHORTAGE",
        job_id="hq-zero-price",
        as_of="2026-08-14",
        config=config,
    )

    row = next(row for row in result["rows"] if row["sku_code"] == "SKU-A")
    assert row["suggested_qty"] > 0
    assert row["unit_price_krw"] == expected_price
    assert row["suggested_amount_krw"] == expected_amount


def test_hq_uses_cms_unit_cost_instead_of_avg_unit_cost() -> None:
    raw = _raw_hq_source()
    raw["inventory"][0]["avg_unit_cost"] = "7509"
    raw["inventory"][0]["unit_cost"] = "3183.11"

    prepared = build_hq_order_logic_source(raw, as_of="2026-08-14")
    row = next(row for row in prepared.rows if row["sku_code"] == "SKU-A")

    assert row["unit_price_local"] == pytest.approx(3183.11)
    assert row["unit_price_krw"] == pytest.approx(3183.11)


def test_hq_data_deficient_sku_uses_inventory_identity() -> None:
    raw = _raw_hq_source()
    raw["inventory"].append(
        {
            "sku": "SKU-NO-DEMAND",
            "prod_nm": "재고에는 있으나 판매 데이터 부족 상품",
            "brand_nm": "테스트브랜드",
            "available_qty": 12,
            "avg_unit_cost": 1000,
            "unit_cost": 1000,
            "stock_status": "normal",
        }
    )

    prepared = build_hq_order_logic_source(raw, as_of="2026-08-14")

    row = next(row for row in prepared.rows if row["sku_code"] == "SKU-NO-DEMAND")
    assert row["product_name"] == "재고에는 있으나 판매 데이터 부족 상품"
    assert row["brand"] == "테스트브랜드"


def test_hq_source_excludes_trouble_stock_from_available_quantity() -> None:
    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")

    trouble = next(
        (row for row in prepared.rows if row["sku_code"] == "SKU-TROUBLE"),
        None,
    )
    assert trouble is None


def test_hq_source_blocks_when_a_required_feed_is_missing() -> None:
    raw = _raw_hq_source()
    del raw["inventory"]

    with pytest.raises(ValueError, match="inventory"):
        build_hq_order_logic_source(raw, as_of="2026-08-14")


def test_hq_source_blocks_when_no_demand_falls_in_window() -> None:
    with pytest.raises(ValueError, match="91개 완료일"):
        build_hq_order_logic_source(
            _raw_hq_source(ship_date="2020-01-01"), as_of="2026-08-14"
        )


def test_hq_result_identity_reports_warehouse_timezone_and_krw() -> None:
    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")
    config = HQOrderLogicConfig(
        lt_days=float(prepared.lead_time["mean_days"]),
        sigma_l_days=float(prepared.lead_time["stdev_days"]),
    )

    result = service.build_hq_order_logic_v2_result(
        prepared,
        cache_info={"created_at": "2026-08-14T00:00:00+00:00", "age_seconds": 10},
        applied_mode="SHORTAGE",
        job_id="hq-test",
        as_of="2026-08-14",
        config=config,
    )

    assert result["entity_code"] == "HQ"
    assert result["warehouse_code"] == "OPO"
    assert result["timezone"] == "Asia/Seoul"
    assert result["currency_code"] == "KRW"
    assert result["exchange_rate_krw"] is None
    assert result["duration_unit"] == "CALENDAR_DAY"
    assert result["period_unit"] == "day"
    assert result["demand_period_count"] == COMPLETED_DAY_COUNT
    assert result["observation_window_days"] == COMPLETED_DAY_COUNT
    assert result["demand_grain"] == "DAY_1D"
    assert result["demand_grain_days"] == 1
    assert result["period_start"] == "2026-05-15"
    assert result["period_end"] == "2026-08-13"
    # Both scenarios must come from one immutable snapshot.
    assert (
        result["scenarios"]["CASH"]["source_snapshot_id"]
        == result["scenarios"]["SHORTAGE"]["source_snapshot_id"]
        == result["source_snapshot_id"]
    )
    settings = result["settings"]
    assert settings["review_days"] == 28.0
    assert settings["ss_floor_days"] == 14.0
    assert settings["ss_cap_days"] == 35.0
    row = result["rows"][0]
    assert row["protection_days"] == pytest.approx(row["protection_weeks"])
    assert row["review_days"] == pytest.approx(
        row["protection_days"] - row["lead_time_days"]
    )
    assert row["lead_time_sigma_days"] == pytest.approx(
        row["lead_time_sigma_weeks"]
    )


@pytest.mark.parametrize("include_product_master", [False, True])
def test_hq_stock_unregistered_sku_is_review_only_regardless_of_product_master(
    include_product_master: bool,
) -> None:
    raw = _raw_hq_source()
    raw["sales_history"].append(
        {
            "prod_cd": "SKU-STOCK-MISSING",
            "prod_nm": "재고 미등록 상품",
            "brand_nm": "테스트브랜드",
            "whouse_nm": "OPO",
            "biz_type": "KR-OVERSEAS",
            "qty": 300,
            "amount_krw": "3000000",
            "ship_dt": "2026-08-13",
        }
    )
    if include_product_master:
        raw["products"].append(
            {
                "prod_cd": "SKU-STOCK-MISSING",
                "prod_nm": "상품마스터 등록 상품",
                "brand_nm": "마스터브랜드",
            }
        )

    prepared = build_hq_order_logic_source(raw, as_of="2026-08-14")
    config = HQOrderLogicConfig(lt_days=12.03, sigma_l_days=16.45)
    result = service.build_hq_order_logic_v2_result(
        prepared,
        cache_info={"created_at": "2026-08-14T00:00:00+00:00", "age_seconds": 10},
        applied_mode="SHORTAGE",
        job_id="hq-stock-review",
        as_of="2026-08-14",
        config=config,
    )

    row = next(
        row for row in result["rows"] if row["sku_code"] == "SKU-STOCK-MISSING"
    )
    assert row["check_required"] is True
    assert "오포창고 재고 마스터에 없는 SKU" in row["check_required_reason"]
    assert row["calculable"] is False
    assert row["suggested_qty"] is None
    assert row["suggested_amount_krw"] is None
    assert row["need_order"] is None
    assert result["summary"]["check_required_skus"] == 1
    assert result["summary"]["order_skus"] == 0


def test_hq_rows_report_day_unit_lead_time_and_domestic_sourcing() -> None:
    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")
    config = HQOrderLogicConfig(lt_days=12.03, sigma_l_days=16.45)

    result = service.build_hq_order_logic_v2_result(
        prepared,
        cache_info={"created_at": "2026-08-14T00:00:00+00:00", "age_seconds": 10},
        applied_mode="SHORTAGE",
        job_id="hq-test",
        as_of="2026-08-14",
        config=config,
    )

    row = next(row for row in result["rows"] if row["sku_code"] == "SKU-A")
    assert row["transport_mode"] == "HQ_DOMESTIC"
    assert row["lead_time_days"] == pytest.approx(12.03)
    assert row["protection_weeks"] == pytest.approx(40.03)
    assert row["reorder_point"] is None
    assert row["inventory_position"] == pytest.approx(165.0)
    assert row["inventory_position_without_incoming"] == pytest.approx(40.0)


class _RequestStub:
    def __init__(self, entity_code: str) -> None:
        self.state = type("State", (), {"entity_code": entity_code})()


def test_router_allows_hq_and_still_blocks_unconfigured_entities() -> None:
    # HQ now has an approved policy and source adapter, so the gate must pass.
    for entity_code in ("HQ", "PL", "USA"):
        router._require_v2_policy_configuration(_RequestStub(entity_code))

    for entity_code in ("UK", "ME", "MX", "MY", "VN"):
        with pytest.raises(Exception) as excinfo:
            router._require_v2_policy_configuration(_RequestStub(entity_code))
        assert getattr(excinfo.value, "status_code", None) == 409


def test_hq_entity_does_not_reuse_the_pl_result_builder() -> None:
    # PL/USA rows carry EUR/USD and a shipping transport mode, so the shared
    # builder must keep refusing an HQ payload.
    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")
    with pytest.raises(service.OrderLogicV2SourceUnavailable, match="결과 원천·통화"):
        service.build_order_logic_v2_result(
            prepared,  # type: ignore[arg-type]
            cache_info={"created_at": "2026-08-14T00:00:00+00:00", "age_seconds": 10},
            applied_mode="SHORTAGE",
            job_id="hq-guard",
            as_of="2026-08-14",
            entity_code="HQ",
        )


def test_hq_http_gate_no_longer_returns_409(monkeypatch: pytest.MonkeyPatch) -> None:
    """The HQ V2 tab must reach the workflow instead of the old 409 block."""

    import backend.auth.user_store as user_store_module
    import backend.routers.auth as auth_router
    import backend.services.auth as auth_service
    from backend.routers import order_logic_v2 as order_logic_router
    from tests.auth_helpers import TEST_PASSWORDS, TestUserStore

    monkeypatch.setattr(user_store_module, "_user_store", TestUserStore())
    monkeypatch.setattr(
        order_logic_router,
        "latest_order_logic_v2_result",
        lambda _client_id, _entity_code: None,
    )
    auth_service.reset_sessions()
    auth_router._login_limiter.reset()

    client = TestClient(main_app)
    login = client.post(
        "/api/auth/login",
        headers={"X-Requested-With": "fetch"},
        json={"id": "adminmaster", "password": TEST_PASSWORDS["adminmaster"]},
    )
    assert login.status_code == 200

    response = client.get(
        "/api/order-logic-v2/latest",
        headers={"X-Entity-Code": "HQ", "X-Client-Id": "hq-browser"},
    )
    assert response.status_code == 200

    blocked = client.get(
        "/api/order-logic-v2/latest",
        headers={"X-Entity-Code": "MY", "X-Client-Id": "hq-browser"},
    )
    assert blocked.status_code == 409


def test_hq_excel_uses_day_unit_labels_and_krw() -> None:
    import openpyxl

    from backend.services.order_logic_v2_excel import (
        CHECK_REQUIRED_HEADER_ROW,
        LOGISTICS_HEADER_ROW,
        generate_order_logic_v2_excel,
    )

    prepared = build_hq_order_logic_source(_raw_hq_source(), as_of="2026-08-14")
    config = HQOrderLogicConfig(lt_days=12.03, sigma_l_days=16.45)
    result = service.build_hq_order_logic_v2_result(
        prepared,
        cache_info={"created_at": "2026-08-14T00:00:00+00:00", "age_seconds": 10},
        applied_mode="SHORTAGE",
        job_id="hq-excel",
        as_of="2026-08-14",
        config=config,
    )

    content = generate_order_logic_v2_excel(
        result, include_amounts=True, entity_code="HQ"
    )
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    try:
        order_sheet = workbook["발주제안"]
        headers = [
            str(order_sheet.cell(7, column).value or "").replace("\n", " ")
            for column in range(1, 13)
        ]
        assert "판매량 (91일)" in headers
        assert "일평균" in headers
        assert "판매량 (13주)" not in headers
        assert "주평균" not in headers
        assert str(order_sheet["A3"].value or "").startswith("일별 수요")
        check_required_sheet = workbook["확인필요"]
        assert (
            check_required_sheet.cell(CHECK_REQUIRED_HEADER_ROW, 4).value
            == "91일 판매수량"
        )
        assert (
            check_required_sheet.cell(CHECK_REQUIRED_HEADER_ROW, 6).value
            == "미입고/입고예정 수량"
        )
        logistics_sheet = workbook["물류전망"]
        assert (
            logistics_sheet.cell(LOGISTICS_HEADER_ROW, 4).value
            == "판매량\n(91일)"
        )
        assert (
            logistics_sheet.cell(LOGISTICS_HEADER_ROW, 6).value
            == "일평균\n판매량"
        )
        all_headers = {
            str(order_sheet.cell(7, column).value or "").replace("\n", " "): column
            for column in range(1, order_sheet.max_column + 1)
        }
        first_row = 8
        assert order_sheet.cell(first_row, all_headers["① 미입고 수량"]).value == 100
        assert order_sheet.cell(first_row, all_headers["② PNFM확정 수량"]).value == 20
        assert order_sheet.cell(first_row, all_headers["③ 입고진행중 수량"]).value == 5
        assert order_sheet.cell(first_row, all_headers["④ 입고완료 수량"]).value == 10
        assert order_sheet.cell(first_row, all_headers["운송중"]).value == 0

        calc_sheet = workbook["_계산기준"]
        basis = {
            calc_sheet.cell(row, 1).value: calc_sheet.cell(row, 2).value
            for row in range(1, 14)
        }
        assert basis["국내조달 L/T(일)"] == pytest.approx(12.03)
        assert basis["정기 검토주기 R(일)"] == pytest.approx(28.0)
        assert basis["안전재고 하한(일분)"] == pytest.approx(14.0)
        assert basis["안전재고 상한(일분)"] == pytest.approx(35.0)
        assert basis["발주금액 통화"] == "KRW"
        # The weekly shipping ladder must not leak into an HQ workbook.
        assert "항공 L/T(일)" not in basis
        assert "해운 L/T(일)" not in basis
    finally:
        workbook.close()
