from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.auth.amount_permissions import (
    can_manage_order_logic_v2_settings,
    can_view_amount_data,
    is_amount_field,
    sanitize_amount_data,
)
from backend.services.amount_safe_workbook import create_amount_safe_workbook


def test_only_master_and_internal_audit_can_view_amount_data() -> None:
    assert can_view_amount_data("adminmaster")
    assert can_view_amount_data(" ADMINMASTER ")
    assert can_view_amount_data("ia")
    for username in ("eu_manager", "bm1", "bm2", "bm3", "my_team", "vn_team", "hnb_team", "sales_team", "", None):
        assert not can_view_amount_data(username)


def test_only_master_and_internal_audit_can_manage_order_logic_v2_settings() -> None:
    assert can_manage_order_logic_v2_settings("adminmaster")
    assert can_manage_order_logic_v2_settings(" IA ")
    for username in (
        "eu_manager",
        "bm1",
        "bm2",
        "bm3",
        "my_team",
        "vn_team",
        "hnb_team",
        "sales_team",
        "",
        None,
    ):
        assert not can_manage_order_logic_v2_settings(username)


def test_amount_classifier_keeps_ratios_quantities_and_ranks() -> None:
    for field in ("매출 비중", "점유율", "sales_share", "growth_rate", "판매수량", "재고수량", "순위"):
        assert not is_amount_field(field), field
    for field in ("매출액", "판매금액", "재고금액", "발주금액", "sales_amount", "inventory_value", "unit_price"):
        assert is_amount_field(field), field


def test_recursive_sanitizer_removes_amounts_but_keeps_allowed_analysis() -> None:
    payload = {
        "매출액": 123_456,
        "판매수량": 42,
        "점유율": 17.5,
        "순위": 2,
        "metric": "판매금액",
        "value": "₩123,456",
        "htmlSnapshot": "<div>₩123,456</div>",
        "rows": [
            {
                "sales_amount": 99,
                "sales_share": 21.2,
                "growth_rate": -3.5,
                "quantity": 7,
            }
        ],
    }

    result = sanitize_amount_data(payload)

    assert "매출액" not in result
    assert "value" not in result
    assert result["htmlSnapshot"] is None
    assert result["판매수량"] == 42
    assert result["점유율"] == 17.5
    assert result["순위"] == 2
    assert result["rows"] == [{"sales_share": 21.2, "growth_rate": -3.5, "quantity": 7}]


def test_general_account_workbook_removes_amount_columns_and_summary_values(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "발주"
    sheet.append(["SKU", "판매수량", "매출액", "매출 비중", "순위"])
    sheet.append(["A-1", 10, 1000, 0.25, 1])
    sheet["A4"] = "총 발주금액"
    sheet["B4"] = "=SUM(C2:C2)"
    hidden = workbook.create_sheet("_source")
    hidden.sheet_state = "hidden"
    hidden.append(["SKU", "재고금액", "재고수량"])
    hidden.append(["A-1", 500, 5])
    workbook.save(source)
    workbook.close()

    safe_path = create_amount_safe_workbook(source)
    try:
        safe = load_workbook(safe_path, data_only=False)
        visible_values = [cell.value for row in safe["발주"].iter_rows() for cell in row]
        hidden_values = [cell.value for row in safe["_source"].iter_rows() for cell in row]
        assert "매출액" not in visible_values
        assert 1000 not in visible_values
        assert "총 발주금액" not in visible_values
        assert not any(isinstance(value, str) and value.startswith("=") for value in visible_values)
        assert "매출 비중" in visible_values
        assert 0.25 in visible_values
        assert "재고금액" not in hidden_values
        assert 500 not in hidden_values
        assert "재고수량" in hidden_values
        assert 5 in hidden_values
        safe.close()
    finally:
        safe_path.unlink(missing_ok=True)
