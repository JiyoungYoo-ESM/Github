from __future__ import annotations

from base64 import b64decode
from collections import Counter
from itertools import permutations
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from backend.services import report_template_export as export_service
from backend.services.report_template_export import ReportBlockPayload, audience_badge, build_report_export


def pptx_text(path: Path) -> str:
    parts: list[str] = []
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not (
                name.startswith("ppt/slides/")
                or name.startswith("ppt/charts/")
                or name.startswith("ppt/embeddings/")
            ):
                continue
            data = archive.read(name)
            try:
                parts.append(data.decode("utf-8", errors="ignore"))
            except UnicodeDecodeError:
                parts.append(data.decode("latin-1", errors="ignore"))
    return "\n".join(parts)


def duplicate_pptx_part_names(path: Path) -> set[str]:
    with ZipFile(path) as archive:
        counts = Counter(archive.namelist())
    return {name for name, count in counts.items() if count > 1}


def slide_text(slide) -> str:
    return "\n".join(
        shape.text.strip()
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def first_pptx_table_rows_with_headers(path: Path, required_headers: set[str]) -> list[list[str]]:
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(path) as archive:
        slide_names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        for name in slide_names:
            root = ET.fromstring(archive.read(name))
            for table in root.findall(".//a:tbl", ns):
                rows: list[list[str]] = []
                for table_row in table.findall("a:tr", ns):
                    row: list[str] = []
                    for cell in table_row.findall("a:tc", ns):
                        row.append("".join(text.text or "" for text in cell.findall(".//a:t", ns)).strip())
                    rows.append(row)
                if rows and required_headers.issubset(set(rows[0])):
                    return rows
    return []


def test_export_normalizes_placeholder_english_report_title(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[country_ranking_block()],
        title="Silicon2 Market Report - Web Card Charts Final",
    )
    prs = Presentation(str(output_path))
    first_slide_text = slide_text(prs.slides[0])
    deck_text = pptx_text(output_path)

    assert "ESM 데이터 분석 리포트" in first_slide_text
    assert "Silicon2 Market Report - Web Card Charts Final" not in deck_text
    assert "ESM 데이터 분석 리포트" in output_path.name
    assert "Silicon2 Market Report - Web Card Charts Final" not in slide_text(prs.slides[-1])


def brand_ranking_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="brand:rank:test",
        title="토리든 브랜드 판매 순위",
        subtitle="아누아 · 토리든 · 롬앤 · 조선미녀",
        meta="브랜드 기준 분석",
        type="brand_rank",
        kind="ranking",
        section="brand",
        size="full",
        params={"brand": "토리든", "원가율": "42%", "eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["순위", "브랜드명", "매출액", "원가율", "원화표시"],
            "rows": [
                {"순위": 1, "브랜드명": "아누아", "brand_original": "ANUA", "brand_role": "self", "is_self": True, "매출액": 9_000_000, "원가율": "40%", "원화표시": "약 ₩158.4억"},
                {"순위": 2, "브랜드명": "토리든", "brand_original": "Torriden", "brand_role": "competitor", "매출액": 7_000_000, "원가율": "41%", "원화표시": "약 ₩123.2억"},
                {"순위": 3, "브랜드명": "롬앤", "brand_original": "Romand", "brand_role": "competitor", "매출액": 5_000_000, "원가율": "39%", "원화표시": "약 ₩88.0억"},
                {"순위": 4, "브랜드명": "조선미녀", "brand_original": "Beauty of Joseon", "brand_role": "competitor", "매출액": 4_000_000, "원가율": "38%", "원화표시": "약 ₩70.4억"},
            ],
        },
    )


def country_ranking_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="country:rank:test",
        title="국가별 판매 순위",
        subtitle="2개국 · 매출 기준",
        meta="국가 · 권역 기준 판매 인사이트",
        type="country_rank",
        kind="ranking",
        section="region",
        size="full",
        params={"eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["순위", "국가", "권역", "매출액", "점유율", "원화표시"],
            "rows": [
                {"순위": 1, "국가": "미국", "권역": "북미", "매출액": 1_200_000, "점유율": "60%", "원화표시": "약 ₩21.1억"},
                {"순위": 2, "국가": "일본", "권역": "아시아", "매출액": 800_000, "점유율": "40%", "원화표시": "약 ₩14.1억"},
            ],
        },
    )


def low_rank_country_ranking_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="country:rank:low-page",
        title="국가별 판매 순위",
        subtitle="18-20위 · 매출 기준",
        meta="국가 · 권역 기준 판매 인사이트",
        type="country_rank",
        kind="ranking",
        section="region",
        size="full",
        params={"eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["순위", "국가", "권역", "매출액", "매출표시", "원화표시", "점유율", "YoY", "MoM"],
            "rows": [
                {"순위": 18, "국가": "Lithuania", "권역": "유럽", "매출액": 239_100, "매출표시": "€239.1K", "원화표시": "약 ₩4.25억", "점유율": "0.7%", "YoY": "-65%"},
                {"순위": 19, "국가": "Bulgaria", "권역": "유럽", "매출액": 144_900, "매출표시": "€144.9K", "원화표시": "약 ₩2.57억", "점유율": "0.4%", "YoY": "-90.3%"},
                {"순위": 20, "국가": "Belgium", "권역": "유럽", "매출액": 135_400, "매출표시": "€135.4K", "원화표시": "약 ₩2.41억", "점유율": "0.4%", "YoY": "-39.7%"},
            ],
        },
    )


def region_share_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="country:region-share:test",
        title="권역별 매출 비중",
        subtitle="유럽 · 점유율 99.9%",
        meta="국가 · 권역 기준 판매 인사이트",
        type="region_share",
        kind="share",
        section="region",
        size="full",
        params={"eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["순위", "권역", "매출액", "매출표시", "원화표시", "점유율"],
            "rows": [
                {"순위": 1, "권역": "유럽", "매출액": 144_300_000, "매출표시": "€144.3M", "원화표시": "약 ₩2,512.1억", "점유율": "99.9%"},
            ],
        },
    )


def country_detail_card_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="country:detail:germany",
        title="Germany 국가 판매 상세",
        subtitle="유럽 · 총 매출 €52.2M · 점유율 16.8%",
        meta="선택 국가 판매 인사이트",
        type="country_detail",
        kind="kpi",
        section="region",
        size="full",
        params={"country": "Germany", "region": "유럽", "amount": 52_200_000, "share": 16.8, "peakMonth": "6월", "eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["구분", "순위", "항목", "매출액", "매출표시", "원화표시", "비고"],
            "rows": [
                {"구분": "국가 요약", "순위": "", "항목": "Germany", "매출액": 52_200_000, "매출표시": "€52.2M", "원화표시": "약 ₩919.0억", "비고": "점유율 16.8%"},
                {"구분": "월 최고 매출", "순위": "", "항목": "6월", "매출액": 7_500_000, "매출표시": "€7.5M", "원화표시": "약 ₩132.0억", "비고": "6월 최고"},
                {"구분": "상위 브랜드", "순위": 1, "항목": "조선미녀", "매출액": 15_200_000, "매출표시": "€15.2M", "원화표시": "약 ₩267.6억", "비고": ""},
                {"구분": "상위 브랜드", "순위": 2, "항목": "메디큐브", "매출액": 10_900_000, "매출표시": "€10.9M", "원화표시": "약 ₩191.9억", "비고": ""},
                {"구분": "Top 5 카테고리", "순위": 1, "항목": "썬케어 > 크림", "매출액": 7_590_000, "매출표시": "€7.59M", "원화표시": "약 ₩133.6억", "비고": ""},
                {"구분": "Top 5 카테고리", "순위": 2, "항목": "스킨케어 > 크림", "매출액": 7_500_000, "매출표시": "€7.5M", "원화표시": "약 ₩132.0억", "비고": ""},
            ],
        },
    )


def sku_detail_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="sku:detail:test",
        title="[EU] 조선미녀크림 50ml SKU 분석",
        subtitle="조선미녀 · 피크 10월 · €2.74M",
        meta="SKU 월별 추이 · 국가별 판매 분포",
        type="sku_detail",
        kind="trend",
        section="sku",
        size="full",
        params={"sku": "EU-JM-50", "brand": "조선미녀", "peakMonth": 10, "eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["구분", "항목", "매출액", "매출표시", "원화표시", "판매수량", "점유율", "비고"],
            "rows": [
                {"구분": "월별 추이", "항목": "8월", "매출액": 620_000, "매출표시": "€0.62M", "원화표시": "약 ₩10.9억", "판매수량": 18_400, "점유율": "", "비고": ""},
                {"구분": "월별 추이", "항목": "10월", "매출액": 1_200_000, "매출표시": "€1.20M", "원화표시": "약 ₩21.1억", "판매수량": 30_100, "점유율": "", "비고": "피크"},
                {"구분": "국가 분포", "항목": "Estonia", "매출액": 1_213_820, "매출표시": "€1.21M", "원화표시": "약 ₩21.4억", "판매수량": 15_200, "점유율": "44.3%", "비고": ""},
            ],
        },
    )


def sku_order_reference_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="sku:order-reference:test",
        title="[EUP]콜라겐 나이트 랩핑 마스크 75ml 발주 참고 정보",
        subtitle="메디큐브 · 피크 6월 · MECUP10-PKREUP",
        meta="SKU 발주 참고 정보",
        type="sku_detail",
        kind="kpi",
        section="sku",
        size="half",
        params={"sku": "MECUP10-PKREUP", "brand": "메디큐브", "peakMonth": 6, "preparationMonth": 3, "eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["항목", "값", "보조값", "비고"],
            "rows": [
                {"항목": "매출 순위", "값": "10위 / 613 SKU", "보조값": "", "비고": ""},
                {"항목": "피크 3개월 전 준비월", "값": "3월", "보조값": "", "비고": "6월 피크 기준"},
                {"항목": "판매수량", "값": "613.6K개", "보조값": "", "비고": ""},
                {"항목": "매출", "값": "€5.1M", "보조값": "약 ₩89.8억", "비고": ""},
                {"항목": "판매 국가", "값": "34개국", "보조값": "", "비고": ""},
                {"항목": "피크월 매출", "값": "€658.3K", "보조값": "약 ₩11.6억", "비고": "6월 · 선택 SKU 월 최대 매출"},
            ],
        },
    )


def cross_matrix_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="cross:country-brand:test",
        title="국가 × 브랜드 교차분석",
        subtitle="United Kingdom, Estonia, Netherlands The × 조선미녀",
        meta="매출액 · 금액",
        type="cross_matrix",
        kind="matrix",
        section="cross",
        size="full",
        params={"axis": "country-brand", "metric": "sales", "scale": "amount", "rowLabel": "국가", "columnLabel": "브랜드", "eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["국가", "브랜드", "표시값", "원화표시", "점유율"],
            "rows": [
                {"국가": "Estonia", "브랜드": "조선미녀", "값": 10067.8, "표시값": "€10,067.8K", "원화표시": "약 ₩177.2억", "점유율": "", "metric": "sales", "scale": "amount"},
                {"국가": "United Kingdom", "브랜드": "조선미녀", "값": 2595.6, "표시값": "€2,595.6K", "원화표시": "약 ₩45.7억", "점유율": "", "metric": "sales", "scale": "amount"},
                {"국가": "Netherlands The", "브랜드": "조선미녀", "값": 2345.4, "표시값": "€2,345.4K", "원화표시": "약 ₩41.3억", "점유율": "", "metric": "sales", "scale": "amount"},
                {"국가": "Estonia", "브랜드": "", "값": 1200.0, "표시값": "€1,200.0K", "원화표시": "약 ₩21.1억", "점유율": "", "metric": "sales", "scale": "amount"},
            ],
        },
    )


def ingredient_country_cross_block() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="cross:ingredient-country:test",
        title="성분 × 국가 교차분석",
        subtitle="스네일뮤신 × Estonia",
        meta="매출액 · 금액",
        type="cross_matrix",
        kind="matrix",
        section="cross",
        size="full",
        params={"axis": "ingredient-country", "metric": "sales", "scale": "amount", "rowLabel": "성분", "columnLabel": "국가", "eur_krw_rate": 1760.49},
        snapshot={
            "columns": ["성분", "국가", "표시값", "원화표시", "점유율"],
            "rows": [
                {"성분": "스네일뮤신", "국가": "Estonia", "값": 676.1, "표시값": "€676.1K", "원화표시": "약 ₩11.9억", "점유율": "", "metric": "sales", "scale": "amount"},
                {"성분": "스네일뮤신", "국가": "Germany", "값": 218.4, "표시값": "€218.4K", "원화표시": "약 ₩3.8억", "점유율": "", "metric": "sales", "scale": "amount"},
            ],
        },
    )


def cross_matrix_block_without_calculated_columns() -> ReportBlockPayload:
    return ReportBlockPayload(
        id="cross:country-brand:fallback",
        title="국가 × 브랜드 교차분석",
        subtitle="Germany × 조선미녀",
        meta="매출액 · 금액",
        type="cross_matrix",
        kind="matrix",
        section="cross",
        size="full",
        params={"axis": "country-brand", "metric": "sales", "scale": "amount", "rowLabel": "국가", "columnLabel": "브랜드"},
        snapshot={
            "columns": ["국가", "브랜드", "표시값"],
            "rows": [
                {"국가": "Germany", "브랜드": "조선미녀", "값": 1914.7, "표시값": "€1,914.7K", "metric": "sales", "scale": "amount"},
                {"국가": "Germany", "브랜드": "아누아", "값": 771.5, "표시값": "€771.5K", "metric": "sales", "scale": "amount"},
            ],
        },
    )


def all_report_block_types() -> list[ReportBlockPayload]:
    specs = [
        ("season_calendar", "season", "matrix"),
        ("sku_detail", "sku", "trend"),
        ("brand_rank", "brand", "ranking"),
        ("brand_growth", "brand", "ranking"),
        ("brand_overview", "brand", "kpi"),
        ("brand_sku_concentration", "brand", "ranking"),
        ("brand_country_distribution", "brand", "share"),
        ("brand_line_composition", "brand", "share"),
        ("brand_comparison", "brand", "trend"),
        ("brand_top_sku_seasonality", "brand", "trend"),
        ("cross_matrix", "cross", "matrix"),
        ("country_growth", "region", "ranking"),
        ("region_share", "region", "share"),
        ("country_rank", "region", "ranking"),
        ("country_detail", "region", "kpi"),
        ("country_sku_season", "region", "trend"),
    ]
    blocks = []
    for index, (block_type, section, kind) in enumerate(specs, start=1):
        rows = [{"항목": f"값 {index}", "매출액": index * 1000, "점유율": f"{index}%"}]
        columns = ["항목", "매출액", "점유율"]
        params = {"eur_krw_rate": 1760.49}
        if block_type == "season_calendar":
            rows = [{"기능군": "선케어", "피크": "7월", "피크강도": "높음", "SKU수": 12, **{f"{month}월": month for month in range(1, 13)}}]
            columns = ["기능군", "피크", "피크강도", "SKU수", *[f"{month}월" for month in range(1, 13)]]
        elif block_type == "cross_matrix":
            rows = [{"국가": "Estonia", "브랜드": "아누아", "값": 1000, "표시값": "€1K", "점유율": "25%", "metric": "sales", "scale": "amount"}]
            columns = ["국가", "브랜드", "표시값", "점유율"]
            params.update({"rowLabel": "국가", "columnLabel": "브랜드", "metric": "sales", "scale": "amount"})
        blocks.append(
            ReportBlockPayload(
                id=f"all-types:{block_type}",
                title=f"TYPE {index:02d} {block_type}",
                subtitle="전체 유형 회귀 테스트",
                meta="report cart",
                type=block_type,
                kind=kind,
                section=section,
                size="full",
                params=params,
                snapshot={"columns": columns, "rows": rows},
            )
        )
    return blocks


def brand_block_with_preview_neighbors(brand: str, amount: int) -> ReportBlockPayload:
    return ReportBlockPayload(
        id=f"brand:rank:{brand}",
        title=f"{brand} 브랜드 판매 순위",
        subtitle="브랜드 기준 분석",
        meta="브랜드 기준 분석",
        type="brand_rank",
        kind="ranking",
        section="brand",
        size="full",
        params={"brand": brand, "metric": "sales"},
        snapshot={
            "columns": ["순위", "브랜드명", "매출액", "성장률MoM", "비교기간"],
            "rows": [
                {"순위": 1, "브랜드명": "조선미녀", "매출액": 24_400_000, "성장률MoM": "-29.3%", "비교기간": "2026.05 → 2026.06"},
                {"순위": 2, "브랜드명": "라운드랩", "매출액": 2_660_000, "성장률MoM": "-21.2%", "비교기간": "2026.05 → 2026.06"},
                {"순위": 3, "브랜드명": "아누아", "매출액": 2_250_000, "성장률MoM": "-3.6%", "비교기간": "2026.05 → 2026.06"},
                {"순위": 4, "브랜드명": "이즈앤트리", "매출액": 1_530_000, "성장률MoM": "-20.7%", "비교기간": "2026.05 → 2026.06"},
                {"순위": 5, "브랜드명": "코스알엑스", "매출액": 1_370_000, "성장률MoM": "+59.5%", "비교기간": "2026.05 → 2026.06"},
            ],
        },
    )


def test_partner_report_masks_competitor_names_in_pptx(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="partner", export_format="ppt", blocks=[brand_ranking_block()], title="마스킹 테스트")
    text = pptx_text(output_path)

    for leaked_name in ["토리든", "Torriden", "롬앤", "Romand", "조선미녀", "Beauty of Joseon"]:
        assert leaked_name not in text
    assert "경쟁사 A" in text
    assert "경쟁사 B" in text
    assert "경쟁사 C" in text
    for confidential_value in ["40%", "41%", "39%", "38%"]:
        assert confidential_value not in text


@pytest.mark.skip(reason="PPT export is card-only; raw snapshot detail rows are not rendered as separate PPT tables.")
def test_snapshot_rows_are_rendered_in_pptx(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[country_ranking_block()], title="실데이터 테스트")
    text = pptx_text(output_path)

    assert "미국" in text
    assert "일본" in text
    assert "1200000" in text or "1,200,000" in text or "€1.2M" in text


def test_legacy_region_blocks_without_snapshot_do_not_render_empty_region_slide(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[
            ReportBlockPayload(
                id="region:유럽",
                title="유럽 권역 매출 비중",
                subtitle="60.0% · €3.61M · 약 ₩64.2억 · 30개국",
                meta="국가 · 권역 기준 판매 인사이트",
            ),
            ReportBlockPayload(
                id="country:rank:유럽:Estonia",
                title="Estonia 판매 순위",
                subtitle="유럽 · 점유율 32.0%",
                meta="국가별 판매 순위",
            ),
        ],
        title="레거시 권역 블록 테스트",
    )
    text = pptx_text(output_path)

    assert "표시할 데이터 없음" not in text
    assert "표시할 국가·권역 데이터가 없습니다." not in text
    assert "유럽" in text
    assert "Estonia" in text
    assert "€3.61M" in text
    assert "60.0%" in text
    assert "32.0%" in text


@pytest.mark.skip(reason="PPT export is card-only; raw KRW detail table cells are not rendered as separate PPT tables.")
def test_internal_report_adds_krw_amounts_to_pptx(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[country_ranking_block(), brand_ranking_block(), sku_detail_block()],
        title="원화 환산 테스트",
    )
    text = pptx_text(output_path)

    assert "원화 환산" in text
    assert "국가별 판매 순위" in text
    assert "토리든 브랜드 판매 순위" in text
    assert "[EU] 조선미녀크림 50ml SKU 분석" in text
    assert "약 ₩21.1억" in text
    assert "약 ₩123.2억" in text
    assert "최상위 매출" not in text
    assert "원가/마진" not in text
    assert "수익성 지표" not in text


def test_report_pptx_adds_slide_numbers(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[country_ranking_block()], title="쪽수 테스트")
    text = pptx_text(output_path)
    total_slides = len(Presentation(str(output_path)).slides)

    assert f"01 / {total_slides:02d}" in text
    assert f"{total_slides:02d} / {total_slides:02d}" in text


def test_report_pptx_adds_analysis_period_to_data_slides_only(tmp_path, monkeypatch):
    """분석 기간 라벨은 데이터 슬라이드에만 표시한다(표지·목차·섹션 인트로·Thank You 제외)."""
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    block = export_service.replace_report_block(
        country_ranking_block(),
        params={"eur_krw_rate": 1760.49, "analysis_period_label": "2024-01-01 ~ 2024-12-31"},
    )
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="분석 기간 테스트")
    prs = Presentation(str(output_path))
    slides = list(prs.slides)
    period_text = "분석 기간: 2024-01-01 ~ 2024-12-31"
    labeled = [index for index, slide in enumerate(slides) if period_text in slide_text(slide)]

    assert labeled == [2]


def test_report_pptx_marks_mixed_analysis_periods(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    first = export_service.replace_report_block(
        country_ranking_block(),
        params={"eur_krw_rate": 1760.49, "analysis_period_label": "2024-01-01 ~ 2024-12-31"},
    )
    second = export_service.replace_report_block(
        brand_ranking_block(),
        params={"eur_krw_rate": 1760.49, "analysis_period_label": "2025-01-01 ~ 2025-12-31"},
    )
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[first, second], title="혼합 기간 테스트")
    prs = Presentation(str(output_path))
    slides = list(prs.slides)
    period_text = "분석 기간: 여러 기간 혼합: 2024-01-01 ~ 2024-12-31 외 1개"
    labeled = [index for index, slide in enumerate(slides) if period_text in slide_text(slide)]

    assert labeled == [2, 3]


def test_region_insights_keep_web_share_basis_for_low_rank_page(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[low_rank_country_ranking_block()], title="낮은 순위 국가 테스트")
    text = pptx_text(output_path)

    assert "CARD 01 · 국가 순위" in text
    assert "Lithuania" in text
    assert "0.7%" in text
    assert "전체의 100%" not in text
    assert "현재 데이터에서 가장 큰 시장" not in text


@pytest.mark.skip(reason="PPT export is card-only; detail table growth headers are no longer generated.")
def test_growth_headers_show_yoy_or_mom_basis(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    region_path = build_report_export(audience="internal", export_format="ppt", blocks=[low_rank_country_ranking_block()], title="성장률 기준 테스트")
    region_text = pptx_text(region_path)
    region_rows = first_pptx_table_rows_with_headers(region_path, {"YoY"})

    assert region_rows
    assert "-65%" in region_text

    brand_path = build_report_export(audience="internal", export_format="ppt", blocks=[brand_block_with_preview_neighbors("조선미녀", 24_400_000)], title="성장률 기준 테스트")
    brand_rows = first_pptx_table_rows_with_headers(brand_path, {"성장률(MoM)"})

    assert brand_rows
    assert "-29.3%" in pptx_text(brand_path)


def test_generic_growth_headers_are_never_ambiguous():
    assert export_service.display_table_header("성장률") == "성장률(MoM)"
    assert export_service.display_table_header("성장률MoM") == "성장률(MoM)"
    assert export_service.display_table_header("YoY") == "YoY"
    assert export_service.display_table_header("growth_yoy") == "성장률(YoY)"


@pytest.mark.skip(reason="PPT export is card-only; long-tail detail rows are intentionally not rendered in PPT.")
def test_brand_rank_keeps_growth_for_rows_outside_growth_card_top_five(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    brands = ["조선미녀", "메디큐브", "바이오던스", "닥터엘시아", "코스알엑스", "아누아", "라운드랩"]
    growth = ["-77.7%", "-61.5%", "-65.1%", "-64.1%", "-57.5%", "-8.3%", "-63.8%"]
    rank_block = ReportBlockPayload(
        id="brand:ranking:all-growth",
        title="전체 브랜드 판매 순위",
        subtitle="7개 브랜드",
        meta="브랜드 기준 분석",
        type="brand_rank",
        kind="ranking",
        section="brand",
        size="full",
        params={"scope": "all_brands", "eur_krw_rate": 1712.88},
        snapshot={
            "columns": ["순위", "브랜드명", "매출액", "성장률MoM"],
            "rows": [
                {"순위": index, "브랜드명": brand, "매출액": (8 - index) * 1_000_000, "성장률MoM": growth[index - 1]}
                for index, brand in enumerate(brands, start=1)
            ],
        },
    )
    growth_block = ReportBlockPayload(
        id="brand:growth:mom:top-five",
        title="브랜드별 성장 변화",
        subtitle="상위 5개만 표시",
        meta="브랜드별 매출 성장률 비교",
        type="brand_growth",
        kind="ranking",
        section="brand",
        size="full",
        params={"comparison_basis": "mom", "comparison_month": "2026-06", "target_month": "2026-07"},
        snapshot={
            "columns": ["순위", "브랜드명", "성장률"],
            "rows": [
                {"순위": index, "브랜드명": brand, "성장률": growth[index - 1]}
                for index, brand in enumerate(brands[:5], start=1)
            ],
        },
    )

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[growth_block, rank_block],
        title="브랜드 성장률 누락 회귀 테스트",
    )
    deck_text = pptx_text(output_path)

    assert "아누아" in deck_text and "-8.3%" in deck_text
    assert "라운드랩" in deck_text and "-63.8%" in deck_text


@pytest.mark.skip(reason="PPT export is card-only; paginated detail tables are no longer generated.")
def test_brand_rank_keeps_growth_column_on_last_page_when_last_row_has_no_growth(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    rows = []
    for index in range(1, 72):
        row = {
            "순위": index,
            "브랜드명": f"BRAND-{index:02d}",
            "매출액": 72_000_000 - index,
            "매출표시": f"€{72 - index / 10:.1f}M",
            "원화표시": f"약 ₩{1200 - index:.1f}억",
            "비교기간": "2026.05 → 2026.06",
        }
        if index < 71:
            row["성장률MoM"] = f"+{index / 10:.1f}%"
        rows.append(row)
    block = ReportBlockPayload(
        id="brand:ranking:last-page-growth-missing",
        title="전체 브랜드 판매 순위",
        subtitle="마지막 페이지 성장률 누락 회귀 테스트",
        meta="브랜드 기준 분석",
        type="brand_rank",
        kind="ranking",
        section="brand",
        size="full",
        params={"scope": "all_brands"},
        snapshot={
            "columns": ["순위", "브랜드명", "매출액", "매출표시", "원화표시", "성장률MoM", "비교기간"],
            "rows": rows,
        },
    )

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[block],
        title="브랜드 마지막 페이지 성장률 컬럼 회귀 테스트",
    )
    prs = Presentation(str(output_path))
    growth_tables = [
        shape.table
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_table", False)
        and "성장률(MoM)" in [cell.text for cell in shape.table.rows[0].cells]
    ]
    assert growth_tables
    table = growth_tables[-1]
    headers = [cell.text for cell in table.rows[0].cells]
    values = [cell.text for cell in table.rows[len(table.rows) - 1].cells]

    assert "성장률(MoM)" in headers
    assert values[headers.index("성장률(MoM)")] == "-"


def test_report_card_numeric_percent_values_drive_bar_width(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    block = ReportBlockPayload(
        id="region:share:numeric-percent",
        title="권역별 매출 비중",
        subtitle="숫자형 점유율 막대 회귀 테스트",
        meta="권역별 매출 비중",
        type="region_share",
        kind="share",
        section="region",
        size="full",
        snapshot={
            "columns": ["순위", "권역", "매출액", "점유율"],
            "rows": [
                {"순위": 1, "권역": "아시아", "매출액": 10, "점유율": 0},
                {"순위": 2, "권역": "기타", "매출액": 10, "점유율": 0},
                {"순위": 3, "권역": "아프리카", "매출액": 10, "점유율": 0},
                {"순위": 4, "권역": "유럽", "매출액": 310_400_000, "점유율": 99.9},
            ],
        },
    )

    columns, rows = export_service.report_snapshot_rows_and_columns(block, "internal")
    card_rows = export_service.report_card_rows(block, columns, rows)
    by_label = {row["label"]: row for row in card_rows}

    assert by_label["유럽"]["value"] == "99.9%"
    assert by_label["유럽"]["ratio"] == 1
    assert by_label["아시아"]["value"] == "0%"
    assert by_label["아시아"]["ratio"] == 0

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="퍼센트 막대 회귀 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · 권역 비중" in slide_text(slide))
    bars = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Report Card Value Bar"]
    widths = [shape.width for shape in bars]

    assert len(widths) == 4
    assert max(widths) > min(widths) * 20


@pytest.mark.skip(reason="PPT export is card-only; structured detail table alignment no longer applies.")
def test_structured_brand_name_with_number_stays_left_aligned(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    block = ReportBlockPayload(
        id="brand:ranking:numbered-name",
        title="전체 브랜드 판매 순위",
        subtitle="숫자 포함 브랜드명 정렬 회귀 테스트",
        meta="브랜드 기준 분석",
        type="brand_rank",
        kind="ranking",
        section="brand",
        size="full",
        params={"scope": "all_brands"},
        snapshot={
            "columns": ["순위", "브랜드명", "매출액", "매출표시", "원화표시", "SKU수"],
            "rows": [
                {
                    "순위": 28,
                    "브랜드명": "센텔리안24",
                    "매출액": 471_158.94,
                    "매출표시": "€471.1K",
                    "원화표시": "약 ₩8.04억",
                    "SKU수": 18,
                },
                {
                    "순위": 63,
                    "브랜드명": "스튜디오 17",
                    "매출액": 1_754.40,
                    "매출표시": "€1.8K",
                    "원화표시": "약 ₩29만",
                    "SKU수": 31,
                }
            ],
        },
    )

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[block],
        title="숫자 포함 브랜드명 정렬 테스트",
    )
    prs = Presentation(str(output_path))
    tables = [
        shape.table
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_table", False)
        and "브랜드명" in [cell.text for cell in shape.table.rows[0].cells]
    ]
    table = tables[0]
    headers = [cell.text for cell in table.rows[0].cells]
    brand_index = headers.index("브랜드명")
    amount_index = headers.index("매출액(EUR)")

    for row_index, expected_brand in ((1, "센텔리안24"), (2, "스튜디오 17")):
        brand_cell = table.cell(row_index, brand_index)
        amount_cell = table.cell(row_index, amount_index)
        assert brand_cell.text == expected_brand
        assert brand_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
        assert amount_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT


@pytest.mark.skip(reason="PPT export is card-only; all-row detail pagination is intentionally removed.")
def test_template_brand_cart_renders_all_detail_rows_in_cart_order(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    brands = [f"RANK-BRAND-{index:02d}" for index in range(1, 16)]
    growth_block = ReportBlockPayload(
        id="brand:growth:all-rows",
        title="GROWTH DETAIL BLOCK",
        subtitle="18 growth rows",
        meta="growth detail",
        type="brand_growth",
        kind="ranking",
        section="brand",
        size="full",
        params={"comparison_basis": "mom", "comparison_month": "2026-05", "target_month": "2026-06"},
        snapshot={
            "columns": ["구분", "순위", "브랜드명", "성장률", "행표식"],
            "rows": [
                {
                    "구분": "성장률 전체",
                    "순위": index,
                    "브랜드명": brands[(index - 1) % len(brands)],
                    "성장률": f"{index:+d}%",
                    "행표식": f"GROWTH-ROW-{index:02d}",
                }
                for index in range(1, 19)
            ],
        },
    )
    rank_block = ReportBlockPayload(
        id="brand:ranking:all-rows",
        title="RANK DETAIL BLOCK",
        subtitle="15 rank rows",
        meta="rank detail",
        type="brand_rank",
        kind="ranking",
        section="brand",
        size="full",
        params={"scope": "all_brands"},
        snapshot={
            "columns": ["순위", "브랜드명", "매출액"],
            "rows": [
                {"순위": index, "브랜드명": brand, "매출액": (16 - index) * 1_000_000}
                for index, brand in enumerate(brands, start=1)
            ],
        },
    )
    comparison_block = ReportBlockPayload(
        id="brand:comparison:all-rows",
        title="COMPARISON DETAIL BLOCK",
        subtitle="24 comparison rows",
        meta="comparison detail",
        type="brand_comparison",
        kind="trend",
        section="brand",
        size="half",
        params={"brands": brands[:2]},
        snapshot={
            "columns": ["브랜드명", "월", "매출액", "브랜드 내 월 매출 비중", "행표식"],
            "rows": [
                {
                    "브랜드명": brands[(index - 1) % 2],
                    "월": f"{((index - 1) % 12) + 1}월",
                    "매출액": index * 10_000,
                    "브랜드 내 월 매출 비중": f"{index / 10:.1f}%",
                    "행표식": f"COMPARE-ROW-{index:02d}",
                }
                for index in range(1, 25)
            ],
        },
    )

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[growth_block, rank_block, comparison_block],
        title="브랜드 장바구니 전체 행 회귀 테스트",
    )
    prs = Presentation(str(output_path))
    slide_texts = [slide_text(slide) for slide in prs.slides]
    deck_text = pptx_text(output_path)

    for brand in brands:
        assert brand in deck_text
    card_pages = [(index, text) for index, text in enumerate(slide_texts) if "CARD " in text]
    assert not any("BLOCK " in text for text in slide_texts)
    assert [any(title in text for _, text in card_pages) for title in ("GROWTH DETAIL BLOCK", "RANK DETAIL BLOCK", "COMPARISON DETAIL BLOCK")] == [True, True, True]
    assert duplicate_pptx_part_names(output_path) == set()


def test_region_share_page_uses_its_own_insight_cards(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[low_rank_country_ranking_block(), region_share_block()],
        title="권역 카드 중복 테스트",
    )
    prs = Presentation(str(output_path))
    slide_texts = [slide_text(slide) for slide in prs.slides]
    share_slide_text = next(text for text in slide_texts if "CARD 02 · 권역 비중" in text)

    assert "권역별 매출 비중" in share_slide_text
    assert "유럽" in share_slide_text
    assert "99.9%" in share_slide_text
    assert "성장률" not in share_slide_text
    assert "Lithuania 매출 비중" not in share_slide_text
    assert "Bulgaria보다" not in share_slide_text
    assert all("01 국가별 판매 순위 및 성장" not in text for text in slide_texts)


def test_country_detail_card_uses_web_like_composite_layout(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[country_detail_card_block()], title="국가 상세 웹 카드 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · 국가 판매 상세" in slide_text(slide))
    text = slide_text(card_slide)
    summary_tiles = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Country Detail Summary Tile"]
    brand_bars = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Country Detail Brand Bar"]
    category_pills = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Country Detail Category Pill"]
    generic_bars = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Report Card Value Bar"]

    assert "선택 국가" in text
    assert "Germany" in text
    assert "총 매출" in text
    assert "€52.2M" in text
    assert "점유율 16.8%" in text
    assert "월 최고 매출" in text
    assert "상위 브랜드 · 미리보기" in text
    assert "조선미녀" in text
    assert "메디큐브" in text
    assert "Top 5 카테고리" in text
    assert "썬케어 > 크림" in text
    assert len(summary_tiles) == 3
    assert len(brand_bars) >= 2
    assert len(category_pills) >= 2
    assert not generic_bars


def test_region_template_keeps_only_country_rows_and_merges_growth():
    blocks = [
        ReportBlockPayload(
            id="country:rank:all",
            title="국가별 판매 순위",
            subtitle="2개국",
            meta="국가 순위",
            type="country_rank",
            kind="ranking",
            section="region",
            size="full",
            params={"yoy_target_month": "2024-12", "yoy_comparison_month": "2023-12"},
            snapshot={
                "columns": ["순위", "국가", "권역", "매출액", "점유율"],
                "rows": [
                    {"순위": 1, "국가": "Germany", "권역": "유럽", "매출액": 38_300_000, "점유율": "18.1%"},
                    {"순위": 2, "국가": "United Kingdom", "권역": "유럽", "매출액": 37_100_000, "점유율": "17.5%"},
                ],
            },
        ),
        ReportBlockPayload(
            id="country:growth:yoy",
            title="국가별 성장 변화",
            subtitle="YoY",
            meta="국가 성장",
            type="country_growth",
            kind="ranking",
            section="region",
            size="full",
            params={
                "comparison_basis": "yoy",
                "target_month": "2024-12",
                "comparison_month": "2023-12",
                "period_label": "2023.12 → 2024.12",
            },
            snapshot={
                "columns": ["국가", "권역", "성장률"],
                "rows": [
                    {"국가": "Germany", "권역": "유럽", "성장률": "-42.9%"},
                    {"국가": "United Kingdom", "권역": "유럽", "성장률": "+72.9%"},
                ],
            },
        ),
        ReportBlockPayload(
            id="country:detail:germany",
            title="Germany 국가 판매 상세",
            subtitle="상세",
            meta="국가 상세",
            type="country_detail",
            kind="kpi",
            section="region",
            size="full",
            params={"country": "Germany", "region": "유럽", "amount": 38_300_000, "share": 18.1},
            snapshot={
                "columns": ["구분", "항목", "매출액"],
                "rows": [
                    {"구분": "국가 요약", "항목": "Germany", "매출액": 38_300_000},
                    {"구분": "상위 브랜드", "항목": "조선미녀", "매출액": 10_000_000},
                ],
            },
        ),
        ReportBlockPayload(
            id="country:sku-season:germany",
            title="상위 SKU · Germany",
            subtitle="시즌성",
            meta="국가 SKU",
            type="country_sku_season",
            kind="trend",
            section="region",
            size="full",
            params={"country": "Germany"},
            snapshot={
                "columns": ["구분", "항목", "매출액"],
                "rows": [
                    {"구분": "상위 SKU", "항목": "상품 A", "매출액": 5_000_000},
                    {"구분": "국가 시즌성", "항목": "12월", "매출액": 3_000_000},
                ],
            },
        ),
    ]

    descriptors, _ = export_service.region_page_descriptors(blocks)
    rows = [row for descriptor in descriptors if descriptor["title"].startswith("01 국가별") for row in descriptor["rows"]]

    assert [row["국가"] for row in rows] == ["Germany", "United Kingdom"]
    assert rows[0]["YoY"] == "-42.9%"
    assert rows[1]["YoY"] == "+72.9%"
    assert rows[0]["_yoy_target_month"] == "2024-12"
    assert rows[0]["_yoy_comparison_month"] == "2023-12"
    assert "YoY 성장률: -42.9% (기준월 2024년 12월 vs 비교월 2023년 12월)" in export_service.build_region_insights(rows)[0]


def test_country_growth_only_snapshot_keeps_sales_share_and_krw_in_region_summary(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    block = ReportBlockPayload(
        id="country:growth:mom:with-amounts",
        title="국가별 성장 변화",
        subtitle="2026.05 → 2026.06",
        meta="국가 성장",
        type="country_growth",
        kind="ranking",
        section="region",
        size="full",
        params={
            "comparison_basis": "mom",
            "target_month": "2026-06",
            "comparison_month": "2026-05",
            "eur_krw_rate": 1_700,
        },
        snapshot={
            "columns": ["구분", "순위", "국가", "권역", "기준월 매출", "비교월 금액", "점유율", "성장률", "비교기간"],
            "rows": [
                {
                    "구분": "성장률 상위",
                    "순위": 1,
                    "국가": "United Kingdom",
                    "권역": "유럽",
                    "기준월 매출": "€1.31M",
                    "비교월 금액": "€1.00M",
                    "점유율": "56.0%",
                    "성장률": "+31%",
                    "비교기간": "2026.05 → 2026.06",
                    "매출액": 1_310_000,
                    "매출표시": "€1.31M",
                    "원화표시": "약 ₩22.3억",
                },
                {
                    "구분": "성장률 상위",
                    "순위": 2,
                    "국가": "Germany",
                    "권역": "유럽",
                    "기준월 매출": "€1.03M",
                    "비교월 금액": "€777K",
                    "점유율": "44.0%",
                    "성장률": "+32.6%",
                    "비교기간": "2026.05 → 2026.06",
                    "매출액": 1_030_000,
                    "매출표시": "€1.03M",
                    "원화표시": "약 ₩17.5억",
                },
            ],
        },
    )

    rows = export_service.region_section_rows([block])
    table_rows = export_service.build_region_table_rows(
        rows,
        include_krw=True,
        eur_krw_rate=1_700,
        growth_key="MoM",
        include_growth=True,
    )

    assert table_rows == [
        ["1", "United Kingdom", "€1.31M", "약 ₩22.3억", "56.0%", "+31%"],
        ["2", "Germany", "€1.03M", "약 ₩17.5억", "44.0%", "+32.6%"],
    ]

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[block],
        title="국가 성장 카드 단독 PPT 회귀 테스트",
    )
    deck_text = pptx_text(output_path)
    for expected in ("United Kingdom", "Germany", "€1.31M", "€1.03M", "+31%", "+32.6%"):
        assert expected in deck_text


def test_country_growth_card_shows_rising_and_declining_like_web_card(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    block = ReportBlockPayload(
        id="country:growth:mom:split-card",
        title="국가별 성장 변화",
        subtitle="2026.05 → 2026.06 · 상승 2개국 · 하락 2개국",
        meta="국가별 매출 성장률 비교",
        type="country_growth",
        kind="ranking",
        section="region",
        size="full",
        params={"comparison_basis": "mom", "target_month": "2026-06", "comparison_month": "2026-05"},
        snapshot={
            "columns": ["구분", "순위", "국가", "권역", "기준월 매출", "비교월 금액", "성장률"],
            "rows": [
                {"구분": "성장률 상위", "순위": 1, "국가": "United Kingdom", "권역": "유럽", "기준월 매출": "€5.1M", "비교월 금액": "€3.89M", "성장률": "+31%"},
                {"구분": "성장률 상위", "순위": 2, "국가": "Germany", "권역": "유럽", "기준월 매출": "€3.5M", "비교월 금액": "€2.64M", "성장률": "+32.6%"},
                {"구분": "감소율 상위", "순위": 1, "국가": "Poland", "권역": "유럽", "기준월 매출": "€1.81M", "비교월 금액": "€2.38M", "성장률": "-24%"},
                {"구분": "감소율 상위", "순위": 2, "국가": "France", "권역": "유럽", "기준월 매출": "€1.22M", "비교월 금액": "€1.72M", "성장률": "-28.8%"},
            ],
        },
    )

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="국가 성장 카드 웹 구조 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · 국가 성장 변화" in slide_text(slide))
    text = slide_text(card_slide)

    assert "성장 상위" in text
    assert "감소 상위" in text
    assert "United Kingdom" in text
    assert "Germany" in text
    assert "Poland" in text
    assert "France" in text
    assert "+31%" in text
    assert "-24%" in text


def test_brand_growth_card_shows_rising_and_declining_like_web_card(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    block = ReportBlockPayload(
        id="brand:growth:mom:split-card",
        title="브랜드별 성장 변화",
        subtitle="2026.05 → 2026.06 · 상승 2개 · 하락 2개",
        meta="브랜드별 매출 성장률 비교",
        type="brand_growth",
        kind="ranking",
        section="brand",
        size="full",
        params={"comparison_basis": "mom", "target_month": "2026-06", "comparison_month": "2026-05"},
        snapshot={
            "columns": ["구분", "순위", "브랜드명", "성장률", "비교기간"],
            "rows": [
                {"구분": "성장률 상위", "순위": 1, "브랜드명": "조선미녀", "성장률": "+12.4%", "비교기간": "2026.05 → 2026.06"},
                {"구분": "성장률 상위", "순위": 2, "브랜드명": "아누아", "성장률": "+8.1%", "비교기간": "2026.05 → 2026.06"},
                {"구분": "감소율 상위", "순위": 1, "브랜드명": "메디큐브", "성장률": "-3.2%", "비교기간": "2026.05 → 2026.06"},
                {"구분": "감소율 상위", "순위": 2, "브랜드명": "토리든", "성장률": "-2.4%", "비교기간": "2026.05 → 2026.06"},
            ],
        },
    )

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="브랜드 성장 카드 웹 구조 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · 브랜드 성장" in slide_text(slide))
    text = slide_text(card_slide)

    assert "성장 상위" in text
    assert "감소 상위" in text
    assert "조선미녀" in text
    assert "아누아" in text
    assert "메디큐브" in text
    assert "토리든" in text
    assert "+12.4%" in text
    assert "-3.2%" in text


def test_brand_overview_card_uses_kpi_tiles_like_web_card(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    block = ReportBlockPayload(
        id="brand:overview:anua",
        title="아누아 브랜드 요약",
        subtitle="매출 €12.3M · 판매수량 12345 · SKU 42개 · 판매 국가 17개국",
        meta="선택 브랜드 핵심 지표",
        type="brand_overview",
        kind="kpi",
        section="brand",
        size="full",
        snapshot={
            "columns": ["브랜드명", "매출액", "매출표시", "원화표시", "판매수량", "SKU수", "국가수", "점유율"],
            "rows": [{"브랜드명": "아누아", "매출액": 12_300_000, "매출표시": "€12.3M", "원화표시": "약 ₩210.0억", "판매수량": 12345, "SKU수": 42, "국가수": 17, "점유율": "8.4%"}],
        },
    )

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="브랜드 요약 KPI 카드 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · 브랜드 개요" in slide_text(slide))
    text = slide_text(card_slide)

    assert "매출" in text
    assert "원화 환산" in text
    assert "판매수량" in text
    assert "SKU" in text
    assert "판매 국가" in text
    assert "€12.3M" in text
    assert "약 ₩210.0억" in text
    assert "12,345" in text
    assert "42개" in text
    assert "17개국" in text
    assert "TOP" not in text


def test_sku_order_reference_card_uses_web_like_kpi_tiles(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[sku_order_reference_block()], title="SKU 발주 참고 카드 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · SKU 상세" in slide_text(slide))
    text = slide_text(card_slide)
    tiles = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 SKU Order Reference Tile"]
    bars = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Report Card Value Bar"]

    assert "매출 순위" in text
    assert "피크 3개월 전 준비월" in text
    assert "판매수량" in text
    assert "매출" in text
    assert "약 ₩89.8억" in text
    assert "판매 국가" in text
    assert "피크월 매출" in text
    assert "6월 피크 기준 약 3개월 전인 3월부터 판매/발주 준비가 필요합니다." in text
    assert len(tiles) == 6
    assert not bars


def test_sku_order_reference_card_uses_tiles_even_when_kind_is_not_kpi(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    block = export_service.replace_report_block(sku_order_reference_block(), kind="trend")

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="SKU 발주 참고 kind 회귀 테스트")
    prs = Presentation(str(output_path))
    card_slide = next(slide for slide in prs.slides if "CARD 01 · SKU 상세" in slide_text(slide))
    text = slide_text(card_slide)
    tiles = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 SKU Order Reference Tile"]
    bars = [shape for shape in card_slide.shapes if getattr(shape, "name", "") == "Silicon2 Report Card Value Bar"]

    assert "발주 참고 정보" in text
    assert "피크 3개월 전 준비월" in text
    assert len(tiles) == 6
    assert not bars


@pytest.mark.skip(reason="PPT export is card-only; country detail sections are no longer split into detail table slides.")
def test_country_detail_splits_mixed_sections_into_separate_detail_pages(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    block = ReportBlockPayload(
        id="country:detail:spain",
        title="Spain 국가 판매 상세",
        subtitle="유럽 · 총 매출 €22.4M · 점유율 7.2%",
        meta="선택 국가 판매 인사이트",
        type="country_detail",
        kind="kpi",
        section="region",
        size="full",
        params={
            "country": "Spain",
            "region": "유럽",
            "amount": 22_435_870.76,
            "share": 7.2,
            "analysis_period_label": "2024-07-14 ~ 2026-07-14",
        },
        snapshot={
            "columns": ["구분", "순위", "항목", "매출액", "매출표시", "원화표시", "비고"],
            "rows": [
                {
                    "구분": "국가 요약",
                    "순위": "",
                    "항목": "Spain",
                    "매출액": 22_435_870.76,
                    "매출표시": "€22.4M",
                    "원화표시": "약 ₩382.4억",
                    "비고": "점유율 7.2%",
                },
                {
                    "구분": "월 최고 매출",
                    "순위": "",
                    "항목": "3월",
                    "매출액": 3_269_752.97,
                    "매출표시": "€3.27M",
                    "원화표시": "약 ₩55.7억",
                    "비고": "3월 최고",
                },
                {
                    "구분": "상위 브랜드",
                    "순위": 1,
                    "항목": "조선미녀",
                    "매출액": 6_874_992.29,
                    "매출표시": "€6.87M",
                    "원화표시": "약 ₩117.2억",
                    "비고": "",
                },
                {
                    "구분": "상위 브랜드",
                    "순위": 2,
                    "항목": "닥터엘시아",
                    "매출액": 3_908_748.44,
                    "매출표시": "€3.91M",
                    "원화표시": "약 ₩66.6억",
                    "비고": "",
                },
                {
                    "구분": "Top 5 카테고리",
                    "순위": 1,
                    "항목": "스킨케어 > 크림",
                    "매출액": 5_248_103.34,
                    "매출표시": "€5.25M",
                    "원화표시": "약 ₩89.4억",
                    "비고": "",
                },
                {
                    "구분": "Top 5 카테고리",
                    "순위": 2,
                    "항목": "스킨케어 > 세럼",
                    "매출액": 3_137_185.66,
                    "매출표시": "€3.14M",
                    "원화표시": "약 ₩53.5억",
                    "비고": "",
                },
            ],
        },
    )

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[block],
        title="국가 상세 혼합 구분 분리 테스트",
    )
    prs = Presentation(str(output_path))
    card_slides = [slide for slide in prs.slides if "CARD 01" in slide_text(slide)]
    detail_slides = [slide for slide in prs.slides if "BLOCK 01" in slide_text(slide)]
    slide_texts = [slide_text(slide) for slide in detail_slides]

    assert len(card_slides) == 1
    card_text = slide_text(card_slides[0])
    assert "Spain 국가 판매 상세" in card_text
    assert "조선미녀" in card_text
    assert "스킨케어 > 크림" in card_text
    assert not any(getattr(shape, "has_table", False) for shape in card_slides[0].shapes)
    assert len(detail_slides) == 4
    assert any("Spain 국가 판매 상세 · 국가 요약" in text for text in slide_texts)
    assert any("Spain 국가 판매 상세 · 월 최고 매출" in text for text in slide_texts)
    assert any("Spain 국가 판매 상세 · 상위 브랜드" in text for text in slide_texts)
    assert any("Spain 국가 판매 상세 · Top 5 카테고리" in text for text in slide_texts)

    table_text_by_group: dict[str, str] = {}
    for slide in detail_slides:
        text = slide_text(slide)
        table_shape = next(shape for shape in slide.shapes if getattr(shape, "has_table", False))
        headers = [cell.text for cell in table_shape.table.rows[0].cells]
        assert "구분" not in headers
        table_text = "\n".join(cell.text for row in table_shape.table.rows for cell in row.cells)
        for group in ("국가 요약", "월 최고 매출", "상위 브랜드", "Top 5 카테고리"):
            if f"Spain 국가 판매 상세 · {group}" in text:
                table_text_by_group[group] = table_text

    assert "Spain" in table_text_by_group["국가 요약"]
    assert "3월" in table_text_by_group["월 최고 매출"]
    assert "조선미녀" in table_text_by_group["상위 브랜드"]
    assert "닥터엘시아" in table_text_by_group["상위 브랜드"]
    assert "스킨케어 > 크림" in table_text_by_group["Top 5 카테고리"]
    assert "스킨케어 > 세럼" in table_text_by_group["Top 5 카테고리"]
    assert "월 최고 매출" not in table_text_by_group["상위 브랜드"]
    assert "상위 브랜드" not in table_text_by_group["Top 5 카테고리"]

    html_path = build_report_export(
        audience="internal",
        export_format="html",
        blocks=[block],
        title="국가 상세 혼합 구분 HTML 분리 테스트",
    )
    html = html_path.read_text(encoding="utf-8")
    assert "<h4>국가 요약</h4>" in html
    assert "<h4>월 최고 매출</h4>" in html
    assert "<h4>상위 브랜드</h4>" in html
    assert "<h4>Top 5 카테고리</h4>" in html
    assert "<th>구분</th>" not in html


@pytest.mark.skip(reason="PPT export is card-only; country SKU season detail tables are no longer generated.")
def test_country_sku_season_detail_pages_are_readable_complete_and_clear_of_period_label(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    rows = [
        {
            "구분": "상위 SKU",
            "순위": index,
            "항목": f"긴 상품명 테스트 SKU {index} 50ml",
            "브랜드": "조선미녀" if index % 2 else "바이오던스",
            "매출액": 4_324_702.82 / index,
            "매출표시": f"€{4.32 / index:.2f}M",
            "원화표시": f"약 ₩{73.7 / index:.1f}억",
            "월": "",
            "매출비중": "",
        }
        for index in range(1, 6)
    ]
    rows.extend(
        {
            "구분": "국가 시즌성",
            "순위": "",
            "항목": f"{month}월",
            "브랜드": "",
            "매출액": 5_326_628.18 / month,
            "매출표시": f"€{5.33 / month:.2f}M",
            "원화표시": f"약 ₩{90.8 / month:.1f}억",
            "월": month,
            "매출비중": f"{month:.1f}%",
        }
        for month in range(1, 14)
    )
    block = ReportBlockPayload(
        id="country:sku-season:Germany",
        title="상위 SKU · Germany",
        subtitle="6월 최고 비중 · 11.4%",
        meta="선택 국가 상위 SKU 및 월별 매출 비중",
        type="country_sku_season",
        kind="trend",
        section="region",
        size="full",
        params={
            "country": "Germany",
            "eur_krw_rate": 1_704.39,
            "analysis_start_date": "2024-07-14",
            "analysis_end_date": "2026-07-14",
            "analysis_period_label": "2024-07-14 ~ 2026-07-14",
        },
        snapshot={
            "columns": ["구분", "순위", "항목", "브랜드", "매출액", "매출표시", "원화표시", "월", "매출비중"],
            "rows": rows,
        },
    )

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[block],
        title="국가 SKU 시즌성 상세 표 회귀 테스트",
    )
    prs = Presentation(str(output_path))
    card_slides = [slide for slide in prs.slides if "CARD 01" in slide_text(slide)]
    detail_slides = [slide for slide in prs.slides if "BLOCK 01" in slide_text(slide)]

    assert len(card_slides) == 1
    assert "담김" not in slide_text(card_slides[0])
    assert "1월" in slide_text(card_slides[0])
    assert "13월" in slide_text(card_slides[0])
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.LINE for shape in card_slides[0].shapes) >= 12
    assert len(detail_slides) == 4
    detail_text = "\n".join(
        cell.text
        for slide in detail_slides
        for shape in slide.shapes
        if getattr(shape, "has_table", False)
        for row in shape.table.rows
        for cell in row.cells
    )
    assert "매출액(EUR)" in detail_text
    assert "매출 규모" in detail_text
    assert "€4,324,702.82" in detail_text
    assert "4324702.82" not in detail_text
    assert "전체 브랜드" in detail_text

    for slide in detail_slides:
        assert block.meta not in slide_text(slide)
        table_shapes = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
        assert len(table_shapes) == 1
        table_shape = table_shapes[0]
        headers = [cell.text for cell in table_shape.table.rows[0].cells]
        assert {"구분", "항목", "브랜드"}.issubset(headers)
        assert all(cell.text.strip() for row in table_shape.table.rows for cell in row.cells)
        assert slide_text(slide).count(audience_badge("internal")) == 1
        period_shape = next(shape for shape in slide.shapes if getattr(shape, "name", "") == "Silicon2 Analysis Period")
        assert table_shape.top + table_shape.height <= period_shape.top - Inches(0.45)

    html_path = build_report_export(
        audience="internal",
        export_format="html",
        blocks=[block],
        title="국가 SKU 시즌성 HTML 표시 테스트",
    )
    html = html_path.read_text(encoding="utf-8")
    assert "<td>€4,324,702.82</td>" in html
    assert "<td>전체 브랜드</td>" in html


@pytest.mark.skip(reason="PPT export is card-only; brand top SKU seasonality detail tables are no longer generated.")
def test_brand_top_sku_seasonality_detail_splits_sku_and_monthly_rows(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    rows = [
        {
            "구분": "상위 SKU",
            "순위": index,
            "SKU": f"cos{index:02d}",
            "상품명": f"코스알엑스 테스트 SKU {index}",
            "카테고리": "스킨케어",
            "관측 완료월 수": "",
            "매출액": 1_000_000 / index,
            "매출표시": f"€{1_000_000 / index:,.2f}",
            "판매수량": 100 * index,
            "월": "",
            "매출비중": "",
        }
        for index in range(1, 6)
    ]
    rows.extend(
        {
            "구분": "월별 매출 비중",
            "순위": "",
            "SKU": "-",
            "상품명": "-",
            "카테고리": "-",
            "관측 완료월 수": 2,
            "매출액": 200_000 * month,
            "매출표시": f"€{200_000 * month:,.2f}",
            "판매수량": "-",
            "월": f"{month}월",
            "매출비중": f"{month * 1.1:.1f}%",
        }
        for month in range(1, 7)
    )
    block = ReportBlockPayload(
        id="brand:top-sku-seasonality:cosrx",
        title="코스알엑스 상위 SKU·월별 판매 시즌성",
        subtitle="상위 5개 SKU · 완료월 연평균 피크월 6월 15.6%",
        meta="선택 브랜드 상위 SKU 및 완료월 월평균 매출 비중",
        type="brand_top_sku_seasonality",
        kind="trend",
        section="brand",
        size="full",
        snapshot={
            "columns": ["구분", "순위", "SKU", "상품명", "카테고리", "관측 완료월 수", "매출액", "매출표시", "판매수량", "월", "매출비중"],
            "rows": rows,
        },
    )

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="브랜드 SKU 시즌성 분리 테스트")
    prs = Presentation(str(output_path))
    detail_slides = [slide for slide in prs.slides if "BLOCK 01" in slide_text(slide)]
    top_sku_slide = next(slide for slide in detail_slides if "상위 SKU" in slide_text(slide))
    monthly_slide = next(slide for slide in detail_slides if "월별 매출 비중" in slide_text(slide))
    top_table = next(shape.table for shape in top_sku_slide.shapes if getattr(shape, "has_table", False))
    monthly_table = next(shape.table for shape in monthly_slide.shapes if getattr(shape, "has_table", False))
    top_headers = [cell.text for cell in top_table.rows[0].cells]
    monthly_headers = [cell.text for cell in monthly_table.rows[0].cells]
    top_text = "\n".join(cell.text for row in top_table.rows for cell in row.cells)
    monthly_text = "\n".join(cell.text for row in monthly_table.rows for cell in row.cells)

    assert "월" not in top_headers
    assert "매출 비중" not in top_headers
    assert "SKU" not in monthly_headers
    assert "상품명" not in monthly_headers
    assert "카테고리" not in monthly_headers
    assert "cos01" in top_text
    assert "1월" in monthly_text
    assert not any(cell.text.strip() == "-" for row in top_table.rows for cell in row.cells)
    assert not any(cell.text.strip() == "-" for row in monthly_table.rows for cell in row.cells)


def test_region_insight_explains_top_three_country_share():
    rows = [
        {"순위": 1, "국가": "Germany", "매출액": 38_300_000, "점유율": "18.1%"},
        {"순위": 2, "국가": "United Kingdom", "매출액": 37_100_000, "점유율": "17.5%"},
        {"순위": 3, "국가": "Spain", "매출액": 19_000_000, "점유율": "8.9%"},
    ]

    insights = export_service.build_region_insights(rows, eur_krw_rate=1717.34)

    assert insights[2].startswith("매출 상위 3개국 합계:")
    assert "분석 대상 전체 국가 매출의 44.5%를 차지합니다." in insights[2]
    assert "표시된 3개 시장" not in insights[2]


@pytest.mark.skip(reason="PPT export is card-only; paginated region tables are no longer generated.")
def test_analysis_period_stays_below_paginated_region_table(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    block = ReportBlockPayload(
        id="country:rank:ten",
        title="국가별 판매 순위",
        subtitle="10개국",
        meta="국가 순위",
        type="country_rank",
        kind="ranking",
        section="region",
        size="full",
        params={"analysis_period_label": "2025-07-13 ~ 2026-07-13", "eur_krw_rate": 1717.34},
        snapshot={
            "columns": ["순위", "국가", "권역", "매출액", "점유율", "YoY"],
            "rows": [
                {"순위": index, "국가": f"Country {index}", "권역": "유럽", "매출액": 20_000_000 - index * 500_000, "점유율": f"{20 - index}%", "YoY": f"{index}%"}
                for index in range(1, 11)
            ],
        },
    )
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[block], title="분석기간 배치 테스트")
    prs = Presentation(str(output_path))

    checked = 0
    for slide in prs.slides:
        period = next((shape for shape in slide.shapes if getattr(shape, "name", "") == "Silicon2 Analysis Period"), None)
        table = next((shape for shape in slide.shapes if getattr(shape, "name", "") == "Silicon2 Data Slide Table"), None)
        if period is None or table is None:
            continue
        checked += 1
        assert table.top + table.height < period.top
    assert checked == 2


def test_internal_report_omits_default_share_criteria_slide(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[brand_ranking_block()], title="내부용 테스트")
    text = pptx_text(output_path)

    assert "공유 외부 공유용 데이터 처리 기준" not in text
    assert "내부용 원본 데이터 취급 기준" not in text
    assert "공유 기준 — 마스킹 및 공유 범위" not in text
    assert "부록 데이터 기준 및 안전 원칙" not in text
    assert "부록 · 공유 및 마스킹 기준" not in text


def test_brand_slide_uses_selected_brand_blocks_not_preview_neighbors(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[
            brand_block_with_preview_neighbors("조선미녀", 24_400_000),
            brand_block_with_preview_neighbors("라운드랩", 2_660_000),
            brand_block_with_preview_neighbors("아누아", 2_250_000),
        ],
        title="내부용 테스트",
    )
    text = pptx_text(output_path)

    assert "조선미녀" in text
    assert "라운드랩" in text
    assert "아누아" in text
    assert "이즈앤트리" not in text
    assert "코스알엑스" not in text


@pytest.mark.skip(reason="PPT export is card-only; SKU detail tables are no longer generated.")
def test_sku_slide_uses_selected_sku_data_without_placeholder_cards(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[sku_detail_block()], title="내부용 테스트")
    text = pptx_text(output_path)
    sku_table_rows = first_pptx_table_rows_with_headers(output_path, {"항목", "원화 환산"})

    assert "[EU] 조선미녀크림 50ml" in text
    assert "Estonia" in text
    assert "44.3%" in text
    assert "원화 환산" in text
    assert "약 ₩21.1억" in text
    assert sku_table_rows
    krw_cell_text = sku_table_rows[1][sku_table_rows[0].index("원화 환산")]
    assert krw_cell_text.startswith("약 ₩")
    assert "조선미녀" not in krw_cell_text
    assert "선택 SKU 핵심 요약" not in text
    assert "선택 SKU 1개를 기준으로 정리했습니다." not in text
    assert "선택 SKU 없음" not in text
    assert "공급망 전략 권고" not in text
    assert "공급 안정성" not in text
    assert "최고 매출 SKU" not in text


@pytest.mark.skip(reason="PPT export is card-only; cross detail tables are no longer generated.")
def test_cross_slide_uses_axis_labels_and_removes_template_overlap_text(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[cross_matrix_block()], title="내부용 테스트")
    text = pptx_text(output_path)
    cross_table_rows = first_pptx_table_rows_with_headers(output_path, {"국가", "브랜드", "원화 환산", "비중"})

    assert "국가 × 브랜드 교차분석" in text
    assert "Estonia" in text
    assert "조선미녀" in text
    assert "€10,067,800.00" in text
    assert "원화 환산" in text
    assert "약 ₩177.2억" in text
    assert "브랜드 미상" not in text
    assert cross_table_rows
    headers = cross_table_rows[0]
    krw_cell_text = cross_table_rows[1][headers.index("원화 환산")]
    share_cell_text = cross_table_rows[1][headers.index("비중")]
    assert krw_cell_text.startswith("약 ₩")
    assert share_cell_text.endswith("%")
    assert share_cell_text != "-"
    assert "교차 분석" in text
    assert "누적 최대 기여 라인" not in text
    assert "어성초 라인" not in text


@pytest.mark.skip(reason="PPT export is card-only; cross detail table backfill no longer applies to PPT.")
def test_cross_slide_backfills_missing_krw_and_share_values(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[cross_matrix_block_without_calculated_columns()], title="교차분석 보강 테스트")
    text = pptx_text(output_path)
    cross_table_rows = first_pptx_table_rows_with_headers(output_path, {"국가", "브랜드", "원화 환산", "비중"})

    assert "국가 × 브랜드 교차분석" in text
    assert cross_table_rows
    headers = cross_table_rows[0]
    krw_cell_text = cross_table_rows[1][headers.index("원화 환산")]
    share_cell_text = cross_table_rows[1][headers.index("비중")]
    assert krw_cell_text.startswith("약 ₩")
    assert share_cell_text.endswith("%")
    assert share_cell_text != "-"


def test_single_cross_row_without_source_share_does_not_force_100_percent(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[
            ReportBlockPayload(
                id="cross:ingredient-country:single",
                title="성분 × 국가 교차분석",
                subtitle="스네일뮤신 × Estonia",
                meta="매출액 · 금액",
                type="cross_matrix",
                kind="matrix",
                section="cross",
                size="full",
                params={"axis": "ingredient-country", "metric": "sales", "scale": "amount", "rowLabel": "성분", "columnLabel": "국가", "eur_krw_rate": 1760.49},
                snapshot={
                    "columns": ["성분", "국가", "표시값", "원화표시", "점유율"],
                    "rows": [
                        {"성분": "스네일뮤신", "국가": "Estonia", "값": 676.1, "표시값": "€676.1K", "원화표시": "약 ₩11.9억", "점유율": "", "metric": "sales", "scale": "amount"},
                    ],
                },
            )
        ],
        title="교차분석 단일행 비중 테스트",
    )
    text = pptx_text(output_path)

    assert "€676.1K" in text
    assert "100%" not in text
    assert "성분 × 국가 교차분석" in text


def test_cross_matrix_card_uses_web_like_matrix_not_top_bars(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[cross_matrix_block()], title="cross matrix card test")
    prs = Presentation(str(output_path))
    card_slides = [slide for slide in prs.slides if "CARD 01" in slide_text(slide)]
    assert card_slides
    text = slide_text(card_slides[0])
    shape_names = [getattr(shape, "name", "") for shape in card_slides[0].shapes]
    assert shape_names.count("Silicon2 Cross Matrix Cell") >= 6
    assert "Silicon2 Report Card Value Bar" not in shape_names
    assert "약 ₩" in text


@pytest.mark.skip(reason="PPT export is card-only; cross blocks render as cards only.")
def test_cross_blocks_are_rendered_as_separate_template_slides(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[cross_matrix_block(), ingredient_country_cross_block()],
        title="교차분석 다중 테스트",
    )
    text = pptx_text(output_path)

    assert "국가 × 브랜드 교차분석" in text
    assert "성분 × 국가 교차분석" in text
    assert "스네일뮤신" in text
    assert "약 ₩11.9억" in text
    assert "브랜드 미상" not in text


def test_full_template_cross_expansion_does_not_duplicate_slide_parts(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    ingredient_block = ingredient_country_cross_block()
    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=[
            country_ranking_block(),
            brand_ranking_block(),
            sku_detail_block(),
            cross_matrix_block(),
            ingredient_block,
        ],
        title="full template cross test",
    )
    prs = Presentation(str(output_path))
    last_slide_text = slide_text(prs.slides[-1])

    assert duplicate_pptx_part_names(output_path) == set()
    assert len(prs.slides) < 16
    assert "01 국가별 판매 순위 및 성장" not in pptx_text(output_path)
    assert "02 브랜드별 판매 순위 및 원화 환산" not in pptx_text(output_path)
    assert last_slide_text
    assert ingredient_block.title not in last_slide_text


def test_soffice_path_can_be_configured_without_path_lookup(tmp_path, monkeypatch):
    fake_soffice = tmp_path / "soffice.exe"
    fake_soffice.write_text("", encoding="utf-8")

    monkeypatch.setenv("SOFFICE_PATH", str(fake_soffice))
    monkeypatch.setattr(export_service.shutil, "which", lambda _name: None)

    assert export_service.find_soffice_executable() == str(fake_soffice)


def test_html_export_is_generated_without_powerpoint_or_pdf_converter(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)

    def fail_office_converter(*_args, **_kwargs):
        raise AssertionError("HTML export should not call the PPTX-to-PDF converter")

    def fake_slide_renderer(_pptx_path, output_dir):
        slide_path = output_dir / "slide_001.png"
        slide_path.write_bytes(b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="))
        return [slide_path]

    monkeypatch.setattr(export_service, "convert_pptx_to_pdf", fail_office_converter)
    monkeypatch.setattr(export_service, "render_pptx_slides_to_png", fail_office_converter)
    output_path = build_report_export(
        audience="partner",
        export_format="html",
        blocks=[country_ranking_block()],
        title="HTML Export Test",
    )
    text = output_path.read_text(encoding="utf-8")

    assert output_path.suffix == ".html"
    assert "<!doctype html>" in text
    assert "HTML Export Test" in text
    assert "block-card" in text
    assert "국가별 판매 순위" in text
    assert "미국" in text
    assert "data:image/png;base64" not in text
    assert not list(output_path.parent.glob("*.pptx"))
    assert not list(output_path.parent.glob("*.pdf"))


def test_all_frontend_report_block_types_keep_titles_and_cart_order_in_pptx(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    blocks = list(reversed(all_report_block_types()))
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=blocks, title="전체 유형 순서 테스트")
    prs = Presentation(str(output_path))
    data_slide_titles = [
        slide_text(slide)
        for slide in prs.slides
        if "CARD " in slide_text(slide) and "전체 유형 회귀 테스트" in slide_text(slide)
    ]

    first_positions = []
    for block in blocks:
        matches = [index for index, text in enumerate(data_slide_titles) if block.title in text]
        assert matches, f"missing report block: {block.type}"
        first_positions.append(matches[0])
    assert first_positions == sorted(first_positions)
    assert "BLOCK " not in "\n".join(slide_text(slide) for slide in prs.slides)
    assert duplicate_pptx_part_names(output_path) == set()


def test_actual_template_keeps_every_frontend_block_as_ordered_lossless_detail(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    blocks = list(reversed(all_report_block_types()))
    output_path = build_report_export(
        audience="internal",
        export_format="ppt",
        blocks=blocks,
        title="실제 템플릿 전체 유형 순서 테스트",
    )
    prs = Presentation(str(output_path))
    detail_slide_texts = [slide_text(slide) for slide in prs.slides if "CARD " in slide_text(slide)]

    first_positions = []
    for expected_index, block in enumerate(blocks, start=1):
        matches = [index for index, text in enumerate(detail_slide_texts) if block.title in text]
        assert matches, f"missing template card block: {block.type}"
        assert any(f"CARD {expected_index:02d}" in detail_slide_texts[index] for index in matches)
        first_positions.append(matches[0])
    assert first_positions == sorted(first_positions)
    assert "BLOCK " not in "\n".join(slide_text(slide) for slide in prs.slides)
    assert duplicate_pptx_part_names(output_path) == set()


def test_html_keeps_all_120_country_card_orders(tmp_path):
    country_blocks = [block for block in all_report_block_types() if block.section == "region"]
    assert len(country_blocks) == 5

    output_path = tmp_path / "country-order.html"
    for ordered_blocks in permutations(country_blocks):
        export_service.write_report_html(
            output_path,
            audience="internal",
            blocks=list(ordered_blocks),
            title="국가 카드 순열 테스트",
        )
        text = output_path.read_text(encoding="utf-8")
        positions = [text.index(block.title) for block in ordered_blocks]
        assert positions == sorted(positions)
        for block in ordered_blocks:
            assert f'id="block-{ordered_blocks.index(block) + 1}"' in text


def test_season_calendar_pptx_outputs_all_twelve_months_across_pages(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    season_block = all_report_block_types()[0]
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[season_block], title="시즌 전체 열 테스트")
    text = pptx_text(output_path)
    prs = Presentation(str(output_path))
    card_slides = [slide for slide in prs.slides if "CARD 01 · 시즌 캘린더" in slide_text(slide)]

    for month in range(1, 13):
        assert f"{month}월" in text
    assert len(card_slides) == 1
    assert not any(getattr(shape, "has_table", False) for shape in card_slides[0].shapes)
    assert len([shape for shape in card_slides[0].shapes if getattr(shape, "name", "") == "Silicon2 Season Calendar Cell"]) == 12
    assert "BLOCK " not in text
    assert "TYPE 01 season_calendar (1/3)" not in text
    assert "TYPE 01 season_calendar (3/3)" not in text


def test_season_calendar_card_uses_calendar_heatmap_not_bar_list(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    season_block = ReportBlockPayload(
        id="season:calendar:function-groups",
        title="시즌 캘린더 — 기능군별 수요 지수",
        subtitle="3개 기능군 · 완료월 12개월 · 판매수량 수요 지수(2~11)",
        meta="시즌 캘린더",
        type="season_calendar",
        kind="matrix",
        section="season",
        size="full",
        params={"rowLabel": "기능군", "columnLabel": "월"},
        snapshot={
            "columns": ["기능군", "피크", "피크강도", "SKU수", *[f"{month}월" for month in range(1, 13)]],
            "rows": [
                {"기능군": "향수/아로마", "피크": "5월", "피크강도": "높음", "SKU수": 12, **{f"{month}월": month for month in range(1, 13)}},
                {"기능군": "네일", "피크": "7월", "피크강도": "중간", "SKU수": 8, **{f"{month}월": 13 - month for month in range(1, 13)}},
                {"기능군": "화장소품", "피크": "3월", "피크강도": "중간", "SKU수": 5, **{f"{month}월": 6 for month in range(1, 13)}},
                *[
                    {"기능군": f"기능군 {index}", "피크": "6월", "피크강도": "중간", "SKU수": index, **{f"{month}월": 2 if month % 2 else 8 for month in range(1, 13)}}
                    for index in range(4, 14)
                ],
            ],
        },
    )
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=[season_block], title="시즌 캘린더 카드 테스트")
    prs = Presentation(str(output_path))
    card_slides = [slide for slide in prs.slides if "CARD 01 · 시즌 캘린더" in slide_text(slide)]
    text = "\n".join(slide_text(slide) for slide in card_slides)
    calendar_cells = [
        shape
        for slide in card_slides
        for shape in slide.shapes
        if getattr(shape, "name", "") == "Silicon2 Season Calendar Cell"
    ]

    assert len(card_slides) == 3
    assert "시즌 캘린더 — 기능군별 수요 지수 (1/3)" in text
    assert "시즌 캘린더 — 기능군별 수요 지수 (3/3)" in text
    assert "향수/아로마" in text
    assert "기능군 13" in text
    assert "1월" in text
    assert "12월" in text
    assert "피크월" in text
    assert "100%" in text
    assert "25%" in text
    assert "0%" not in {line.strip() for line in text.splitlines()}
    assert len(calendar_cells) == 156


def test_season_calendar_values_are_relative_to_each_row_peak():
    assert export_service.format_season_calendar_value(10, 10) == "100%"
    assert export_service.format_season_calendar_value(7, 10) == "70%"
    assert export_service.format_season_calendar_value(5, 10) == "50%"
    assert export_service.format_season_calendar_value(10, 10, is_peak=False) == "99%"
    assert export_service.format_season_calendar_value(0, 10) == "-"


@pytest.mark.skip(reason="PPT export is card-only; template brand detail tables are no longer generated.")
def test_original_brand_template_uses_brand_rows_not_card_titles(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    blocks = [
        ReportBlockPayload(
            id="brand:rank:all",
            title="전체 브랜드 판매 순위",
            subtitle="매출 기준",
            meta="브랜드 기준 분석",
            type="brand_rank",
            kind="ranking",
            section="brand",
            size="full",
            params={"scope": "all_brands", "eur_krw_rate": 1717.34},
            snapshot={
                "columns": ["순위", "브랜드명", "매출액", "매출표시", "원화표시"],
                "rows": [
                    {"순위": 1, "브랜드명": "조선미녀", "매출액": 68_900_000, "매출표시": "€68.9M", "원화표시": "약 ₩1,182.5억"},
                    {"순위": 2, "브랜드명": "메디큐브", "매출액": 32_200_000, "매출표시": "€32.2M", "원화표시": "약 ₩552.3억"},
                ],
            },
        ),
        ReportBlockPayload(
            id="brand:growth:mom",
            title="브랜드별 성장 변화",
            subtitle="2024.11 → 2024.12",
            meta="브랜드별 매출 성장률 비교",
            type="brand_growth",
            kind="ranking",
            section="brand",
            size="full",
            params={"comparison_basis": "mom"},
            snapshot={
                "columns": ["구분", "순위", "브랜드명", "성장률"],
                "rows": [
                    {"구분": "성장률 상위", "순위": 1, "브랜드명": "조선미녀", "성장률": "+12.4%"},
                    {"구분": "감소율 상위", "순위": 1, "브랜드명": "메디큐브", "성장률": "-3.2%"},
                ],
            },
        ),
        ReportBlockPayload(
            id="brand:comparison",
            title="브랜드 월별 매출 비중 비교",
            subtitle="2개 브랜드 월별 추이",
            meta="월별 비교",
            type="brand_comparison",
            kind="trend",
            section="brand",
            size="half",
            params={"brands": ["조선미녀", "메디큐브"]},
            snapshot={
                "columns": ["브랜드명", "월", "매출액", "브랜드 내 월 매출 비중"],
                "rows": [
                    {"브랜드명": "조선미녀", "월": "11월", "매출액": 5_000_000, "브랜드 내 월 매출 비중": "45%"},
                    {"브랜드명": "조선미녀", "월": "12월", "매출액": 6_000_000, "브랜드 내 월 매출 비중": "55%"},
                    {"브랜드명": "메디큐브", "월": "11월", "매출액": 3_000_000, "브랜드 내 월 매출 비중": "50%"},
                    {"브랜드명": "메디큐브", "월": "12월", "매출액": 3_000_000, "브랜드 내 월 매출 비중": "50%"},
                ],
            },
        ),
    ]

    output_path = build_report_export(audience="internal", export_format="ppt", blocks=blocks, title="보고서 장바구니")
    table_rows = first_pptx_table_rows_with_headers(output_path, {"브랜드명", "원화 환산"})
    table_text = "\n".join(" | ".join(row) for row in table_rows)
    deck_text = pptx_text(output_path)

    assert "02 브랜드별 판매 순위 및 원화 환산" not in deck_text
    assert "조선미녀" in table_text and "€68.9M" in table_text and "약 ₩1,182.5억" in table_text
    assert "메디큐브" in table_text and "€32.2M" in table_text and "약 ₩552.3억" in table_text
    assert "브랜드별 성장 변화" not in table_text
    assert "브랜드 월별 매출 비중 비교" not in table_text


def test_html_keeps_all_block_types_and_cart_order(tmp_path, monkeypatch):
    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    blocks = list(reversed(all_report_block_types()))
    output_path = build_report_export(audience="internal", export_format="html", blocks=blocks, title="HTML 전체 유형 테스트")
    text = output_path.read_text(encoding="utf-8")

    positions = [text.find(f'<h3>{block.title}</h3>') for block in blocks]
    assert all(position >= 0 for position in positions)
    assert positions == sorted(positions)
    assert "APPENDIX" not in text
    assert "공유 및 마스킹 기준" not in text
    for month in range(1, 13):
        assert f">{month}월<" in text


def test_thirty_blocks_create_paginated_agenda_without_dropping_blocks(tmp_path, monkeypatch):
    from pptx import Presentation

    monkeypatch.setattr(export_service, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(export_service, "REPORT_CART_TEMPLATE_PPTX", tmp_path / "missing-template.pptx")
    source = all_report_block_types()
    blocks = [
        export_service.replace_report_block(source[index % len(source)], id=f"max:{index}", title=f"MAX BLOCK {index + 1:02d}")
        for index in range(30)
    ]
    output_path = build_report_export(audience="internal", export_format="ppt", blocks=blocks, title="최대 블록 테스트")
    prs = Presentation(str(output_path))
    text = pptx_text(output_path)

    assert sum("CONTENTS" in slide_text(slide) for slide in prs.slides) == 3
    for index in range(1, 31):
        assert f"MAX BLOCK {index:02d}" in text


def test_audience_badges_match_policy():
    assert audience_badge("internal") == "INTERNAL ONLY · 대외비"
    assert audience_badge("partner") == "PARTNER VIEW · 민감정보 제외"
    assert audience_badge("sales") == "공개용"
