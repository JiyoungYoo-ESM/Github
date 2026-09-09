import pandas as pd
from openpyxl import Workbook

from core import common, export_excel_util, inbound, order_review_report, session, transport


def _check_required_df():
    return pd.DataFrame(
        [
            {
                "\ud655\uc778 \uad6c\ubd84": "SKU \ub9e4\uce6d \ud655\uc778",
                "\uc0c1\ud488\ucf54\ub4dc": "SKU-1",
                "\uc0c1\ud488\uba85": "Sample 1",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\ubc1c\uacac \uc6d0\ubcf8": "\ud310\ub9e4\ub0b4\uc5ed/\uc6b4\uc1a1\uc911",
                "\ud310\ub9e4\uc218\ub7c9": 120,
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 800,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": 0,
                "\uc218\ub7c9": 800,
                "\ucd9c\uace0\uc77c": "2026-05-10",
                "\uc608\uc0c1 \uc785\uace0\uc77c": "2026-05-25",
                "\uc6b4\uc1a1\uc218\ub2e8": "\ud574\uc6b4",
                "\ud655\uc778 \ub0b4\uc6a9": "SKU \ub9e4\uce6d \ud544\uc694",
                "\uad8c\uc7a5 \ud655\uc778 \uc561\uc158": "\ub2e8\uc885/\ucf54\ub4dc \ubcc0\uacbd \ud655\uc778",
                "\ub2f4\ub2f9\uc790\ud310\ub2e8": "",
                "\ud310\ub2e8\uc77c\uc2dc": "",
                "\ub2f4\ub2f9\uc790": "",
            },
            {
                "\ud655\uc778 \uad6c\ubd84": "\ub2e8\uac00 \ubbf8\ub4f1\ub85d \ud655\uc778",
                "\uc0c1\ud488\ucf54\ub4dc": "SKU-2",
                "\uc0c1\ud488\uba85": "Sample 2",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\ubc1c\uacac \uc6d0\ubcf8": "EU \ud604\uc9c0 \uc7ac\uace0",
                "\ud310\ub9e4\uc218\ub7c9": 10,
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 0,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": 0,
                "\uc218\ub7c9": 10,
                "\ucd9c\uace0\uc77c": "",
                "\uc608\uc0c1 \uc785\uace0\uc77c": "",
                "\uc6b4\uc1a1\uc218\ub2e8": "",
                "\ud655\uc778 \ub0b4\uc6a9": "\ub2e8\uac00 \ud655\uc778 \ud544\uc694",
                "\uad8c\uc7a5 \ud655\uc778 \uc561\uc158": "FOC/\ub2e8\uac00 \ub204\ub77d \ud655\uc778",
                "\ub2f4\ub2f9\uc790\ud310\ub2e8": "\ud655\uc778 \uc644\ub8cc",
                "\ud310\ub2e8\uc77c\uc2dc": "2026-05-20",
                "\ub2f4\ub2f9\uc790": "Owner",
            },
        ],
        columns=common.CHECK_REQUIRED_DISPLAY_COLUMNS,
    )


def test_check_required_sheet_uses_shared_reference_layout():
    wb = Workbook()
    ws = wb.active
    export_excel_util.append_df(ws, _check_required_df(), start_row=8)

    order_review_report.format_check_required_sheet(ws, {"base_date": "2026-05-20"}, header_row=8)
    export_excel_util.apply_date_picker_validation(
        ws,
        "\ud310\ub2e8\uc77c\uc2dc",
        "_\ud310\ub2e8\uc77c\uc2dc_\uc120\ud0dd\uac12",
        "\ud310\ub2e8\uc77c\uc2dc",
        "\ub4dc\ub86d\ub2e4\uc6b4\uc5d0\uc11c \ud310\ub2e8\uc77c\uc2dc\ub97c \uc120\ud0dd\ud558\uac70\ub098 yyyy-mm-dd \ud615\uc2dd\uc73c\ub85c \uc785\ub825\ud558\uc138\uc694.",
        header_row=8,
    )

    assert ws["A3"].value == "\u25bc \uc6d0\ubcf8 \ub370\uc774\ud130 \ud655\uc778\ud544\uc694 | \uc785\ub825 \uc624\ub958 \uc810\uac80"
    assert ws["I1"].value is None
    assert ws["J1"].value is None
    assert ws["A4"].value == "★ 법인별 운송수단·ETA와 입력 이상값을 표시합니다. 담당자판단·판단일시·담당자는 입력 영역입니다."
    assert ws["C4"].value is None
    assert ws["E4"].value is None
    merged_ranges = {str(merged_range) for merged_range in ws.merged_cells.ranges}
    assert "A3:E3" in merged_ranges
    assert "A4:E4" in merged_ranges
    assert "A3:F3" not in merged_ranges
    assert "A4:H4" not in merged_ranges
    for coord in ("F3", "G3", "F4", "G4"):
        assert ws[coord].fill.fill_type is None
    assert ws.cell(8, 1).value == "\uc0c1\ud488\ucf54\ub4dc"
    assert ws.cell(8, 14).value == "\ub2f4\ub2f9\uc790\ud310\ub2e8"
    assert ws.cell(8, 11).value == "원본 비고"
    assert ws.cell(9, 12).alignment.wrap_text is True
    assert ws.cell(9, 12).alignment.horizontal == "left"
    assert ws.cell(9, 14).alignment.wrap_text is False
    assert ws.cell(9, 14).alignment.horizontal == "center"
    assert ws.row_dimensions[9].height >= 48
    assert ws.freeze_panes == "D9"
    assert ws.auto_filter.ref == "A8:P10"
    assert "_\ud310\ub2e8\uc77c\uc2dc_\uc120\ud0dd\uac12" in wb.sheetnames


def test_empty_check_required_sheet_shows_no_data_message():
    ws = Workbook().active
    ws.title = "\ud655\uc778\ud544\uc694"
    export_excel_util.append_df(ws, pd.DataFrame(columns=common.CHECK_REQUIRED_DISPLAY_COLUMNS), start_row=8)

    order_review_report.format_check_required_sheet(ws, {"base_date": "2026-05-20"}, header_row=8)

    assert ws["A9"].value == "\ud655\uc778\ud544\uc694 \ub370\uc774\ud130 \uc5c6\uc74c"


def test_unrecognized_transport_keeps_eta_and_separates_original_remark():
    unrecognized = pd.DataFrame(
        {
            "상품코드": ["SKU-USA"],
            "상품명": ["Sample"],
            "수량": [100],
            "출고일": ["2026-07-31"],
            "예상 입고일": ["2026-08-15"],
            "표준 운송수단": ["확인필요"],
            "원본값": ["HJ 73rd (ETA 8/15)"],
            "확인필요 사유": ["입고 예정일은 확인되었으나 운송수단을 분류할 수 없습니다."],
            "권장 확인 액션": ["법인 표준 운송수단(해운/항공) 확인"],
        }
    )

    result = order_review_report.build_check_required_sheet_df(
        pd.DataFrame(),
        unrecognized,
        pd.DataFrame(),
        pd.DataFrame(),
    )

    row = result.iloc[0]
    assert row["예상 입고일"] == "2026-08-15"
    assert row["운송수단"] == "확인필요"
    assert row["원본 비고"] == "HJ 73rd (ETA 8/15)"
    assert row["확인 내용"] == "입고 예정일은 확인되었으나 운송수단을 분류할 수 없습니다."


def test_unrecognized_transport_keeps_the_source_brand():
    unrecognized = pd.DataFrame(
        {
            "상품코드": ["AbibS11-C", "ANS11-E"],
            "상품명": ["Sample A", "Sample B"],
            "브랜드": ["아비브", "아누아"],
            "수량": [480, 2000],
            "출고일": ["2026-08-05", "2026-08-05"],
            "예상 입고일": ["2026-08-24", "2026-08-24"],
            "표준 운송수단": ["확인필요", "확인필요"],
            "원본값": ["ER 34th (ETA 8/24)", "ER 34th (ETA 8/24)"],
        }
    )

    result = order_review_report.build_check_required_sheet_df(
        pd.DataFrame(),
        unrecognized,
        pd.DataFrame(),
        pd.DataFrame(),
    )

    assert result["브랜드"].tolist() == ["아비브", "아누아"]


def test_human_data_issue_allows_comma_formatted_shipping_qty():
    context = session.SessionContext(
        uploaded_data={
            "shipping": pd.DataFrame(
                {
                    "SKU": ["BODP04-MEU"],
                    "\uc0c1\ud488\uba85": ["[EU][4\ub9e4\uc785]\ubc14\uc774\uc624 \ucf5c\ub77c\uac90 \ub9ac\uc5bc \ub525 \ub9c8\uc2a4\ud06c"],
                    "\ube0c\ub79c\ub4dc": ["\ubc14\uc774\uc624\ub358\uc2a4"],
                    "\uc218\ub7c9": ["33,600"],
                    "\ucd9c\uace0\uc77c": ["2026-06-01"],
                    "\uc6b4\uc1a1\uc218\ub2e8": ["\ud574\uc6b4"],
                }
            )
        }
    )

    issues = order_review_report.build_human_data_issue_df({}, context)

    assert issues.empty


def test_human_data_issue_does_not_flag_sku_format_variants():
    context = session.SessionContext(
        uploaded_data={
            "eu_stock": pd.DataFrame(
                {
                    "SKU": ["MBM01-EU17(20ml)R", " CODE WITH SPACE "],
                    "\uc0c1\ud488\uba85": ["M\ud37c\ud399\ud2b8 \ucee4\ubc84 \ube44\ube44 17\ud638 20ml", "Custom SKU"],
                    "\uc7ac\uace0\uc218\ub7c9": [10, 5],
                }
            )
        }
    )

    issues = order_review_report.build_human_data_issue_df({}, context)

    assert issues.empty


def test_human_data_issue_does_not_flag_case_variant_skus():
    context = session.SessionContext(
        uploaded_data={
            "eu_stock": pd.DataFrame(
                {
                    "SKU": ["A-001", "a-001"],
                    "\uc0c1\ud488\uba85": ["Upper Item", "Lower Item"],
                    "\uc7ac\uace0\uc218\ub7c9": [10, 5],
                }
            )
        }
    )

    issues = order_review_report.build_human_data_issue_df({}, context)

    assert issues.empty


def test_human_data_issue_ignores_delivery_charge_sales_rows():
    context = session.SessionContext(
        uploaded_data={
            "sales_detail": pd.DataFrame(
                {
                    "SKU": ["Delivery Charge"],
                    "\uc0c1\ud488\uba85": ["Delivery Charge"],
                    "\uc218\ub7c9": ["not-a-product-qty"],
                }
            )
        }
    )

    issues = order_review_report.build_human_data_issue_df({}, context)

    assert issues.empty


def test_human_data_issue_treats_negative_open_po_as_zero():
    context = session.SessionContext(
        uploaded_data={
            "open_po": pd.DataFrame(
                {
                    "SKU": ["TIRM30-CR3W"],
                    "\uc0c1\ud488\uba85": ["Sample SKU"],
                    "\ubbf8\uc785\uace0\uc218\ub7c9": [-80],
                }
            )
        }
    )

    issues = order_review_report.build_human_data_issue_df({}, context)

    assert issues.empty


def test_unrecognized_transport_report_carries_the_shipping_brand():
    context = session.SessionContext(
        uploaded_data={
            "shipping": pd.DataFrame(
                {
                    "SKU": ["AbibS11-C"],
                    "\uc0c1\ud488\uba85": ["\uc5b4\uc131\ucd08 \ud06c\ub9bc \uce74\ubc0d \ud29c\ubcf8 75ml"],
                    "\ube0c\ub79c\ub4dc": ["\uc544\ubc14\ube0c"],
                    "\uc218\ub7c9": [480],
                    "\ucd9c\uace0\uc77c": ["2026-08-05"],
                    "\uc6b4\uc1a1\uc218\ub2e8": ["ER 34th (ETA 8/24)"],
                }
            )
        }
    )

    issues = transport.unrecognized_transport_report_df({}, context)

    assert not issues.empty
    assert issues.loc[0, "\ube0c\ub79c\ub4dc"] == "\uc544\ubc14\ube0c"


def test_shipping_transport_mode_can_be_embedded_in_raw_text():
    context = session.SessionContext(
        uploaded_data={
            "shipping": pd.DataFrame(
                {
                    "SKU": ["NLS05-A"],
                    "\uc0c1\ud488\uba85": ["\uc5d0\uc774\ucee8\ud2b8\ub864 10% \uc138\ub7fc"],
                    "\uc218\ub7c9": [12],
                    "\ucd9c\uace0\uc77c": ["2026-06-01"],
                    "\uc6b4\uc1a1\uc218\ub2e8": ["FR \ud56d\uacf5 4\ucc28(\uc5d0\uc774\ud53c\uc5d8\ube44)"],
                }
            )
        }
    )

    shipping = transport.get_shipping({}, context)
    issues = transport.unrecognized_transport_report_df({}, context)

    assert shipping.loc[0, "\uc6b4\uc1a1\uc218\ub2e8"] == "\ud56d\uacf5"
    assert issues.empty


def test_check_required_includes_open_po_and_shipping_eta_reference_rows():
    master_unregistered = pd.DataFrame(
        [
            {
                "\uc0c1\ud488\ucf54\ub4dc": "NEW-PO",
                "\uc0c1\ud488\uba85": "New PO SKU",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 0,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": 5,
            },
            {
                "\uc0c1\ud488\ucf54\ub4dc": "NEW-SHIP",
                "\uc0c1\ud488\uba85": "New Shipping SKU",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 7,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": 0,
                "\ucd9c\uace0\uc77c": "2026-06-01",
                "\uc608\uc0c1 \uc785\uace0\uc77c": "2026-08-10",
                "\uc6b4\uc1a1\uc218\ub2e8": "\ud574\uc6b4",
            },
        ]
    )

    result = order_review_report.build_check_required_sheet_df(
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        master_unregistered,
    )

    assert "\ud655\uc778 \uad6c\ubd84" not in result.columns
    assert set(result["\ubc1c\uacac \uc6d0\ubcf8"]) == {"\ubbf8\uc785\uace0\ud604\ud669", "\uc6b4\uc1a1\uc911"}
    po_row = result.loc[result["\uc0c1\ud488\ucf54\ub4dc"].eq("NEW-PO")].iloc[0]
    shipping_row = result.loc[result["\uc0c1\ud488\ucf54\ub4dc"].eq("NEW-SHIP")].iloc[0]
    assert po_row["\ubc1c\uacac \uc6d0\ubcf8"] == "\ubbf8\uc785\uace0\ud604\ud669"
    assert shipping_row["\ubc1c\uacac \uc6d0\ubcf8"] == "\uc6b4\uc1a1\uc911"
    assert po_row["\ucd9c\uace0\uc77c"] == ""
    assert po_row["\uc608\uc0c1 \uc785\uace0\uc77c"] == ""
    assert po_row["\uc6b4\uc1a1\uc218\ub2e8"] == ""
    assert shipping_row["\ucd9c\uace0\uc77c"] == "2026-06-01"
    assert shipping_row["\uc608\uc0c1 \uc785\uace0\uc77c"] == "2026-08-10"
    assert shipping_row["\uc6b4\uc1a1\uc218\ub2e8"] == "\ud574\uc6b4"
    assert "\ub300\uc0c1 SKU: NEW-PO" in po_row["\ud655\uc778 \ub0b4\uc6a9"]
    assert "\ub300\uc0c1 SKU: NEW-SHIP" in shipping_row["\ud655\uc778 \ub0b4\uc6a9"]
    assert "\ud604\uc9c0 \uc7ac\uace0" in po_row["\ud655\uc778 \ub0b4\uc6a9"]
    assert "\ud604\uc9c0 \uc7ac\uace0" in shipping_row["\ud655\uc778 \ub0b4\uc6a9"]


def test_backend_fast_check_required_includes_eta_reference_rows(monkeypatch):
    import backend.analysis as analysis

    master_unregistered = pd.DataFrame(
        [
            {
                "\uc0c1\ud488\ucf54\ub4dc": "FAST-PO",
                "\uc0c1\ud488\uba85": "Fast PO SKU",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 0,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": 3,
            },
            {
                "\uc0c1\ud488\ucf54\ub4dc": "FAST-SHIP",
                "\uc0c1\ud488\uba85": "Fast Shipping SKU",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 4,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": 0,
            },
        ]
    )
    monkeypatch.setattr(analysis, "build_master_unregistered_sku_df", lambda settings, context: master_unregistered)
    monkeypatch.setattr(analysis, "unrecognized_transport_report_df", lambda settings, context: pd.DataFrame())
    monkeypatch.setattr(analysis, "build_human_data_issue_df", lambda settings, context: pd.DataFrame())
    monkeypatch.setattr(analysis, "build_excluded_order_review_df", lambda excluded_review, settings: pd.DataFrame())

    result = analysis.build_check_required_df({}, pd.DataFrame(), pd.DataFrame(), session.SessionContext())

    assert "\ud655\uc778 \uad6c\ubd84" not in result.columns
    assert set(result["\ubc1c\uacac \uc6d0\ubcf8"]) == {"\ubbf8\uc785\uace0\ud604\ud669", "\uc6b4\uc1a1\uc911"}


def test_master_unregistered_shipping_dates_show_first_eta_only(monkeypatch):
    context = session.SessionContext(
        uploaded_data={
            "eu_stock": pd.DataFrame({"SKU": ["KNOWN"]}),
            "shipping": pd.DataFrame(
                {
                    "SKU": ["NEW-SHIP", "NEW-SHIP"],
                    "\uc0c1\ud488\uba85": ["New Shipping SKU", "New Shipping SKU"],
                    "\ube0c\ub79c\ub4dc": ["Brand", "Brand"],
                    "\uc218\ub7c9": [3, 4],
                    "\ucd9c\uace0\uc77c": ["2026-06-05", "2026-06-01"],
                    "\uc6b4\uc1a1\uc218\ub2e8": ["\ud574\uc6b4", "\ud574\uc6b4"],
                }
            ),
        }
    )
    settings = {"lead_times": {"\ud574\uc6b4": 70, "\ud56d\uacf5": 15, "\ucca0\uc1a1": 30, "\ud2b8\ub7ed": 30}}

    result = inbound.build_master_unregistered_sku_df(settings, context)

    row = result.loc[result["\uc0c1\ud488\ucf54\ub4dc"].eq("NEW-SHIP")].iloc[0]
    assert row["\ucd9c\uace0\uc77c"] == "2026-06-01"
    assert row["\uc608\uc0c1 \uc785\uace0\uc77c"] == "2026-08-10"


def test_check_required_treats_negative_open_po_reference_qty_as_zero():
    master_unregistered = pd.DataFrame(
        [
            {
                "\uc0c1\ud488\ucf54\ub4dc": "NEG-PO",
                "\uc0c1\ud488\uba85": "Negative PO SKU",
                "\ube0c\ub79c\ub4dc": "Brand",
                "\uc6b4\uc1a1\uc911 \uc218\ub7c9": 0,
                "\ubbf8\uc785\uace0 \uc218\ub7c9": -80,
            }
        ]
    )

    result = order_review_report.build_check_required_sheet_df(
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        master_unregistered,
    )

    assert result.empty


def test_check_required_excludes_blank_sku_transport_noise():
    unrecognized = pd.DataFrame(
        [
            {
                "\uc6d0\ubcf8\uac12": "",
                "\uc0c1\ud488\ucf54\ub4dc": "",
                "\uc0c1\ud488\uba85": "",
                "\uc218\ub7c9": 0,
                "\ucd9c\uace0\uc77c": "13624399",
            }
        ]
    )

    result = order_review_report.build_check_required_sheet_df(
        pd.DataFrame(),
        unrecognized,
        pd.DataFrame(),
        pd.DataFrame(),
    )

    assert result.empty


def test_unrecognized_transport_report_ignores_blank_sku_rows():
    context = session.SessionContext(
        uploaded_data={
            "shipping": pd.DataFrame(
                {
                    "SKU": [""],
                    "\uc0c1\ud488\uba85": [""],
                    "\uc218\ub7c9": [0],
                    "\ucd9c\uace0\uc77c": ["13624399"],
                    "\uc6b4\uc1a1\uc218\ub2e8": [""],
                }
            )
        }
    )

    issues = transport.unrecognized_transport_report_df({}, context)

    assert issues.empty
