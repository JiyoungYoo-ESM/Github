from __future__ import annotations

from datetime import date

import numpy as np
import pandas as pd

from core.common import (
    ALERT_LOCAL_SHORT_LABEL,
    ALERT_SHIPPING_SHORT_LABEL,
    ALERT_URGENT_LABEL,
    DEFAULT_LEAD_TIME_DAYS,
    korea_today,
    ORDER_REVIEW_ELIGIBILITY_HELPER,
    ORDER_SHEET_COLUMNS,
    UNIFIED_ESM_ORDER_REVIEW_COLUMNS,
)
from core import kpi as kpi_mod, order_review as order_review_mod, transport as transport_mod
from core.export_excel_calculations import order_review_excluded_mask, template_order_metrics
from core.export_excel_sheets import _clean_merge_key, _first_non_blank



def build_stock_eta_sheet_df(
    review_df: pd.DataFrame,
    arrival_calendar_df: pd.DataFrame | None = None,
    settings: dict | None = None,
) -> pd.DataFrame:
    settings = settings or {}
    base = pd.DataFrame(review_df).copy()
    sku_code = order_review_mod.report_col(base, ["상품코드"]).map(_clean_merge_key)
    product_name = order_review_mod.report_col(base, ["상품명"])
    brand = order_review_mod.report_col(base, ["브랜드"])
    sales_3m = pd.to_numeric(order_review_mod.report_col(base, ["최근 3개월 판매수량", "기준_3M_판매수량", "PA_CA_3M_판매수량"], 0), errors="coerce").fillna(0)
    monthly_sales = pd.to_numeric(order_review_mod.report_col(base, ["월평균 판매수량"], 0), errors="coerce").fillna(0)
    safety_stock = pd.to_numeric(order_review_mod.report_col(base, ["안전재고 목표수량", "안전재고수량"], 0), errors="coerce").fillna(0)
    eu_stock = pd.to_numeric(order_review_mod.report_col(base, ["EU 현지 가용수량", "유럽 가용재고"], 0), errors="coerce").fillna(0)
    shipping_stock = pd.to_numeric(order_review_mod.report_col(base, ["운송중 수량"], 0), errors="coerce").fillna(0)
    combined_stock = eu_stock + shipping_stock
    open_po_qty = pd.to_numeric(order_review_mod.report_col(base, ["미입고수량", "미입고 수량"], 0), errors="coerce").fillna(0)
    hq_eu_stock = pd.to_numeric(order_review_mod.report_col(base, ["본사 EU창고 가용수량", "본사 EU창고"], 0), errors="coerce").fillna(0)
    metrics = template_order_metrics(base, settings)
    monthly_sales = metrics["월평균_계산"]
    safety_stock = metrics["안전재고_계산"]
    eu_stock = metrics["유럽재고_계산"]
    shipping_stock = metrics["운송재고_계산"]
    combined_stock = metrics["유럽운송합산_계산"]
    order_required_qty = metrics["발주필요수량_계산"]

    def source_or_fallback(candidates: list[str], fallback):
        for col in candidates:
            if col in base.columns:
                return base[col]
        return fallback

    def cover_months(qty: pd.Series) -> pd.Series:
        display_monthly = monthly_sales.round(0)
        return np.where(display_monthly.gt(0), qty / display_monthly.replace(0, np.nan), 0)

    order_required_numeric = pd.to_numeric(
        source_or_fallback(["최종 발주 필요수량", "발주필요수량", "추가 발주 필요 수량"], order_required_qty),
        errors="coerce",
    ).fillna(0)
    is_oos = eu_stock.le(0)
    is_order_needed = order_required_numeric.gt(0)
    status = pd.Series(
        np.select(
            [is_oos, is_order_needed],
            ["OOS", "발주필요"],
            default="-",
        ),
        index=base.index,
    )
    order_required_display = order_required_qty
    open_po_display = source_or_fallback(
        ["미입고 현황", "미입고수량", "미입고 수량"],
        open_po_qty.round(0).astype(int),
    )
    hq_eu_stock_display = source_or_fallback(
        ["현지 창고 재고", "EU 창고 재고", "본사 EU창고 가용수량", "본사 EU창고"],
        hq_eu_stock.round(0).astype(int),
    )
    out = pd.DataFrame(
        {
            "코드": sku_code,
            "No": range(1, len(base) + 1),
            "SKU": _first_non_blank(product_name, sku_code),
            "브랜드": brand,
            "3개월 판매량": sales_3m.round(0).astype(int),
            "월평균 판매량": monthly_sales.round(0).astype(int),
            "안전재고": safety_stock.round(0).astype(int),
            "1. 현지 재고": eu_stock.round(0).astype(int),
            "2. 운송 재고": shipping_stock.round(0).astype(int),
            "3. 현지+운송": combined_stock.round(0).astype(int),
            "1. 현지 재고(M)": pd.Series(cover_months(eu_stock), index=base.index).round(1),
            "2. 운송중(M)": pd.Series(cover_months(shipping_stock), index=base.index).round(1),
            "3. 합산(M)": pd.Series(cover_months(combined_stock), index=base.index).round(1),
            "알림": status,
            "미입고 현황": open_po_display,
            "현지 창고 재고": hq_eu_stock_display,
            "발주필요수량": order_required_display,
        }
    )

    if arrival_calendar_df is not None and not arrival_calendar_df.empty and {"SKU", "도착일", "수량"}.issubset(arrival_calendar_df.columns):
        arrival_work = arrival_calendar_df[["SKU", "도착일", "수량"]].copy()
        arrival_work["SKU"] = arrival_work["SKU"].map(_clean_merge_key)
        arrival_work["도착일"] = kpi_mod.parse_date_series(arrival_work["도착일"]).dt.date
        arrival_work["수량"] = pd.to_numeric(arrival_work["수량"], errors="coerce").fillna(0)
        arrival_work = arrival_work[arrival_work["SKU"].ne("") & arrival_work["도착일"].notna() & arrival_work["수량"].gt(0)]
        if not arrival_work.empty:
            pivot = arrival_work.pivot_table(index="SKU", columns="도착일", values="수량", aggfunc="sum", fill_value=0)
            pivot = pivot.reindex(columns=sorted(pivot.columns))
            for arrival_date in pivot.columns:
                label = f"{arrival_date.year}-{arrival_date.month}-{arrival_date.day}"
                mapped_qty = sku_code.map(pivot[arrival_date]).fillna(0).round(0).astype(int)
                if mapped_qty.sum() > 0:
                    out[label] = mapped_qty.astype(object).where(mapped_qty.ne(0), "-").to_numpy()

    # 엑셀에서 안전재고 개월 수를 바꿔도 발주제외/확인필요 행이 다시 발주 대상으로
    # 살아나지 않도록 수식이 참조하는 판정값을 마지막 숨김 열로 전달한다.
    out[ORDER_REVIEW_ELIGIBILITY_HELPER] = (~order_review_excluded_mask(base)).astype(int)
    fixed_columns = [col for col in UNIFIED_ESM_ORDER_REVIEW_COLUMNS if col in out.columns]
    dynamic_columns = [
        col
        for col in out.columns
        if col not in fixed_columns and col != ORDER_REVIEW_ELIGIBILITY_HELPER
    ]
    return out[fixed_columns + dynamic_columns + [ORDER_REVIEW_ELIGIBILITY_HELPER]].reset_index(drop=True)


def _cover_month_display(qty: pd.Series, monthly_sales: pd.Series) -> pd.Series:
    display_monthly = pd.to_numeric(monthly_sales, errors="coerce").fillna(0).round(0)
    values = pd.Series(np.where(display_monthly.gt(0), qty / display_monthly.replace(0, np.nan), np.nan), index=monthly_sales.index)
    return values.round(1).astype(object).where(display_monthly.gt(0), "—")


def _order_sheet_transport_recommendation(
    order_qty: pd.Series,
    cover_days: pd.Series,
    lead_times: dict,
    *,
    lead_times_by_code: dict[str, int] | None = None,
    entity_code: str | None = None,
) -> pd.Series:
    qty = pd.to_numeric(order_qty, errors="coerce").fillna(0)
    cover = pd.to_numeric(cover_days, errors="coerce").fillna(999999)
    recommendations = [
        transport_mod.transport_recommendation_by_depletion_days(
            days,
            lead_times,
            order_needed=float(required_qty) > 0,
            lead_times_by_code=lead_times_by_code,
            entity_code=entity_code,
        )[0]
        for required_qty, days in zip(qty, cover)
    ]
    return pd.Series(recommendations, index=order_qty.index, dtype=object).replace("● 발주불필요", "-")


def _supported_order_sheet_transport_columns(settings: dict | None = None) -> list[str]:
    """Return the transport-quantity columns the selected entity actually uses.

    The order sheet always carries 항공/철송/해운 필요량 columns, but an entity
    only ships some of those modes.  We key off the same per-entity lead-time
    data the recommendation legend uses (`lead_time_methods` /
    `lead_times_by_code`), and fall back to the legacy name-keyed `lead_times`
    so existing PL behaviour is unchanged.
    """

    settings = settings or {}
    group_to_column = {
        "AIR": "항공\n필요량",
        "RAIL": "철송\n필요량",
        "OCEAN": "해운\n필요량",
    }
    column_order = ["항공\n필요량", "철송\n필요량", "해운\n필요량"]

    supported_groups: set[str] = set()
    lead_times_by_code = settings.get("lead_times_by_code") or {}
    for method in settings.get("lead_time_methods") or []:
        code = str(method.get("transport_code") or method.get("code") or "").strip().upper()
        if code in {"AIR_DIR", "AIR_TS", "AIR"}:
            supported_groups.add("AIR")
        elif code == "RAIL":
            supported_groups.add("RAIL")
        elif code == "OCEAN":
            supported_groups.add("OCEAN")
    for code in lead_times_by_code:
        upper = str(code).strip().upper()
        if upper in {"AIR_DIR", "AIR_TS", "AIR"}:
            supported_groups.add("AIR")
        elif upper == "RAIL":
            supported_groups.add("RAIL")
        elif upper == "OCEAN":
            supported_groups.add("OCEAN")

    if not supported_groups:
        legacy = settings.get("lead_times") if "lead_times" in settings else DEFAULT_LEAD_TIME_DAYS
        legacy = legacy or {}
        for group, label in (("AIR", "항공"), ("RAIL", "철송"), ("OCEAN", "해운")):
            if label in legacy:
                supported_groups.add(group)

    if not supported_groups:
        return column_order

    return [
        group_to_column[group]
        for group in ("AIR", "RAIL", "OCEAN")
        if group in supported_groups
    ]


def build_order_sheet_df(report_df: pd.DataFrame, settings: dict | None = None) -> pd.DataFrame:
    settings = settings or {}
    base = pd.DataFrame(report_df).copy()
    metrics = template_order_metrics(base, settings)
    sku_code = order_review_mod.report_col(base, ["상품코드", "SKU"], "").map(_clean_merge_key)
    product_name = order_review_mod.report_col(base, ["상품명", "SKU명", "제품명"], "")
    brand = order_review_mod.report_col(base, ["브랜드"])
    current_stock = pd.to_numeric(order_review_mod.report_col(base, ["EU 현지 재고", "유럽 가용재고", "EU 가용재고", "EU 현지 가용수량"], 0), errors="coerce").fillna(0)
    safety_stock = pd.to_numeric(order_review_mod.report_col(base, ["안전재고", "안전재고수량", "안전재고 목표수량"], 0), errors="coerce").fillna(0)
    monthly_sales = pd.to_numeric(order_review_mod.report_col(base, ["월평균", "월평균 판매수량"], 0), errors="coerce").fillna(0)
    shipping_stock = pd.to_numeric(order_review_mod.report_col(base, ["운송중", "운송중 수량", "운송재고"], 0), errors="coerce").fillna(0)
    current_stock = metrics["유럽재고_계산"]
    safety_stock = metrics["안전재고_계산"]
    monthly_sales = metrics["월평균_계산"]
    shipping_stock = metrics["운송재고_계산"]
    open_po_qty = pd.to_numeric(order_review_mod.report_col(base, ["미입고 수량", "미입고", "미입고수량"], 0), errors="coerce").fillna(0)
    hq_eu_stock = pd.to_numeric(order_review_mod.report_col(base, ["본사 EU창고", "본사 EU창고 가용수량", "EU창고 재고", "현지창고 재고"], 0), errors="coerce").fillna(0)
    safety_months = float(settings.get("safety_months", 3.0))
    order_qty = metrics["발주필요수량_계산"]
    order_amount_krw = metrics["발주금액_KRW_계산"]
    cover_days = pd.to_numeric(order_review_mod.report_col(base, ["EU 현지 커버일수"], 999999), errors="coerce").fillna(999999)
    depletion_raw = order_review_mod.report_col(base, ["예상 소진일", "고갈 예상일"], "판매없음")
    depletion_dates = kpi_mod.parse_date_series(depletion_raw)
    depletion = depletion_raw.astype(object).copy()
    depletion.loc[depletion_dates.notna()] = depletion_dates.loc[depletion_dates.notna()].dt.date
    lead_times = (
        settings.get("lead_times")
        if "lead_times" in settings
        else DEFAULT_LEAD_TIME_DAYS
    )
    recommendation = _order_sheet_transport_recommendation(
        order_qty,
        cover_days,
        lead_times or {},
        lead_times_by_code=settings.get("lead_times_by_code"),
        entity_code=str(settings.get("entity_code") or "PL"),
    )
    air_qty = pd.to_numeric(order_review_mod.report_col(base, ["항공 필요수량", "시간기준_항공필요수량"], 0), errors="coerce").fillna(0)
    rail_qty = pd.to_numeric(order_review_mod.report_col(base, ["철송 필요수량", "시간기준_철송필요수량"], 0), errors="coerce").fillna(0)
    sea_qty = pd.to_numeric(order_review_mod.report_col(base, ["해운 필요수량", "시간기준_해운필요수량"], 0), errors="coerce").fillna(0)
    distribution_missing = (air_qty + rail_qty + sea_qty).le(0) & order_qty.gt(0)
    recommendation_text = recommendation.astype(str)
    air_qty = air_qty.mask(distribution_missing & recommendation_text.str.contains("항공", na=False), order_qty)
    rail_qty = rail_qty.mask(distribution_missing & recommendation_text.str.contains("철송", na=False), order_qty)
    sea_qty = sea_qty.mask(distribution_missing & recommendation_text.str.contains("해운", na=False), order_qty)
    order_qty_display = order_qty

    out = pd.DataFrame(
        {
            "No": range(1, len(base) + 1),
            "SKU": _first_non_blank(sku_code, product_name),
            "상품명": product_name,
            "브랜드": brand,
            "현지재고": current_stock.round(0).astype(int),
            "월평균판매량": monthly_sales.round(0).astype(int),
            "운송재고": shipping_stock.round(0).astype(int),
            "① 현지\n(M)": _cover_month_display(current_stock, monthly_sales),
            "② 운송중\n(M)": _cover_month_display(shipping_stock, monthly_sales),
            "③ 합산\n(M)": _cover_month_display(current_stock + shipping_stock, monthly_sales),
            "안전재고 기준(M)": safety_months,
            "확보재고 수량": (current_stock + shipping_stock).round(0).astype(int),
            "안전재고 목표수량": safety_stock.round(0).astype(int),
            "⚠ 알람": "-",
            "미입고\n현황": open_po_qty.round(0).astype(int),
            "현지창고 재고": hq_eu_stock.round(0).astype(int),
            "발주총\n필요수량": order_qty_display,
            "총발주금액": order_amount_krw,
            "고갈\n예정일": depletion,
            "운송수단\n추천": recommendation,
            "항공\n필요량": air_qty.round(0).astype(int),
            "철송\n필요량": rail_qty.round(0).astype(int),
            "해운\n필요량": sea_qty.round(0).astype(int),
            "★담당자\n발주": "",
            "발주\n신호등": "",
            "★발주 메모": "",
        }
    )
    transport_columns = _supported_order_sheet_transport_columns(settings)
    optional_transport_columns = {"항공\n필요량", "철송\n필요량", "해운\n필요량"}
    output_columns = [
        column
        for column in ORDER_SHEET_COLUMNS
        if column not in optional_transport_columns or column in transport_columns
    ]
    return out[output_columns].reset_index(drop=True)


def build_order_sheet_top_sales_df(report_df: pd.DataFrame, top_n: int = 5) -> pd.DataFrame:
    base = pd.DataFrame(report_df).copy()
    sales_col = "최근 3개월 판매량"
    sales = pd.to_numeric(
        order_review_mod.report_col(base, ["PA+CA 판매수량", "기준_3M_판매수량", "최근 3개월 수요", "최근 3개월 판매수량"], 0),
        errors="coerce",
    ).fillna(0)
    total_sales = float(sales.sum())
    top = pd.DataFrame(
        {
            "SKU": _first_non_blank(order_review_mod.report_col(base, ["상품코드", "SKU"], "").map(_clean_merge_key), order_review_mod.report_col(base, ["상품명"], "")),
            "브랜드": order_review_mod.report_col(base, ["브랜드"]),
            sales_col: sales,
        }
    )
    top = top[top[sales_col].gt(0)].sort_values(sales_col, ascending=False).head(top_n)
    top[sales_col] = top[sales_col].round(0).astype(int)
    top["비중(%)"] = top[sales_col] / total_sales if total_sales > 0 else 0
    rows = top[["SKU", "브랜드", sales_col, "비중(%)"]].to_dict("records")
    top_total = int(top[sales_col].sum()) if not top.empty else 0
    rows.append({"SKU": f"합계 (상위 {top_n}개)", "브랜드": "", sales_col: top_total, "비중(%)": (top_total / total_sales if total_sales > 0 else 0)})
    return pd.DataFrame(rows, columns=["SKU", "브랜드", sales_col, "비중(%)"])


def _order_sheet_transport_criteria(settings: dict | None = None) -> list[tuple[str, str]]:
    """Build the visible recommendation legend from the selected entity's methods."""

    settings = settings or {}
    lead_times_by_code = settings.get("lead_times_by_code") or {}
    lead_time_methods = settings.get("lead_time_methods") or []
    supported_codes = {"AIR_DIR", "AIR_TS", "AIR", "RAIL", "OCEAN"}
    code_priority = {code: index for index, code in enumerate(("AIR_DIR", "AIR_TS", "AIR", "RAIL", "OCEAN"))}
    number_labels = ("①", "②", "③", "④", "⑤", "⑥")
    candidates: list[tuple[str, str, int]] = []

    for method in lead_time_methods:
        code = str(method.get("transport_code") or method.get("code") or "").strip().upper()
        if code not in supported_codes or code not in lead_times_by_code:
            continue
        label = str(method.get("display_name") or method.get("label") or "").strip()
        if not label:
            continue
        candidates.append((code, label, int(lead_times_by_code[code])))

    if not candidates:
        legacy = settings.get("lead_times") if "lead_times" in settings else DEFAULT_LEAD_TIME_DAYS
        for code, label in (("AIR", "항공"), ("RAIL", "철송"), ("OCEAN", "해운")):
            if legacy and label in legacy:
                candidates.append((code, label, int(legacy[label])))

    candidates.sort(key=lambda item: (item[2], code_priority[item[0]]))
    criteria: list[tuple[str, str]] = [("운송수단 추천 기준", "조건")]
    if not candidates:
        criteria.append(("판단불가", "선택 법인에 지원되는 운송수단 없음"))
    else:
        fastest = candidates[0]
        criteria.append((f"{number_labels[0]} {fastest[1]} 긴급", f"고갈까지 ≤ {fastest[1]} L/T일"))
        for index, current in enumerate(candidates[:-1], start=2):
            next_method = candidates[index - 1]
            criteria.append(
                (
                    f"{number_labels[index - 1]} {current[1]}",
                    f"{current[1]} L/T < 고갈 ≤ {next_method[1]} L/T",
                )
            )
        slowest = candidates[-1]
        criteria.append(
            (
                f"{number_labels[len(candidates)]} {slowest[1]} 가능",
                f"{slowest[1]} L/T < 고갈",
            )
        )
    criteria.append(("● 발주불필요", "안전재고 충족"))
    return criteria


def format_order_sheet(ws, settings: dict | None = None, top_sales_df: pd.DataFrame | None = None) -> None:
    from copy import copy
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter

    settings = settings or {}
    top_sales_df = pd.DataFrame(top_sales_df) if top_sales_df is not None else pd.DataFrame(columns=["SKU", "브랜드", "최근 3개월 판매량", "비중(%)"])
    header_row = 10
    data_start_row = header_row + 1
    dark_fill = PatternFill("solid", fgColor="4F6228")
    light_fill = PatternFill("solid", fgColor="E2F0D9")
    param_value_fill = PatternFill("solid", fgColor="FFF2CC")
    note_fill = PatternFill("solid", fgColor="F0F5E8")
    summary_value_fill = PatternFill("solid", fgColor="EBF1DE")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    alert_header_fill = PatternFill("solid", fgColor="FFFF00")
    risk_fill = PatternFill("solid", fgColor="FFC7CE")
    gray_fill = PatternFill("solid", fgColor="D9D9D9")
    white_font = Font(color="FFFFFF", bold=True)
    green_font = Font(color="4F6228", bold=True)
    parameter_title_font = Font(color="FFFFFF", bold=True, size=11)
    parameter_label_font = Font(color="4F6228", bold=True, size=9)
    parameter_value_font = Font(color="4F6228", bold=True, size=10)
    red_font = Font(color="C00000", bold=True)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.sheet_view.showGridLines = True
    lead_times = (
        settings.get("lead_times")
        if "lead_times" in settings
        else DEFAULT_LEAD_TIME_DAYS
    ) or {}
    base_date = settings.get("base_date") or korea_today()
    safety_months = float(settings.get("safety_months", 3))

    ws.merge_cells("A3:F3")
    title_cell = ws["A3"]
    title_cell.value = "▼ 파라미터 설정  (읽기전용)"
    title_cell.fill = dark_fill
    title_cell.font = parameter_title_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    parameter_cells = [
        ("A4", "안전재고(M)", "A5", safety_months),
        ("B4", "항공 L/T(일)", "B5", int(lead_times["항공"]) if "항공" in lead_times else "-"),
        ("C4", "해운 L/T(일)", "C5", int(lead_times["해운"]) if "해운" in lead_times else "-"),
        ("D4", "철송\nL/T(일)", "D5", int(lead_times["철송"]) if "철송" in lead_times else "-"),
        ("E4", "트럭\nL/T(일)", "E5", int(lead_times["트럭"]) if "트럭" in lead_times else "-"),
        ("F4", "기준일", "F5", base_date),
    ]
    for label_ref, label, value_ref, value in parameter_cells:
        label_cell = ws[label_ref]
        value_cell = ws[value_ref]
        label_cell.value = label
        value_cell.value = value
        label_cell.fill = light_fill
        value_cell.fill = param_value_fill
        label_cell.font = parameter_label_font
        value_cell.font = parameter_value_font
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.border = value_cell.border = border
        if isinstance(value, date):
            value_cell.number_format = "yyyy-mm-dd"
    for col_idx in range(1, 9):
        for row_idx in (6,):
            cell = ws.cell(row_idx, col_idx)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    criteria = _order_sheet_transport_criteria(settings)
    for r_offset, (left, right) in enumerate(criteria, start=3):
        left_cell = ws.cell(r_offset, 11, left)
        right_cell = ws.cell(r_offset, 12, right)
        ws.merge_cells(start_row=r_offset, start_column=12, end_row=r_offset, end_column=13)
        left_cell.fill = dark_fill if r_offset == 3 else light_fill
        right_cell.fill = dark_fill if r_offset == 3 else light_fill
        left_cell.font = white_font if r_offset == 3 else green_font
        right_cell.font = white_font if r_offset == 3 else Font(color="4F6228")
        left_cell.alignment = Alignment(horizontal="center", vertical="center")
        right_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        left_cell.border = right_cell.border = border
        for col_idx in (12, 13):
            ws.cell(r_offset, col_idx).border = border

    top_start_col = 19
    top_headers = ["SKU", "브랜드", "최근 3개월 판매량", "비중(%)"]
    for idx, header in enumerate(top_headers, start=top_start_col):
        cell = ws.cell(2, idx, header)
        cell.fill = gray_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    top_detail_df = top_sales_df.iloc[:-1].head(5) if not top_sales_df.empty else top_sales_df
    total_values = top_sales_df.iloc[-1].tolist() if not top_sales_df.empty else ["합계 (상위 5개)", "", 0, 0]
    for row_offset in range(5):
        excel_row = 3 + row_offset
        if row_offset < len(top_detail_df):
            values = top_detail_df.iloc[row_offset].tolist()
        else:
            values = ["", "", "", ""]
        for col_offset, value in enumerate(values):
            cell = ws.cell(excel_row, top_start_col + col_offset, value)
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if col_offset == 0 else "center", vertical="center")
            if col_offset == 2:
                cell.number_format = "#,##0"
            elif col_offset == 3:
                cell.number_format = "0.0%"
    for col_offset, value in enumerate(total_values):
        cell = ws.cell(8, top_start_col + col_offset, value)
        cell.fill = gray_fill
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="left" if col_offset == 0 else "center", vertical="center")
        if col_offset == 2:
            cell.number_format = "#,##0"
        elif col_offset == 3:
            cell.number_format = "0.0%"

    data_max_column = ws.max_column
    headers = {str(ws.cell(header_row, col).value or ""): col for col in range(1, data_max_column + 1)}
    data_end_row = ws.max_row
    data_row_count = max(data_end_row - data_start_row + 1, 0)
    fast_data_style = data_row_count > 300
    orange_headers = {
        "미입고\n현황",
        "EU창고 재고",
        "발주총\n필요수량",
        "총발주금액",
        "항공\n필요량",
        "철송\n필요량",
        "해운\n필요량",
        "★담당자\n발주",
        "발주\n신호등",
        "★발주 메모",
    }
    decimal_headers = {"① 현지\n(M)", "② 운송중\n(M)", "③ 합산\n(M)"}
    text_headers = {"SKU", "브랜드", "⚠ 알람", "운송수단\n추천", "발주\n신호등", "★발주 메모"}
    date_headers = {"고갈\n예정일"}
    for col_idx in range(1, ws.max_column + 1):
        header_cell = ws.cell(header_row, col_idx)
        header = str(header_cell.value or "")
        header_cell.fill = alert_header_fill if header == "⚠ 알람" else orange_fill if header in orange_headers else light_fill
        header_cell.font = green_font
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cell.border = border

    required_col = headers.get("발주총\n필요수량")
    amount_col = headers.get("총발주금액")
    recommendation_col = headers.get("운송수단\n추천")
    manager_col = headers.get("★담당자\n발주")
    signal_col = headers.get("발주\n신호등")
    alert_col = headers.get("⚠ 알람")
    local_stock_col = headers.get("현지재고") or headers.get("유럽현지재고")
    monthly_sales_col = headers.get("월평균판매량")
    combined_months_col = headers.get("③ 합산\n(M)")
    last_col_letter = get_column_letter(ws.max_column)
    if required_col:
        total_cell = ws.cell(header_row - 1, required_col)
        total_cell.value = f"=SUM({get_column_letter(required_col)}{data_start_row}:{get_column_letter(required_col)}{data_end_row})"
        total_cell.fill = orange_fill
        total_cell.font = Font(color="C00000", bold=True)
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.border = border
        total_cell.number_format = "#,##0"
    if amount_col:
        total_cell = ws.cell(header_row - 1, amount_col)
        total_cell.value = f"=SUM({get_column_letter(amount_col)}{data_start_row}:{get_column_letter(amount_col)}{data_end_row})"
        total_cell.fill = orange_fill
        total_cell.font = Font(color="C00000", bold=True)
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.border = border
        total_cell.number_format = "#,##0"
    if fast_data_style and required_col:
        required_letter = get_column_letter(required_col)
        ws.conditional_formatting.add(
            f"A{data_start_row}:{last_col_letter}{data_end_row}",
            FormulaRule(
                formula=[f'OR(${required_letter}{data_start_row}>0,${required_letter}{data_start_row}="-")'],
                fill=risk_fill,
            ),
        )
    if recommendation_col:
        recommendation_letter = get_column_letter(recommendation_col)
        ws.conditional_formatting.add(
            f"{recommendation_letter}{data_start_row}:{recommendation_letter}{data_end_row}",
            FormulaRule(
                formula=[f'AND(${recommendation_letter}{data_start_row}<>"",${recommendation_letter}{data_start_row}<>"-")'],
                font=red_font,
            ),
        )
    if alert_col and local_stock_col and required_col:
        local_stock_letter = get_column_letter(local_stock_col)
        required_letter = get_column_letter(required_col)
        for row_idx in range(data_start_row, data_end_row + 1):
            ws.cell(row_idx, alert_col).value = (
                f'=IF({local_stock_letter}{row_idx}<=0,"{ALERT_URGENT_LABEL}",'
                f'IF({required_letter}{row_idx}>0,"{ALERT_LOCAL_SHORT_LABEL}","-"))'
            )
    sku_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="center")
    if fast_data_style:
        for row_idx in range(data_start_row, data_end_row + 1):
            for col_idx in range(1, data_max_column + 1):
                ws.cell(row_idx, col_idx).border = border
            if manager_col:
                ws.cell(row_idx, manager_col).value = None
    else:
        for row_idx in range(data_start_row, data_end_row + 1):
            required_value = ws.cell(row_idx, required_col).value if required_col else 0
            required_number = pd.to_numeric(required_value, errors="coerce")
            highlight = (str(required_value).strip() == "-") or (pd.notna(required_number) and float(required_number) > 0)
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row_idx, col_idx)
                header = str(ws.cell(header_row, col_idx).value or "")
                cell.border = border
                cell.alignment = sku_alignment if header == "SKU" else center_alignment
                if highlight:
                    cell.fill = risk_fill
                if header in text_headers:
                    cell.number_format = "@"
                    if header == "운송수단\n추천" and str(cell.value or "").strip() not in {"", "-"}:
                        cell.font = red_font
                elif header in date_headers:
                    cell.number_format = "yyyy-mm-dd"
                elif header in decimal_headers:
                    cell.number_format = "#,##0.##"
                elif header == "총발주금액":
                    cell.number_format = "#,##0"
                else:
                    cell.number_format = "#,##0"
            if manager_col:
                ws.cell(row_idx, manager_col).value = None
                ws.cell(row_idx, manager_col).number_format = "#,##0"

    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(header_row, col_idx).value or "")
        if header in decimal_headers:
            for row_idx in range(data_start_row, data_end_row + 1):
                ws.cell(row_idx, col_idx).number_format = "#,##0.##"

    if required_col and manager_col and signal_col:
        required_letter = get_column_letter(required_col)
        manager_letter = get_column_letter(manager_col)
        signal_letter = get_column_letter(signal_col)
        for row_idx in range(data_start_row, data_end_row + 1):
            signal_cell = ws.cell(row_idx, signal_col)
            signal_cell.value = (
                f'=IF(NOT(ISNUMBER({required_letter}{row_idx})),"● 미발주",'
                f'IF({manager_letter}{row_idx}="",'
                f'IF({required_letter}{row_idx}<=0,"● 발주불필요","● 미발주"),'
                f'IF(NOT(ISNUMBER({manager_letter}{row_idx})),"● 확인필요",'
                f'IF({manager_letter}{row_idx}<0,"● 확인필요",'
                f'IF({manager_letter}{row_idx}>={required_letter}{row_idx},"● 안전","● 부족")))))'
            )
        signal_range = f"{signal_letter}{data_start_row}:{signal_letter}{ws.max_row}"
        signal_fonts = {
            "● 안전": Font(color="2E7D32", bold=True),
            "● 미발주": red_font,
            "● 확인필요": red_font,
            "● 부족": red_font,
            "● 발주불필요": Font(color="000000", bold=True),
        }
        for signal_text, signal_font in signal_fonts.items():
            ws.conditional_formatting.add(signal_range, CellIsRule(operator="equal", formula=[f'"{signal_text}"'], font=signal_font))
        for header_text, col_idx in [("★담당자\n발주", manager_col), ("발주\n신호등", signal_col)]:
            header_cell = ws.cell(header_row, col_idx)
            new_font = copy(header_cell.font)
            new_font.color = "C00000" if header_text == "★담당자\n발주" else "4F6228"
            header_cell.font = new_font

    # Keep the order table visibly boxed even when fast styling skips per-cell formatting.
    for row_idx in range(header_row, data_end_row + 1):
        for col_idx in range(1, data_max_column + 1):
            ws.cell(row_idx, col_idx).border = border

    widths = {
        "No": 6,
        "SKU": 18,
        "상품명": 42,
        "브랜드": 15,
        "유럽현지재고": 12,
        "월평균판매량": 12,
        "운송재고": 10,
        "① 현지\n(M)": 10,
        "② 운송중\n(M)": 11,
        "③ 합산\n(M)": 14,
        "안전재고 기준(M)": 14,
        "확보재고 수량": 14,
        "안전재고 목표수량": 15,
        "⚠ 알람": 11,
        "미입고\n현황": 9,
        "EU창고 재고": 10,
        "발주총\n필요수량": 12,
        "총발주금액": 16,
        "고갈\n예정일": 18,
        "운송수단\n추천": 13,
        "항공\n필요량": 10,
        "철송\n필요량": 10,
        "해운\n필요량": 10,
        "★담당자\n발주": 12,
        "발주\n신호등": 12,
        "★발주 메모": 15,
    }
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(header_row, col_idx).value or "")
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, 12)

    # The parameter table lives on A:H, so keep these columns visually even.
    # Otherwise B inherits the SKU column width and makes the top controls look broken.
    for col_letter in "ABCDEFGH":
        ws.column_dimensions[col_letter].width = 15
    ws.column_dimensions["K"].width = max(ws.column_dimensions["K"].width or 0, 19)
    ws.column_dimensions["L"].width = max(ws.column_dimensions["L"].width or 0, 24)
    ws.column_dimensions["M"].width = max(ws.column_dimensions["M"].width or 0, 18)
    header_cols = {str(ws.cell(header_row, col).value or ""): col for col in range(1, ws.max_column + 1)}
    for header, width in {"SKU": 18, "상품명": 42, "브랜드": 15}.items():
        col_idx = header_cols.get(header)
        if col_idx:
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    for header, width in {
        "안전재고 기준(M)": 18,
        "확보재고 수량": 18,
        "운송수단\n추천": 28,
        "항공\n필요량": 12,
        "철송\n필요량": 13,
        "해운\n필요량": 11,
    }.items():
        col_idx = header_cols.get(header)
        if col_idx:
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
    for col_idx, width in {
        top_start_col: 28,
        top_start_col + 1: 12,
        top_start_col + 2: 20,
        top_start_col + 3: 11,
    }.items():
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
    for row_idx, height in {3: 18, 4: 34, 5: 20, 6: 20, 8: 20, header_row: 34}.items():
        ws.row_dimensions[row_idx].height = height
    for row_idx in range(3, 9):
        ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 0, 22)
    ws.row_dimensions[6].hidden = False
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{data_end_row}"
    ws.freeze_panes = f"D{data_start_row}"
    for row_idx in range(1, data_end_row + 1):
        for col_idx in range(1, data_max_column + 1):
            ws.cell(row_idx, col_idx).protection = Protection(locked=False)
    if alert_col:
        for row_idx in range(data_start_row, data_end_row + 1):
            ws.cell(row_idx, alert_col).protection = Protection(locked=True, hidden=True)
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.sort = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
