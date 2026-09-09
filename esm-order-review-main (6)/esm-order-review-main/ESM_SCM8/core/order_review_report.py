from __future__ import annotations

from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd

from core.common import (
    _DEFAULT_EUR_KRW_RATE,
    CHECK_REQUIRED_COLUMNS,
    CHECK_REQUIRED_DISPLAY_COLUMNS,
    CHECK_REQUIRED_INTERNAL_COLUMNS,
    CHECK_REQUIRED_OUTPUT_COLUMNS,
)
from core.preprocess import product_sku_mask
from core.session import SessionContext
from core.export_excel_calculations import template_order_metrics
from core import export_excel as export_excel_mod, kpi as kpi_mod, loaders as loaders_mod, order_review as order_review_mod, preprocess as preprocess_mod


def build_order_review_report_df(review: pd.DataFrame, settings: dict | None = None) -> pd.DataFrame:
    currency_code = str((settings or {}).get("currency_code") or "EUR").strip().upper()
    if currency_code not in {"EUR", "USD"}:
        currency_code = "EUR"
    out = pd.DataFrame(
        {
            "No": range(1, len(review) + 1),
            "구분": order_review_mod.report_col(review, ["발주검토여부"]),
            "최종 액션": order_review_mod.report_col(review, ["최종 액션", "우선 액션"]),
            "상태": order_review_mod.report_col(review, ["상태"], "정상"),
            "조치 요약": order_review_mod.report_col(review, ["조치 요약"]),
            "근거": order_review_mod.report_col(review, ["근거", "판단 사유"]),
            "예외 플래그": order_review_mod.report_col(review, ["예외 플래그"]),
            "예외 사유": order_review_mod.report_col(review, ["예외 사유"]),
            "우선 액션": order_review_mod.report_col(review, ["우선 액션"]),
            "SKU등록상태": order_review_mod.report_col(review, ["SKU등록상태"]),
            "재고운영상태": order_review_mod.report_col(review, ["재고운영상태"]),
            "판매이력여부": order_review_mod.report_col(review, ["판매이력여부"]),
            "운송중여부": order_review_mod.report_col(review, ["운송중여부"]),
            "미입고여부": order_review_mod.report_col(review, ["미입고여부"]),
            "판단메모": order_review_mod.report_col(review, ["판단메모"]),
            "판매수량 기준": order_review_mod.report_col(review, ["판매수량 기준"]),
            "브랜드": order_review_mod.report_col(review, ["브랜드"]),
            "상품명": order_review_mod.report_col(review, ["상품명"]),
            "바코드": order_review_mod.report_col(review, ["바코드"]),
            "상품코드": order_review_mod.report_col(review, ["상품코드"]),
            "커버개월": pd.to_numeric(order_review_mod.report_col(review, ["커버가능 개월수"], 0), errors="coerce").fillna(0),
            "기준_3M_판매수량": pd.to_numeric(order_review_mod.report_col(review, ["기준_3M_판매수량"], 0), errors="coerce").fillna(0),
            "PA_CA_3M_판매수량": pd.to_numeric(order_review_mod.report_col(review, ["PA_CA_3M_판매수량"], 0), errors="coerce").fillna(0),
            "PA+CA 판매수량": pd.to_numeric(order_review_mod.report_col(review, ["PA+CA 판매수량"], 0), errors="coerce").fillna(0),
            "최근 3개월 수요": pd.to_numeric(order_review_mod.report_col(review, ["최근 3개월 판매수량"], 0), errors="coerce").fillna(0),
            "월평균": pd.to_numeric(order_review_mod.report_col(review, ["월평균 판매수량"], 0), errors="coerce").fillna(0),
            "일평균 판매수량": pd.to_numeric(order_review_mod.report_col(review, ["일평균 판매수량"], 0), errors="coerce").fillna(0),
            "EU 현지 커버일수": pd.to_numeric(order_review_mod.report_col(review, ["EU 현지 커버일수"], 0), errors="coerce").fillna(0),
            "안전재고 개월 수": pd.to_numeric(order_review_mod.report_col(review, ["안전재고 개월 수"], 0), errors="coerce").fillna(0),
            "안전재고": pd.to_numeric(order_review_mod.report_col(review, ["안전재고 목표수량"], 0), errors="coerce").fillna(0),
            "안전재고수량": pd.to_numeric(order_review_mod.report_col(review, ["안전재고수량"], 0), errors="coerce").fillna(0),
            "EU 현지 재고": pd.to_numeric(order_review_mod.report_col(review, ["EU 현지 가용수량"], 0), errors="coerce").fillna(0),
            "유럽 가용재고": pd.to_numeric(order_review_mod.report_col(review, ["유럽 가용재고"], 0), errors="coerce").fillna(0),
            "본사 EU창고": pd.to_numeric(order_review_mod.report_col(review, ["본사 EU창고 가용수량"], 0), errors="coerce").fillna(0),
            "운송중": pd.to_numeric(order_review_mod.report_col(review, ["운송중 수량"], 0), errors="coerce").fillna(0),
            "입고 전 예상 결품량": pd.to_numeric(order_review_mod.report_col(review, ["입고 전 예상 결품량"], 0), errors="coerce").fillna(0),
            "미입고": pd.to_numeric(order_review_mod.report_col(review, ["미입고수량"], 0), errors="coerce").fillna(0),
            "미입고 수량": pd.to_numeric(order_review_mod.report_col(review, ["미입고수량"], 0), errors="coerce").fillna(0),
            "유럽+운송": pd.to_numeric(order_review_mod.report_col(review, ["유럽+운송 수량"], 0), errors="coerce").fillna(0),
            "재고 ETA 상태": order_review_mod.report_col(review, ["재고 ETA 상태"], ""),
            "부족수량 산정 기준 보유량": pd.to_numeric(order_review_mod.report_col(review, ["부족수량 산정 기준 보유량"], 0), errors="coerce").fillna(0),
            "운송포함 보유개월": pd.to_numeric(order_review_mod.report_col(review, ["운송 포함 보유개월수"], 0), errors="coerce").fillna(0),
            "안전재고 목표수량": pd.to_numeric(order_review_mod.report_col(review, ["안전재고 목표수량", "안전재고수량"], 0), errors="coerce").fillna(0),
            "부족수량": pd.to_numeric(order_review_mod.report_col(review, ["최종 부족수량"], 0), errors="coerce").fillna(0),
            "1차 부족수량": pd.to_numeric(order_review_mod.report_col(review, ["1차 부족수량"], 0), errors="coerce").fillna(0),
            "본사 이동 수량": pd.to_numeric(order_review_mod.report_col(review, ["본사 이동 수량"], 0), errors="coerce").fillna(0),
            "추가 발주 필요 수량": pd.to_numeric(order_review_mod.report_col(review, ["추가 발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0),
            "발주필요수량": pd.to_numeric(order_review_mod.report_col(review, ["발주필요수량"], 0), errors="coerce").fillna(0),
            "긴급 보충 필요 수량": pd.to_numeric(order_review_mod.report_col(review, ["긴급 보충 필요 수량"], 0), errors="coerce").fillna(0),
            "입고 전 결품 위험 여부": order_review_mod.report_col(review, ["입고 전 결품 위험 여부"], "N"),
            "권장 긴급 액션": order_review_mod.report_col(review, ["권장 긴급 액션"]),
            "논의필요": order_review_mod.report_col(review, ["논의필요"]),
            f"현지 입고단가({currency_code})": pd.to_numeric(order_review_mod.report_col(review, [f"현지 입고단가_{currency_code}", "현지 입고단가", "EU 입고단가"], 0), errors="coerce").fillna(0),
            f"부족금액_{currency_code}": pd.to_numeric(order_review_mod.report_col(review, [f"부족금액_{currency_code}", "부족금액_EUR"], 0), errors="coerce").fillna(0),
            "부족금액_KRW": pd.to_numeric(order_review_mod.report_col(review, ["부족금액_KRW"], 0), errors="coerce").fillna(0),
            "추가 발주 필요 금액": pd.to_numeric(order_review_mod.report_col(review, ["추가 발주 필요 금액", "발주필요금액"], 0), errors="coerce").fillna(0),
            "발주필요금액_KRW": pd.to_numeric(order_review_mod.report_col(review, ["발주필요금액"], 0), errors="coerce").fillna(0),
            "ETA 지연 플래그": order_review_mod.report_col(review, ["ETA 지연 플래그"]),
            "운송 검토안": order_review_mod.report_col(review, ["운송수단 검토안"]),
            "권장 운송안": order_review_mod.report_col(review, ["시간기준_권장운송안"], "추가 운송 불필요"),
            "권장 수량 요약": order_review_mod.report_col(review, ["시간기준_권장수량요약"], "추가 운송 불필요"),
            "항공 필요수량": pd.to_numeric(order_review_mod.report_col(review, ["시간기준_항공필요수량"], 0), errors="coerce").fillna(0),
            "철송 필요수량": pd.to_numeric(order_review_mod.report_col(review, ["시간기준_철송필요수량"], 0), errors="coerce").fillna(0),
            "해운 필요수량": pd.to_numeric(order_review_mod.report_col(review, ["시간기준_해운필요수량"], 0), errors="coerce").fillna(0),
            "판단 사유": order_review_mod.report_col(review, ["판단 사유"]),
            "예상 소진일": order_review_mod.report_col(review, ["고갈 예상일"]),
            "첫 입고 예정일": order_review_mod.report_col(review, ["최초 ETA"]),
            "기준일": order_review_mod.report_col(review, ["기준일"]),
            "적용 리드타임": pd.to_numeric(order_review_mod.report_col(review, ["적용 리드타임"], 0), errors="coerce").fillna(0),
            "예상 도착 가능일": order_review_mod.report_col(review, ["예상 도착 가능일"]),
        }
    )
    # 구형 PL Excel 조립 코드와 과거 결과 호환용 별칭이다.
    out["EU 입고단가(EUR)"] = out[f"현지 입고단가({currency_code})"]
    out["부족금액_EUR"] = out[f"부족금액_{currency_code}"]
    stockout_df = export_excel_mod.build_stockout_calendar_report_df(review)
    if not stockout_df.empty and "상품코드" in stockout_df.columns and "쇼티지 예상 일수" in stockout_df.columns:
        shortage_map = stockout_df.drop_duplicates("상품코드").set_index("상품코드")["쇼티지 예상 일수"]
        out["쇼티지 예상 일수"] = out["상품코드"].map(shortage_map)
    else:
        out["쇼티지 예상 일수"] = np.nan
    metrics = template_order_metrics(review, settings)
    canonical_monthly = pd.to_numeric(metrics["월평균_계산"], errors="coerce").fillna(0)
    canonical_safety = pd.to_numeric(metrics["안전재고_계산"], errors="coerce").fillna(0)
    canonical_eu_stock = pd.to_numeric(metrics["유럽재고_계산"], errors="coerce").fillna(0)
    canonical_shipping = pd.to_numeric(metrics["운송재고_계산"], errors="coerce").fillna(0)
    canonical_combined = pd.to_numeric(metrics["유럽운송합산_계산"], errors="coerce").fillna(0)
    canonical_target = canonical_safety
    canonical_order_qty = pd.to_numeric(metrics["발주필요수량_계산"], errors="coerce").fillna(0)
    canonical_order_amount_eur = pd.to_numeric(metrics["발주금액_EUR_계산"], errors="coerce").fillna(0)
    canonical_order_amount = pd.to_numeric(metrics["발주금액_KRW_계산"], errors="coerce").fillna(0)
    out["월평균"] = canonical_monthly.to_numpy()
    out["안전재고"] = canonical_safety.to_numpy()
    out["안전재고수량"] = canonical_safety.to_numpy()
    out["EU 현지 재고"] = canonical_eu_stock.to_numpy()
    out["유럽 가용재고"] = canonical_eu_stock.to_numpy()
    out["운송중"] = canonical_shipping.to_numpy()
    out["유럽+운송"] = canonical_combined.to_numpy()
    out["부족수량 산정 기준 보유량"] = canonical_combined.to_numpy()
    # 커버개월은 리뷰 단계의 반올림 전 월평균으로 계산돼 있어, 같은 시트에 표시되는
    # 월평균·EU 현지 재고(canonical 반올림 값)와 나누어 떨어지지 않았다. 재고ETA 시트의
    # 커버개월 계산과도 기준이 달라 한 결과 안에 커버기간이 두 종류 존재했다.
    out["커버개월"] = np.where(
        canonical_monthly.gt(0),
        canonical_eu_stock / canonical_monthly.where(canonical_monthly.gt(0)),
        99,
    )
    out["재고 ETA 상태"] = np.select(
        [
            canonical_eu_stock.le(0),
            canonical_order_qty.gt(0),
        ],
        [
            "OOS",
            "발주필요",
        ],
        default="-",
    )
    out["안전재고 목표수량"] = canonical_target.to_numpy()
    out["부족수량"] = canonical_order_qty.to_numpy()
    out["1차 부족수량"] = canonical_order_qty.to_numpy()
    out["최종 부족수량"] = canonical_order_qty.to_numpy()
    out["추가 발주 필요 수량"] = canonical_order_qty.to_numpy()
    out["부족금액_EUR"] = canonical_order_amount_eur.to_numpy()
    out["부족금액_KRW"] = canonical_order_amount.to_numpy()
    out["추가 발주 필요 금액"] = canonical_order_amount.to_numpy()
    out["발주필요수량"] = canonical_order_qty.to_numpy()
    out["발주 필요 수량"] = canonical_order_qty.to_numpy()
    out["발주필요금액_KRW"] = canonical_order_amount.to_numpy()
    out["발주 필요 금액"] = canonical_order_amount.to_numpy()
    normal_status = out["상태"].astype(str).eq("정상")
    canonical_order_needed = canonical_order_qty.gt(0) & normal_status
    canonical_no_order = canonical_order_qty.le(0) & normal_status
    out.loc[canonical_order_needed, "최종 액션"] = "발주 필요"
    out.loc[canonical_order_needed, "우선 액션"] = "발주 필요"
    out.loc[canonical_order_needed, "구분"] = "필요"
    out.loc[canonical_order_needed, "발주검토여부"] = "필요"
    out.loc[canonical_order_needed, "근거"] = "안전재고 목표수량 대비 현지 가용수량+운송중 수량 부족"
    out.loc[canonical_order_needed, "판단 사유"] = out.loc[canonical_order_needed, "근거"]
    out.loc[canonical_order_needed, "조치 요약"] = canonical_order_qty.loc[canonical_order_needed].round(0).astype(int).map(
        lambda qty: f"신규 발주 {qty:,}개"
    )
    out.loc[canonical_no_order, "최종 액션"] = "발주 필요 없음"
    out.loc[canonical_no_order, "우선 액션"] = "발주 필요 없음"
    out.loc[canonical_no_order, "구분"] = "불필요"
    out.loc[canonical_no_order, "발주검토여부"] = "불필요"
    out.loc[canonical_no_order, "근거"] = "현지 가용수량+운송중 수량으로 안전재고 목표수량 충족"
    out.loc[canonical_no_order, "판단 사유"] = out.loc[canonical_no_order, "근거"]
    out.loc[canonical_no_order, "조치 요약"] = "발주 필요 없음"
    out["논의필요"] = np.where(
        out["최종 액션"].astype(str).eq("발주 필요") | out["상태"].astype(str).isin(["확인필요"]),
        "Y",
        "",
    )
    out["EU 가용재고"] = out["유럽 가용재고"]
    out["예상 입고일"] = out["첫 입고 예정일"]
    base_date_value = None
    if "기준일" in out.columns and not out.empty:
        parsed_base = kpi_mod.parse_date_series(out["기준일"]).dropna()
        if not parsed_base.empty:
            base_date_value = parsed_base.iloc[0].date()
    out = order_review_mod.apply_order_review_reference_flags(out, base_date=base_date_value)
    out["운송수단 추천 사유"] = out.apply(order_review_mod.describe_transport_recommendation, axis=1)
    out = order_review_mod.sort_order_review_output(out).reset_index(drop=True)
    out["No"] = range(1, len(out) + 1)
    front_cols = [
        "브랜드", "상품코드", "상품명", "우선 액션", "참고 플래그", "판단 사유",
        "즉시OOS여부", "입고지연위험", "ETA미확인", "쇼티지 예상 일수",
        "발주 필요 금액", "발주 필요 수량", "긴급 보충 필요 수량", "입고 전 결품 위험 여부", "권장 긴급 액션",
        "미입고 수량", "운송중", "재고 ETA 상태", "입고 전 예상 결품량", "EU 가용재고",
        "예상 입고일", "예상 소진일",
    ]
    ordered_cols = [col for col in front_cols if col in out.columns] + [col for col in out.columns if col not in front_cols]
    out = out[ordered_cols]
    return out


def order_review_dashboard_kpi_values(report_df: pd.DataFrame) -> dict[str, float]:
    report = pd.DataFrame(report_df).copy()
    action = order_review_mod.report_col(report, ["우선 액션", "최종 액션"]).astype(str)
    status = order_review_mod.report_col(report, ["상태"], "정상").astype(str)
    pre_arrival_risk = order_review_mod.report_col(report, ["입고 전 결품 위험 여부"], "N").astype(str).eq("Y")
    order_needed = order_review_mod.order_needed_action_mask(action)
    return {
        "review_needed_sku": float((order_needed | pre_arrival_risk | status.isin(["본사이동", "운송대기", "확인필요"])).sum()),
        "order_needed_sku": float(order_needed.sum()),
        "hq_move_sku": float(status.eq("본사이동").sum()),
        "transport_wait_sku": float(status.eq("운송대기").sum()),
        "order_needed_qty": float(pd.to_numeric(order_review_mod.report_col(report, ["추가 발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0).sum()),
        "order_needed_amount": float(pd.to_numeric(order_review_mod.report_col(report, ["추가 발주 필요 금액", "발주필요금액_KRW", "발주필요금액"], 0), errors="coerce").fillna(0).sum()),
    }


def stock_eta_dashboard_summary_rows(report_df: pd.DataFrame) -> list[tuple[str, object]]:
    values = order_review_dashboard_kpi_values(report_df)
    return [
        ("검토필요 SKU", int(values["review_needed_sku"])),
        ("발주필요 SKU", int(values["order_needed_sku"])),
        ("본사이동 SKU", int(values["hq_move_sku"])),
        ("운송대기 SKU", int(values["transport_wait_sku"])),
        ("발주필요 수량", int(round(values["order_needed_qty"]))),
        ("발주필요 금액", kpi_mod.fmt_krw_compact(values["order_needed_amount"])),
    ]


def build_order_review_top10(report_df: pd.DataFrame) -> pd.DataFrame:
    action = order_review_mod.report_col(report_df, ["우선 액션", "최종 액션"]).astype(str)
    order_qty = pd.to_numeric(order_review_mod.report_col(report_df, ["추가 발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0)
    top = report_df[(order_qty > 0) | order_review_mod.order_needed_action_mask(action)].copy()
    top = top.sort_values(["추가 발주 필요 금액", "추가 발주 필요 수량"], ascending=[False, False]).head(10)
    return pd.DataFrame(
        {
            "순위": list(range(1, len(top) + 1)),
            "브랜드": top["브랜드"].tolist(),
            "상품명": top["상품명"].tolist(),
            "상품코드": top["상품코드"].tolist(),
            "커버개월": top["커버개월"].tolist(),
            "부족수량": top["부족수량"].tolist(),
            "부족금액_KRW": top["부족금액_KRW"].tolist(),
            "운송 검토안": top["운송 검토안"].tolist(),
            "최종 액션": top["최종 액션"].tolist() if "최종 액션" in top.columns else top["우선 액션"].tolist(),
            "우선 액션": top["우선 액션"].tolist(),
        }
    )


def build_sales_concentration_top10_df(report_df: pd.DataFrame) -> pd.DataFrame:
    top = report_df.copy()
    top["최근 3개월 PA+CA 판매수량"] = pd.to_numeric(top["PA+CA 판매수량"], errors="coerce").fillna(0)
    total_sales = float(top["최근 3개월 PA+CA 판매수량"].sum())
    top = top.sort_values(["최근 3개월 PA+CA 판매수량", "월평균"], ascending=[False, False]).head(10)
    recent_sales_display = top["최근 3개월 PA+CA 판매수량"].round(0).astype(int)
    monthly_sales_display = pd.to_numeric(top["월평균"], errors="coerce").fillna(0).round(0).astype(int)
    return pd.DataFrame(
        {
            "순위": range(1, len(top) + 1),
            "브랜드": top["브랜드"].tolist(),
            "상품코드": top["상품코드"].tolist(),
            "상품명": top["상품명"].tolist(),
            "최근 3개월 PA+CA 판매수량": recent_sales_display.tolist(),
            "월평균 판매수량": monthly_sales_display.tolist(),
            "판매 집중도": [float(value) / total_sales if total_sales > 0 else 0 for value in top["최근 3개월 PA+CA 판매수량"]],
            "우선 액션": top["우선 액션"].tolist(),
        }
    )


def build_order_risk_top10_df(report_df: pd.DataFrame) -> pd.DataFrame:
    action = order_review_mod.report_col(report_df, ["우선 액션", "최종 액션"]).astype(str)
    order_qty = pd.to_numeric(order_review_mod.report_col(report_df, ["추가 발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0)
    top = report_df[(order_qty > 0) | order_review_mod.order_needed_action_mask(action)].copy()
    top = top.sort_values(["추가 발주 필요 금액", "추가 발주 필요 수량", "커버개월"], ascending=[False, False, True]).head(10)
    return pd.DataFrame(
        {
            "순위": range(1, len(top) + 1),
            "브랜드": top["브랜드"].tolist(),
            "상품코드": top["상품코드"].tolist(),
            "상품명": top["상품명"].tolist(),
            "발주필요수량": top["발주필요수량"].tolist(),
            "부족금액_KRW": top["부족금액_KRW"].tolist(),
            "커버개월": top["커버개월"].tolist(),
            "운송 검토안": top["운송 검토안"].tolist(),
            "최종 액션": top["최종 액션"].tolist() if "최종 액션" in top.columns else top["우선 액션"].tolist(),
            "우선 액션": top["우선 액션"].tolist(),
        }
    )


def build_pre_eta_stockout_top10_df(stockout_df: pd.DataFrame) -> pd.DataFrame:
    if stockout_df.empty:
        return pd.DataFrame(
            columns=["순위", "상품코드", "상품명", "예상 소진일", "첫 입고 예정일", "쇼티지 예상 일수", "추천 대응", "우선 액션"]
        )
    top = stockout_df.copy()
    shortage_days = top.get("쇼티지 예상 일수", top.get("ETA 전 공백일수", pd.Series([0] * len(top), index=top.index)))
    if not isinstance(shortage_days, pd.Series):
        shortage_days = pd.Series([shortage_days] * len(top), index=top.index)
    top["_공백일수"] = pd.to_numeric(shortage_days, errors="coerce").fillna(0)
    top["_예상소진일"] = pd.to_datetime(top.get("예상 소진일", ""), errors="coerce")
    top["_첫입고예정일"] = pd.to_datetime(top.get("첫 입고 예정일", ""), errors="coerce")
    risk_mask = top["입고 전 품절 여부"].astype(str).eq("Y") | top["우선 액션"].astype(str).str.contains("ETA 지연 위험", na=False)
    top = top[risk_mask].sort_values(
        ["_예상소진일", "_첫입고예정일", "_공백일수"],
        ascending=[True, True, False],
        na_position="last",
    ).head(10)
    top = top.drop(columns=["_공백일수", "_예상소진일", "_첫입고예정일"], errors="ignore")
    top.insert(0, "순위", range(1, len(top) + 1))
    keep_cols = ["순위", "상품코드", "상품명", "예상 소진일", "첫 입고 예정일", "쇼티지 예상 일수", "추천 대응", "우선 액션"]
    return top[[col for col in keep_cols if col in top.columns]]


def build_order_review_summary_df(report_df: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "브랜드",
        "상품명",
        "상품코드",
        "최종 액션",
        "상태",
        "조치 요약",
        "발주필요수량",
        "추가 발주 필요 수량",
        "본사 이동 수량",
        "발주필요금액_KRW",
        "추가 발주 필요 금액",
        "운송 검토안",
        "우선 액션",
        "판단메모",
        "근거",
        "판단 사유",
        "일평균 판매수량",
        "PA+CA 판매수량",
        "월평균",
        "안전재고수량",
        "안전재고 목표수량",
        "유럽 가용재고",
        "운송중",
        "미입고 수량",
        "커버개월",
        "운송포함 보유개월",
        "EU 현지 커버일수",
        "예상 소진일",
        "첫 입고 예정일",
        "쇼티지 예상 일수",
        "운송중여부",
        "미입고여부",
        "운송수단 추천 사유",
    ]
    out = report_df[[col for col in columns if col in report_df.columns]].copy()
    out = out.rename(
        columns={
            "PA+CA 판매수량": "기준 3개월 판매수량(재고 PA+CA)",
            "월평균": "월평균 판매수량(기준/3)",
        }
    )
    monthly_col = "월평균 판매수량(기준/3)"
    if monthly_col in out.columns:
        out[monthly_col] = pd.to_numeric(out[monthly_col], errors="coerce").fillna(0).round(0).astype(int)
    sort_cols = [col for col in ["발주필요수량", "발주필요금액_KRW", "부족수량", monthly_col] if col in out.columns]
    if sort_cols:
        out = out.sort_values(sort_cols, ascending=[False] * len(sort_cols)).reset_index(drop=True)
    return out


def build_excluded_order_review_df(excluded_review: pd.DataFrame, settings: dict | None = None) -> pd.DataFrame:
    currency_code = str((settings or {}).get("currency_code") or "EUR").strip().upper()
    if currency_code not in {"EUR", "USD"}:
        currency_code = "EUR"
    if excluded_review.empty:
        return pd.DataFrame(
            columns=[
                "제외유형", "제외사유", "브랜드", "상품명", "제품상태", "상품코드", "PA+CA 판매수량",
                "월평균 판매수량", "안전재고수량", "유럽 가용재고", "본사 EU창고 가용수량", "유럽+본사 가용재고",
                "운송중 수량", "미입고 수량",
                "제외 수량", "제외 금액_KRW", "최종 액션", "상태", "조치 요약", "근거", "예외 플래그", "예외 사유", "우선 액션", "판단 사유",
            ]
        )
    eu_available = pd.to_numeric(order_review_mod.report_col(excluded_review, ["유럽 가용재고"], 0), errors="coerce").fillna(0)
    hq_available = pd.to_numeric(order_review_mod.report_col(excluded_review, ["본사 EU창고 가용수량"], 0), errors="coerce").fillna(0)
    out = pd.DataFrame(
        {
            "제외유형": order_review_mod.report_col(excluded_review, ["제외유형"]),
            "제외사유": order_review_mod.report_col(excluded_review, ["제외사유"]),
            "브랜드": order_review_mod.report_col(excluded_review, ["브랜드"]),
            "상품명": order_review_mod.report_col(excluded_review, ["상품명"]),
            "제품상태": order_review_mod.report_col(excluded_review, ["제품상태"]),
            "상품코드": order_review_mod.report_col(excluded_review, ["상품코드"]),
            "PA+CA 판매수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["PA+CA 판매수량"], 0), errors="coerce").fillna(0),
            "월평균 판매수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["월평균 판매수량"], 0), errors="coerce").fillna(0),
            "안전재고수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["안전재고수량"], 0), errors="coerce").fillna(0),
            "유럽 가용재고": eu_available,
            "본사 EU창고 가용수량": hq_available,
            "유럽+본사 가용재고": eu_available + hq_available,
            "운송중 수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["운송중 수량"], 0), errors="coerce").fillna(0),
            "미입고 수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["미입고수량"], 0), errors="coerce").fillna(0),
            "제외 수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["발주필요수량"], 0), errors="coerce").fillna(0),
            "제외 금액_KRW": pd.to_numeric(order_review_mod.report_col(excluded_review, ["발주필요금액"], 0), errors="coerce").fillna(0),
            "현지 입고단가": pd.to_numeric(order_review_mod.report_col(excluded_review, [f"현지 입고단가_{currency_code}", "현지 입고단가", "EU 입고단가", "입고단가"], 0), errors="coerce").fillna(0),
            "재고금액": pd.to_numeric(order_review_mod.report_col(excluded_review, ["재고금액"], 0), errors="coerce").fillna(0),
            "재고수량": pd.to_numeric(order_review_mod.report_col(excluded_review, ["재고수량"], 0), errors="coerce").fillna(0),
            "최종 액션": "발주 필요 없음",
            "상태": order_review_mod.report_col(excluded_review, ["상태"], "발주제외"),
            "조치 요약": order_review_mod.report_col(excluded_review, ["조치 요약"], "발주 검토 제외"),
            "근거": order_review_mod.report_col(excluded_review, ["근거", "판단 사유", "제외사유"]),
            "예외 플래그": order_review_mod.report_col(excluded_review, ["예외 플래그"]),
            "예외 사유": order_review_mod.report_col(excluded_review, ["예외 사유"]),
        }
    )
    out["EU 입고단가"] = out["현지 입고단가"]
    out["상태"] = np.where(out["예외 플래그"].astype(str).eq("Y"), "확인필요", "발주제외")
    out["우선 액션"] = out["최종 액션"]
    out["판단 사유"] = out["근거"]

    local_krw_rate = float(
        (settings or {}).get(
            "currency_krw_rate",
            (settings or {}).get("eur_krw_rate", _DEFAULT_EUR_KRW_RATE),
        )
    )

    def _compute_row_amount(r: pd.Series) -> tuple[float, str]:
        qty = float(r.get("제외 수량", 0) or 0)
        if qty <= 0:
            return 0.0, "수량 0"
        unit_local = float(r.get("현지 입고단가", 0) or 0)
        if unit_local > 0:
            return qty * unit_local * local_krw_rate, f"현지 입고단가({currency_code}) × 환율"
        inv_amount = float(r.get("재고금액", 0) or 0)
        inv_qty = float(r.get("재고수량", 0) or 0)
        if inv_amount > 0 and inv_qty > 0:
            unit_krw = inv_amount / inv_qty
            return qty * unit_krw, "재고금액/재고수량"
        return float("nan"), "계산불가: 단가 없음"

    computed = out.apply(lambda r: _compute_row_amount(r), axis=1)
    out["제외 금액_KRW"] = [c[0] for c in computed]
    out["제외 금액 사유"] = [c[1] for c in computed]

    return out.sort_values(["제외유형", "제외 금액_KRW", "제외 수량"], ascending=[True, False, False]).reset_index(drop=True)


def _empty_check_required_df() -> pd.DataFrame:
    return pd.DataFrame(columns=CHECK_REQUIRED_COLUMNS)


def _normalize_check_required_df(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return _empty_check_required_df()
    out = df.copy()
    for col in CHECK_REQUIRED_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out[CHECK_REQUIRED_COLUMNS]


def _strip_check_required_internal_columns(df: pd.DataFrame) -> pd.DataFrame:
    return df.drop(columns=CHECK_REQUIRED_INTERNAL_COLUMNS, errors="ignore")


def price_missing_review_df(review: pd.DataFrame) -> pd.DataFrame:
    if review is None or review.empty:
        return pd.DataFrame()
    flag = order_review_mod.report_col(review, ["단가 미등록 플래그"]).astype(str).str.upper().eq("Y")
    return review[flag].copy()


def build_human_data_issue_df(settings: dict | None = None, context: SessionContext | None = None) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    max_rows_per_issue = int((settings or {}).get("human_data_issue_max_rows_per_type", 100))

    def _to_number_for_input_check(series: pd.Series) -> pd.Series:
        text = series.fillna("").astype(str).str.strip()
        normalized = text.str.replace(",", "", regex=False)
        return pd.to_numeric(normalized, errors="coerce")

    def _source_df(key: str, sample_fn) -> pd.DataFrame:
        try:
            return pd.DataFrame(loaders_mod.get_data_or_sample(key, sample_fn, context))
        except Exception:
            return pd.DataFrame()

    datasets = [
        {
            "source": "현지 재고",
            "key": "eu_stock",
            "sample": loaders_mod.sample_eu_stock,
            "sku": ["상품코드", "SKU", "품목코드", "아이템코드", "itemcode"],
            "name": ["상품명", "제품명", "품목명", "itemname", "Item Name"],
            "brand": ["브랜드", "brand"],
            "qty": ["재고수량", "현재재고", "수량", "qty"],
            "date": [],
            "mode": [],
        },
        {
            "source": "판매내역상세",
            "key": "sales_detail",
            "sample": loaders_mod.sample_sales_detail,
            "sku": ["SKU", "상품코드", "품목코드", "itemcode"],
            "name": ["상품명", "제품명", "품목명", "itemname", "Item Name"],
            "brand": ["브랜드", "brand"],
            "qty": ["수량", "판매수량", "판매 기준기간 판매량", "기준기간 판매량", "최근 판매량", "qty"],
            "date": ["판매일", "판매일자", "출고일", "출고일자", "일자", "date"],
            "mode": [],
        },
        {
            "source": "운송중",
            "key": "shipping",
            "sample": loaders_mod.sample_shipping,
            "sku": ["SKU", "상품코드", "품목코드", "아이템코드", "itemcode"],
            "name": ["상품명", "제품명", "품목명", "itemname", "Item Name"],
            "brand": ["브랜드", "brand"],
            "qty": ["수량", "운송수량", "qty"],
            "date": ["출고일", "선적일", "shipdate"],
            "mode": ["운송수단", "운송 수단", "배송수단", "mode"],
        },
        {
            "source": "미입고 PO",
            "key": "open_po",
            "sample": loaders_mod.sample_open_po,
            "sku": ["SKU", "상품코드", "품목코드", "아이템코드", "itemcode"],
            "name": ["상품명", "제품명", "품목명", "itemname", "Item Name"],
            "brand": ["브랜드", "brand"],
            "qty": ["미입고수량", "미입고 수량", "미입고", "수량", "미량", "openqty"],
            "date": [],
            "mode": [],
            "negative_qty_as_zero": True,
        },
    ]

    def _row(
        *,
        source: str,
        check_type: str,
        reason: str,
        action: str,
        sku: object = "",
        name: object = "",
        brand: object = "",
        qty: object = 0,
        date_value: object = "",
        mode_value: object = "",
        raw_value: object = "",
    ) -> dict[str, object]:
        numeric_qty = _to_number_for_input_check(pd.Series([qty])).fillna(0).iloc[0]
        return {
            "확인 구분": check_type,
            "상품코드": sku,
            "상품명": name,
            "브랜드": brand,
            "발견 원본": source,
            "판매수량": numeric_qty if source == "판매내역상세" else 0,
            "운송중 수량": numeric_qty if source == "운송중" else 0,
            "미입고 수량": numeric_qty if source == "미입고 PO" else 0,
            "확인대상 수량": numeric_qty,
            "출고일": date_value if source == "운송중" else "",
            "예상 입고일": "",
            "운송수단": mode_value,
            "원본값": raw_value,
            "확인필요 사유": reason,
            "권장 확인 액션": action,
        }

    for spec in datasets:
        df = _source_df(spec["key"], spec["sample"])
        if df.empty:
            continue

        source = str(spec["source"])
        sku_col = kpi_mod.find_column(df, spec["sku"])
        name_col = kpi_mod.find_column(df, spec["name"])
        brand_col = kpi_mod.find_column(df, spec["brand"])
        qty_col = kpi_mod.find_column(df, spec["qty"])
        date_col = kpi_mod.find_column(df, spec["date"]) if spec["date"] else None
        mode_col = kpi_mod.find_column(df, spec["mode"]) if spec["mode"] else None

        if sku_col is None:
            rows.append(
                _row(
                    source=source,
                    check_type="필수 컬럼 확인",
                    reason="상품코드/SKU 컬럼을 찾을 수 없습니다.",
                    action="원본 파일의 상품코드 컬럼명을 확인",
                )
            )
            continue

        raw_sku = df[sku_col].fillna("").astype(str)
        sku = preprocess_mod.clean_identifier_series(df[sku_col])
        name = df[name_col].fillna("").astype(str) if name_col else pd.Series([""] * len(df), index=df.index)
        brand = df[brand_col].fillna("").astype(str) if brand_col else pd.Series([""] * len(df), index=df.index)
        qty_raw = df[qty_col] if qty_col else pd.Series([0] * len(df), index=df.index)
        qty = _to_number_for_input_check(qty_raw)
        valid_product = product_sku_mask(sku)

        empty_sku = pd.Series(False, index=df.index) & valid_product
        for idx in df[empty_sku].index[:max_rows_per_issue]:
            rows.append(
                _row(
                    source=source,
                    check_type="상품코드 공란",
                    reason="상품코드가 비어 있어 발주/재고/판매 데이터 연결이 불가능합니다.",
                    action="원본 파일에서 상품코드 입력",
                    sku="",
                    name=name.loc[idx],
                    brand=brand.loc[idx],
                    qty=qty_raw.loc[idx],
                    raw_value=raw_sku.loc[idx],
                )
            )

        stripped_sku = raw_sku.str.strip()
        whitespace_sku = pd.Series(False, index=df.index) & valid_product
        for idx in df[whitespace_sku].index[:max_rows_per_issue]:
            rows.append(
                _row(
                    source=source,
                    check_type="상품코드 공백",
                    reason="상품코드 앞뒤에 불필요한 공백이 있습니다.",
                    action="원본 파일에서 상품코드 앞뒤 공백 제거",
                    sku=sku.loc[idx],
                    name=name.loc[idx],
                    brand=brand.loc[idx],
                    qty=qty_raw.loc[idx],
                    raw_value=raw_sku.loc[idx],
                )
            )

        suspicious_sku = pd.Series(False, index=df.index) & valid_product
        for idx in df[suspicious_sku].index[:max_rows_per_issue]:
            rows.append(
                _row(
                    source=source,
                    check_type="상품코드 문자 확인",
                    reason="상품코드에 일반적이지 않은 문자나 중간 공백이 포함되어 있습니다.",
                    action="원본 파일에서 상품코드 오타 여부 확인",
                    sku=sku.loc[idx],
                    name=name.loc[idx],
                    brand=brand.loc[idx],
                    qty=qty_raw.loc[idx],
                    raw_value=raw_sku.loc[idx],
                )
            )

        if qty_col is None:
            rows.append(
                _row(
                    source=source,
                    check_type="필수 컬럼 확인",
                    reason="수량 컬럼을 찾을 수 없습니다.",
                    action="원본 파일의 수량 컬럼명을 확인",
                )
            )
        else:
            bad_qty = qty.isna() & qty_raw.fillna("").astype(str).str.strip().ne("")
            negative_qty = pd.Series(False, index=df.index) if spec.get("negative_qty_as_zero") else qty.fillna(0).lt(0)
            for idx in df[(bad_qty | negative_qty) & valid_product].index[:max_rows_per_issue]:
                rows.append(
                    _row(
                        source=source,
                        check_type="수량 입력 확인",
                        reason="수량이 숫자가 아니거나 음수로 입력되어 있습니다.",
                        action="원본 파일에서 수량 입력값 확인",
                        sku=sku.loc[idx],
                        name=name.loc[idx],
                        brand=brand.loc[idx],
                        qty=qty_raw.loc[idx],
                        raw_value=qty_raw.loc[idx],
                    )
                )

        if date_col is not None and source == "운송중":
            raw_date = df[date_col].fillna("").astype(str).str.strip()
            parsed_date = kpi_mod.parse_date_series(df[date_col])
            bad_date = raw_date.ne("") & parsed_date.isna()
            for idx in df[bad_date & valid_product].index[:max_rows_per_issue]:
                rows.append(
                    _row(
                        source=source,
                        check_type="날짜 입력 확인",
                        reason="날짜 형식을 인식할 수 없습니다.",
                        action="원본 파일에서 날짜 형식을 yyyy-mm-dd 등으로 수정",
                        sku=sku.loc[idx],
                        name=name.loc[idx],
                        brand=brand.loc[idx],
                        qty=qty_raw.loc[idx],
                        date_value=raw_date.loc[idx],
                        raw_value=raw_date.loc[idx],
                    )
                )

        if bool((settings or {}).get("enable_deep_name_consistency_check", False)) and name_col is not None:
            names_by_sku = (
                pd.DataFrame({"상품코드": sku, "상품명": name.astype(str).str.strip()})
                .loc[lambda frame: frame["상품코드"].ne("") & frame["상품명"].ne("")]
                .groupby("상품코드")["상품명"]
                .nunique()
            )
            inconsistent_codes = set(names_by_sku[names_by_sku > 1].index.astype(str))
            for code in sorted(inconsistent_codes):
                first_idx = sku[sku.astype(str).eq(code)].index[0]
                rows.append(
                    _row(
                        source=source,
                        check_type="상품명 입력 확인",
                        reason="같은 상품코드에 서로 다른 상품명이 입력되어 있습니다.",
                        action="원본 파일에서 상품명 오타 여부 확인",
                        sku=code,
                        name=name.loc[first_idx],
                        brand=brand.loc[first_idx],
                        qty=0,
                    )
                )

    if not rows:
        return pd.DataFrame(columns=CHECK_REQUIRED_COLUMNS)
    return _normalize_check_required_df(pd.DataFrame(rows))


def build_check_required_sheet_df(
    eta_unmatched_df: pd.DataFrame,
    unrecognized_transport_df: pd.DataFrame,
    excluded_df: pd.DataFrame,
    master_unregistered_df: pd.DataFrame,
    price_missing_df: pd.DataFrame | None = None,
    human_data_issue_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    sections: list[pd.DataFrame] = []

    if human_data_issue_df is not None and not human_data_issue_df.empty:
        sections.append(_normalize_check_required_df(human_data_issue_df))

    if not unrecognized_transport_df.empty:
        sections.append(
            _normalize_check_required_df(
                pd.DataFrame(
                    {
                        "확인 구분": "운송정보 확인",
                        "상품코드": order_review_mod.report_col(unrecognized_transport_df, ["상품코드"]),
                        "상품명": order_review_mod.report_col(unrecognized_transport_df, ["상품명"]),
                        "브랜드": order_review_mod.report_col(unrecognized_transport_df, ["브랜드"]),
                        "발견 원본": "운송중",
                        "판매수량": 0,
                        "운송중 수량": pd.to_numeric(order_review_mod.report_col(unrecognized_transport_df, ["수량"], 0), errors="coerce").fillna(0),
                        "미입고 수량": 0,
                        "확인대상 수량": pd.to_numeric(order_review_mod.report_col(unrecognized_transport_df, ["수량"], 0), errors="coerce").fillna(0),
                        "출고일": order_review_mod.report_col(unrecognized_transport_df, ["출고일"]),
                        "예상 입고일": order_review_mod.report_col(unrecognized_transport_df, ["예상 입고일"]),
                        "운송수단": order_review_mod.report_col(
                            unrecognized_transport_df,
                            ["표준 운송수단"],
                            "확인필요",
                        ),
                        "원본값": order_review_mod.report_col(unrecognized_transport_df, ["원본값"]),
                        "확인필요 사유": order_review_mod.report_col(
                            unrecognized_transport_df,
                            ["확인필요 사유"],
                            "운송수단/출고일/ETA 계산 기준 확인 필요",
                        ),
                        "권장 확인 액션": order_review_mod.report_col(
                            unrecognized_transport_df,
                            ["권장 확인 액션"],
                            "선택 법인의 표준 운송수단 및 출고일 확인",
                        ),
                    }
                )
            )
        )

    if not master_unregistered_df.empty:
        sku = order_review_mod.report_col(master_unregistered_df, ["상품코드"])
        name = order_review_mod.report_col(master_unregistered_df, ["상품명"])
        brand = order_review_mod.report_col(master_unregistered_df, ["브랜드"])
        shipping_qty = pd.to_numeric(order_review_mod.report_col(master_unregistered_df, ["운송중 수량"], 0), errors="coerce").fillna(0)
        open_po_qty = pd.to_numeric(order_review_mod.report_col(master_unregistered_df, ["미입고 수량"], 0), errors="coerce").fillna(0)
        shipping_date = order_review_mod.report_col(master_unregistered_df, ["출고일"])
        shipping_eta = order_review_mod.report_col(master_unregistered_df, ["예상 입고일", "ETA"])
        shipping_mode = order_review_mod.report_col(master_unregistered_df, ["운송수단"])
        open_po_qty = open_po_qty.clip(lower=0)
        promo_goodie_bag_mask = (
            brand.fillna("").astype(str).str.strip().str.casefold().eq("stylekorean uk limited")
            & name.fillna("").astype(str).str.strip().str.casefold().eq("goodie bag")
        )
        shipping_qty = shipping_qty.mask(promo_goodie_bag_mask, 0)
        open_po_qty = open_po_qty.mask(promo_goodie_bag_mask, 0)
        sku_text = sku.fillna("").astype(str).str.strip()
        open_po_reason = (
            "대상 SKU: "
            + sku_text
            + " / 미입고현황에는 있으나 현지 재고 데이터에는 없는 SKU입니다. 신규 SKU이거나 현지 재고 이력이 없는 상품일 수 있으므로 확인이 필요합니다."
        )
        shipping_reason = (
            "대상 SKU: "
            + sku_text
            + " / 운송상세내역에는 있으나 현지 재고 데이터에는 없는 SKU입니다. 현재 이동 중인 신규 SKU 또는 현지 재고 이력이 없는 상품일 수 있으므로 확인이 필요합니다."
        )
        eta_sections: list[pd.DataFrame] = []
        if open_po_qty.gt(0).any():
            eta_sections.append(
                pd.DataFrame(
                    {
                        "확인 구분": "미입고 ETA",
                        "상품코드": sku,
                        "상품명": name,
                        "브랜드": brand,
                        "발견 원본": "미입고현황",
                        "판매수량": 0,
                        "운송중 수량": 0,
                        "미입고 수량": open_po_qty,
                        "확인대상 수량": open_po_qty,
                        "출고일": "",
                        "예상 입고일": "",
                        "운송수단": "",
                        "확인필요 사유": "미입고 현황에는 존재하지만, 현지 재고 데이터에는 아직 존재하지 않는 SKU입니다. 신규 SKU이거나 현지 재고 이력이 없는 상품일 수 있으므로 확인이 필요합니다.",
                        "권장 확인 액션": "신규 SKU 여부, 현지 재고 이력/마스터 등록 여부 확인",
                    }
                ).loc[open_po_qty.gt(0)]
            )
        if shipping_qty.gt(0).any():
            eta_sections.append(
                pd.DataFrame(
                    {
                        "확인 구분": "운송 ETA",
                        "상품코드": sku,
                        "상품명": name,
                        "브랜드": brand,
                        "발견 원본": "운송중",
                        "판매수량": 0,
                        "운송중 수량": shipping_qty,
                        "미입고 수량": 0,
                        "확인대상 수량": shipping_qty,
                        "출고일": shipping_date,
                        "예상 입고일": shipping_eta,
                        "운송수단": shipping_mode,
                        "확인필요 사유": "운송 중 데이터에는 존재하지만, 현지 재고 데이터에는 아직 존재하지 않는 SKU입니다. 현재 이동 중인 신규 SKU 또는 현지 재고 이력이 없는 상품일 수 있으므로 확인이 필요합니다.",
                        "권장 확인 액션": "신규 SKU 여부, 현지 재고 이력/마스터 등록 여부 확인",
                    }
                ).loc[shipping_qty.gt(0)]
            )
        if eta_sections:
            sections.append(_normalize_check_required_df(pd.concat(eta_sections, ignore_index=True)))

    if not sections:
        return pd.DataFrame(columns=CHECK_REQUIRED_DISPLAY_COLUMNS)
    out = pd.concat(sections, ignore_index=True)
    check_col = CHECK_REQUIRED_COLUMNS[0]
    sku_col = CHECK_REQUIRED_COLUMNS[1]
    reason_col = CHECK_REQUIRED_COLUMNS[13]
    eta_sku_text = out[sku_col].fillna("").astype(str).str.strip()
    open_po_eta_mask = out[check_col].astype(str).eq("미입고 ETA")
    shipping_eta_mask = out[check_col].astype(str).eq("운송 ETA")
    out.loc[open_po_eta_mask, reason_col] = (
        "대상 SKU: "
        + eta_sku_text.loc[open_po_eta_mask]
        + " / 미입고현황에는 있으나 현지 재고 데이터에는 없는 SKU입니다. 신규 SKU이거나 현지 재고 이력이 없는 상품일 수 있으므로 확인이 필요합니다."
    )
    out.loc[shipping_eta_mask, reason_col] = (
        "대상 SKU: "
        + eta_sku_text.loc[shipping_eta_mask]
        + " / 운송상세내역에는 있으나 현지 재고 데이터에는 없는 SKU입니다. 현재 이동 중인 신규 SKU 또는 현지 재고 이력이 없는 상품일 수 있으므로 확인이 필요합니다."
    )
    out["확인필요 사유"] = (
        out["확인필요 사유"]
        .astype(str)
        .str.replace("현지 재고 마스터 미등록", "현지 재고 목록 미존재", regex=False)
        .str.replace(", ", "/", regex=False)
        .str.replace(" 원본에 존재", " 내역에 존재", regex=False)
        .str.replace("판매내역 내역에 존재", "판매내역에 존재", regex=False)
    )

    def _row_to_keys(r: pd.Series) -> tuple[str, str]:
        row = pd.Series({
            "상품코드": r.get("상품코드", ""),
            "SKU": r.get("상품코드", ""),
            "거래처명": r.get("브랜드", ""),
            "원본 거래처명": r.get("브랜드", ""),
            "Invoice 번호": r.get("Invoice 번호", ""),
            "수량": r.get("확인대상 수량", r.get("수량", 0)),
            "금액": r.get("금액", 0),
            "Invoice 비고": r.get("원본값", r.get("확인필요 사유", "")),
            "원본값": r.get("원본값", ""),
            "거래유형": r.get("거래유형", ""),
        })
        try:
            return loaders_mod.make_instance_key(row), loaders_mod.make_pattern_key(row)
        except Exception:
            return "", ""

    keys = out.apply(_row_to_keys, axis=1)
    out["instance_key"] = [k[0] for k in keys]
    out["pattern_key"] = [k[1] for k in keys]

    decisions = loaders_mod.load_exception_decisions()
    inst_map = {r.instance_key: r for _, r in decisions[decisions["instance_key"].astype(bool)].iterrows()} if not decisions.empty else {}
    pat_map = {r.pattern_key: r for _, r in decisions[decisions["pattern_key"].astype(bool)].iterrows()} if not decisions.empty else {}

    def _apply_decision(idx: int, r: pd.Series) -> pd.Series:
        inst = r.get("instance_key", "")
        pat = r.get("pattern_key", "")
        applied = r.copy()
        if inst and inst in inst_map:
            dec = inst_map[inst]
        elif pat and pat in pat_map:
            dec = pat_map[pat]
        else:
            dec = None
        if dec is not None:
            applied["담당자판단"] = dec.get("담당자 판단", "")
            applied["판단일시"] = dec.get("판단일시", "")
            applied["담당자"] = dec.get("담당자", "") if "담당자" in dec.index else ""
        else:
            applied["담당자판단"] = ""
            applied["판단일시"] = ""
            applied["담당자"] = ""
        return applied

    out = out.apply(lambda r: _apply_decision(r.name, r), axis=1)
    out = _strip_check_required_internal_columns(out)
    out = out.drop(columns=["담당자 판단", "판단 사유"], errors="ignore")
    for col in CHECK_REQUIRED_OUTPUT_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    total_like_name = out["상품명"].astype(str).str.strip().str.lower().isin({"total", "subtotal", "grand total", "합계", "총계", "소계"})
    qty = pd.to_numeric(out["확인대상 수량"], errors="coerce").fillna(0)
    check_type = out["확인 구분"].astype(str)
    reason = out["확인필요 사유"].astype(str)
    sku_output_col = CHECK_REQUIRED_OUTPUT_COLUMNS[1]
    check_type_output_col = CHECK_REQUIRED_OUTPUT_COLUMNS[0]
    valid_sku = product_sku_mask(out[sku_output_col])
    eta_reference_issue = out[check_type_output_col].astype(str).str.contains("ETA", na=False)
    required_column_issue = out[check_type_output_col].astype(str).eq("필수 컬럼 확인")
    empty_eta_noise = (
        check_type.eq("입고예정 확인")
        & reason.str.contains("입고 예정", na=False)
        & qty.eq(0)
    )
    out["확인 내용"] = out["확인필요 사유"].astype(str).str.strip()
    raw_value = out["원본값"].astype(str).str.strip()
    out["원본 비고"] = raw_value.where(raw_value.ne("nan"), "")
    out["수량"] = out["확인대상 수량"]
    for col in CHECK_REQUIRED_DISPLAY_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out.loc[((valid_sku | required_column_issue | eta_reference_issue) & ~total_like_name & ~empty_eta_noise), CHECK_REQUIRED_DISPLAY_COLUMNS].reset_index(drop=True)


def format_check_required_sheet(ws, settings: dict | None = None, header_row: int = 8) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    settings = settings or {}
    data_start_row = header_row + 1
    dark_fill = PatternFill("solid", fgColor="4F6228")
    light_fill = PatternFill("solid", fgColor="E2F0D9")
    note_fill = PatternFill("solid", fgColor="F3F8EC")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    no_fill = PatternFill(fill_type=None)
    white_font = Font(color="FFFFFF", bold=True)
    green_font = Font(color="4F6228", bold=True)
    red_font = Font(color="C00000", bold=True)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = {str(ws.cell(header_row, col).value or ""): col for col in range(1, ws.max_column + 1)}
    decision_col = headers.get("담당자판단")
    decision_date_col = headers.get("판단일시")
    owner_col = headers.get("담당자")

    for row_idx in range(1, 3):
        for col_idx in range(9, 11):
            cell = ws.cell(row_idx, col_idx)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.font = Font()
            cell.alignment = Alignment()

    ws.merge_cells("A3:E3")
    title_cell = ws["A3"]
    title_cell.value = "▼ 원본 데이터 확인필요 | 입력 오류 점검"
    title_cell.fill = dark_fill
    title_cell.font = white_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A4:E4")
    note_cell = ws["A4"]
    note_cell.value = "★ 법인별 운송수단·ETA와 입력 이상값을 표시합니다. 담당자판단·판단일시·담당자는 입력 영역입니다."
    note_cell.fill = note_fill
    note_cell.font = Font(color="4F6228", italic=True, size=9)
    note_cell.alignment = Alignment(horizontal="left", vertical="center")

    method_labels = [
        str(method.get("display_name") or method.get("label") or "").strip()
        for method in list(settings.get("lead_time_methods") or [])
        if str(method.get("display_name") or method.get("label") or "").strip()
    ]
    allowed_transport_text = "/".join(dict.fromkeys(method_labels)) or "해운/항공/철송/트럭"
    guide_rows = [
        ("확인 처리 기준", "권장 입력"),
        ("미입고 ETA", "미입고에는 있으나 현지 재고 이력 없는 SKU 확인"),
        ("운송 ETA", "운송중에는 있으나 현지 재고 이력 없는 SKU 확인"),
        ("입력 오류", "필수 컬럼/수량/날짜 이상값 확인"),
        ("운송수단 오류", f"{allowed_transport_text} 미인식값 확인"),
    ]
    guide_end_col = max(ws.max_column, len(CHECK_REQUIRED_DISPLAY_COLUMNS))
    guide_start_col = max(9, guide_end_col - 6)
    guide_text_start_col = guide_start_col + 1
    for row_offset, (left, right) in enumerate(guide_rows, start=3):
        ws.merge_cells(
            start_row=row_offset,
            start_column=guide_text_start_col,
            end_row=row_offset,
            end_column=guide_end_col,
        )
        ws.cell(row_offset, guide_start_col, left)
        ws.cell(row_offset, guide_text_start_col, right)
        for col_idx in range(guide_start_col, guide_end_col + 1):
            cell = ws.cell(row_offset, col_idx)
            cell.fill = dark_fill if row_offset == 3 else light_fill
            cell.font = white_font if row_offset == 3 else (green_font if col_idx == guide_start_col else Font(color="4F6228"))
            cell.alignment = Alignment(
                horizontal="center" if col_idx == guide_start_col else "left",
                vertical="center",
                wrap_text=col_idx >= guide_text_start_col,
            )
            cell.border = border

    editable_headers = {"담당자판단", "판단일시", "담당자"}
    text_headers = {"확인 구분", "상품코드", "상품명", "브랜드", "발견 원본", "운송수단", "원본 비고", "확인 내용", "권장 확인 액션", "담당자판단", "담당자"}
    date_headers = {"출고일", "예상 입고일", "판단일시"}
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(header_row, col_idx).value or "")
        cell = ws.cell(header_row, col_idx)
        cell.fill = orange_fill if header in editable_headers else light_fill
        cell.font = red_font if header in editable_headers else green_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_idx in range(data_start_row, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(header_row, col_idx).value or "")
            cell = ws.cell(row_idx, col_idx)
            cell.border = border
            long_text_headers = {"상품명", "원본 비고", "확인 내용", "권장 확인 액션"}
            cell.alignment = Alignment(
                horizontal="left" if header in long_text_headers else "center",
                vertical="center",
                wrap_text=header in long_text_headers,
            )
            cell.fill = no_fill
            if header in text_headers:
                cell.number_format = "@"
            elif header in date_headers:
                cell.number_format = "yyyy-mm-dd"
            else:
                cell.number_format = "#,##0"

    has_data_rows = ws.max_row >= data_start_row and any(
        any(ws.cell(row_idx, col_idx).value not in (None, "") for col_idx in range(1, ws.max_column + 1))
        for row_idx in range(data_start_row, ws.max_row + 1)
    )
    if not has_data_rows:
        message_row = data_start_row
        merge_end_col = min(ws.max_column, 14)
        ws.merge_cells(start_row=message_row, start_column=1, end_row=message_row, end_column=merge_end_col)
        message_cell = ws.cell(message_row, 1, "확인필요 데이터 없음")
        message_cell.fill = note_fill
        message_cell.font = Font(color="4F6228", bold=True)
        message_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in range(1, merge_end_col + 1):
            ws.cell(message_row, col_idx).border = border
        ws.row_dimensions[message_row].height = 24

    if decision_col:
        decision_letter = get_column_letter(decision_col)
        validation = DataValidation(
            type="list",
            formula1='"수정 완료,원본 확인 완료,수정 불필요,보류"',
            allow_blank=True,
        )
        validation.promptTitle = "담당자판단"
        validation.prompt = "처리 결과를 선택하거나 직접 입력하세요."
        ws.add_data_validation(validation)
        if ws.max_row >= data_start_row:
            validation.add(f"{decision_letter}{data_start_row}:{decision_letter}{ws.max_row}")
    if decision_date_col:
        for row_idx in range(data_start_row, ws.max_row + 1):
            ws.cell(row_idx, decision_date_col).number_format = "yyyy-mm-dd"

    widths = {
        "확인 구분": 16,
        "상품코드": 18,
        "상품명": 42,
        "브랜드": 14,
        "발견 원본": 16,
        "판매수량": 12,
        "운송중 수량": 12,
        "미입고 수량": 12,
        "수량": 12,
        "출고일": 14,
        "예상 입고일": 14,
        "운송수단": 14,
        "원본 비고": 38,
        "확인 내용": 48,
        "권장 확인 액션": 36,
        "담당자판단": 16,
        "판단일시": 14,
        "담당자": 12,
    }
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(header_row, col_idx).value or "")
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, 14)
    ws.column_dimensions["I"].width = max(ws.column_dimensions["I"].width or 0, 18)
    for col_letter in ("J", "K", "L", "M", "N", "O", "P"):
        ws.column_dimensions[col_letter].width = max(ws.column_dimensions[col_letter].width or 0, 16)
    for row_idx, height in {3: 20, 4: 28, 5: 30, 6: 28, 7: 28, header_row: 36}.items():
        ws.row_dimensions[row_idx].height = height
    if has_data_rows:
        for row_idx in range(data_start_row, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 0, 48)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.freeze_panes = f"D{data_start_row}"


def order_action_priority(action: object) -> int:
    text = str(action)
    if text == "발주 필요":
        return 1
    if text == "발주 필요 없음":
        return 90
    if text == "발주 불필요":
        return 90
    if "ETA 지연 위험" in text:
        return 2
    if "본사 EU창고" in text:
        return 3
    if "본사 일부 이동" in text:
        return 3
    if "운송 도착 후 검토" in text:
        return 3
    if "미입고 반영 후 모니터링" in text:
        return 4
    if "운송중 도착 대기" in text:
        return 5
    if "발주 불필요" in text:
        return 90
    if "판매없음" in text or "발주 제외" in text:
        return 99
    return 50


def apply_low_unit_price_review_label(
    df: pd.DataFrame,
    *,
    price_columns: list[str],
    qty_columns: list[str],
) -> tuple[pd.DataFrame, int]:
    out = df.copy()
    if "우선 액션" not in out.columns:
        return out, 0
    if "판단 사유" not in out.columns:
        out["판단 사유"] = ""

    price = pd.Series(0, index=out.index, dtype="float64")
    for col in price_columns:
        if col in out.columns:
            price = pd.to_numeric(out[col], errors="coerce").fillna(0)
            break

    order_qty = pd.Series(0, index=out.index, dtype="float64")
    for col in qty_columns:
        if col in out.columns:
            order_qty = pd.to_numeric(out[col], errors="coerce").fillna(0)
            break

    action = out["우선 액션"].astype(str)
    protected_action = action.str.contains("FOC|단가 확인|제외|저단가", na=False)
    low_unit_price = (
        (price > 0)
        & (price <= 0.5)
        & (order_qty > 0)
        & ~protected_action
    )
    count = int(low_unit_price.sum())
    if count:
        low_price_note = "저단가 확인: 단가 0.5 EUR 이하(무상/최소과세금액/임시단가 가능성)"
        existing_reason = out.loc[low_unit_price, "판단 사유"].fillna("").astype(str).str.strip()
        needs_note = ~existing_reason.str.contains("저단가 확인|단가 0.5 EUR 이하", na=False)
        target_index = existing_reason[needs_note].index
        out.loc[target_index, "판단 사유"] = np.where(
            existing_reason.loc[target_index].eq(""),
            low_price_note,
            existing_reason.loc[target_index] + " / " + low_price_note,
        )
    return out, count


def build_order_review_display_df(review: pd.DataFrame, only_needed: bool, show_detail: bool) -> pd.DataFrame:
    display = build_order_review_report_df(review)
    detail = pd.DataFrame(
        {
            "최근 3개월 판매수량": pd.to_numeric(order_review_mod.report_col(review, ["최근 3개월 판매수량"], 0), errors="coerce").fillna(0),
            "월평균 판매수량": pd.to_numeric(order_review_mod.report_col(review, ["월평균 판매수량"], 0), errors="coerce").fillna(0),
            "안전재고 목표수량": pd.to_numeric(order_review_mod.report_col(review, ["안전재고 목표수량"], 0), errors="coerce").fillna(0),
            "부족수량 산정 기준 보유량": pd.to_numeric(order_review_mod.report_col(review, ["부족수량 산정 기준 보유량"], 0), errors="coerce").fillna(0),
            "적용 리드타임": pd.to_numeric(order_review_mod.report_col(review, ["적용 리드타임"], 0), errors="coerce").fillna(0),
            "예상 도착 가능일": order_review_mod.report_col(review, ["예상 도착 가능일"]),
        }
    )
    display = pd.concat([display, detail], axis=1)
    display["현재 가용재고"] = display["EU 현지 재고"]
    display["운송중 수량"] = display["운송중"]
    display["월평균 판매수량 정렬"] = display["월평균 판매수량"]

    if only_needed:
        action = order_review_mod.report_col(display, ["우선 액션", "최종 액션"]).astype(str)
        status = order_review_mod.report_col(display, ["상태"], "정상").astype(str)
        pre_arrival_risk = order_review_mod.report_col(display, ["입고 전 결품 위험 여부"], "N").astype(str).eq("Y")
        display = display[order_review_mod.order_needed_action_mask(action) | pre_arrival_risk | status.isin(["본사이동", "운송대기", "확인필요"])]

    display = order_review_mod.sort_order_review_output(display).reset_index(drop=True)
    display, _ = apply_low_unit_price_review_label(
        display,
        price_columns=["EU 입고단가(EUR)", "EU 입고단가"],
        qty_columns=["발주필요수량", "추가 발주 필요 수량"],
    )
    display["No"] = range(1, len(display) + 1)

    base_cols = [
        "브랜드", "상품명", "상품코드", "최종 액션", "상태", "조치 요약", "일평균 판매수량", "현재 가용재고",
        "EU 현지 커버일수", "예상 소진일", "운송중 수량", "입고 전 예상 결품량", "첫 입고 예정일",
        "쇼티지 예상 일수", "본사 이동 수량", "추가 발주 필요 수량", "발주필요수량",
        "긴급 보충 필요 수량", "입고 전 결품 위험 여부", "권장 긴급 액션", "추가 발주 필요 금액", "운송 검토안",
    ]
    detail_cols = [
        "구분", "우선 액션", "근거", "예외 플래그", "예외 사유", "운송중여부", "미입고여부",
        "커버개월", "부족수량", "부족금액_KRW",
        "EU 현지 재고", "본사 EU창고", "운송중", "미입고", "바코드", "판단 사유",
        "기준_3M_판매수량", "PA_CA_3M_판매수량",
        "판단메모", "최근 3개월 판매수량", "월평균 판매수량", "고갈 예상일",
        "안전재고 목표수량", "부족수량 산정 기준 보유량", "EU 입고단가(EUR)", "부족금액_EUR", "적용 리드타임", "예상 도착 가능일",
    ]
    cols = base_cols + (detail_cols if show_detail else [])
    return display[[col for col in cols if col in display.columns]]
