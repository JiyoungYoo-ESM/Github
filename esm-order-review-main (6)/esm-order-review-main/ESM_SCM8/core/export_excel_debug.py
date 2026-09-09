from __future__ import annotations

from numbers import Number

import pandas as pd

from core.common import _ORDER_REVIEW_EXPORT_VERSION, _ORDER_REVIEW_LOGIC_VERSION, REPORT_TABLE_HEADER_FILL
from core.export_excel_util import append_df, apply_report_column_formats, autosize_columns, style_table
from core.session import SessionContext, ensure_session_context
import os
from datetime import datetime
from core.common import PROJECT_ROOT, TRANSPORT_REVIEW_REQUIRED, format_months
from core import (
    eta as eta_mod,
    inventory as inventory_mod,
    kpi as kpi_mod,
    loaders as loaders_mod,
    order_review as order_review_mod,
    preprocess as preprocess_mod,
    transport as transport_mod,
    validation as validation_mod,
)

def apply_workbook_thousands_number_formats(wb) -> None:
    """Normalize numeric cells that still use comma-less Excel formats."""
    comma_less_formats = {"General", "0", "0.0"}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, Number) or isinstance(value, bool):
                    continue
                if cell.is_date or cell.number_format not in comma_less_formats:
                    continue
                if cell.number_format == "0.0":
                    cell.number_format = "#,##0.0"
                elif isinstance(value, float) and not float(value).is_integer():
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"


def build_validation_log_df(settings: dict, review: pd.DataFrame, signal_counts: dict[str, int] | None = None, context: SessionContext | None = None) -> pd.DataFrame:
    rows: list[dict[str, object]] = []

    def add(level: str, item: str, detail: str, count: int | float | str = "") -> None:
        rows.append({"구분": level, "항목": item, "내용": detail, "건수/값": count})

    stock = inventory_mod.get_eu_stock()
    hq_stock = inventory_mod.get_hq_eu_stock()
    sales = loaders_mod.get_data_or_sample("sales_detail", loaders_mod.sample_sales_detail, context)
    shipping = loaders_mod.get_data_or_sample("shipping", loaders_mod.sample_shipping, context)
    open_po_raw = loaders_mod.get_data_or_sample("open_po", loaders_mod.sample_open_po, context)
    required_checks = [
        ("EU 현지 재고", stock, [["상품코드", "SKU", "품목코드"], ["재고수량", "현재재고"], ["Hold수량", "홀드수량"]]),
        ("본사 EU창고 재고", hq_stock, [["상품코드", "SKU", "품목코드"]]),
        ("판매내역상세(검증용)", sales, [["상품코드", "SKU", "품목코드"], ["수량", "판매수량", "판매 기준기간 판매량", "기준기간 판매량", "최근 판매량", "qty"]]),
        ("운송중", shipping, [["상품코드", "SKU", "품목코드"], ["수량", "운송수량", "qty"], ["출고일", "선적일", "shipdate"], ["운송수단", "운송 수단", "배송수단", "mode"]]),
        ("미입고 PO", open_po_raw, [["상품코드", "SKU", "품목코드"]]),
    ]

    for name, df, required in required_checks:
        missing = [candidates[0] for candidates in required if kpi_mod.find_column(df, candidates) is None]
        if missing:
            add("확인필요", f"{name} 컬럼", f"인식되지 않은 컬럼: {', '.join(missing)}", len(missing))

    pa_ca_col = settings.get("pa_ca_sales_column_override") or preprocess_mod.detect_pa_ca_sales_column(stock)
    if pa_ca_col is None:
        add("확인필요", "PA+CA 판매수량 컬럼", f"자동 탐지 실패. 사용 가능한 컬럼: {', '.join(map(str, stock.columns))}")
    else:
        add("정보", "판매수량 계산 기준", f"판매수량 계산 기준: 재고 파일 PA+CA 3개월 판매수량 (컬럼: {pa_ca_col})")
    add("정보", "판매내역상세 검증", "발주 검토 산출물에서는 PA+CA 기준 단일화를 위해 판매내역상세 검증값을 표시하지 않습니다.")

    if {"재고수량", "Hold수량", "가용수량"}.issubset(set(stock.columns)):
        expected = kpi_mod.to_number_series(stock["재고수량"]) - kpi_mod.to_number_series(stock["Hold수량"])
        actual = kpi_mod.to_number_series(stock["가용수량"])
        mismatch = (expected - actual).abs() > 1
        if mismatch.any():
            add("확인필요", "가용수량 검증", "가용수량과 재고수량-Hold수량이 다른 행이 있습니다.", int(mismatch.sum()))

    po_qty_col = kpi_mod.find_column(open_po_raw, ["미입고수량", "미입고 수량", "미입고", "수량", "잔량", "openqty"])
    if po_qty_col is not None:
        negative_count = int((kpi_mod.to_number_series(open_po_raw[po_qty_col]) < 0).sum())
        if negative_count:
            add("보정", "미입고 PO 음수", "산출표에서는 0으로 보정했습니다.", negative_count)
    else:
        po_col = kpi_mod.find_column(open_po_raw, ["PO 수량", "발주수량", "poqty", "orderqty"])
        received_col = kpi_mod.find_column(open_po_raw, ["입고 수량", "입고수량", "receivedqty"])
        if po_col is not None and received_col is not None:
            negative_count = int(((kpi_mod.to_number_series(open_po_raw[po_col]) - kpi_mod.to_number_series(open_po_raw[received_col])) < 0).sum())
            if negative_count:
                add("보정", "미입고 PO 음수", "PO 수량-입고 수량이 음수인 행은 산출표에서 0으로 보정했습니다.", negative_count)

    prepared_shipping = transport_mod.get_shipping(settings, context)
    unrecognized_transport_df = transport_mod.unrecognized_transport_report_df(settings, context)
    if not unrecognized_transport_df.empty:
        raw_values = ", ".join(
            sorted(
                value
                for value in unrecognized_transport_df["원본값"].astype(str).str.strip().drop_duplicates().tolist()
                if value
            )
        )
        add(
            "확인필요",
            TRANSPORT_REVIEW_REQUIRED,
            f"표준 운송수단(해운/항공/철송/트럭)이 원본값에 포함되지 않은 경우 ETA 계산 및 운송수단별 집계에서 제외합니다. 원본값: {raw_values or '(빈값)'}",
            len(unrecognized_transport_df),
        )
    estimated_count = int(prepared_shipping.get("ETA 구분", pd.Series(dtype=str)).astype(str).str.contains("계산ETA", na=False).sum())
    if estimated_count:
        add("정보", "ETA 산정", "출고일 + 운송 L/T 기준으로 ETA를 산정한 행입니다.", estimated_count)
    eta_unmatched_df = eta_mod.build_eta_unmatched_shipping_report_df(review, settings, shipping_df=prepared_shipping)
    add(
        "정보",
        "ETA 표시 기준",
        "입고 예정은 있으나 발주 검토에는 없는 SKU는 발주필요수량 산식에는 자동 반영하지 않습니다.",
        len(eta_unmatched_df),
    )
    zero_sales = int((pd.to_numeric(order_review_mod.report_col(review, ["월평균 판매수량"], 0), errors="coerce").fillna(0) <= 0).sum())
    if zero_sales:
        add("정보", "판매속도 0", "월평균 판매수량이 0 이하인 SKU입니다.", zero_sales)

    excluded_review = validation_mod.apply_order_review_filters(order_review_mod.excluded_order_review_df(settings), settings)
    excluded_qty = float(pd.to_numeric(order_review_mod.report_col(excluded_review, ["발주필요수량"], 0), errors="coerce").fillna(0).sum())
    excluded_amount = float(pd.to_numeric(order_review_mod.report_col(excluded_review, ["발주필요금액"], 0), errors="coerce").fillna(0).sum())
    add("정보", "제외 SKU 수", "상품명/상태에서 명확히 확인되는 샘플/FOC/무상/단종 SKU는 정상 본품 발주검토 합계에서 제외하고 정보성 요약으로만 표시합니다.", len(excluded_review))
    add("정보", "제외 수량", "제외 SKU의 발주필요수량 합계입니다.", f"{excluded_qty:,.0f}")
    add("정보", "제외 금액", "제외 SKU의 발주필요금액 합계입니다.", kpi_mod.fmt_krw(excluded_amount))
    detail_short_qty = float(pd.to_numeric(order_review_mod.report_col(review, ["발주필요수량"], 0), errors="coerce").fillna(0).sum())
    detail_short_amount = float(pd.to_numeric(order_review_mod.report_col(review, ["발주필요금액"], 0), errors="coerce").fillna(0).sum())
    add("정보", "요약-상세 합계 검증", "요약 시트 정상 본품 부족수량/부족금액은 발주 검토 상세 합계와 동일한 원천 데이터로 산출합니다.", f"{detail_short_qty:,.0f}개 / {kpi_mod.fmt_krw(detail_short_amount)}")

    add("정보", "부족수량 계산 기준", "MAX(0, 안전재고 목표수량 - (EU 현지 가용수량 + 운송중 수량))")
    add("정보", "적용 리드타임", transport_mod.lead_time_summary(settings))
    add("정보", "ETA 산정 방식", "ETA = 출고일 + 운송 L/T")
    add("정보", "추천 기준 산정 방식", "운송수단 추천/발주 검토에는 운송 L/T 사용")
    add("정보", "파일별 통화 기준", "EU 현지 창고 평균단가는 원본 컬럼/값으로 EUR·KRW 자동 판별, 본사 EU창고 평균단가=KRW, EU 현지 판매내역상세 금액=EUR, 본사→유럽법인향 판매내역상세 금액=EUR, 해상컨테이너 금액=EUR")
    add("정보", "운송중 도착 예정 금액 기준", "메인 KPI는 해상/운송중 원본 파일의 EUR 금액 합계 × 적용 환율입니다. 발주 산식의 운송재고에는 SKU별 전체 운송수량을 반영하고, ETA/운송수단별 집계는 운송수단 인식 성공 행을 기준으로 표시합니다.")
    add("정보", "안전재고 기준 개월 수", format_months(settings.get("safety_months", 3.0)))
    add("정보", "파라미터 조정 위치", "조건 변경은 웹 업로드 화면의 설정에서 수행하며 엑셀은 결과 보고/검토용 산출물로 유지")
    add("정보", "엑셀 파라미터 편집 여부", "엑셀 산출물에는 파라미터 수정 기능을 넣지 않고 적용 기준만 표시")
    add("정보", "미입고 반영 기준", "미입고수량은 브랜드사 발주 후 본사 창고 미입고된 수량입니다. 기본 발주 필요량에는 포함하지 않고, 보고서 시트의 미입고 해소 후 합계에서 수량 차감 비교합니다.")
    add("정보", "제외 기준", f"샘플/FOC/무상 명칭 제외={settings.get('exclude_sample', True)}, 단종 제외={settings.get('exclude_discontinued', True)}, 단가 0 발주검토 제외=True")
    if signal_counts is not None:
        add("정보", "신호등 발주불필요 건수", "ESM 내부 검토 시트의 발주 신호등 기준입니다.", signal_counts.get("발주불필요", 0))
        add("정보", "신호등 미입력 건수", "ESM 내부 검토 시트의 발주 신호등 기준입니다.", signal_counts.get("미입력", 0))
        add("정보", "신호등 안전 건수", "ESM 내부 검토 시트의 발주 신호등 기준입니다.", signal_counts.get("안전", 0))
        add("정보", "신호등 부족 건수", "ESM 내부 검토 시트의 발주 신호등 기준입니다.", signal_counts.get("부족", 0))
        add("정보", "신호등 확인필요 건수", "ESM 내부 검토 시트의 발주 신호등 기준입니다.", signal_counts.get("확인필요", 0))

    return pd.DataFrame(rows, columns=["구분", "항목", "내용", "건수/값"])


def current_git_hash() -> str:
    git_dir = PROJECT_ROOT / ".git"
    try:
        head = (git_dir / "HEAD").read_text(encoding="utf-8").strip()
        if head.startswith("ref:"):
            ref_path = head.split(":", 1)[1].strip()
            return (git_dir / ref_path).read_text(encoding="utf-8").strip()[:12]
        return head[:12]
    except Exception:
        return os.environ.get("GIT_COMMIT", "-")[:12] or "-"


def uploaded_file_build_info_df(context: SessionContext | None = None) -> pd.DataFrame:
    context = ensure_session_context(context)
    upload_labels = {
        "eu_stock": "EU 현지 재고",
        "hq_eu_stock": "본사 EU창고 재고",
        "sales_detail": "EU 현지 B2B 판매내역상세",
        "open_po": "미입고현황",
        "shipping": "해상/운송중",
        "hq_to_eu_sales_detail": "본사→유럽법인향 판매내역상세",
        "past_sales": "과거 판매내역",
    }
    metadata = context.uploaded_file_metadata
    uploaded_data = context.uploaded_data
    rows: list[dict[str, object]] = []
    for key, label in upload_labels.items():
        meta = metadata.get(key, {}) if isinstance(metadata, dict) else {}
        df = uploaded_data.get(key)
        shape = meta.get("shape") or getattr(df, "shape", None)
        row_count, col_count = ("-", "-")
        if isinstance(shape, tuple) and len(shape) >= 2:
            row_count, col_count = shape[0], shape[1]
        shape_text = f"{row_count} x {col_count}" if row_count != "-" and col_count != "-" else "-"
        rows.append(
            {
                "파일키": key,
                "입력 파일": label,
                "상태": "업로드 완료" if df is not None else "미업로드",
                "파일명": meta.get("file_name") or context.uploaded_files.get(key, "미업로드"),
                "크기(bytes)": meta.get("size", "-"),
                "last_modified": meta.get("last_modified", "-"),
                "읽은 시각": meta.get("loaded_at", "-"),
                "행 수": row_count,
                "열 수": col_count,
                "DataFrame shape": shape_text,
            }
        )
    return pd.DataFrame(rows)


def add_input_file_audit_sheet(wb, build_timestamp: datetime | None = None, context: SessionContext | None = None):
    """생성 엑셀에 실제 반영된 입력 파일 현황을 숨김 시트로 남긴다."""
    ws = wb.create_sheet("_입력파일현황")
    info_df = uploaded_file_build_info_df(context)
    append_df(ws, info_df)
    if len(info_df.columns):
        style_table(ws, 1, max(1, len(info_df) + 1), len(info_df.columns), freeze=False)
        apply_report_column_formats(ws)
    autosize_columns(ws, max_width=60)
    ws.sheet_state = "hidden"
    return ws


def build_info_df(
    settings: dict,
    review: pd.DataFrame,
    build_timestamp: datetime,
    context: SessionContext | None = None,
) -> pd.DataFrame:
    try:
        shipping = transport_mod.get_shipping(settings, context)
        shipping_qty_total = float(pd.to_numeric(shipping.get("수량", 0), errors="coerce").fillna(0).sum())
        shipping_recognized_qty = float(pd.to_numeric(shipping.loc[transport_mod.recognized_transport_mask(shipping), "수량"], errors="coerce").fillna(0).sum())
    except Exception:
        shipping_qty_total = 0.0
        shipping_recognized_qty = 0.0
    review_shipping_total = float(pd.to_numeric(order_review_mod.report_col(review, ["운송중 수량"], 0), errors="coerce").fillna(0).sum()) if isinstance(review, pd.DataFrame) and not review.empty else 0.0
    _, hq_note = order_review_mod.hq_eu_stock_column_status(review) if isinstance(review, pd.DataFrame) else (None, "본사 EU창고 재고 컬럼 미확인")
    return pd.DataFrame(
        [
            {"항목": "추출 시각", "값": build_timestamp.strftime("%Y-%m-%d %H:%M:%S")},
            {"항목": "기준일", "값": settings.get("base_date")},
            {"항목": "발주검토 로직 버전", "값": _ORDER_REVIEW_LOGIC_VERSION},
            {"항목": "엑셀 산출 버전", "값": _ORDER_REVIEW_EXPORT_VERSION},
            {"항목": "git hash", "값": current_git_hash()},
            {"항목": "운송중 수량 합계(prepare_shipping 전체)", "값": f"{shipping_qty_total:,.0f}"},
            {"항목": "운송중 수량 합계(운송수단 인식 성공)", "값": f"{shipping_recognized_qty:,.0f}"},
            {"항목": "ESM 운송중 수량 합계(최종 review)", "값": f"{review_shipping_total:,.0f}"},
            {"항목": "주의", "값": "미입고수량은 참고값이며 발주필요수량 산식에서 직접 차감하지 않음"},
            {"항목": "발주 필요 수량", "값": "MAX(0, 안전재고 목표수량 - (EU 현지 가용수량 + 운송중 수량))"},
            {"항목": "신규 추가", "값": "입고 전 예상 결품량 / 긴급 보충 필요 수량 / 입고 전 결품 위험 여부 / 권장 긴급 액션"},
            {"항목": "긴급 보충 필요 수량", "값": "신규 발주 확정 수량이 아니라 입고 전 공백을 메우기 위한 검토 수량"},
            {"항목": "미입고수량 정책", "값": "미입고수량은 발주 필요 수량 산식에서 직접 차감하지 않고 참고/확인필요 데이터로만 표시"},
            {"항목": "예상입고일 기준", "값": "SKU별 가장 빠른 운송중 ETA 기준"},
            {"항목": "본사 EU창고 재고 정책", "값": f"{hq_note} / 발주필요수량 산식에서는 차감하지 않고 참고값으로만 표시"},
        ]
    )


def _excel_safe_df(df: pd.DataFrame | None) -> pd.DataFrame:
    if df is None:
        return pd.DataFrame()
    out = pd.DataFrame(df).copy()
    for col in out.columns:
        out[col] = out[col].map(lambda value: repr(value) if isinstance(value, (tuple, list, dict, set)) else value)
    return out


def write_excel_debug_section(ws, title: str, df: pd.DataFrame | None, start_row: int) -> int:
    from openpyxl.styles import Font, PatternFill

    section_df = _excel_safe_df(df)
    if section_df.empty:
        section_df = pd.DataFrame([{"내용": "데이터 없음"}])
    end_col = max(len(section_df.columns), 1)
    title_fill = PatternFill("solid", fgColor=REPORT_TABLE_HEADER_FILL)
    for col_idx in range(1, end_col + 1):
        cell = ws.cell(start_row, col_idx)
        cell.fill = title_fill
        cell.font = Font(color="FFFFFF", bold=True, size=12)
    ws.cell(start_row, 1, title)
    append_df(ws, section_df, start_row=start_row + 1)
    header_row = start_row + 1
    for col_idx in range(1, end_col + 1):
        cell = ws.cell(header_row, col_idx)
        cell.fill = title_fill
        cell.font = Font(color="FFFFFF", bold=True)
    apply_report_column_formats(ws, header_row=header_row)
    return start_row + len(section_df) + 4


def add_build_and_shipping_debug_sheets(wb, settings: dict, review: pd.DataFrame, build_timestamp: datetime, context: SessionContext | None = None) -> list:
    sheets = []

    ws_build = wb.create_sheet("_빌드정보")
    row = 1
    row = write_excel_debug_section(ws_build, "빌드 정보", build_info_df(settings, review, build_timestamp, context), row)
    row = write_excel_debug_section(ws_build, "입력 파일 정보", uploaded_file_build_info_df(context), row)
    write_excel_debug_section(ws_build, "세션 캐시 현황", transport_mod.session_cache_inventory_df(), row)
    autosize_columns(ws_build, max_width=60)
    ws_build.sheet_state = "hidden"
    sheets.append(ws_build)

    debug = transport_mod.shipping_pipeline_debug(settings, esm_review=review, context=context)
    ws_debug = wb.create_sheet("_운송중디버그")
    row = 1
    for title, df in [
        ("Stage 0. 파일 read 방식 확인", debug.get("stage0_summary")),
        ("Stage 0. read 시도 이력/fallback", debug.get("read_attempts")),
        ("Stage 0. table/candidate shape", debug.get("candidate_shapes")),
        ("Stage 1. raw 업로드 직후", debug.get("stage1_summary")),
        ("Stage 1. 추적 SKU raw", debug.get("raw_tracking")),
        ("Stage 2. 컬럼 표준화 후", debug.get("stage2_summary")),
        ("Stage 2. 추적 SKU 표준화", debug.get("stage2_tracking")),
        ("Stage 2. SKU 정규화 전후 매칭", debug.get("stage2_sku_normalization_tracking")),
        ("Stage 2-1. 필터 여부", debug.get("stage2_filters")),
        ("Stage 3. 운송수단 정규화 후", debug.get("stage3_summary")),
        ("Stage 3. 정규화 실패 원본값 Top 20", debug.get("transport_failure_counts")),
        ("Stage 3. 운송수단별 행 수/수량", debug.get("transport_mode_summary")),
        ("Stage 4. SKU groupby 후", debug.get("stage4_summary")),
        ("Stage 4. 수량 기준 상위 20 SKU", debug.get("top_skus")),
        ("추적 SKU 단계별 요약", debug.get("tracked_wide")),
        ("복사용 요약", pd.DataFrame({"요약": str(debug.get("summary_text", "")).splitlines()})),
    ]:
        row = write_excel_debug_section(ws_debug, title, df, row)
    autosize_columns(ws_debug, max_width=70)
    ws_debug.sheet_state = "hidden"
    sheets.append(ws_debug)

    ws_compare = wb.create_sheet("_운송중비교")
    row = 1
    for title, df in [
        ("Stage 5. raw vs 최종 ESM 차이 큰 순 Top 30", debug.get("compare_top_diff")),
        ("Stage 5. raw에는 있는데 ESM 운송중이 0인 SKU Top 30", debug.get("compare_raw_only")),
        ("Stage 5. ESM에는 있는데 raw에는 없는 SKU", debug.get("compare_esm_only")),
    ]:
        row = write_excel_debug_section(ws_compare, title, df, row)
    autosize_columns(ws_compare, max_width=70)
    ws_compare.sheet_state = "hidden"
    sheets.append(ws_compare)

    return sheets
