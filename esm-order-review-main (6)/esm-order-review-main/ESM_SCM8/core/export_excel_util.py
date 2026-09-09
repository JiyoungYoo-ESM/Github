from __future__ import annotations

from datetime import date, datetime, timedelta
from numbers import Number

import numpy as np
import pandas as pd

from core.common import (
    REPORT_TABLE_HEADER_FILL,
    korea_today,
)


def autosize_columns(ws, max_width: int = 38) -> None:
    for col in ws.columns:
        letter = col[0].column_letter
        sample = col[:50]
        width = max((len(str(cell.value)) for cell in sample if cell.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(width + 2, 8), max_width)


def clean_excel_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, float) and (np.isnan(value) or np.isinf(value)):
        return ""
    if pd.isna(value):
        return ""
    return value


def excel_text_identifier(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none"}:
        return ""
    return text[:-2] if text.endswith(".0") else text


def excel_exact_match_formula(lookup_value_ref: str, lookup_range_ref: str) -> str:
    """Return an Excel MATCH formula that compares text with case sensitivity."""
    return f"MATCH(TRUE,EXACT({lookup_value_ref},{lookup_range_ref}),0)"


def excel_exact_index_formula(
    lookup_value_ref: str,
    lookup_range_ref: str,
    return_range_ref: str,
    default: str = '""',
) -> str:
    """Return an INDEX/MATCH formula that uses EXACT for SKU-safe lookups."""
    return f'=IFERROR(INDEX({return_range_ref},{excel_exact_match_formula(lookup_value_ref, lookup_range_ref)}),{default})'


def display_text(value: object, header: object = "") -> object:
    if not isinstance(value, str):
        return value
    text = value
    header_text = str(header or "")
    preserve_review_reason = header_text in {"판단 사유", "추천 사유", "운송수단 추천 사유"}
    if not preserve_review_reason:
        text = text.replace("신규 발주 검토", "발주 필요")
        text = text.replace("추가 발주 검토", "발주 필요")
    if "ETA 지연 위험 SKU" not in text:
        text = text.replace("입고 전 품절 위험", "ETA 지연 위험")
    if text == "운송 검토안":
        text = "권장 대응 / 운송 검토안"
    text = text.replace("일평균 판매수량(기준/90)", "일평균 판매수량")
    text = text.replace("고갈 예정일", "예상 소진일")
    text = text.replace("쇼티지 예상 일수", "쇼티지 예상일수")
    if "SKU" in header_text or "SKU" in text or header_text in {"SKU등록상태", "확인 구분", "확인필요 사유", "권장 확인 액션"}:
        text = text.replace("SKU 확인필요", "상품코드 확인 필요")
        text = text.replace("SKU 확인 필요", "상품코드 확인 필요")
        text = text.replace("코드 확인필요", "상품코드 확인 필요")
    return text


def append_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> None:
    text_headers = {
        "상품코드", "SKU", "바코드", "컨테이너번호",
        "우선 액션", "운송 검토안", "권장 대응 / 운송 검토안", "추천 운송수단", "권장 운송안", "권장 수량 요약", "판단 사유", "운송수단 추천 사유", "추천 사유",
        "권장 긴급 액션", "입고 전 결품 위험 여부",
        "발주 신호등", "발주 메모", "담당자판단", "담당자", "브랜드", "상품명", "제품상태", "구분", "판매수량 기준",
        "보조 검증 기준", "판매수량 기준 확인 필요", "원본 출처", "확인 필요 사유",
        "SKU등록상태", "재고운영상태", "판매이력여부", "운송중여부", "미입고여부", "판단메모",
        "재고파일 존재여부", "판매내역 존재여부", "운송중 존재여부", "미입고 존재여부",
        "발견 원본", "확인 구분", "확인필요 사유", "권장 확인 액션",
    }
    col_names = [str(c) for c in df.columns]
    text_offsets = [i for i, c in enumerate(col_names) if c in text_headers]

    for c_idx, col in enumerate(df.columns, start_col):
        ws.cell(start_row, c_idx, display_text(col))

    if start_col == 1:
        for row_tuple in df.itertuples(index=False):
            ws.append([
                excel_text_identifier(display_text(v, col_names[i])) if i in text_offsets
                else clean_excel_value(display_text(v, col_names[i]))
                for i, v in enumerate(row_tuple)
            ])
        text_offset_set = set(text_offsets)
        for r_idx in range(start_row + 1, start_row + 1 + len(df)):
            for offset in text_offset_set:
                ws.cell(r_idx, offset + 1).number_format = "@"
    else:
        for r_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
            for offset, value in enumerate(row):
                c_idx = start_col + offset
                header = col_names[offset]
                cell = ws.cell(r_idx, c_idx)
                if header in text_headers:
                    cell.value = excel_text_identifier(display_text(value, header))
                    cell.number_format = "@"
                else:
                    cell.value = clean_excel_value(display_text(value, header))


def style_table(ws, start_row: int, end_row: int, end_col: int, header_fill: str = REPORT_TABLE_HEADER_FILL, freeze: bool = True) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D9DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor=header_fill)
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in {5, 17})
    for row in ws.iter_rows(min_row=start_row, max_row=start_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    if freeze:
        ws.freeze_panes = ws.cell(start_row + 1, 1)


def format_number_columns(ws, header_row: int = 1) -> None:
    one_decimal_keywords = ["개월", "일평균", "커버"]
    header_map = {cell.column: str(cell.value or "") for cell in ws[header_row]}
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if not isinstance(cell.value, (int, float)):
                continue
            header = header_map.get(cell.column, "")
            cell.number_format = "#,##0.0" if any(keyword in header for keyword in one_decimal_keywords) else "#,##0"


def apply_report_column_formats(ws, header_row: int = 1) -> None:
    from openpyxl.styles import Alignment

    text_headers = {
        "상품코드", "SKU", "바코드", "컨테이너번호",
        "우선 액션", "운송 검토안", "권장 대응 / 운송 검토안", "추천 운송수단", "권장 운송안", "권장 수량 요약", "판단 사유", "운송수단 추천 사유", "추천 사유",
        "발주 신호등", "발주 메모", "담당자판단", "담당자", "브랜드", "상품명", "제품상태", "구분", "판매수량 기준",
        "보조 검증 기준", "판매수량 기준 확인 필요", "원본 출처", "확인 필요 사유",
        "SKU등록상태", "재고운영상태", "판매이력여부", "운송중여부", "미입고여부", "판단메모",
        "재고파일 존재여부", "판매내역 존재여부", "운송중 존재여부", "미입고 존재여부",
        "발견 원본", "확인 구분", "확인필요 사유", "권장 확인 액션",
    }
    date_headers = {
        "도착일", "출고일", "발주일", "납기예정일", "납기 예정일", "예상 도착 가능일",
        "고갈 예상일", "고갈 예정일", "기준일", "예상 소진일", "첫 입고 예정일",
        "최초 도착 예정일", "최초 ETA", "검토일", "최초 도착일", "마지막 도착일", "예상 입고일",
    }
    date_keywords = ["도착일", "출고일", "발주일", "납기예정일", "납기 예정일", "예정일", "소진일", "검토일", "기준일"]
    non_date_keywords = ["공백일수", "경과/잔여일", "리드타임"]
    percent_headers = {"판매 집중도", "판매수량 차이율"}
    krw_keywords = ["KRW"]
    eur_keywords = ["EUR"]
    amount_keywords = ["금액"]
    unit_price_keywords = ["단가", "입고단가"]
    month_keywords = ["개월", "커버"]
    qty_headers = {
        "최근 3개월 수요", "3개월 수요", "월평균", "월평균 판매수량", "안전재고",
        "EU 현지 재고", "유럽 재고", "본사 EU창고", "운송중", "미입고",
        "유럽+운송", "EU+운송중 합계", "부족수량 산정 기준 보유량",
        "부족수량", "발주필요수량", "입고 전 예상 결품량", "긴급 보충 필요 수량", "파이프라인 합계", "SKU 수",
        "기준_3M_판매수량", "PA_CA_3M_판매수량",
    }
    rounded_qty_headers = {
        "월평균", "월평균 판매량", "월평균 판매수량",
        "최근 3개월 PA+CA 판매수량", "최근 3개월 판매량", "일평균 판매량",
        "기준_3M_판매수량", "PA_CA_3M_판매수량",
        "기준 3개월 판매수량(재고 PA+CA)",
        "월평균 판매수량(기준/3)", "일평균 판매수량(기준/90)",
    }
    average_qty_headers: set[str] = set()
    qty_keywords = ["수량", "수요", "재고", "운송중", "운송", "미입고", "월평균", "안전재고", "파이프라인", "보유량", "합계", "결품량"]

    headers = {cell.column: str(cell.value or "") for cell in ws[header_row]}
    header_replacements = {
        "첫 ETA 도착일": "첫 입고 예정일",
        "최초 ETA": "첫 입고 예정일",
        "최초 ETA 수량": "첫 입고 예정 수량",
    }
    for cell in ws[header_row]:
        replacement = header_replacements.get(str(cell.value or ""))
        if replacement:
            cell.value = replacement
            headers[cell.column] = replacement

    def is_date_header(header: str) -> bool:
        if header in date_headers:
            return True
        if any(keyword in header for keyword in non_date_keywords):
            return False
        return any(keyword in header for keyword in date_keywords)

    _TEXT, _DATE, _OTHER = 0, 1, 2
    col_type: dict[int, int] = {}
    col_fmt: dict[int, str] = {}
    col_is_numeric: set[int] = set()
    col_skip_align: set[int] = set()
    wrap_cols = frozenset({5, 17, 21})

    for col_idx, header in headers.items():
        if header in text_headers:
            col_type[col_idx] = _TEXT
            col_fmt[col_idx] = "@"
        elif is_date_header(header):
            col_type[col_idx] = _DATE
            col_fmt[col_idx] = "yyyy-mm-dd"
        else:
            col_type[col_idx] = _OTHER
            if header == "판매수량 차이율":
                col_fmt[col_idx] = '0.0"%"'
                col_skip_align.add(col_idx)
            elif header in percent_headers:
                col_fmt[col_idx] = "0.0%"
                col_skip_align.add(col_idx)
            elif any(kw in header for kw in unit_price_keywords):
                col_fmt[col_idx] = "#,##0.00"
                col_is_numeric.add(col_idx)
            elif any(kw in header for kw in krw_keywords + eur_keywords + amount_keywords):
                col_fmt[col_idx] = "#,##0"
                col_is_numeric.add(col_idx)
            elif header in rounded_qty_headers:
                col_fmt[col_idx] = "#,##0"
                col_is_numeric.add(col_idx)
            elif header in average_qty_headers:
                col_fmt[col_idx] = "#,##0.0"
                col_is_numeric.add(col_idx)
            elif "판매수량" in header or " 수량" in header or header.endswith("수량"):
                col_fmt[col_idx] = "#,##0"
                col_is_numeric.add(col_idx)
            elif any(kw in header for kw in month_keywords):
                col_fmt[col_idx] = "0.0"
                col_is_numeric.add(col_idx)
            elif header in qty_headers or any(kw in header for kw in qty_keywords) or header.startswith("ETA "):
                col_fmt[col_idx] = "#,##0"
                col_is_numeric.add(col_idx)

    align_wrap = Alignment(vertical="center", wrap_text=True)
    align_no_wrap = Alignment(vertical="center", wrap_text=False)

    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            col = cell.column
            ctype = col_type.get(col, _OTHER)
            if ctype == _TEXT:
                cell.number_format = "@"
                if cell.value not in (None, ""):
                    cell.value = excel_text_identifier(cell.value)
                continue
            if ctype == _DATE:
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = align_wrap if col in wrap_cols else align_no_wrap
                continue
            if col in col_is_numeric and isinstance(cell.value, str):
                numeric_value = pd.to_numeric(cell.value.replace(",", ""), errors="coerce")
                if pd.notna(numeric_value):
                    cell.value = float(numeric_value)
            fmt = col_fmt.get(col)
            if fmt:
                cell.number_format = fmt
            if col not in col_skip_align:
                cell.alignment = align_wrap if col in wrap_cols else align_no_wrap


def apply_fast_excel_column_layout(ws, header_row: int = 1) -> None:
    header_widths = {
        "상품코드": 18,
        "상품명": 42,
        "브랜드": 16,
        "우선 액션": 20,
        "권장 긴급 액션": 34,
        "발주필요금액(KRW)": 20,
        "발주 필요 수량": 16,
        "긴급 보충 필요 수량": 18,
        "입고 전 예상 결품량": 18,
        "입고 전 결품 위험 여부": 18,
        "미입고 수량": 16,
        "운송중 수량": 16,
        "유럽+운송 수량": 18,
        "운송 포함 보유개월": 18,
        "EU 현지 커버일수": 18,
        "유럽 가용재고": 16,
        "유럽 현재 재고": 16,
        "항공 검토수량": 16,
        "철송 검토수량": 16,
        "해운 검토수량": 16,
        "월평균 판매수량": 18,
        "일평균 판매수량": 18,
        "기준 3개월 판매수량": 20,
        "안전재고 필요 수량": 20,
        "부족 수량": 16,
        "판단 사유": 42,
        "권장 운송안": 18,
        "권장 수량 요약": 28,
        "추천 운송수단": 18,
        "예상 소진일": 14,
        "예상 입고일": 14,
        "첫 입고 예정일": 14,
        "예상 도착 가능일": 16,
        "고갈 예상일": 14,
        "최초 ETA": 14,
        "검토일": 14,
        "출고일": 14,
        "도착일": 14,
    }
    date_headers = {"예상 소진일", "예상 입고일", "검토일", "출고일", "도착일", "첫 입고 예정일", "예상 도착 가능일"}
    integer_headers = {
        "발주필요금액(KRW)",
        "발주 필요 수량",
        "긴급 보충 필요 수량",
        "입고 전 예상 결품량",
        "미입고 수량",
        "유럽 가용재고",
        "운송중 수량",
        "유럽+운송 수량",
        "월평균 판매수량",
        "일평균 판매수량",
        "기준 3개월 판매수량",
        "안전재고 필요 수량",
        "부족 수량",
        "유럽 현재 재고",
        "항공 검토수량",
        "철송 검토수량",
        "해운 검토수량",
    }
    decimal_headers = {
        "운송 포함 보유개월",
        "EU 현지 커버일수",
        "쇼티지 예상일수",
    }
    integer_keywords = [
        "수량", "재고", "금액", "월평균", "일평균", "기준 3개월", "안전재고",
        "유럽", "운송중", "미입고", "부족",
    ]
    decimal_keywords = ["커버일수", "보유개월", "쇼티지", "단가", "포함 보유개월"]
    pct_like_headers = {"운송 포함 보유개월", "EU 현지 커버일수", "쇼티지 예상일수"}

    def coerce_numeric_cell(cell) -> None:
        if cell.value in (None, ""):
            return
        if isinstance(cell.value, Number) and not isinstance(cell.value, bool):
            return
        if isinstance(cell.value, str):
            text = cell.value.strip().replace(",", "")
            if not text:
                return
            numeric_value = pd.to_numeric(text, errors="coerce")
            if pd.notna(numeric_value):
                value = float(numeric_value)
                cell.value = int(value) if value.is_integer() else value

    for cell in ws[header_row]:
        header = str(cell.value or "")
        letter = cell.column_letter
        ws.column_dimensions[letter].width = header_widths.get(header, max(12, min(24, len(header) + 4)))
        is_date_col = header in date_headers or any(keyword in header for keyword in ["소진일", "입고일", "도착일", "ETA", "출고일", "검토일"])
        if is_date_col:
            target_format = "yyyy-mm-dd"
        elif header in decimal_headers or any(keyword in header for keyword in decimal_keywords):
            target_format = "#,##0.00" if "단가" in header else "#,##0.0"
        elif header in integer_headers or any(keyword in header for keyword in integer_keywords):
            target_format = "#,##0"
        else:
            target_format = ""

        if target_format:
            for col_cells in ws.iter_cols(min_col=cell.column, max_col=cell.column, min_row=header_row + 1, max_row=ws.max_row):
                for item in col_cells:
                    if target_format == "yyyy-mm-dd" and isinstance(item.value, str):
                        parsed = pd.to_datetime(item.value, errors="coerce")
                        if pd.notna(parsed):
                            item.value = parsed.to_pydatetime()
                        item.number_format = target_format
                    elif target_format != "yyyy-mm-dd":
                        coerce_numeric_cell(item)
                        item.number_format = target_format
                    elif isinstance(item.value, (datetime, date)):
                        item.number_format = target_format
            if is_date_col:
                ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, 14)


def apply_readable_excel_layout(ws, header_rows: list[int] | None = None) -> None:
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_rows = header_rows or [1]
    thin_border = Border(
        left=Side(style="thin", color="D9DEE8"),
        right=Side(style="thin", color="D9DEE8"),
        top=Side(style="thin", color="D9DEE8"),
        bottom=Side(style="thin", color="D9DEE8"),
    )
    header_fill = PatternFill("solid", fgColor=REPORT_TABLE_HEADER_FILL)
    date_keywords = ["일", "ETA", "검토일"]
    numeric_keywords = ["수량", "재고", "운송중", "미입고", "금액", "KRW", "EUR", "개월", "커버", "수요", "합계"]

    for header_row in header_rows:
        if header_row > ws.max_row:
            continue
        for cell in ws[header_row]:
            if isinstance(cell, MergedCell):
                continue
            if cell.value in (None, ""):
                continue
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.border = thin_border
            if cell.row not in header_rows:
                cell.alignment = Alignment(vertical="center", wrap_text=cell.alignment.wrap_text)

    best_headers: dict[int, str] = {}
    for header_row in header_rows:
        if header_row > ws.max_row:
            continue
        for cell in ws[header_row]:
            if isinstance(cell, MergedCell):
                continue
            header = str(cell.value or "")
            if header:
                best_headers.setdefault(cell.column, header)

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        header = best_headers.get(col_idx, "")
        if header in {"상품명", "주요 상품", "판단메모", "판단 사유", "추천 사유", "권장 수량 요약"}:
            width = 34
        elif header in {"상품코드", "SKU", "바코드"}:
            width = 16
        elif header == "브랜드":
            width = 16
        elif any(keyword in header for keyword in numeric_keywords):
            width = 14
        elif any(keyword in header for keyword in date_keywords):
            width = 14
        else:
            width = min(max(ws.column_dimensions[letter].width or 10, 10), 24)
        ws.column_dimensions[letter].width = width


def apply_action_row_highlights(ws, header_rows: list[int] | None = None) -> None:
    from copy import copy
    from openpyxl.styles import PatternFill

    header_rows = header_rows or [1]
    header_row_set = set(header_rows)
    risk_fill = PatternFill("solid", fgColor="FCE4D6")
    for header_row in header_rows:
        if header_row > ws.max_row:
            continue
        headers = {str(cell.value or ""): cell.column for cell in ws[header_row]}
        action_col = headers.get("우선 액션")
        qty_col = next(
            (headers[col_name] for col_name in ["발주필요수량", "발주 필요 수량", "부족수량", "부족 수량"] if col_name in headers),
            None,
        )
        if not action_col and not qty_col:
            continue
        for row_idx in range(header_row + 1, ws.max_row + 1):
            if row_idx in header_row_set or any(str(cell.value or "") == "우선 액션" for cell in ws[row_idx]):
                break
            action_text = str(ws.cell(row_idx, action_col).value or "") if action_col else ""
            if "입고 전 결품 위험" in action_text or "입고 전 품절 위험 SKU" in action_text or "입고 전 품절 위험" in action_text or "ETA 지연 위험" in action_text:
                ws.cell(row_idx, action_col).fill = risk_fill
            elif action_text == "발주 필요":
                action_cell = ws.cell(row_idx, action_col)
                bold_font = copy(action_cell.font)
                bold_font.bold = True
                action_cell.font = bold_font
            if qty_col:
                qty_value = pd.to_numeric(ws.cell(row_idx, qty_col).value, errors="coerce")
                if pd.notna(qty_value) and float(qty_value) > 0:
                    qty_cell = ws.cell(row_idx, qty_col)
                    bold_font = copy(qty_cell.font)
                    bold_font.bold = True
                    qty_cell.font = bold_font


def apply_internal_review_formulas(ws, row_count: int) -> None:
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    headers = {str(cell.value or ""): cell.column for cell in ws[1]}
    manager_qty_col = headers.get("담당자 발주량")
    required_qty_col = headers.get("발주 필요 수량")
    diff_col = headers.get("발주 차이")
    signal_col = headers.get("발주 신호등")
    review_date_col = headers.get("검토일")
    if not all([manager_qty_col, required_qty_col, diff_col, signal_col]):
        return

    manager_letter = get_column_letter(manager_qty_col)
    required_letter = get_column_letter(required_qty_col)
    diff_letter = get_column_letter(diff_col)
    signal_letter = get_column_letter(signal_col)
    review_date_letter = get_column_letter(review_date_col) if review_date_col else ""
    for row in range(2, row_count + 2):
        ws[f"{manager_letter}{row}"].value = None
        ws[f"{manager_letter}{row}"].number_format = "#,##0"
        ws[f"{diff_letter}{row}"] = (
            f'=IF(NOT(ISNUMBER({required_letter}{row})),"",'
            f'IF({manager_letter}{row}="","",'
            f'IF(ISNUMBER({manager_letter}{row}),{manager_letter}{row}-{required_letter}{row},"")))'
        )
        ws[f"{signal_letter}{row}"] = (
            f'=IF(NOT(ISNUMBER({required_letter}{row})),"● 확인필요",'
            f'IF({manager_letter}{row}="",'
            f'IF({required_letter}{row}<=0,"● 발주불필요","● 미입력"),'
            f'IF(NOT(ISNUMBER({manager_letter}{row})),"● 확인필요",'
            f'IF({manager_letter}{row}<0,"● 확인필요",'
            f'IF({manager_letter}{row}>={required_letter}{row},"● 안전","● 부족")))))'
        )
        if review_date_letter:
            ws[f"{review_date_letter}{row}"].number_format = "yyyy-mm-dd"
    ws[f"{manager_letter}1"].comment = Comment("실제로 담당자가 발주하기로 결정한 수량을 입력하세요.", "Codex")
    ws[f"{diff_letter}1"].comment = Comment("담당자 발주량과 발주 필요 수량이 모두 숫자일 때만 자동 계산됩니다.", "Codex")
    ws[f"{signal_letter}1"].comment = Comment("발주 필요 수량이 0 이하라도 담당자 발주량을 입력하면 입력 수량 기준으로 판정됩니다.", "Codex")
    if row_count <= 0:
        return
    signal_range = f"{signal_letter}2:{signal_letter}{row_count + 1}"
    signal_fonts = {
        "● 안전": Font(color="2E7D32", bold=True),
        "● 미입력": Font(color="ED7D31", bold=True),
        "● 확인필요": Font(color="C00000", bold=True),
        "● 부족": Font(color="C00000", bold=True),
        "● 발주불필요": Font(color="7F7F7F", bold=True),
    }
    for signal_text, signal_font in signal_fonts.items():
        ws.conditional_formatting.add(signal_range, CellIsRule(operator="equal", formula=[f'"{signal_text}"'], font=signal_font))

    if review_date_letter:
        wb = ws.parent
        helper_title = "_검토일_선택값"
        if helper_title in wb.sheetnames:
            del wb[helper_title]
        helper_ws = wb.create_sheet(helper_title)
        first_date = ws[f"{review_date_letter}2"].value if row_count else None
        if isinstance(first_date, datetime):
            start_date = first_date.date()
        elif isinstance(first_date, date):
            start_date = first_date
        else:
            start_date = korea_today()
        for idx in range(366):
            cell = helper_ws.cell(idx + 1, 1, start_date + timedelta(days=idx))
            cell.number_format = "yyyy-mm-dd"
        helper_ws.sheet_state = "hidden"

        date_validation = DataValidation(
            type="list",
            formula1=f"'{helper_title}'!$A$1:$A$366",
            allow_blank=True,
        )
        date_validation.promptTitle = "검토일"
        date_validation.prompt = "드롭다운에서 검토일을 선택하거나 yyyy-mm-dd 형식으로 입력하세요."
        date_validation.errorTitle = "날짜 선택 확인"
        date_validation.error = "검토일은 드롭다운 날짜 또는 yyyy-mm-dd 형식으로 입력해야 합니다."
        ws.add_data_validation(date_validation)
        date_validation.add(f"{review_date_letter}2:{review_date_letter}{row_count + 1}")


def apply_date_picker_validation(
    ws,
    header_name: str,
    helper_title: str,
    prompt_title: str,
    prompt: str,
    header_row: int = 1,
) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    headers = {str(cell.value or ""): cell.column for cell in ws[header_row]}
    target_col = headers.get(header_name)
    row_count = max(ws.max_row - header_row, 0)
    if not target_col or row_count <= 0:
        return

    target_letter = get_column_letter(target_col)
    data_start_row = header_row + 1
    data_end_row = header_row + row_count
    for row in range(data_start_row, data_end_row + 1):
        ws[f"{target_letter}{row}"].number_format = "yyyy-mm-dd"

    wb = ws.parent
    if helper_title in wb.sheetnames:
        del wb[helper_title]
    helper_ws = wb.create_sheet(helper_title)
    start_date = korea_today() - timedelta(days=30)
    for idx in range(427):
        cell = helper_ws.cell(idx + 1, 1, start_date + timedelta(days=idx))
        cell.number_format = "yyyy-mm-dd"
    helper_ws.sheet_state = "hidden"

    date_validation = DataValidation(
        type="list",
        formula1=f"'{helper_title}'!$A$1:$A$427",
        allow_blank=True,
    )
    date_validation.promptTitle = prompt_title
    date_validation.prompt = prompt
    date_validation.errorTitle = "날짜 선택 확인"
    date_validation.error = f"{header_name}은 드롭다운 날짜 또는 yyyy-mm-dd 형식으로 입력해야 합니다."
    ws.add_data_validation(date_validation)
    date_validation.add(f"{target_letter}{data_start_row}:{target_letter}{data_end_row}")


def enable_excel_auto_calculation(wb) -> None:
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcOnSave = True
    wb.calculation.calcId = 0
