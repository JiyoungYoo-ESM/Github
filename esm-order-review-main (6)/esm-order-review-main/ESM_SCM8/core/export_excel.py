from __future__ import annotations

from datetime import date

import pandas as pd

from core.common import _DEFAULT_EUR_KRW_RATE, REPORT_TABLE_HEADER_FILL
from core.export_excel_util import append_df, autosize_columns, enable_excel_auto_calculation
from core.session import SessionContext
from datetime import datetime
from io import BytesIO
from core.common import format_months, korea_now
from core import (
    eta as eta_mod,
    export_excel_util as export_excel_util_mod,
    inbound as inbound_mod,
    kpi as kpi_mod,
    loaders as loaders_mod,
    order_review as order_review_mod,
    order_review_report as order_review_report_mod,
    transport as transport_mod,
    validation as validation_mod,
)
from core.export_excel_sheets import (
    _clean_merge_key,
    _series_from_candidates,
    _build_unique_lookup_df,
    _lookup_order_values,
    _first_non_blank,
    build_unified_esm_order_review_sheet,
    validate_unified_esm_order_review_sheet,
    format_unified_esm_order_review_sheet,
)
from core.export_excel_order_sheet import (
    build_stock_eta_sheet_df,
    _cover_month_display,
    _order_sheet_transport_recommendation,
    build_order_sheet_df,
    build_order_sheet_top_sales_df,
    format_order_sheet,
)
from core.export_excel_report import (
    build_internal_work_review_df,
    build_order_review_report_summary,
    format_report_sheet,
    build_monthly_arrival_summary_df,
    build_stockout_calendar_report_df,
    pre_arrival_shortage_logic_summary_text,
    internal_review_signal_counts,
)
from core.export_excel_debug import (
    apply_workbook_thousands_number_formats,
    build_validation_log_df,
    current_git_hash,
    uploaded_file_build_info_df,
    add_input_file_audit_sheet,
    build_info_df,
    _excel_safe_df,
    write_excel_debug_section,
    add_build_and_shipping_debug_sheets,
)



def generate_order_review_excel(settings: dict, review: pd.DataFrame, build_timestamp: datetime | None = None, context: SessionContext | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    build_timestamp = build_timestamp or korea_now()
    report_df = order_review_report_mod.build_order_review_report_df(review, settings)
    excluded_review = validation_mod.apply_order_review_filters(order_review_mod.excluded_order_review_df(settings), settings)
    excluded_df = order_review_report_mod.build_excluded_order_review_df(excluded_review, settings)
    unrecognized_transport_df = transport_mod.unrecognized_transport_report_df(settings, context)
    master_unregistered_df = inbound_mod.build_master_unregistered_sku_df(settings, context)
    human_data_issue_df = order_review_report_mod.build_human_data_issue_df(settings, context)
    check_required_df = order_review_report_mod._strip_check_required_internal_columns(
        order_review_report_mod.build_check_required_sheet_df(
            pd.DataFrame(),
            unrecognized_transport_df,
            excluded_df,
            master_unregistered_df,
            pd.DataFrame(),
            human_data_issue_df,
        )
    )
    calendar_df = eta_mod.build_arrival_calendar_report_df(settings, context)
    monthly_arrival_df = build_monthly_arrival_summary_df(settings, context)
    transport_arrival_df = eta_mod.build_transport_arrival_summary_df(settings, context)
    target_month_arrival_df = eta_mod.build_target_month_arrival_summary_df(settings, context=context)
    stockout_df = build_stockout_calendar_report_df(review)
    sales_concentration_top10_df = order_review_report_mod.build_sales_concentration_top10_df(report_df)
    order_risk_top10_df = order_review_report_mod.build_order_risk_top10_df(report_df)
    pre_eta_stockout_top10_df = order_review_report_mod.build_pre_eta_stockout_top10_df(stockout_df)
    stock_eta_df = build_stock_eta_sheet_df(review, calendar_df, settings)
    order_sheet_df = build_order_sheet_df(report_df, settings)
    order_sheet_top_sales_df = build_order_sheet_top_sales_df(report_df)
    if not calendar_df.empty:
        eta_summary_rows: list[dict[str, object]] = []
        for (arrival_date, transport_mode), group in calendar_df.groupby(["도착일", "운송수단"], dropna=False):
            sku_summary = (
                group.groupby("SKU", dropna=False)
                .agg(
                    수량=("수량", "sum"),
                    상품명=("상품명", lambda values: next((str(value).strip() for value in values if str(value).strip()), "")),
                )
                .sort_values("수량", ascending=False)
            )
            top_items = []
            for sku, item in sku_summary.head(3).iterrows():
                sku_text = str(sku).strip()
                name_text = str(item.get("상품명", "")).strip()
                if sku_text and name_text:
                    top_items.append(f"{sku_text} / {name_text}")
                elif sku_text:
                    top_items.append(sku_text)
            if not top_items:
                top_items = [""]
            if len(sku_summary) > 3:
                top_items.append(f"외 {len(sku_summary) - 3}개")
            for product_text in top_items:
                eta_summary_rows.append(
                    {
                        "도착일": arrival_date,
                        "운송수단": transport_mode,
                        "SKU 수": int(group["SKU"].nunique()),
                        "주요 상품": product_text,
                        "총 수량": float(group["수량"].sum()),
                    }
                )
        eta_summary_df = pd.DataFrame(eta_summary_rows).sort_values(["도착일", "운송수단"], na_position="last")
    else:
        eta_summary_df = pd.DataFrame(columns=["도착일", "운송수단", "SKU 수", "주요 상품", "총 수량"])

    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "ESM 발주 검토 요약"
    ws["A1"].font = Font(size=18, bold=True, color="1F2937")
    ws["A3"] = "기준일"
    ws["B3"] = f'{settings["base_date"]:%Y-%m-%d}'
    ws["A4"] = "판매 기준기간"
    ws["B4"] = f'{settings["period_start"]:%Y-%m-%d} ~ {settings["period_end"]:%Y-%m-%d}'
    ws["A5"] = "안전재고 기준"
    ws["B5"] = f'월평균 판매량 x {format_months(settings.get("safety_months", 3.0))}'
    ws["A6"] = "환율"
    ws["B6"] = f'{float(settings.get("eur_krw_rate", _DEFAULT_EUR_KRW_RATE)):,.2f} KRW/EUR'

    action = order_review_mod.report_col(report_df, ["우선 액션", "최종 액션"]).astype(str)
    status = order_review_mod.report_col(report_df, ["상태"], "정상").astype(str)
    total_short_qty = float(pd.to_numeric(order_review_mod.report_col(report_df, ["추가 발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0).sum())
    total_short_eur = float(report_df["부족금액_EUR"].sum())
    total_short_krw = float(pd.to_numeric(order_review_mod.report_col(report_df, ["추가 발주 필요 금액", "부족금액_KRW"], 0), errors="coerce").fillna(0).sum())
    excluded_qty = float(excluded_df["제외 수량"].sum()) if "제외 수량" in excluded_df.columns else 0
    excluded_krw = float(excluded_df["제외 금액_KRW"].sum()) if "제외 금액_KRW" in excluded_df.columns else 0
    base_month = date(settings["base_date"].year, settings["base_date"].month, 1)
    next_month = eta_mod.add_months(base_month, 1)
    current_month_arrival_qty = eta_mod.arrival_qty_for_month(calendar_df, base_month)
    next_month_arrival_qty = eta_mod.arrival_qty_for_month(calendar_df, next_month)
    current_month_arrival_amount = eta_mod.arrival_amount_for_month(calendar_df, base_month)
    next_month_arrival_amount = eta_mod.arrival_amount_for_month(calendar_df, next_month)
    order_needed_sku_count = int(order_review_mod.order_needed_action_mask(action).sum())
    eta_stockout_risk_count = int(stockout_df["입고 전 품절 여부"].astype(str).eq("Y").sum()) if "입고 전 품절 여부" in stockout_df.columns else 0
    kpis = [
        ("정상 본품 SKU 수", f"{len(report_df):,}"),
        ("발주필요 SKU 수", f"{order_needed_sku_count:,}"),
        ("데이터 확인필요 SKU 수", f"{len(check_required_df):,}"),
        ("본사이동 SKU 수", f'{int(status.eq("본사이동").sum()):,}'),
        ("운송대기 SKU 수", f'{int(status.eq("운송대기").sum()):,}'),
        ("입고 전 품절위험 SKU 수", f"{eta_stockout_risk_count:,}"),
        ("추가 발주 필요 수량(정상 본품)", f"{total_short_qty:,.0f}"),
        ("추가 발주 필요 금액_EUR(정상 본품)", f"{total_short_eur:,.0f} EUR"),
        ("추가 발주 필요 금액_KRW(정상 본품)", kpi_mod.fmt_krw(total_short_krw)),
        ("이번 달 도착 예정 수량", f"{current_month_arrival_qty:,.0f}"),
        ("이번 달 도착 예정 금액", kpi_mod.fmt_krw(current_month_arrival_amount)),
        ("다음 달 도착 예정 수량", f"{next_month_arrival_qty:,.0f}"),
        ("다음 달 도착 예정 금액", kpi_mod.fmt_krw(next_month_arrival_amount)),
        ("제외 SKU 수", f"{len(excluded_df):,}"),
        ("제외 수량", f"{excluded_qty:,.0f}"),
        ("제외 금액_KRW", kpi_mod.fmt_krw(excluded_krw)),
    ]
    ws["A8"] = "구분"
    ws["B8"] = "값"
    for idx, (label, value) in enumerate(kpis, 9):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
        ws.cell(idx, 2).font = Font(bold=True)
    summary_header_fill = PatternFill("solid", fgColor=REPORT_TABLE_HEADER_FILL)
    for cell in ws[8][0:2]:
        cell.fill = summary_header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    summary_header_rows: list[int] = []
    target_month_title = "/".join(f"{eta_mod.add_months(base_month, offset).month}월" for offset in (1, 2, 3))

    for row in ws.iter_rows(min_row=3, max_row=8 + len(kpis), min_col=1, max_col=2):
        for cell in row:
            cell.border = Border(bottom=Side(style="thin", color="D9DEE8"))
            if cell.column == 1:
                cell.font = Font(bold=cell.row == 8)
            if cell.column == 2 and cell.row > 8:
                cell.font = Font(bold=True)
    autosize_columns(ws)
    export_excel_util_mod.apply_readable_excel_layout(ws, summary_header_rows)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 34

    ws_check_required = wb.create_sheet("확인필요")
    append_df(ws_check_required, check_required_df, start_row=8)
    order_review_report_mod.format_check_required_sheet(ws_check_required, settings, header_row=8)
    export_excel_util_mod.apply_date_picker_validation(
        ws_check_required,
        "판단일시",
        "_판단일시_선택값",
        "판단일시",
        "드롭다운에서 판단일시를 선택하거나 yyyy-mm-dd 형식으로 입력하세요.",
        header_row=8,
    )

    ws_report = wb.create_sheet("보고서")
    format_report_sheet(ws_report, settings, report_df)

    ws_arrival_plan = wb.create_sheet("입고예정")
    eta_mod.write_arrival_reference_sheet(ws_arrival_plan, settings, calendar_df)

    ws_stock_eta = wb.create_sheet("재고 ETA")
    append_df(ws_stock_eta, stock_eta_df)
    format_unified_esm_order_review_sheet(ws_stock_eta, settings)
    ws_stock_eta.freeze_panes = "D10"

    ws_order = wb.create_sheet("발주")
    append_df(ws_order, order_sheet_df, start_row=10)
    format_order_sheet(ws_order, settings, order_sheet_top_sales_df)

    input_audit_sheet = add_input_file_audit_sheet(wb, build_timestamp, context)
    debug_sheets = add_build_and_shipping_debug_sheets(wb, settings, review, build_timestamp, context)
    enable_excel_auto_calculation(wb)
    helper_sheets = [
        wb[name]
        for name in ["_검토일_선택값", "_판단일시_선택값", "_보고서원본", "_브랜드_선택값"]
        if name in wb.sheetnames
    ]

    wb._sheets = [
        ws_stock_eta,
        ws_order,
        ws_report,
        ws_arrival_plan,
        ws_check_required,
    ] + [input_audit_sheet] + debug_sheets + helper_sheets

    output = BytesIO()
    apply_workbook_thousands_number_formats(wb)
    wb.save(output)
    return output.getvalue()


def generate_order_review_excel_fast(
    settings: dict,
    review: pd.DataFrame,
    build_timestamp: datetime | None = None,
    excluded_review: pd.DataFrame | None = None,
    report_df: pd.DataFrame | None = None,
    check_required_df: pd.DataFrame | None = None,
    calendar_df: pd.DataFrame | None = None,
    stock_eta_df: pd.DataFrame | None = None,
    context: SessionContext | None = None,
) -> bytes:
    from openpyxl import Workbook

    build_timestamp = build_timestamp or korea_now()
    report_df = pd.DataFrame(report_df).copy() if report_df is not None else order_review_report_mod.build_order_review_report_df(review, settings)
    if check_required_df is None:
        excluded_review_df = (
            pd.DataFrame(excluded_review).copy()
            if excluded_review is not None
            else validation_mod.apply_order_review_filters(loaders_mod.cached_order_review_df(settings, excluded_only=True, context=context), settings)
        )
        excluded_df = order_review_report_mod.build_excluded_order_review_df(excluded_review_df, settings)
        unrecognized_transport_df = transport_mod.unrecognized_transport_report_df(settings, context)
        master_unregistered_df = inbound_mod.build_master_unregistered_sku_df(settings, context)
        human_data_issue_df = order_review_report_mod.build_human_data_issue_df(settings, context)
        check_required_df = order_review_report_mod._strip_check_required_internal_columns(
            order_review_report_mod.build_check_required_sheet_df(
                pd.DataFrame(),
                unrecognized_transport_df,
                excluded_df,
                master_unregistered_df,
                pd.DataFrame(),
                human_data_issue_df,
            )
        )
    else:
        check_required_df = pd.DataFrame(check_required_df).copy()
    calendar_df = pd.DataFrame(calendar_df).copy() if calendar_df is not None else eta_mod.build_arrival_calendar_report_df(settings, context)
    stock_eta_df = (
        pd.DataFrame(stock_eta_df).copy()
        if stock_eta_df is not None
        else build_stock_eta_sheet_df(review, calendar_df, settings)
    )
    order_sheet_df = build_order_sheet_df(report_df, settings)
    order_sheet_top_sales_df = build_order_sheet_top_sales_df(report_df)

    wb = Workbook()
    ws_stock_eta = wb.active
    ws_stock_eta.title = "재고 ETA"
    append_df(ws_stock_eta, stock_eta_df)
    format_unified_esm_order_review_sheet(ws_stock_eta, settings)
    ws_order = wb.create_sheet("발주")
    append_df(ws_order, order_sheet_df, start_row=10)
    format_order_sheet(ws_order, settings, order_sheet_top_sales_df)
    ws_check_required = wb.create_sheet("확인필요")
    append_df(ws_check_required, check_required_df, start_row=8)
    order_review_report_mod.format_check_required_sheet(ws_check_required, settings, header_row=8)
    export_excel_util_mod.apply_date_picker_validation(
        ws_check_required,
        "판단일시",
        "_판단일시_선택값",
        "판단일시",
        "드롭다운에서 판단일시를 선택하거나 yyyy-mm-dd 형식으로 입력하세요.",
        header_row=8,
    )

    ws_arrival = wb.create_sheet("입고예정")
    eta_mod.write_arrival_reference_sheet(ws_arrival, settings, calendar_df)

    ws_report = wb.create_sheet("보고서")
    format_report_sheet(ws_report, settings, report_df)

    input_audit_sheet = add_input_file_audit_sheet(wb, build_timestamp, context)
    debug_sheets = add_build_and_shipping_debug_sheets(wb, settings, review, build_timestamp, context) if settings.get("include_excel_debug_sheets") else []
    season_ingredient_sheets = []
    if settings.get("include_season_sheets", False):
        try:
            from core.season_calendar import append_season_sheets
            from core.ingredient_trend import append_ingredient_sheets

            sales_history_df = loaders_mod.get_order_review_data_or_empty(
                "sales_history",
                loaders_mod.sample_sales_history,
                context,
            )
            prod_df = loaders_mod.get_order_review_data_or_empty(
                "prod_list",
                loaders_mod.sample_prod_list,
                context,
            )
            if not sales_history_df.empty:
                before_sheet_names = set(wb.sheetnames)
                append_season_sheets(wb, sales_history_df, prod_df, settings)
                append_ingredient_sheets(wb, sales_history_df, prod_df, settings)
                season_ingredient_sheets = [
                    wb[name]
                    for name in wb.sheetnames
                    if name not in before_sheet_names
                ]
        except Exception as exc:  # noqa: BLE001
            print(f"[season/ingredient sheets skipped] {exc}", flush=True)
    helper_sheets = [
        wb[name]
        for name in ["_검토일_선택값", "_판단일시_선택값", "_보고서원본", "_브랜드_선택값"]
        if name in wb.sheetnames
    ]
    wb._sheets = [
        wb["재고 ETA"],
        wb["발주"],
        wb["보고서"],
        wb["입고예정"],
        wb["확인필요"],
    ] + season_ingredient_sheets + [input_audit_sheet] + debug_sheets + helper_sheets
    enable_excel_auto_calculation(wb)
    output = BytesIO()
    apply_workbook_thousands_number_formats(wb)
    wb.save(output)
    return output.getvalue()
