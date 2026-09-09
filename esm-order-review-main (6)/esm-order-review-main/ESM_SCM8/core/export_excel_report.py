from __future__ import annotations

from copy import copy as copy_style

import numpy as np
import pandas as pd

from core.common import (
    _DEFAULT_EUR_KRW_RATE,
    ALERT_LOCAL_SHORT_LABEL,
    ALERT_SHIPPING_SHORT_LABEL,
    ALERT_URGENT_LABEL,
    DEFAULT_LEAD_TIME_DAYS,
    korea_today,
)
from core.export_excel_calculations import template_order_metrics
from core.session import SessionContext
from core import (
    eta as eta_mod,
    inventory as inventory_mod,
    kpi as kpi_mod,
    order_review as order_review_mod,
    preprocess as preprocess_mod,
)



def build_internal_work_review_df(review: pd.DataFrame, settings: dict) -> pd.DataFrame:
    eu_stock_source = preprocess_mod.prepare_eu_stock(inventory_mod.get_eu_stock())
    if "상품코드" in eu_stock_source.columns and "재고수량" in eu_stock_source.columns:
        current_stock_map = (
            pd.to_numeric(eu_stock_source["재고수량"], errors="coerce")
            .fillna(0)
            .groupby(eu_stock_source["상품코드"])
            .sum()
            .to_dict()
        )
    else:
        current_stock_map = {}
    sku = order_review_mod.report_col(review, ["상품코드"])
    recent_sales = pd.to_numeric(order_review_mod.report_col(review, ["최근 3개월 판매수량"], 0), errors="coerce").fillna(0)
    monthly_sales = pd.to_numeric(order_review_mod.report_col(review, ["월평균 판매수량"], 0), errors="coerce").fillna(0)
    daily_sales = pd.to_numeric(order_review_mod.report_col(review, ["일평균 판매수량"], 0), errors="coerce").fillna(0)
    stockout_df = build_stockout_calendar_report_df(review)
    if not stockout_df.empty and "상품코드" in stockout_df.columns and "쇼티지 예상 일수" in stockout_df.columns:
        shortage_map = stockout_df.drop_duplicates("상품코드").set_index("상품코드")["쇼티지 예상 일수"]
    else:
        shortage_map = {}
    out = pd.DataFrame(
        {
            "상품코드": sku,
            "상품명": order_review_mod.report_col(review, ["상품명"]),
            "브랜드": order_review_mod.report_col(review, ["브랜드"]),
            "제품상태": order_review_mod.report_col(review, ["제품상태"]),
            "기준 3개월 판매수량(재고 PA+CA)": recent_sales.round(0).astype(int),
            "월평균 판매수량(기준/3)": monthly_sales.round(0).astype(int),
            "일평균 판매수량(기준/90)": daily_sales.round(0).astype(int),
            "EU 현지 커버일수": pd.to_numeric(order_review_mod.report_col(review, ["EU 현지 커버일수"], 0), errors="coerce").fillna(0),
            "유럽 현재 재고": sku.map(current_stock_map).fillna(pd.to_numeric(order_review_mod.report_col(review, ["EU 현지 가용수량"], 0), errors="coerce").fillna(0)),
            "유럽 가용재고": pd.to_numeric(order_review_mod.report_col(review, ["EU 현지 가용수량", "유럽 가용재고"], 0), errors="coerce").fillna(0),
            "운송중 수량": pd.to_numeric(order_review_mod.report_col(review, ["운송중 수량"], 0), errors="coerce").fillna(0),
            "입고 전 예상 결품량": pd.to_numeric(order_review_mod.report_col(review, ["입고 전 예상 결품량"], 0), errors="coerce").fillna(0),
            "미입고 수량": pd.to_numeric(order_review_mod.report_col(review, ["미입고수량"], 0), errors="coerce").fillna(0),
            "유럽+운송 수량": pd.to_numeric(order_review_mod.report_col(review, ["유럽+운송 수량"], 0), errors="coerce").fillna(0),
            "운송 포함 보유개월": pd.to_numeric(order_review_mod.report_col(review, ["운송 포함 보유개월수"], 0), errors="coerce").fillna(0),
            "안전재고 기준 개월 수": pd.to_numeric(order_review_mod.report_col(review, ["안전재고 개월 수"], settings.get("safety_months", 3.0)), errors="coerce").fillna(settings.get("safety_months", 3.0)),
            "목표 운영 개월 수": pd.to_numeric(order_review_mod.report_col(review, ["목표 운영개월 수"], float(settings.get("safety_months", 3.0))), errors="coerce").fillna(float(settings.get("safety_months", 3.0))),
            "안전재고 필요 수량": pd.to_numeric(order_review_mod.report_col(review, ["안전재고 목표수량", "안전재고수량"], 0), errors="coerce").fillna(0),
            "부족 수량": pd.to_numeric(order_review_mod.report_col(review, ["최종 부족수량"], 0), errors="coerce").fillna(0),
            "발주 필요 수량": pd.to_numeric(order_review_mod.report_col(review, ["발주필요수량"], 0), errors="coerce").fillna(0),
            "긴급 보충 필요 수량": pd.to_numeric(order_review_mod.report_col(review, ["긴급 보충 필요 수량"], 0), errors="coerce").fillna(0),
            "입고 전 결품 위험 여부": order_review_mod.report_col(review, ["입고 전 결품 위험 여부"], "N"),
            "권장 긴급 액션": order_review_mod.report_col(review, ["권장 긴급 액션"]),
            "고갈 예정일": order_review_mod.report_col(review, ["고갈 예상일"]),
            "쇼티지 예상 일수": sku.map(shortage_map),
            "추천 운송수단": order_review_mod.report_col(review, ["운송수단 검토안"]),
            "추천 사유": order_review_mod.report_col(review, ["판단 사유"]),
            "권장 운송안": order_review_mod.report_col(review, ["시간기준_권장운송안"], "추가 운송 불필요"),
            "권장 수량 요약": order_review_mod.report_col(review, ["시간기준_권장수량요약"], "추가 운송 불필요"),
            "항공 검토수량": pd.to_numeric(order_review_mod.report_col(review, ["시간기준_항공필요수량"], 0), errors="coerce").fillna(0),
            "철송 검토수량": pd.to_numeric(order_review_mod.report_col(review, ["시간기준_철송필요수량"], 0), errors="coerce").fillna(0),
            "해운 검토수량": pd.to_numeric(order_review_mod.report_col(review, ["시간기준_해운필요수량"], 0), errors="coerce").fillna(0),
            "우선 액션": order_review_mod.report_col(review, ["우선 액션"]),
            "담당자 발주량": "",
            "발주 차이": "",
            "발주 신호등": "미입력",
            "발주 메모": "",
            "검토일": settings["base_date"],
            "담당자": "",
        }
    )
    out["추천 사유"] = out.rename(columns={"추천 운송수단": "운송 검토안", "추천 사유": "판단 사유"}).apply(order_review_mod.describe_transport_recommendation, axis=1)
    out["_입력우선순위"] = np.where(
        (pd.to_numeric(out["발주 필요 수량"], errors="coerce").fillna(0) > 0)
        | out["입고 전 결품 위험 여부"].astype(str).eq("Y"),
        0,
        1,
    )
    out["_월평균정렬"] = monthly_sales.to_numpy()
    out = out.sort_values(
        ["_입력우선순위", "발주 필요 수량", "부족 수량", "_월평균정렬"],
        ascending=[True, False, False, False],
    ).drop(columns=["_입력우선순위", "_월평균정렬"])
    return out.reset_index(drop=True)


def build_order_review_report_summary(report_df: pd.DataFrame, settings: dict | None = None) -> dict[str, object]:
    settings = settings or {}
    base = pd.DataFrame(report_df)
    safety_months = float(settings.get("safety_months", 3))
    lead_times = (
        settings.get("lead_times")
        if "lead_times" in settings
        else DEFAULT_LEAD_TIME_DAYS
    ) or {}
    eur_rate = float(settings.get("eur_krw_rate", _DEFAULT_EUR_KRW_RATE))

    metrics = template_order_metrics(base, settings)
    order_qty = pd.to_numeric(metrics["발주필요수량_계산"], errors="coerce").fillna(0)
    order_amount = pd.to_numeric(metrics["발주금액_EUR_계산"], errors="coerce").fillna(0)
    total_order_qty = float(order_qty.sum())
    total_order_eur = float(order_amount.sum())
    total_order_krw = total_order_eur * eur_rate
    total_with_open_qty = float(pd.to_numeric(metrics["미입고해소후_발주필요수량_계산"], errors="coerce").fillna(0).sum())
    total_with_open_eur = float(pd.to_numeric(metrics["미입고해소후_발주금액_EUR_계산"], errors="coerce").fillna(0).sum())
    total_with_open_krw = total_with_open_eur * eur_rate

    return {
        "safety_months": safety_months,
        "air_lt": int(lead_times["항공"]) if "항공" in lead_times else "-",
        "sea_lt": int(lead_times["해운"]) if "해운" in lead_times else "-",
        "rail_lt": int(lead_times["철송"]) if "철송" in lead_times else "-",
        "truck_lt": int(lead_times["트럭"]) if "트럭" in lead_times else "-",
        "base_date": settings.get("base_date") or korea_today(),
        "currency_code": str(settings.get("currency_code") or "EUR"),
        "eur_rate": eur_rate,
        "total_order_qty": total_order_qty,
        "total_order_eur": total_order_eur,
        "total_order_krw": total_order_krw,
        "total_with_open_qty": total_with_open_qty,
        "total_with_open_eur": total_with_open_eur,
        "total_with_open_krw": total_with_open_krw,
    }


def format_report_sheet(ws, settings: dict | None = None, report_df: pd.DataFrame | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    settings = settings or {}
    report_df = report_df if report_df is not None else pd.DataFrame()
    summary = build_order_review_report_summary(report_df, settings)
    wb = ws.parent
    source_sheet_name = "_보고서원본"
    brand_sheet_name = "_브랜드_선택값"
    source_ws = wb[source_sheet_name] if source_sheet_name in wb.sheetnames else wb.create_sheet(source_sheet_name)
    brand_ws = wb[brand_sheet_name] if brand_sheet_name in wb.sheetnames else wb.create_sheet(brand_sheet_name)
    source_ws.sheet_state = "hidden"
    brand_ws.sheet_state = "hidden"
    source_ws.delete_rows(1, max(source_ws.max_row, 1))
    brand_ws.delete_rows(1, max(brand_ws.max_row, 1))

    def sheet_ref(name: str) -> str:
        return "'" + name.replace("'", "''") + "'"

    table_side = Side(style="thin", color="000000")
    table_border = Border(left=table_side, right=table_side, top=table_side, bottom=table_side)
    dark_fill = PatternFill("solid", fgColor="00205B")
    green_fill = PatternFill("solid", fgColor="E2EFD9")
    alert_header_fill = PatternFill("solid", fgColor="FFFF00")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    param_border_side = Side(style="thin", color="000000")
    param_border = Border(left=param_border_side, right=param_border_side, top=param_border_side, bottom=param_border_side)
    param_title_fill = PatternFill("solid", fgColor="4F6228")
    param_label_fill = PatternFill("solid", fgColor="E2F0D9")
    param_value_fill = PatternFill("solid", fgColor="FFF2CC")
    param_summary_fill = PatternFill("solid", fgColor="92D050")
    number_format = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A13"
    ws.merge_cells("A10:Y10")
    ws.merge_cells("U1:Y1")
    ws.merge_cells("H11:J11")
    ws.merge_cells("K11:M11")
    ws.merge_cells("P11:S11")

    column_widths = {
        "A": 16.0, "B": 13.0, "C": 53.57, "D": 13.71, "E": 14.0,
        "F": 16.0, "G": 11.0, "H": 16.0, "I": 18.0, "J": 18.0, "K": 24.0, "L": 26.0, "M": 18.0,
        "U": 14.71, "V": 24.71, "W": 26.71, "X": 27.57, "Y": 16.0,
    }
    hidden_columns: set[str] = set("NOPQRST")
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
        ws.column_dimensions[col].hidden = col in hidden_columns
    for col in hidden_columns:
        ws.column_dimensions[col].hidden = True
        ws.column_dimensions[col].width = 1
    for row_idx in range(1, 13):
        ws.row_dimensions[row_idx].height = 18
    ws.row_dimensions[10].height = 4.5
    ws.row_dimensions[11].hidden = True

    selector_fill = PatternFill("solid", fgColor="4F6228")
    selector_value_fill = PatternFill("solid", fgColor="E2F0D9")
    ws.merge_cells("B2:G2")
    ws.merge_cells("C3:G3")
    ws["B2"].value = "브랜드 조회"
    ws["B2"].fill = selector_fill
    ws["B2"].font = Font(color="FFFFFF", bold=True)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
    ws["B2"].border = param_border
    ws["B3"].value = "브랜드 검색"
    ws["B3"].fill = selector_value_fill
    ws["B3"].font = Font(bold=True, color="375623")
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B3"].border = param_border
    ws["C3"].value = "전체"
    ws["C3"].fill = selector_value_fill
    ws["C3"].font = Font(bold=True, color="375623")
    ws["C3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["C3"].border = param_border
    for col_idx in range(4, 6):
        cell = ws.cell(3, col_idx)
        cell.fill = selector_value_fill
        cell.border = param_border
    ws.merge_cells("B4:G4")
    ws["B4"].value = "검색 후보"
    ws["B4"].fill = selector_fill
    ws["B4"].font = Font(color="FFFFFF", bold=True)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center")
    ws["B4"].border = param_border
    for row_idx in range(5, 10):
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=7)
        cell = ws.cell(row_idx, 2)
        cell.fill = selector_value_fill
        cell.font = Font(color="375623")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = param_border
        for col_idx in range(3, 8):
            merged_cell = ws.cell(row_idx, col_idx)
            merged_cell.fill = selector_value_fill
            merged_cell.border = param_border
    for row_idx in range(2, 10):
        for col_idx in range(8, 21):
            cell = ws.cell(row_idx, col_idx)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.font = Font()
            cell.alignment = Alignment()

    currency_code = str(summary.get("currency_code") or "EUR")
    param_cells = {
        "U1": "▼ 파라미터 설정  | 솔팅 가능",
        "U2": "안전재고(M)", "V2": "AirL/T(일)", "W2": "해운L/T(일)", "X2": "철송L/T(일)", "Y2": "트럭L/T(일)",
        "U3": summary["safety_months"], "V3": summary["air_lt"], "W3": summary["sea_lt"], "X3": summary["rail_lt"], "Y3": summary["truck_lt"],
        "U4": "기준일", "V4": f"환율(KRW/{currency_code})",
        "U5": summary["base_date"], "V5": summary["eur_rate"],
        "V6": "총 발주 필요수량", "W6": f"총 발주금액({currency_code})", "X6": "총 발주 금액(KRW)",
        "V7": summary["total_order_qty"], "W7": summary["total_order_eur"], "X7": summary["total_order_krw"],
        "V8": "미입고 해소 후 총 발주 필요수량", "W8": f"미입고 해소 후 총 발주금액({currency_code})", "X8": "(미입고 해소 후)총 발주 금액(KRW)",
        "V9": summary["total_with_open_qty"], "W9": summary["total_with_open_eur"], "X9": summary["total_with_open_krw"],
    }
    for row_idx in range(1, 10):
        for col_idx in range(21, 26):
            cell = ws.cell(row_idx, col_idx)
            cell.fill = param_label_fill
            cell.border = param_border
            cell.font = Font(bold=True, color="375623")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    for coord, value in param_cells.items():
        cell = ws[coord]
        cell.value = value
        if coord in {"U5"}:
            cell.number_format = "yyyy-mm-dd"
        elif coord in {"V5"}:
            cell.number_format = "#,##0.00"
        elif coord in {"W7", "W9"}:
            cell.number_format = "#,##0.00"
        elif coord in {"X7", "X9"}:
            cell.number_format = "#,##0"
        elif coord[0] in {"V", "X"} and coord[1:] in {"7", "9"}:
            cell.number_format = "#,##0"
    for col_idx in range(21, 26):
        cell = ws.cell(1, col_idx)
        cell.fill = param_title_fill
        cell.border = param_border
    ws["U1"].font = Font(color="FFFFFF", bold=True)
    ws["U1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    for coord in ("U3", "V3", "W3", "X3", "Y3"):
        ws[coord].fill = param_value_fill
    for row_idx in (6, 8):
        ws.cell(row_idx, 21).fill = white_fill
        for col_idx in range(22, 25):
            cell = ws.cell(row_idx, col_idx)
            cell.fill = param_summary_fill
            cell.font = Font(bold=True, color="000000", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    for row_idx in (7, 9):
        ws.cell(row_idx, 21).fill = white_fill
        for col_idx in range(22, 25):
            cell = ws.cell(row_idx, col_idx)
            cell.fill = param_label_fill
            cell.font = Font(bold=True, color="375623")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    for row_idx in range(6, 10):
        cell = ws.cell(row_idx, 21)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()
        cell.font = Font()
        cell.alignment = Alignment()
    for col_idx in range(21, 26):
        cell = ws.cell(1, col_idx)
        cell.fill = param_title_fill
        cell.border = param_border

    headers = [
        "코드", "No", "SKU", "브랜드", "미입고현황", "현지 창고재고",
        "⚠ 알람", "발주 필요수량", f"=발주 금액({currency_code})", "=발주 금액(KRW)",
        "(미입고 해소 후)발주 필요수량", f"=(미입고 해소 후)발주 금액({currency_code})", "=(미입고 해소 후)발주 금액(KRW)",
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(12, col_idx, header)
        if isinstance(header, str) and header.startswith("="):
            cell.data_type = "s"
        if header == "⚠ 알람":
            cell.fill = alert_header_fill
        else:
            cell.fill = green_fill
        cell.font = Font(bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = table_border
        source_cell = source_ws.cell(1, col_idx, header)
        if isinstance(header, str) and header.startswith("="):
            source_cell.data_type = "s"
        source_cell.fill = copy_style(cell.fill)
        source_cell.font = copy_style(cell.font)
        source_cell.alignment = copy_style(cell.alignment)
        source_cell.border = copy_style(cell.border)
    ws.row_dimensions[12].height = 36

    base = pd.DataFrame(report_df).reset_index(drop=True)
    sku = order_review_mod.report_col(base, ["상품코드", "SKU", "sku"], "")
    product_name = order_review_mod.report_col(base, ["상품명", "SKU명", "제품명"], "")
    brand = order_review_mod.report_col(base, ["브랜드"], "").fillna("").astype(str).str.strip()
    sales_3m = pd.to_numeric(order_review_mod.report_col(base, ["기준_3M_판매수량", "최근 3개월 수요", "PA+CA 판매수량"], 0), errors="coerce").fillna(0)
    metrics = template_order_metrics(base, settings)
    monthly = metrics["월평균_계산"]
    safety_stock = metrics["안전재고_계산"]
    eu_stock = metrics["유럽재고_계산"]
    shipping_stock = metrics["운송재고_계산"]
    eu_shipping = metrics["유럽운송합산_계산"]
    open_po = pd.to_numeric(order_review_mod.report_col(base, ["미입고 현황", "미입고 수량", "미입고"], 0), errors="coerce").fillna(0)
    hq_stock = pd.to_numeric(order_review_mod.report_col(base, ["본사 EU창고"], 0), errors="coerce").fillna(0)
    eur_rate = float(summary["eur_rate"])
    order_qty = metrics["발주필요수량_계산"]
    order_amount_eur = metrics["발주금액_EUR_계산"]
    order_amount_krw = metrics["발주금액_KRW_계산"]
    order_with_open_qty = metrics["미입고해소후_발주필요수량_계산"]
    order_with_open_eur = metrics["미입고해소후_발주금액_EUR_계산"]
    order_with_open_krw = metrics["미입고해소후_발주금액_KRW_계산"]
    rounded_monthly = monthly.round(0)
    local_months = pd.Series(
        np.where(rounded_monthly.gt(0), eu_stock / rounded_monthly.replace(0, np.nan), 0),
        index=base.index,
    ).round(1)
    shipping_months = pd.Series(
        np.where(rounded_monthly.gt(0), shipping_stock / rounded_monthly.replace(0, np.nan), 0),
        index=base.index,
    ).round(1)
    alert_values = pd.Series("-", index=base.index, dtype=object)
    has_sales = rounded_monthly.gt(0)
    combined_months = local_months + shipping_months
    safety_months = float(summary["safety_months"])
    alert_values.loc[has_sales & combined_months.lt(safety_months)] = ALERT_URGENT_LABEL
    alert_values.loc[has_sales & combined_months.ge(safety_months) & combined_months.lt(safety_months * 1.5)] = ALERT_LOCAL_SHORT_LABEL
    target_stock = safety_stock
    available_stock = eu_shipping
    brand_values = ["전체"] + sorted(
        {
            str(value).strip()
            for value in brand.fillna("").astype(str).tolist()
            if str(value).strip() and str(value).strip() not in {"-", "nan", "None"}
        }
    )
    for idx, value in enumerate(brand_values, start=1):
        brand_ws.cell(idx, 1, value)
    brand_ws.cell(1, 2, "_검색순번")
    brand_ws.column_dimensions["B"].hidden = True
    brand_last_row = max(len(brand_values), 2)
    for idx in range(2, brand_last_row + 1):
        if idx > len(brand_values):
            brand_ws.cell(idx, 1, "")
        brand_ws.cell(
            idx,
            2,
            (
                f'=IF(OR(TRIM({sheet_ref(ws.title)}!$C$3)="전체",'
                f'TRIM({sheet_ref(ws.title)}!$C$3)="",'
                f'TRIM($A{idx})=""),"",'
                f'IF(ISNUMBER(SEARCH(SUBSTITUTE(TRIM({sheet_ref(ws.title)}!$C$3)," ",""),'
                f'SUBSTITUTE(TRIM($A{idx})," ",""))),COUNT($B$1:B{idx - 1})+1,""))'
            ),
        )
    brand_suggestion_value_range = f"{sheet_ref(brand_sheet_name)}!$A$2:$A${brand_last_row}"
    brand_suggestion_helper_range = f"{sheet_ref(brand_sheet_name)}!$B$2:$B${brand_last_row}"
    for row_idx in range(5, 10):
        ws.cell(
            row_idx,
            2,
            (
                f'=IFERROR(INDEX({brand_suggestion_value_range},'
                f'MATCH(ROWS($B$5:B{row_idx}),{brand_suggestion_helper_range},0)),"")'
            ),
        )
    data_start = 13
    src_last_row = len(base) + 1
    source_helper_col_idx = len(headers) + 1
    source_helper_col = get_column_letter(source_helper_col_idx)
    source_ws.cell(1, source_helper_col_idx, "_브랜드조회순번")
    source_ws.column_dimensions[source_helper_col].hidden = True
    fast_data_style = len(base) > 300
    left_wrap_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_alignment = Alignment(horizontal="right", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")
    for idx, row in base.iterrows():
        excel_row = data_start + idx
        source_row = idx + 2
        row_values = [
            sku.iloc[idx],
            idx + 1 if str(sku.iloc[idx]).strip() else "",
            product_name.iloc[idx],
            brand.iloc[idx],
            open_po.iloc[idx],
            hq_stock.iloc[idx],
            alert_values.iloc[idx],
            order_qty.iloc[idx],
            order_amount_eur.iloc[idx],
            order_amount_krw.iloc[idx],
            order_with_open_qty.iloc[idx],
            order_with_open_eur.iloc[idx],
            order_with_open_krw.iloc[idx],
        ]
        source_ws.cell(
            source_row,
            source_helper_col_idx,
            (
                f'=IF(OR(TRIM({sheet_ref(ws.title)}!$C$3)="전체",'
                f'TRIM({sheet_ref(ws.title)}!$C$3)="",'
                f'ISNUMBER(SEARCH(SUBSTITUTE(TRIM({sheet_ref(ws.title)}!$C$3)," ",""),'
                f'SUBSTITUTE(TRIM($D{source_row})," ","")))),'
                f'COUNT(${source_helper_col}$1:{source_helper_col}{source_row - 1})+1,"")'
            ),
        )
        for col_idx, value in enumerate(row_values, 1):
            source_ws.cell(source_row, col_idx, value)
            source_col = get_column_letter(col_idx)
            cell = ws.cell(
                excel_row,
                col_idx,
                (
                    f'=IFERROR(INDEX({sheet_ref(source_sheet_name)}!${source_col}$2:${source_col}${src_last_row},'
                    f'MATCH(ROWS($A${data_start}:A{excel_row}),'
                    f'{sheet_ref(source_sheet_name)}!${source_helper_col}$2:${source_helper_col}${src_last_row},0)),"")'
                ),
            )
            cell.border = table_border
            cell.alignment = left_wrap_alignment if col_idx in {1, 3, 4} else center_alignment if col_idx == 7 else right_alignment
            if col_idx == 7:
                cell.number_format = "@"
            elif col_idx in {5, 6, 8, 11}:
                cell.number_format = number_format
            elif col_idx in {9, 12}:
                cell.number_format = "#,##0.00"
            elif col_idx in {10, 13}:
                cell.number_format = "#,##0"

    last_row = max(data_start, data_start + len(base) - 1)
    source_ws.sheet_view.showGridLines = False
    if not fast_data_style:
        for row in range(data_start, last_row + 1):
            ws.row_dimensions[row].height = 18
    ws.auto_filter.ref = f"A12:M{last_row}"
    if len(base) > 0:
        report_table = Table(displayName="ReportBrandFilter", ref=f"A12:M{last_row}")
        report_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        ws.add_table(report_table)
    source_ref = sheet_ref(source_sheet_name)
    source_helper_range = f"{source_ref}!${source_helper_col}$2:${source_helper_col}${src_last_row}"

    def brand_total_formula(sum_range: str) -> str:
        return f'=SUMIF({source_helper_range},">0",{sum_range})'

    visible_total_formulas = {
        "V7": brand_total_formula(f"{source_ref}!$H$2:$H${src_last_row}"),
        "W7": brand_total_formula(f"{source_ref}!$I$2:$I${src_last_row}"),
        "X7": brand_total_formula(f"{source_ref}!$J$2:$J${src_last_row}"),
        "V9": brand_total_formula(f"{source_ref}!$K$2:$K${src_last_row}"),
        "W9": brand_total_formula(f"{source_ref}!$L$2:$L${src_last_row}"),
        "X9": brand_total_formula(f"{source_ref}!$M$2:$M${src_last_row}"),
    }
    for coord, formula in visible_total_formulas.items():
        ws[coord].value = formula
    for row_idx in range(1, last_row + 1):
        for col_idx in range(1, 26):
            ws.cell(row_idx, col_idx).protection = Protection(locked=False)
    for coord in visible_total_formulas:
        ws[coord].protection = Protection(locked=True, hidden=True)
    for row_idx in range(data_start, last_row + 1):
        ws.cell(row_idx, 7).protection = Protection(locked=True, hidden=True)
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def build_monthly_arrival_summary_df(settings: dict, context: SessionContext | None = None) -> pd.DataFrame:
    calendar_df = eta_mod.build_arrival_calendar_report_df(settings, context)
    if calendar_df.empty:
        return pd.DataFrame(
            columns=[
                "도착 예정월", "해운 수량", "항공 수량", "철송 수량", "트럭 수량", "총 도착 수량",
                "총 도착 금액_KRW", "누적 도착 수량", "누적 도착 금액_KRW",
            ]
        )
    out = calendar_df.copy()
    out["수량"] = pd.to_numeric(out["수량"], errors="coerce").fillna(0)
    out["도착 예정 금액_KRW"] = pd.to_numeric(out.get("도착 예정 금액_KRW", 0), errors="coerce").fillna(0)
    out = out[out["도착 예정월"].notna()]
    rows = []
    for month, group in out.groupby("도착 예정월"):
        row: dict[str, object] = {"도착 예정월": month}
        total_qty = 0.0
        for mode in ["해운", "항공", "철송", "트럭"]:
            mode_group = group[group["운송수단"].astype(str).eq(mode)]
            qty = float(mode_group["수량"].sum())
            row[f"{mode} 수량"] = qty
            total_qty += qty
        row["총 도착 수량"] = total_qty
        row["총 도착 금액_KRW"] = float(group["도착 예정 금액_KRW"].sum())
        rows.append(row)
    summary = pd.DataFrame(rows).sort_values("도착 예정월").reset_index(drop=True)
    summary["누적 도착 수량"] = summary["총 도착 수량"].cumsum()
    summary["누적 도착 금액_KRW"] = summary["총 도착 금액_KRW"].cumsum()
    columns = [
        "도착 예정월", "해운 수량", "항공 수량", "철송 수량", "트럭 수량", "총 도착 수량",
        "총 도착 금액_KRW", "누적 도착 수량", "누적 도착 금액_KRW",
    ]
    return summary[columns]


def build_stockout_calendar_report_df(review: pd.DataFrame) -> pd.DataFrame:
    out = pd.DataFrame(
        {
            "상품코드": order_review_mod.report_col(review, ["상품코드"]),
            "상품명": order_review_mod.report_col(review, ["상품명"]),
            "현재 가용재고": pd.to_numeric(order_review_mod.report_col(review, ["EU 현지 가용수량"], 0), errors="coerce").fillna(0),
            "월평균 판매량": pd.to_numeric(order_review_mod.report_col(review, ["월평균 판매수량"], 0), errors="coerce").fillna(0),
            "일평균 판매량": pd.to_numeric(order_review_mod.report_col(review, ["일평균 판매수량"], 0), errors="coerce").fillna(0),
            "예상 소진일": order_review_mod.report_col(review, ["고갈 예상일"]),
            "첫 입고 예정일": order_review_mod.report_col(review, ["최초 ETA"]),
            "추천 운송수단": order_review_mod.report_col(review, ["운송수단 검토안"]),
            "우선 액션": order_review_mod.report_col(review, ["우선 액션"]),
            "판단 사유": order_review_mod.report_col(review, ["판단 사유"]),
        }
    )
    depletion = kpi_mod.parse_date_series(out["예상 소진일"])
    first_eta = kpi_mod.parse_date_series(out["첫 입고 예정일"])
    raw_gap = (first_eta - depletion).dt.days
    out["쇼티지 예상 일수"] = np.where(first_eta.notna() & depletion.notna(), np.maximum(raw_gap.fillna(0), 0), np.nan)
    action_risk = out["우선 액션"].astype(str).str.contains("ETA 지연 위험", na=False)
    out["입고 전 품절 여부"] = np.where((pd.to_numeric(out["쇼티지 예상 일수"], errors="coerce").fillna(0) > 0) | action_risk, "Y", "")
    out["지연 여부"] = np.where(out["입고 전 품절 여부"].eq("Y"), "지연/품절위험", "")
    out["추천 대응"] = [
        order_review_mod.recommended_response_from_action(action, recommendation)
        for action, recommendation in zip(out["우선 액션"], out["추천 운송수단"])
    ]
    out["운송수단 추천 사유"] = out.apply(order_review_mod.describe_transport_recommendation, axis=1)
    out["_위험정렬"] = out["입고 전 품절 여부"].astype(str).eq("Y").astype(int)
    out["_공백일수정렬"] = pd.to_numeric(out["쇼티지 예상 일수"], errors="coerce").fillna(0)
    out["_예상소진일정렬"] = kpi_mod.parse_date_series(out["예상 소진일"])
    out["_월평균정렬"] = pd.to_numeric(out["월평균 판매량"], errors="coerce").fillna(0)
    out = out.sort_values(
        ["_위험정렬", "_공백일수정렬", "_월평균정렬", "_예상소진일정렬"],
        ascending=[False, False, False, True],
    )
    return out[
        [
            "상품코드", "상품명", "현재 가용재고", "월평균 판매량", "일평균 판매량", "예상 소진일",
            "첫 입고 예정일", "쇼티지 예상 일수", "입고 전 품절 여부", "지연 여부", "추천 대응",
            "추천 운송수단", "운송수단 추천 사유", "우선 액션",
        ]
    ].reset_index(drop=True)


def pre_arrival_shortage_logic_summary_text(report_df: pd.DataFrame) -> str:
    report = pd.DataFrame(report_df).copy()
    if report.empty:
        return "[입고 전 결품 위험 로직 요약]\n- 전체 SKU 수: 0"

    action = order_review_mod.report_col(report, ["우선 액션", "최종 액션"]).astype(str)
    urgent_qty = pd.to_numeric(order_review_mod.report_col(report, ["긴급 보충 필요 수량"], 0), errors="coerce").fillna(0)
    shortage_flag = order_review_mod.report_col(report, ["입고 전 결품 위험 여부"], "N").astype(str).eq("Y")
    order_qty = pd.to_numeric(order_review_mod.report_col(report, ["발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0)
    order_qty_alias = pd.to_numeric(order_review_mod.report_col(report, ["추가 발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0)
    order_amount = pd.to_numeric(order_review_mod.report_col(report, ["발주 필요 금액", "발주필요금액(KRW)", "발주필요금액_KRW", "추가 발주 필요 금액"], 0), errors="coerce").fillna(0)
    order_amount_alias = pd.to_numeric(order_review_mod.report_col(report, ["추가 발주 필요 금액", "발주필요금액_KRW", "발주필요금액"], 0), errors="coerce").fillna(0)
    exact_sku = order_review_mod.report_col(report, ["상품코드"]).astype(str).str.strip()
    ans_mask = exact_sku.eq("ANS10-MPDRN100EU")
    if ans_mask.any():
        ans_idx = ans_mask[ans_mask].index[0]
        ans_qty = urgent_qty.loc[ans_idx]
        ans_action = action.loc[ans_idx]
    else:
        ans_qty = "-"
        ans_action = "-"
    ans_qty_text = ans_qty if isinstance(ans_qty, str) else f"{ans_qty:,.0f}"

    qty_unchanged = bool(np.allclose(order_qty.fillna(0), order_qty_alias.fillna(0)))
    amount_unchanged = bool(np.allclose(order_amount.fillna(0), order_amount_alias.fillna(0)))
    no_order_shortage_count = int((action.eq("발주 불필요") & shortage_flag).sum())
    order_shortage_count = int((action.eq("발주 필요") & shortage_flag).sum())
    lines = [
        "[입고 전 결품 위험 로직 요약]",
        f"- 전체 SKU 수: {len(report):,}",
        f"- 기존 발주 필요 수량 합계: {order_qty.sum():,.0f}",
        f"- 기존 발주 필요 금액 합계: {order_amount.sum():,.0f}",
        f"- 기존 발주 필요 수량 산식 컬럼 일치: {'Y' if qty_unchanged else 'N'}",
        f"- 기존 발주 필요 금액 산식 컬럼 일치: {'Y' if amount_unchanged else 'N'}",
        f"- 입고 전 결품 위험 Y 건수: {int(shortage_flag.sum()):,}",
        f"- 긴급 보충 필요 수량 합계: {urgent_qty.sum():,.0f}",
        f"- 발주 불필요 중 입고 전 결품 위험 Y: {no_order_shortage_count:,}",
        f"- 발주 필요 중 입고 전 결품 위험 Y: {order_shortage_count:,}",
        f"- ANS10-MPDRN100EU 긴급 보충 필요 수량: {ans_qty_text}",
        f"- ANS10-MPDRN100EU 우선 액션: {ans_action}",
    ]
    return "\n".join(lines)


def internal_review_signal_counts(internal_review_df: pd.DataFrame) -> dict[str, int]:
    required = pd.to_numeric(internal_review_df.get("발주 필요 수량", pd.Series(dtype=object)), errors="coerce")
    manager_raw = internal_review_df.get("담당자 발주량", pd.Series([""] * len(internal_review_df), index=internal_review_df.index))
    manager = pd.to_numeric(manager_raw, errors="coerce")
    manager_blank = manager_raw.isna() | manager_raw.astype(str).str.strip().eq("")

    counts = {"발주불필요": 0, "미입력": 0, "안전": 0, "부족": 0, "확인필요": 0}
    invalid_required = required.isna()
    valid_required = required.notna()
    manager_numeric = manager.notna()
    valid_manager = ~manager_blank & manager_numeric & (manager >= 0)

    counts["확인필요"] += int(invalid_required.sum())
    counts["발주불필요"] = int((valid_required & (required <= 0) & manager_blank).sum())
    counts["미입력"] = int((valid_required & (required > 0) & manager_blank).sum())
    counts["안전"] = int((valid_required & valid_manager & (manager >= required)).sum())
    counts["부족"] = int((valid_required & valid_manager & (manager < required)).sum())
    counts["확인필요"] += int((valid_required & ~manager_blank & (~manager_numeric | (manager < 0))).sum())
    return counts
