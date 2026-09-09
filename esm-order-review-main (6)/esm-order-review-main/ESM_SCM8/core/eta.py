from __future__ import annotations

import calendar
import html as html_lib

import numpy as np
import pandas as pd
from core.session import SessionContext

from core.common import (
    _DEFAULT_EUR_KRW_RATE,
    ARRIVAL_REFERENCE_COLUMNS,
    ARRIVAL_REFERENCE_DETAIL_COLUMNS,
    ARRIVAL_REFERENCE_MODES,
    COLORS,
    korea_today,
    MODE_COLORS,
    RISK_BG,
    TRANSPORT_REVIEW_REQUIRED,
)
from core.export_excel_util import append_df, display_text
from core.lead_times import TRANSPORT_GROUP_DISPLAY_NAMES
from datetime import date
from numbers import Number
from core import inbound as inbound_mod, kpi as kpi_mod, order_review as order_review_mod, order_review_report as order_review_report_mod, preprocess as preprocess_mod, transport as transport_mod, validation as validation_mod

def tooltip_html(rows: pd.DataFrame, eta_date: str) -> str:
    detail_rows = []
    for _, row in rows.iterrows():
        qty = pd.to_numeric(row.get("수량", 0), errors="coerce")
        amount = pd.to_numeric(row.get("도착 예정 금액(M원)", 0), errors="coerce")
        detail_rows.append(
            '<div class="cal-tooltip-item">'
            f'<b>{html_lib.escape(str(row.get("운송수단", "-")))}</b> · {html_lib.escape(str(row.get("브랜드", "-")))}'
            f'<br><b>SKU</b> {html_lib.escape(str(row.get("SKU", "-")))}'
            f'<br><b>상품명</b> {html_lib.escape(str(row.get("상품명", "-")))}'
            f'<br><b>수량</b> {0 if pd.isna(qty) else qty:,.0f} · <b>입고금액</b> {format_calendar_amount_million(amount)}'
            f'<br><b>출고일</b> {html_lib.escape(str(row.get("출고일", "-")))} · <b>적용 리드타임</b> {html_lib.escape(str(row.get("리드타임", "-")))}'
            f'<br><b>위험도</b> {html_lib.escape(str(row.get("위험도", "-")))}'
            '</div>'
        )
    return (
        '<div class="cal-tooltip">'
        f'<div class="cal-tooltip-title">ETA {html_lib.escape(eta_date)}</div>'
        f'{"".join(detail_rows)}'
        '</div>'
    )


def add_shipping_arrival_amount_columns(shipping: pd.DataFrame, settings: dict) -> pd.DataFrame:
    out = shipping.copy()
    zero = pd.Series(0, index=out.index, dtype="float64")
    qty = pd.to_numeric(out.get("수량", zero), errors="coerce").fillna(0)
    eur_amount = pd.to_numeric(out.get("금액", zero), errors="coerce").fillna(0)
    unit_krw = pd.to_numeric(out.get("입고가(KRW)", zero), errors="coerce").fillna(0)
    fallback_krw = unit_krw * qty
    if "CMS 원화 환산금액" in out.columns:
        # CMS 응답에서는 인보이스 계약환율로 계산된 값을 그대로 사용한다.
        # 값이 비어 있어도 현재 환율로 재환산하지 않는 것이 안전하다.
        amount_krw = pd.to_numeric(
            out["CMS 원화 환산금액"], errors="coerce"
        ).fillna(0)
    else:
        # 업로드 파일 등 구 계약에는 API 원화 필드가 없으므로 기존 fallback을
        # 유지한다. CMS mapped frame에는 항상 위 컬럼이 존재한다.
        amount_krw = np.where(
            eur_amount > 0,
            eur_amount * float(settings.get("eur_krw_rate", _DEFAULT_EUR_KRW_RATE)),
            fallback_krw,
        )
    amount_krw = np.where(qty > 0, amount_krw, 0)
    out["도착 예정 금액_KRW"] = amount_krw
    out["도착 예정 금액(M원)"] = pd.Series(amount_krw, index=out.index) / 1_000_000
    return out


def format_calendar_amount_million(value: object) -> str:
    amount = pd.to_numeric(value, errors="coerce")
    if pd.isna(amount):
        amount = 0
    amount = float(amount)
    return f"{amount:,.1f}M" if abs(amount) < 10 and amount != 0 else f"{amount:,.0f}M"


def tooltip_title(rows: pd.DataFrame, eta_date: str) -> str:
    lines = [f"ETA 날짜: {eta_date}"]
    for _, row in rows.iterrows():
        qty = pd.to_numeric(row.get("수량", 0), errors="coerce")
        amount = pd.to_numeric(row.get("도착 예정 금액(M원)", 0), errors="coerce")
        lines.extend(
            [
                "",
                f"운송수단: {row.get('운송수단', '-')}",
                f"브랜드: {row.get('브랜드', '-')}",
                f"SKU: {row.get('SKU', '-')}",
                f"상품명: {row.get('상품명', '-')}",
                f"수량: {0 if pd.isna(qty) else qty:,.0f}",
                f"입고금액: {format_calendar_amount_million(amount)}",
                f"출고일: {row.get('출고일', '-')}",
                f"적용 리드타임: {row.get('리드타임', '-')}",
                f"위험도: {row.get('위험도', '-')}",
            ]
        )
    return html_lib.escape("\n".join(lines), quote=True).replace("\n", "&#10;")


def calendar_html(shipping: pd.DataFrame, year: int, month: int, today_day: int = 11, selected_day: int = 20) -> str:
    month_df = shipping[pd.to_datetime(shipping["ETA"]).dt.to_period("M") == pd.Period(f"{year}-{month:02d}")]
    grouped = (
        month_df.groupby(["ETA"], as_index=False)
        .agg({"도착 예정 금액(M원)": "sum", "SKU": "nunique"})
        .sort_values("ETA")
    )
    events: dict[int, pd.Series] = {}
    details: dict[int, pd.DataFrame] = {}
    for _, row in grouped.iterrows():
        day = pd.to_datetime(row["ETA"]).day
        events[day] = row
        details[day] = month_df[month_df["ETA"] == row["ETA"]].copy()

    weeks = calendar.Calendar(firstweekday=0).monthdayscalendar(year, month)
    html_rows = []
    for week in weeks:
        cells = []
        for col_idx, day in enumerate(week):
            if day == 0:
                cells.append('<td><div class="day-num subtle">&nbsp;</div></td>')
                continue
            event = events.get(day)
            classes = []
            if day == today_day:
                classes.append("cal-today")
            style = ""
            event_html = ""
            if event is not None:
                day_detail = details.get(day, pd.DataFrame())
                risk = day_detail["위험도"].iloc[0] if not day_detail.empty and "위험도" in day_detail.columns else ""
                style = f' style="background:{RISK_BG.get(risk, "#ffffff")};"'
                tooltip = tooltip_html(day_detail, str(event["ETA"]))
                amount_text = format_calendar_amount_million(event["도착 예정 금액(M원)"])
                mode_tags = []
                for mode in day_detail["운송수단"].dropna().astype(str).unique().tolist() if "운송수단" in day_detail.columns else []:
                    dot = MODE_COLORS.get(mode, "#98a2b3")
                    mode_tags.append(
                        f'<span class="cal-mode-tag" title="{html_lib.escape(mode)}">'
                        f'<span class="dot" style="background:{dot}"></span>{html_lib.escape(mode)}</span>'
                    )
                tooltip_side_class = " cal-tooltip-left" if col_idx >= 4 else ""
                event_html = (
                    f'<div class="cal-tooltip-wrap{tooltip_side_class}">'
                    '<div class="cal-event">'
                    f'<div class="cal-mode-row">{"".join(mode_tags)}</div>'
                    f'<span class="cal-event-amount">{amount_text}</span>'
                    f'<small>{int(event["SKU"])} SKU</small>'
                    "</div>"
                    f"{tooltip}"
                    "</div>"
                )
            day_color = COLORS["red"] if day == today_day else COLORS["text"]
            cells.append(f'<td class="{" ".join(classes)}"{style}><div class="day-num" style="color:{day_color}">{day}</div>{event_html}</td>')
        html_rows.append("<tr>" + "".join(cells) + "</tr>")

    return f"""
    <table class="calendar">
        <thead><tr><th>월</th><th>화</th><th>수</th><th>목</th><th>금</th><th>토</th><th>일</th></tr></thead>
        <tbody>{''.join(html_rows)}</tbody>
    </table>
    """


def calendar_selected_detail_html(shipping: pd.DataFrame, selected_dt: pd.Timestamp) -> str:
    selected_date = selected_dt.strftime("%Y-%m-%d")
    detail = shipping[shipping["ETA"].astype(str).eq(selected_date)].copy()
    if detail.empty:
        return ""
    rows = []
    for _, row in detail.iterrows():
        mode = html_lib.escape(str(row.get("운송수단", "-")))
        brand = html_lib.escape(str(row.get("브랜드", "-")))
        sku = html_lib.escape(str(row.get("SKU", "-")))
        name = html_lib.escape(str(row.get("상품명", "-")))
        risk = html_lib.escape(str(row.get("위험도", "-")))
        ship_date = html_lib.escape(str(row.get("출고일", "-")))
        lead_time = html_lib.escape(str(row.get("리드타임", "-")))
        qty = pd.to_numeric(row.get("수량", 0), errors="coerce")
        amount = pd.to_numeric(row.get("도착 예정 금액(M원)", 0), errors="coerce")
        rows.append(
            "<div class=\"cal-detail-item\">"
            f"<b>{mode}</b> · {brand}<br>"
            f"<b>SKU</b> {sku}<br>"
            f"<b>상품명</b> {name}<br>"
            f"<b>수량</b> {0 if pd.isna(qty) else qty:,.0f} · "
            f"<b>입고금액</b> {format_calendar_amount_million(amount)}<br>"
            f"<b>출고일</b> {ship_date} · <b>적용 리드타임</b> {lead_time} · <b>위험도</b> {risk}"
            "</div>"
        )
    return (
        '<div class="cal-detail-panel">'
        f'<div class="cal-detail-header">{selected_date} ETA 상세 ({len(detail)}건)</div>'
        f"{''.join(rows)}"
        "</div>"
    )



def detail_for_date(
    settings: dict,
    selected_date: str = "2026-05-20",
    context: SessionContext | None = None,
) -> pd.DataFrame:
    df = add_shipping_arrival_amount_columns(validation_mod.apply_common_filters(transport_mod.get_shipping(settings, context), settings), settings)
    amount_eur_source = df["금액"] if "금액" in df.columns else pd.Series([0] * len(df), index=df.index)
    amount_eur = pd.to_numeric(amount_eur_source, errors="coerce").fillna(0)
    df["금액(EUR)"] = amount_eur
    columns = [
        "운송수단", "브랜드", "SKU", "상품명", "수량", "금액(EUR)", "도착 예정 금액(M원)",
        "출고일", "ETA", "리드타임", "운송 L/T", "위험도",
    ]
    return df[df["ETA"] == selected_date][[col for col in columns if col in df.columns]].copy()


def eta_date_options(settings: dict, context: SessionContext | None = None) -> list[str]:
    shipping = validation_mod.apply_common_filters(transport_mod.get_shipping(settings, context), settings)
    dates = pd.to_datetime(shipping["ETA"], errors="coerce").dropna().sort_values()
    return dates.dt.strftime("%Y-%m-%d").drop_duplicates().tolist()


def default_eta_detail_date(
    settings: dict,
    preferred_date: str = "2026-05-20",
    context: SessionContext | None = None,
) -> str | None:
    options = eta_date_options(settings, context)
    if not options:
        return None
    base = pd.Timestamp(settings["base_date"])
    base_text = base.strftime("%Y-%m-%d")
    if base_text in options:
        return base_text
    if preferred_date in options and pd.Timestamp(preferred_date) >= base:
        return preferred_date
    upcoming = [option for option in options if pd.Timestamp(option) >= base]
    return upcoming[0] if upcoming else options[0]


def eta_detail_display_df(detail: pd.DataFrame) -> pd.DataFrame:
    columns = ["SKU", "상품명", "수량", "금액(M원)", "금액(EUR)", "운송수단", "출고일", "적용 L/T", "위험도"]
    if detail.empty:
        return pd.DataFrame(columns=columns)

    work = detail.copy()
    work["수량"] = pd.to_numeric(work.get("수량", 0), errors="coerce").fillna(0)
    work["금액(EUR)"] = pd.to_numeric(work.get("금액(EUR)", 0), errors="coerce").fillna(0)
    work["도착 예정 금액(M원)"] = pd.to_numeric(work.get("도착 예정 금액(M원)", 0), errors="coerce").fillna(0)
    work["리드타임"] = pd.to_numeric(work.get("리드타임", 0), errors="coerce").fillna(0)
    if "상품명" not in work.columns:
        work["상품명"] = ""
    group_cols = [col for col in ["SKU", "운송수단", "위험도"] if col in work.columns]
    if not group_cols:
        return pd.DataFrame(columns=columns)

    grouped = (
        work.groupby(group_cols, dropna=False)
        .agg(
            상품명=("상품명", inbound_mod.first_non_empty),
            수량=("수량", "sum"),
            금액_M원=("도착 예정 금액(M원)", "sum"),
            금액_EUR=("금액(EUR)", "sum"),
            출고일=("출고일", lambda values: ", ".join(sorted({str(value) for value in values if str(value).strip()}))),
            적용_LT=("리드타임", "max"),
        )
        .reset_index()
        .sort_values("수량", ascending=False)
    )
    grouped = grouped.rename(columns={"금액_M원": "금액(M원)", "금액_EUR": "금액(EUR)", "적용_LT": "적용 L/T"})
    return grouped[[col for col in columns if col in grouped.columns]]


def eta_detail_summary_html(detail: pd.DataFrame, display_detail: pd.DataFrame, selected_date: str) -> str:
    total_qty = float(pd.to_numeric(detail.get("수량", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
    total_m = float(pd.to_numeric(detail.get("도착 예정 금액(M원)", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
    total_eur = float(pd.to_numeric(detail.get("금액(EUR)", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
    raw_rows = len(detail)
    sku_count = display_detail["SKU"].nunique() if "SKU" in display_detail.columns else 0
    modes = ", ".join(display_detail["운송수단"].dropna().astype(str).drop_duplicates().tolist()) if "운송수단" in display_detail.columns else "-"
    top_sku = "-"
    if not display_detail.empty and {"SKU", "수량"}.issubset(display_detail.columns):
        top_row = display_detail.sort_values("수량", ascending=False).iloc[0]
        top_sku = f'{top_row["SKU"]} / {float(top_row["수량"]):,.0f}개'

    items = [
        ("ETA", selected_date),
        ("SKU", f"{sku_count:,}개"),
        ("원본 행", f"{raw_rows:,}건"),
        ("수량", f"{total_qty:,.0f}개"),
        ("금액", f"{total_m:,.0f}M 원"),
        ("EUR", f"{total_eur:,.0f}"),
        ("운송", modes or "-"),
        ("최대 수량", top_sku),
    ]
    chips = "".join(
        f'<div class="eta-summary-chip"><span>{html_lib.escape(label)}</span><b>{html_lib.escape(value)}</b></div>'
        for label, value in items
    )
    return f'<div class="eta-summary-grid">{chips}</div>'



def build_eta_unmatched_shipping_report_df(
    review: pd.DataFrame,
    settings: dict | None = None,
    shipping_df: pd.DataFrame | None = None,
    context: SessionContext | None = None,
) -> pd.DataFrame:
    shipping = shipping_df.copy() if shipping_df is not None else transport_mod.get_shipping(settings, context)
    if settings is not None:
        shipping = validation_mod.apply_common_filters(shipping, settings)

    shipping = shipping[preprocess_mod.product_sku_mask(shipping["SKU"])].copy()
    eta = pd.to_datetime(shipping["ETA"], errors="coerce")
    shipping = shipping[eta.notna()].copy()
    if shipping.empty:
        return pd.DataFrame(
            columns=[
                "상품코드", "상품명", "브랜드", "운송수단 코드", "운송수단", "운송수단 출처", "출고일", "계산 ETA",
                "수량", "입고가(KRW)", "ETA 구분", "운송 L/T",
                "적용 L/T", "ETA 출처", "확인 필요 사유", "처리 기준",
            ]
        )

    review_codes = set(preprocess_mod.clean_identifier_series(order_review_mod.report_col(review, ["상품코드"], "")).astype(str))
    shipping = shipping[~shipping["SKU"].astype(str).isin(review_codes)].copy()
    if shipping.empty:
        return pd.DataFrame(
            columns=[
                "상품코드", "상품명", "브랜드", "운송수단 코드", "운송수단", "운송수단 출처", "출고일", "계산 ETA",
                "수량", "입고가(KRW)", "ETA 구분", "운송 L/T",
                "적용 L/T", "ETA 출처", "확인 필요 사유", "처리 기준",
            ]
        )

    out = pd.DataFrame(
        {
            "상품코드": shipping["SKU"],
            "상품명": shipping["상품명"],
            "브랜드": shipping["브랜드"],
            "운송수단 코드": shipping.get("운송수단 코드", ""),
            "운송수단": shipping["운송수단"],
            "운송수단 출처": shipping.get("운송수단 출처", ""),
            "출고일": shipping["출고일"],
            "계산 ETA": shipping["ETA"],
            "수량": pd.to_numeric(shipping["수량"], errors="coerce").fillna(0),
            "입고가(KRW)": pd.to_numeric(shipping["입고가(KRW)"], errors="coerce").fillna(0),
            "ETA 구분": shipping.get("ETA 구분", "적용LT계산ETA"),
            "운송 L/T": pd.to_numeric(shipping.get("운송 L/T", 0), errors="coerce").fillna(0),
            "적용 L/T": pd.to_numeric(shipping.get("리드타임", 0), errors="coerce").fillna(0),
            "ETA 출처": shipping.get("ETA 출처", shipping.get("ETA 구분", "")),
            "확인 필요 사유": "입고 예정은 있으나 발주 검토에는 없음",
            "처리 기준": "ETA 타임라인에는 표시, 발주 필요 수량 산식에는 자동 반영하지 않음",
        }
    )
    return out.sort_values(["계산 ETA", "운송수단", "상품코드"], na_position="last").reset_index(drop=True)


def build_eta_timeline_report_df(
    review: pd.DataFrame,
    settings: dict | None = None,
    context: SessionContext | None = None,
) -> pd.DataFrame:
    shipping = transport_mod.get_shipping(settings, context)
    if settings is not None:
        shipping = validation_mod.apply_common_filters(shipping, settings)
    eta_dates = (
        pd.to_datetime(shipping["ETA"], errors="coerce")
        .dropna()
        .sort_values()
        .dt.strftime("%Y-%m-%d")
        .drop_duplicates()
        .tolist()
    )
    base = order_review_report_mod.build_order_review_report_df(review, settings).rename(
        columns={
            "상품코드": "SKU",
            "최근 3개월 수요": "3개월 수요",
            "EU 현지 재고": "유럽 재고",
            "부족수량": "발주필요수량",
            "부족금액_KRW": "발주필요금액(KRW)",
        }
    )
    base = base.loc[:, ~base.columns.duplicated()]
    keep_cols = [
        "No", "브랜드", "SKU", "상품명",
        "우선 액션", "발주필요금액(KRW)", "발주필요수량",
        "운송중", "유럽+운송", "운송포함 보유개월",
        "유럽 재고", "커버개월", "월평균", "안전재고",
        "EU 입고단가(EUR)", "판단 사유",
    ]
    base = base[[col for col in keep_cols if col in base.columns]].copy()
    base["확인 필요 사유"] = ""
    for eta in eta_dates:
        label = pd.to_datetime(eta).strftime("ETA %m/%d")
        qty = pd.to_numeric(shipping.loc[shipping["ETA"] == eta, "수량"], errors="coerce").fillna(0)
        qty_by_sku = qty.groupby(shipping.loc[shipping["ETA"] == eta, "SKU"]).sum()
        qty_by_sku = qty_by_sku[qty_by_sku > 0]
        mapped_qty = pd.to_numeric(base["SKU"].map(qty_by_sku), errors="coerce")
        base[label] = mapped_qty.where(mapped_qty > 0, "")
    date_labels = [pd.to_datetime(eta).strftime("ETA %m/%d") for eta in eta_dates]
    if date_labels:
        base["파이프라인 합계"] = base[date_labels].replace("", 0).sum(axis=1)
    else:
        base["파이프라인 합계"] = 0
    if {"우선 액션", "발주필요수량", "발주필요금액(KRW)"}.issubset(base.columns):
        no_order_eta_risk = (
            base["우선 액션"].astype(str).str.contains("ETA 지연 위험", na=False)
            & (pd.to_numeric(base["발주필요수량"], errors="coerce").fillna(0) <= 0)
            & (pd.to_numeric(base["발주필요금액(KRW)"], errors="coerce").fillna(0) <= 0)
        )
        base.loc[no_order_eta_risk, "우선 액션"] = "ETA 일정 확인"
        zero_amount_order = (
            base["우선 액션"].astype(str).eq("발주 필요")
            & (pd.to_numeric(base["발주필요수량"], errors="coerce").fillna(0) > 0)
            & (pd.to_numeric(base["발주필요금액(KRW)"], errors="coerce").fillna(0) <= 0)
        )
        base.loc[zero_amount_order, "우선 액션"] = "FOC/단가 확인 필요"

    unmatched = build_eta_unmatched_shipping_report_df(review, settings, shipping_df=shipping, context=context)
    if not unmatched.empty:
        extra_rows: list[dict[str, object]] = []
        for sku, group in unmatched.groupby("상품코드", dropna=False):
            row = {col: "" for col in base.columns}
            row["브랜드"] = inbound_mod.first_non_empty(group["브랜드"]) if "브랜드" in group.columns else ""
            row["SKU"] = sku
            row["상품명"] = inbound_mod.first_non_empty(group["상품명"]) if "상품명" in group.columns else ""
            row["운송중"] = float(pd.to_numeric(group["수량"], errors="coerce").fillna(0).sum())
            row["우선 액션"] = "입고예정 확인필요"
            row["확인 필요 사유"] = "입고 예정은 있으나 발주 검토에는 없음"
            for eta in eta_dates:
                eta_text = pd.to_datetime(eta).strftime("%Y-%m-%d")
                label = pd.to_datetime(eta).strftime("ETA %m/%d")
                eta_qty = pd.to_numeric(group.loc[group["계산 ETA"].astype(str).eq(eta_text), "수량"], errors="coerce").fillna(0).sum()
                row[label] = float(eta_qty) if eta_qty else ""
            row["파이프라인 합계"] = row["운송중"]
            extra_rows.append(row)
        base = pd.concat([base, pd.DataFrame(extra_rows)], ignore_index=True)
    def eta_priority(value: object) -> int:
        text = str(display_text(value) or "")
        if "ETA 지연 위험" in text:
            return 0
        if text == "발주 필요":
            return 1
        if "본사 EU창고" in text or "본사 일부 이동" in text:
            return 2
        if "운송중 도착 대기" in text:
            return 3
        if "FOC/단가 확인 필요" in text or "발주금액 확인 필요" in text or "ETA 일정 확인" in text or "입고예정 확인필요" in text or "확인 필요" in text:
            return 4
        return 9

    base["_정렬_우선액션"] = base["우선 액션"].map(eta_priority) if "우선 액션" in base.columns else 9
    base["_정렬_발주금액"] = pd.to_numeric(base.get("발주필요금액(KRW)", 0), errors="coerce").fillna(0)
    base["_정렬_발주수량"] = pd.to_numeric(base.get("발주필요수량", 0), errors="coerce").fillna(0)
    if "우선 액션" in base.columns:
        zero_order_eta_risk = (
            base["우선 액션"].astype(str).str.contains("입고 전 결품 위험|입고 전 품절 위험|ETA 지연 위험", na=False)
            & (base["_정렬_발주금액"] <= 0)
            & (base["_정렬_발주수량"] <= 0)
        )
        base.loc[zero_order_eta_risk, "_정렬_우선액션"] = 4
    base = base.sort_values(
        ["_정렬_발주금액", "_정렬_발주수량", "_정렬_우선액션"],
        ascending=[False, False, True],
        kind="mergesort",
    ).drop(columns=["_정렬_우선액션", "_정렬_발주금액", "_정렬_발주수량"])
    if "No" in base.columns:
        base["No"] = range(1, len(base) + 1)
    base, _ = order_review_report_mod.apply_low_unit_price_review_label(
        base,
        price_columns=["EU 입고단가(EUR)", "EU 입고단가"],
        qty_columns=["발주필요수량"],
    )
    base = base.drop(columns=["EU 입고단가(EUR)", "EU 입고단가"], errors="ignore")
    return base


def build_arrival_calendar_report_df(settings: dict, context: SessionContext | None = None) -> pd.DataFrame:
    shipping = validation_mod.apply_common_filters(
        transport_mod.get_shipping(settings, context),
        settings,
        include_transport=False,
    )
    configured_groups = {
        TRANSPORT_GROUP_DISPLAY_NAMES.get(str(method.get("transport_group") or "").strip().upper(), "")
        for method in list(settings.get("lead_time_methods") or [])
    }
    configured_groups.discard("")
    selected_modes = set(validation_mod.selected_transport_modes(settings))
    if configured_groups and selected_modes != configured_groups:
        shipping = shipping[shipping["운송수단"].astype(str).isin(selected_modes)].copy()
    eta = pd.to_datetime(shipping["ETA"], errors="coerce")
    ship_date = kpi_mod.parse_date_series(shipping["출고일"])
    qty = pd.to_numeric(shipping.get("수량", 0), errors="coerce").fillna(0)
    local_amount_source = shipping["금액"] if "금액" in shipping.columns else pd.Series([0] * len(shipping), index=shipping.index)
    local_amount = pd.to_numeric(local_amount_source, errors="coerce").fillna(0)
    unit_krw = pd.to_numeric(shipping.get("입고가(KRW)", 0), errors="coerce").fillna(0)
    fallback_krw = unit_krw * qty
    currency_code = str(settings.get("currency_code") or "EUR").strip().upper()
    currency_krw_rate = float(
        settings.get("currency_krw_rate", settings.get("eur_krw_rate", _DEFAULT_EUR_KRW_RATE))
    )
    if "CMS 원화 환산금액" in shipping.columns:
        amount_krw = pd.to_numeric(
            shipping["CMS 원화 환산금액"], errors="coerce"
        ).fillna(0)
    else:
        amount_krw = np.where(
            local_amount > 0,
            local_amount * currency_krw_rate,
            fallback_krw,
        )
    amount_krw = np.where(qty > 0, amount_krw, 0)
    days_until_arrival = (eta - pd.to_datetime(settings["base_date"])).dt.days
    mode_label_by_code = {
        str(definition["code"]): str(definition["label"])
        for definition in arrival_reference_mode_definitions(settings)
        if str(definition.get("code") or "")
    }
    transport_codes = order_review_mod.report_col(
        shipping,
        ["운송수단 코드"],
        "",
    ).astype(str).str.strip().str.upper()
    display_mode = transport_codes.map(mode_label_by_code)
    display_mode = display_mode.where(
        display_mode.notna(),
        shipping["운송수단"].fillna("").astype(str),
    ).replace({TRANSPORT_REVIEW_REQUIRED: "확인필요", "트럭킹": "트럭"})

    def dday_label(value: object) -> str:
        if pd.isna(value):
            return ""
        days = int(value)
        if days == 0:
            return "D-Day"
        if days > 0:
            return f"D-{days}"
        return f"D+{abs(days)}"

    out = pd.DataFrame(
        {
            "도착일": eta.dt.date,
            "도착 예정월": eta.dt.strftime("%Y-%m"),
            "ETA 구분": shipping.get("ETA 구분", "실제ETA"),
            "ETA 출처": shipping.get("ETA 출처", shipping.get("ETA 구분", "실제ETA")),
            "운송수단 코드": transport_codes,
            "운송수단": display_mode,
            "운송수단 출처": shipping.get("운송수단 출처", ""),
            "브랜드": shipping["브랜드"],
            "SKU": shipping["SKU"],
            "상품명": shipping["상품명"],
            "수량": qty,
            "도착 예정 금액_현지통화": local_amount,
            f"도착 예정 금액_{currency_code}": local_amount,
            "도착 예정 금액_KRW": amount_krw,
            "출고일": ship_date.dt.date,
            "운송 L/T": pd.to_numeric(shipping.get("운송 L/T", 0), errors="coerce"),
            "적용 L/T": pd.to_numeric(shipping.get("리드타임", 0), errors="coerce").fillna(0),
            "ETA D-Day": days_until_arrival.map(dday_label),
        }
    )
    out = out[preprocess_mod.product_sku_mask(out["SKU"]) & out["도착일"].notna() & (pd.to_numeric(out["수량"], errors="coerce").fillna(0) > 0)]
    return out.sort_values(
        ["도착일", "도착 예정 금액_KRW", "수량", "운송수단", "SKU"],
        ascending=[True, False, False, True, True],
        na_position="last",
    ).reset_index(drop=True)


def add_months(base: date, months: int) -> date:
    month_index = base.month - 1 + months
    year = base.year + month_index // 12
    month = month_index % 12 + 1
    day = min(base.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def arrival_reference_mode_definitions(settings: dict | None = None) -> list[dict[str, object]]:
    """Return legal-entity transport columns without inventing unsupported modes."""

    settings = settings or {}
    methods = list(settings.get("lead_time_methods") or [])
    if not methods:
        return [
            {
                "code": "",
                "group": raw_mode,
                "label": display_mode,
                "match_labels": {raw_mode, display_mode, "트럭킹" if raw_mode == "트럭" else display_mode},
                "allow_label_fallback": True,
            }
            for raw_mode, display_mode in ARRIVAL_REFERENCE_MODES
        ]

    group_counts: dict[str, int] = {}
    for method in methods:
        group = str(method.get("transport_group") or "").strip().upper()
        if group:
            group_counts[group] = group_counts.get(group, 0) + 1

    definitions: list[dict[str, object]] = []
    used_labels: set[str] = set()
    for method in methods:
        code = str(method.get("transport_code") or method.get("code") or "").strip().upper()
        group = str(method.get("transport_group") or "").strip().upper()
        if not code or not group:
            continue
        group_label = TRANSPORT_GROUP_DISPLAY_NAMES.get(group, group)
        method_label = str(method.get("display_name") or method.get("label") or group_label).strip()
        label = group_label if group_counts.get(group, 0) == 1 else method_label
        if label in used_labels:
            continue
        used_labels.add(label)
        definitions.append(
            {
                "code": code,
                "group": group,
                "label": label,
                "match_labels": {label, method_label, group_label, "트럭킹" if group == "TRUCKING" else group_label},
                "allow_label_fallback": group_counts.get(group, 0) == 1,
            }
        )
    return definitions


def arrival_qty_for_month(calendar_df: pd.DataFrame, target_month: date) -> float:
    if calendar_df.empty or "도착 예정월" not in calendar_df.columns:
        return 0.0
    month_key = f"{target_month.year:04d}-{target_month.month:02d}"
    month_mask = calendar_df["도착 예정월"].astype(str).eq(month_key)
    return float(pd.to_numeric(calendar_df.loc[month_mask, "수량"], errors="coerce").fillna(0).sum())


def arrival_amount_for_month(calendar_df: pd.DataFrame, target_month: date) -> float:
    if calendar_df.empty or "도착 예정월" not in calendar_df.columns or "도착 예정 금액_KRW" not in calendar_df.columns:
        return 0.0
    month_key = f"{target_month.year:04d}-{target_month.month:02d}"
    month_mask = calendar_df["도착 예정월"].astype(str).eq(month_key)
    return float(pd.to_numeric(calendar_df.loc[month_mask, "도착 예정 금액_KRW"], errors="coerce").fillna(0).sum())


def build_transport_arrival_summary_df(settings: dict, context: SessionContext | None = None) -> pd.DataFrame:
    calendar_df = build_arrival_calendar_report_df(settings, context)
    if calendar_df.empty:
        return pd.DataFrame(columns=["운송수단", "도착 예정월", "SKU 수", "총 도착 수량", "도착 예정 금액_KRW"])
    out = calendar_df.copy()
    grouped = (
        out.groupby(["운송수단", "도착 예정월"], dropna=False)
        .agg(
            **{
                "SKU 수": ("SKU", "nunique"),
                "총 도착 수량": ("수량", "sum"),
                "도착 예정 금액_KRW": ("도착 예정 금액_KRW", "sum"),
            }
        )
        .reset_index()
    )
    mode_order = {mode: idx for idx, mode in enumerate(["해운", "철송", "트럭", "항공"])}
    grouped["_운송순서"] = grouped["운송수단"].map(mode_order).fillna(99)
    return grouped.sort_values(["_운송순서", "도착 예정월"]).drop(columns=["_운송순서"]).reset_index(drop=True)


def build_target_month_arrival_summary_df(
    settings: dict,
    month_offsets: tuple[int, ...] = (1, 2, 3),
    context: SessionContext | None = None,
) -> pd.DataFrame:
    calendar_df = build_arrival_calendar_report_df(settings, context)
    base_month = date(settings["base_date"].year, settings["base_date"].month, 1)
    mode_definitions = arrival_reference_mode_definitions(settings)
    rows = []
    for offset in month_offsets:
        target = add_months(base_month, offset)
        month_key = f"{target.year:04d}-{target.month:02d}"
        month_df = calendar_df[calendar_df["도착 예정월"].astype(str).eq(month_key)] if not calendar_df.empty else pd.DataFrame()
        row: dict[str, object] = {
            "입고 예정월": f"{target.month}월",
            "연월": month_key,
            "총 도착 수량": float(pd.to_numeric(month_df.get("수량", 0), errors="coerce").fillna(0).sum()) if not month_df.empty else 0,
            "총 도착 금액_KRW": float(pd.to_numeric(month_df.get("도착 예정 금액_KRW", 0), errors="coerce").fillna(0).sum()) if not month_df.empty else 0,
            "SKU 수": int(month_df["SKU"].nunique()) if not month_df.empty and "SKU" in month_df.columns else 0,
        }
        matched = pd.Series(False, index=month_df.index)
        for definition in mode_definitions:
            label = str(definition["label"])
            if month_df.empty:
                mask = pd.Series(False, index=month_df.index)
            else:
                code = str(definition.get("code") or "")
                code_values = order_review_mod.report_col(
                    month_df, ["운송수단 코드"], ""
                ).astype(str).str.strip().str.upper()
                mask = code_values.eq(code) if code else pd.Series(False, index=month_df.index)
                if bool(definition.get("allow_label_fallback", False)):
                    labels = {str(value) for value in definition.get("match_labels", set())}
                    mask = mask | month_df["운송수단"].astype(str).isin(labels)
            matched = matched | mask
            row[f"{label} 수량"] = float(
                pd.to_numeric(month_df.loc[mask, "수량"], errors="coerce").fillna(0).sum()
            ) if not month_df.empty else 0
        row["확인필요 수량"] = float(
            pd.to_numeric(month_df.loc[~matched, "수량"], errors="coerce").fillna(0).sum()
        ) if not month_df.empty else 0
        rows.append(row)
    return pd.DataFrame(rows)


def build_arrival_reference_summary_df(
    settings: dict,
    calendar_df: pd.DataFrame | None = None,
    context: SessionContext | None = None,
) -> pd.DataFrame:
    calendar = build_arrival_calendar_report_df(settings, context) if calendar_df is None else pd.DataFrame(calendar_df).copy()
    base_date = settings.get("base_date") or korea_today()
    year = int(base_date.year)
    currency_krw_rate = float(
        settings.get("currency_krw_rate", settings.get("eur_krw_rate", _DEFAULT_EUR_KRW_RATE))
    )
    currency_code = str(settings.get("currency_code") or "EUR").strip().upper()
    amount_col = f"금액({currency_code})"
    local_amount_source_col = f"도착 예정 금액_{currency_code}"
    mode_definitions = arrival_reference_mode_definitions(settings)
    mode_labels = [str(definition["label"]) for definition in mode_definitions]
    columns = [
        "입고 예정",
        *mode_labels,
        "합계",
        amount_col,
        "금액(KRW)",
        "금액(억원)",
    ]

    if calendar.empty:
        work = pd.DataFrame(
            columns=[
                "도착 예정월", "운송수단 코드", "운송수단", "수량",
                "도착 예정 금액_현지통화", "도착 예정 금액_KRW",
            ]
        )
    else:
        work = calendar.copy()
        if "도착 예정월" not in work.columns:
            work["도착 예정월"] = kpi_mod.parse_date_series(order_review_mod.report_col(work, ["도착일"])).dt.strftime("%Y-%m")
        work["운송수단 코드"] = order_review_mod.report_col(work, ["운송수단 코드"], "").astype(str).str.strip().str.upper()
        work["운송수단"] = order_review_mod.report_col(work, ["운송수단"], "")
        work["수량"] = pd.to_numeric(order_review_mod.report_col(work, ["수량"], 0), errors="coerce").fillna(0)
        work["도착 예정 금액_KRW"] = pd.to_numeric(order_review_mod.report_col(work, ["도착 예정 금액_KRW"], 0), errors="coerce").fillna(0)
        if local_amount_source_col in work.columns:
            work["도착 예정 금액_현지통화"] = pd.to_numeric(work[local_amount_source_col], errors="coerce").fillna(0)
        elif "도착 예정 금액_현지통화" in work.columns:
            work["도착 예정 금액_현지통화"] = pd.to_numeric(work["도착 예정 금액_현지통화"], errors="coerce").fillna(0)
        elif currency_code == "EUR" and "도착 예정 금액_EUR" in work.columns:
            work["도착 예정 금액_현지통화"] = pd.to_numeric(work["도착 예정 금액_EUR"], errors="coerce").fillna(0)
        else:
            work["도착 예정 금액_현지통화"] = np.where(
                currency_krw_rate > 0,
                work["도착 예정 금액_KRW"] / currency_krw_rate,
                0,
            )

    def mode_mask(df: pd.DataFrame, definition: dict[str, object]) -> pd.Series:
        if df.empty:
            return pd.Series(False, index=df.index)
        code = str(definition.get("code") or "")
        code_match = df["운송수단 코드"].astype(str).eq(code) if code else pd.Series(False, index=df.index)
        if not bool(definition.get("allow_label_fallback", False)):
            return code_match
        labels = {str(value) for value in definition.get("match_labels", set())}
        return code_match | df["운송수단"].astype(str).isin(labels)

    def classified_masks(df: pd.DataFrame) -> tuple[list[pd.Series], pd.Series]:
        masks = [mode_mask(df, definition) for definition in mode_definitions]
        matched = pd.Series(False, index=df.index)
        for mask in masks:
            matched = matched | mask
        return masks, ~matched

    # Unclassified transport rows belong exclusively to the 확인필요 sheet.
    # Exclude them from both the visible mode columns and the arrival totals so
    # this sheet always reconciles to the transport columns it displays.
    _, review_mask = classified_masks(work)
    work = work.loc[~review_mask].copy()

    rows: list[dict[str, object]] = []
    for month in range(1, 13):
        month_key = f"{year:04d}-{month:02d}"
        month_df = work[work["도착 예정월"].astype(str).eq(month_key)] if not work.empty else work
        row: dict[str, object] = {"입고 예정": month_key}
        mode_masks, _ = classified_masks(month_df)
        for definition, mask in zip(mode_definitions, mode_masks):
            qty = float(pd.to_numeric(month_df.loc[mask, "수량"], errors="coerce").fillna(0).sum()) if not month_df.empty else 0.0
            row[str(definition["label"])] = int(round(qty))
        total_qty = float(pd.to_numeric(month_df.get("수량", 0), errors="coerce").fillna(0).sum()) if not month_df.empty else 0.0
        amount_local = float(pd.to_numeric(month_df.get("도착 예정 금액_현지통화", 0), errors="coerce").fillna(0).sum()) if not month_df.empty else 0.0
        amount_krw = float(pd.to_numeric(month_df.get("도착 예정 금액_KRW", 0), errors="coerce").fillna(0).sum()) if not month_df.empty else 0.0
        row["합계"] = int(round(total_qty))
        row[amount_col] = int(round(amount_local)) if amount_local > 0 else "-"
        row["금액(KRW)"] = int(round(amount_krw)) if amount_krw > 0 else "-"
        row["금액(억원)"] = f"{int(round(amount_krw / 100_000_000))}억" if amount_krw > 0 else "-"
        rows.append(row)

    total_row: dict[str, object] = {"입고 예정": "합계"}
    amount_eur_row: dict[str, object] = {"입고 예정": amount_col}
    amount_krw_row: dict[str, object] = {"입고 예정": "금액(KRW)"}
    all_mode_masks, _ = classified_masks(work)
    for definition, mask in zip(mode_definitions, all_mode_masks):
        mode_df = work[mask] if not work.empty else work
        qty = float(pd.to_numeric(mode_df.get("수량", 0), errors="coerce").fillna(0).sum()) if not mode_df.empty else 0.0
        amount_local = float(pd.to_numeric(mode_df.get("도착 예정 금액_현지통화", 0), errors="coerce").fillna(0).sum()) if not mode_df.empty else 0.0
        amount_krw = float(pd.to_numeric(mode_df.get("도착 예정 금액_KRW", 0), errors="coerce").fillna(0).sum()) if not mode_df.empty else 0.0
        label = str(definition["label"])
        total_row[label] = int(round(qty))
        amount_eur_row[label] = int(round(amount_local)) if amount_local > 0 else "-"
        amount_krw_row[label] = int(round(amount_krw)) if amount_krw > 0 else "-"
    grand_qty = float(pd.to_numeric(work.get("수량", 0), errors="coerce").fillna(0).sum()) if not work.empty else 0.0
    total_row["합계"] = int(round(grand_qty))
    for row in (total_row, amount_eur_row, amount_krw_row):
        row.setdefault("합계", "")
        row.setdefault(amount_col, "")
        row.setdefault("금액(KRW)", "")
        row.setdefault("금액(억원)", "")
    rows.extend([total_row, amount_eur_row, amount_krw_row])
    return pd.DataFrame(rows, columns=columns)


def build_arrival_reference_detail_df(
    settings: dict,
    calendar_df: pd.DataFrame | None = None,
    context: SessionContext | None = None,
) -> pd.DataFrame:
    calendar = build_arrival_calendar_report_df(settings, context) if calendar_df is None else pd.DataFrame(calendar_df).copy()
    if calendar.empty:
        return pd.DataFrame(columns=ARRIVAL_REFERENCE_DETAIL_COLUMNS)

    work = calendar.copy()
    mode_definitions = arrival_reference_mode_definitions(settings)
    transport_codes = order_review_mod.report_col(
        work, ["운송수단 코드"], ""
    ).astype(str).str.strip().str.upper()
    transport_labels = order_review_mod.report_col(
        work, ["운송수단"], ""
    ).astype(str)
    classified = pd.Series(False, index=work.index)
    for definition in mode_definitions:
        code = str(definition.get("code") or "")
        mask = transport_codes.eq(code) if code else pd.Series(False, index=work.index)
        if bool(definition.get("allow_label_fallback", False)):
            labels = {str(value) for value in definition.get("match_labels", set())}
            mask = mask | transport_labels.isin(labels)
        classified = classified | mask
    work = work.loc[classified].copy()
    if work.empty:
        return pd.DataFrame(columns=ARRIVAL_REFERENCE_DETAIL_COLUMNS)
    shipped_at = kpi_mod.parse_date_series(order_review_mod.report_col(work, ["출고일"]))
    source_lead_time_days = pd.to_numeric(
        order_review_mod.report_col(work, ["운송 L/T", "적용 L/T", "리드타임"], ""),
        errors="coerce",
    ).round(0)
    configured_lead_times = {
        str(code).strip().upper(): int(days)
        for code, days in dict(settings.get("lead_times_by_code") or {}).items()
    }
    configured_lead_time_days = transport_codes.loc[work.index].map(
        configured_lead_times
    )
    lead_time_days = pd.to_numeric(
        configured_lead_time_days.where(
            configured_lead_time_days.notna(),
            source_lead_time_days,
        ),
        errors="coerce",
    ).round(0)
    lead_time_display = lead_time_days.astype(object).where(lead_time_days.notna(), "")
    chosen_arrival = kpi_mod.parse_date_series(order_review_mod.report_col(work, ["도착일"]))
    base_date = pd.Timestamp(settings.get("base_date") or korea_today()).normalize()
    out = pd.DataFrame(
        {
            "도착일": chosen_arrival.dt.date,
            "운송수단": order_review_mod.report_col(work, ["운송수단"], "").astype(str).replace({"트럭킹": "트럭"}),
            "브랜드": order_review_mod.report_col(work, ["브랜드"], ""),
            "SKU": order_review_mod.report_col(work, ["SKU", "상품코드"], ""),
            "상품명": order_review_mod.report_col(work, ["상품명"], ""),
            "수량": pd.to_numeric(order_review_mod.report_col(work, ["수량"], 0), errors="coerce").fillna(0).round(0).astype(int),
            "도착예정금액(KRW)": pd.to_numeric(
                order_review_mod.report_col(work, ["도착 예정 금액_KRW", "도착예정금액(KRW)", "도착 예정금액(KRW)"], 0),
                errors="coerce",
            ).fillna(0).round(0).astype(int),
            "출고일": shipped_at.dt.date,
            "운송 L/T": lead_time_display,
            # Reuse the calendar's chosen final arrival so a confirmed source
            # ETA is never replaced by a ship-date + default-L/T recomputation.
            "도착예정일": chosen_arrival.dt.date,
            "D-Day": (chosen_arrival.dt.normalize() - base_date).dt.days,
        }
    )
    out = out[out["도착일"].notna() & (pd.to_numeric(out["수량"], errors="coerce").fillna(0) > 0)]
    return out.sort_values(
        ["도착일", "도착예정금액(KRW)", "수량", "운송수단", "SKU"],
        ascending=[True, False, False, True, True],
        na_position="last",
    ).reset_index(drop=True)[ARRIVAL_REFERENCE_DETAIL_COLUMNS]


def write_arrival_reference_sheet(ws, settings: dict, calendar_df: pd.DataFrame | None = None) -> None:
    summary_df = build_arrival_reference_summary_df(settings, calendar_df)
    detail_df = build_arrival_reference_detail_df(settings, calendar_df)
    append_df(ws, summary_df)

    detail_header_row = len(summary_df) + 5
    detail_title_row = detail_header_row - 1
    ws.cell(detail_title_row, 1, "SKU별 도착 일정")
    append_df(ws, detail_df, start_row=detail_header_row)
    format_arrival_reference_summary_sheet(ws, settings, detail_header_row=detail_header_row)


def _find_arrival_reference_detail_header_row(ws) -> int | None:
    for row_idx in range(2, ws.max_row + 1):
        headers = {
            str(ws.cell(row_idx, col_idx).value or "")
            for col_idx in range(1, min(ws.max_column, len(ARRIVAL_REFERENCE_DETAIL_COLUMNS)) + 1)
        }
        if str(ws.cell(row_idx, 1).value or "") == "도착일" and "운송수단" in headers:
            return row_idx
    return None


def format_arrival_reference_summary_sheet(ws, settings: dict | None = None, detail_header_row: int | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    settings = settings or {}
    header_fill = PatternFill("solid", fgColor="4F6228")
    light_green_fill = PatternFill("solid", fgColor="E2F0D9")
    light_orange_fill = PatternFill("solid", fgColor="FCE4D6")
    detail_header_fill = PatternFill("solid", fgColor="D9EAD3")
    detail_title_fill = PatternFill("solid", fgColor="4F6228")
    total_fill = PatternFill("solid", fgColor="BFBFBF")
    label_fill = PatternFill("solid", fgColor="4F6228")
    dark_green_font = Font(color="375623", bold=True)
    white_font = Font(color="FFFFFF", bold=True, size=12)
    header_font = Font(color="FFFFFF", bold=True)
    total_font = Font(color="000000", bold=True, size=12)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    detail_header_row = detail_header_row or _find_arrival_reference_detail_header_row(ws)
    detail_last_col = len(ARRIVAL_REFERENCE_DETAIL_COLUMNS) if detail_header_row else 0
    summary_headers: list[str] = []
    for col_idx in range(1, ws.max_column + 1):
        value = str(ws.cell(1, col_idx).value or "").strip()
        if not value:
            break
        summary_headers.append(value)
    summary_last_col = len(summary_headers)
    total_col = summary_headers.index("합계") + 1 if "합계" in summary_headers else summary_last_col
    krw_col = summary_headers.index("금액(KRW)") + 1 if "금액(KRW)" in summary_headers else None
    hundred_million_col = summary_headers.index("금액(억원)") + 1 if "금액(억원)" in summary_headers else None

    for col_idx in range(1, summary_last_col + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row_idx in range(2, 14):
        for col_idx in range(1, summary_last_col + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "right", vertical="center")
            if col_idx == 1:
                cell.fill = light_green_fill
                cell.font = dark_green_font
            elif col_idx in {krw_col, hundred_million_col}:
                cell.fill = light_orange_fill
                cell.font = Font(color="9C4A00", bold=col_idx == krw_col)
            if col_idx > 1 and isinstance(cell.value, Number) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0"
            elif col_idx == hundred_million_col:
                cell.number_format = "0억"

    for col_idx in range(1, total_col + 1):
        cell = ws.cell(14, col_idx)
        cell.fill = total_fill
        cell.font = total_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "right", vertical="center")
        if col_idx > 1:
            cell.number_format = "#,##0"

    for row_idx in (15, 16):
        label_cell = ws.cell(row_idx, 1)
        label_cell.fill = label_fill
        label_cell.font = header_font
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.border = border
        for col_idx in range(2, total_col):
            cell = ws.cell(row_idx, col_idx)
            cell.border = border
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(cell.value, Number) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0"

    if detail_header_row:
        detail_title_row = detail_header_row - 1
        try:
            ws.merge_cells(start_row=detail_title_row, start_column=1, end_row=detail_title_row, end_column=detail_last_col)
        except ValueError:
            pass
        title_cell = ws.cell(detail_title_row, 1)
        title_cell.value = title_cell.value or "SKU별 도착 일정"
        title_cell.fill = detail_title_fill
        title_cell.font = white_font
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        title_cell.border = border

        for col_idx in range(1, detail_last_col + 1):
            cell = ws.cell(detail_header_row, col_idx)
            cell.fill = detail_header_fill
            cell.font = dark_green_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        detail_date_headers = {"도착일", "출고일", "도착예정일"}
        detail_number_headers = {"수량", "도착예정금액(KRW)", "운송 L/T"}
        for row_idx in range(detail_header_row + 1, ws.max_row + 1):
            for col_idx in range(1, detail_last_col + 1):
                header = str(ws.cell(detail_header_row, col_idx).value or "")
                cell = ws.cell(row_idx, col_idx)
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="left" if header == "상품명" else "right" if header in detail_number_headers else "center",
                    vertical="center",
                    wrap_text=header == "상품명",
                )
                if header in detail_date_headers:
                    cell.number_format = "yyyy-mm-dd"
                elif header in detail_number_headers:
                    cell.number_format = "#,##0"
                elif header == "D-Day":
                    cell.number_format = '"D-"0;"D+"0;"D-Day"'
        if ws.max_row >= detail_header_row:
            ws.auto_filter.ref = f"A{detail_header_row}:{get_column_letter(detail_last_col)}{ws.max_row}"

    clear_start_col = max(summary_last_col, detail_last_col) + 1
    for col_idx in range(clear_start_col, ws.max_column + 1):
        for row_idx in range(1, ws.max_row + 1):
            ws.cell(row_idx, col_idx).value = None
    widths: dict[str, float] = {}
    for col_idx, header in enumerate(summary_headers, start=1):
        width = 14
        if header == "금액(KRW)":
            width = 20
        elif header == "금액(억원)":
            width = 13
        elif header.startswith("항공("):
            width = 18
        widths[get_column_letter(col_idx)] = width
    for col_idx in range(summary_last_col + 1, detail_last_col + 1):
        widths.setdefault(get_column_letter(col_idx), 14)
    if detail_header_row:
        detail_width_by_header = {
            "도착일": 13,
            "운송수단": 16,
            "브랜드": 18,
            "SKU": 20,
            "상품명": 58,
            "수량": 14,
            "도착예정금액(KRW)": 22,
            "출고일": 13,
            "운송 L/T": 13,
            "도착예정일": 16,
            "D-Day": 11,
        }
        for col_idx in range(1, detail_last_col + 1):
            header = str(ws.cell(detail_header_row, col_idx).value or "")
            widths[get_column_letter(col_idx)] = detail_width_by_header.get(header, 14)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for row_idx in range(1, 17):
        ws.row_dimensions[row_idx].height = 24 if row_idx == 1 else 20
    if detail_header_row:
        ws.row_dimensions[detail_header_row - 1].height = 24
        ws.row_dimensions[detail_header_row].height = 30
        for row_idx in range(detail_header_row + 1, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 36
    ws.freeze_panes = None

