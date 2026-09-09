from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd

from core.common import (
    ALERT_LOCAL_SHORT_LABEL,
    ALERT_SHIPPING_SHORT_LABEL,
    ALERT_URGENT_LABEL,
    DEFAULT_LEAD_TIME_DAYS,
    korea_today,
    ORDER_REVIEW_ELIGIBILITY_HELPER,
    UNIFIED_ESM_ORDER_REVIEW_COLUMNS,
)
from core.export_excel_util import display_text, excel_text_identifier
from core import (
    kpi as kpi_mod,
    order_review as order_review_mod,
    order_review_report as order_review_report_mod,
)



def _clean_merge_key(value: object) -> str:
    return excel_text_identifier(value).strip()


def _series_from_candidates(df: pd.DataFrame, candidates: list[str], default: object = "") -> pd.Series:
    return order_review_mod.report_col(df, candidates, default)


def _build_unique_lookup_df(source_df: pd.DataFrame, key_cols: list[str]) -> tuple[pd.DataFrame, int]:
    if source_df.empty or any(col not in source_df.columns for col in key_cols):
        return pd.DataFrame(columns=key_cols), 0
    lookup = source_df.copy()
    for col in key_cols:
        lookup[col] = lookup[col].map(_clean_merge_key)
    valid = lookup[key_cols[0]].astype(str).ne("")
    lookup = lookup[valid].copy()
    duplicated = lookup.duplicated(key_cols, keep=False)
    duplicate_count = int(duplicated.sum())
    lookup = lookup[~duplicated].drop_duplicates(key_cols, keep="first")
    return lookup, duplicate_count


def _lookup_order_values(
    base_df: pd.DataFrame,
    order_df: pd.DataFrame,
    source_candidates: list[str],
    default: object = "",
) -> pd.Series:
    source_col = order_review_mod.first_existing_col(order_df, source_candidates)
    if source_col is None or "상품코드" not in base_df.columns or "상품코드" not in order_df.columns:
        return pd.Series([default] * len(base_df), index=base_df.index)

    order_work = order_df[["상품코드", "브랜드", source_col] if "브랜드" in order_df.columns else ["상품코드", source_col]].copy()
    base_keys = pd.DataFrame({"상품코드": base_df["상품코드"].map(_clean_merge_key)}, index=base_df.index)
    if "브랜드" in base_df.columns and "브랜드" in order_work.columns:
        base_keys["브랜드"] = base_df["브랜드"].map(_clean_merge_key)

    sku_lookup, _ = _build_unique_lookup_df(order_work[["상품코드", source_col]], ["상품코드"])
    result = base_keys["상품코드"].map(sku_lookup.set_index("상품코드")[source_col]) if not sku_lookup.empty else pd.Series(index=base_df.index, dtype=object)

    if "브랜드" in base_keys.columns and "브랜드" in order_work.columns:
        missing = result.isna()
        sku_brand_lookup, _ = _build_unique_lookup_df(order_work[["상품코드", "브랜드", source_col]], ["상품코드", "브랜드"])
        if not sku_brand_lookup.empty and missing.any():
            brand_map = sku_brand_lookup.set_index(["상품코드", "브랜드"])[source_col]
            brand_values = [
                brand_map.get((sku, brand), np.nan)
                for sku, brand in zip(base_keys.loc[missing, "상품코드"], base_keys.loc[missing, "브랜드"])
            ]
            result.loc[missing] = brand_values

    return result.fillna(default)


def _first_non_blank(primary: pd.Series, fallback: pd.Series) -> pd.Series:
    if primary.empty:
        return primary.astype(object)
    primary_text = primary.astype(str).str.strip()
    mask = primary.isna() | primary_text.eq("") | primary_text.str.lower().isin(["nan", "none"])
    out = primary.astype(object).copy()
    if mask.any():
        out.loc[mask] = fallback.loc[mask] if isinstance(fallback, pd.Series) else fallback
    return out


def render_order_review_workbook_for_scenario(workbook_path: str | Path, scenario: str) -> bytes:
    """Return the full order workbook with its summary set to the selected inbound scenario."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(workbook_path)
    worksheet = workbook["재고 ETA"]

    def compact(value: object) -> str:
        return "".join(str(value or "").split())

    header_row = 0
    sku_column = 0
    order_qty_column = 0
    inbound_column = 0
    for row_number in range(1, min(worksheet.max_row, 20) + 1):
        current_sku_column = 0
        current_order_qty_column = 0
        current_inbound_column = 0
        for column_number in range(1, worksheet.max_column + 1):
            header = compact(worksheet.cell(row_number, column_number).value)
            if header == "SKU":
                current_sku_column = column_number
            elif header in {"발주필요수량", "최종발주필요수량"}:
                current_order_qty_column = column_number
            elif header == "미입고현황":
                current_inbound_column = column_number
        if current_sku_column and current_order_qty_column and current_inbound_column:
            header_row = row_number
            sku_column = current_sku_column
            order_qty_column = current_order_qty_column
            inbound_column = current_inbound_column
            break
    if not header_row:
        raise ValueError("원본 파일에서 시나리오 계산에 필요한 열을 찾지 못했습니다.")

    summary_label_row = 0
    summary_label_column = 0
    for row_number in range(1, min(header_row - 1, 8) + 1):
        for column_number in range(1, worksheet.max_column + 1):
            if compact(worksheet.cell(row_number, column_number).value) == "발주필요SKU":
                summary_label_row = row_number
                summary_label_column = column_number
                break
        if summary_label_row:
            break
    if not summary_label_row:
        raise ValueError("원본 파일에서 발주필요 SKU 요약 셀을 찾지 못했습니다.")

    scenario_label = "미입고 해소 후" if str(scenario).strip().lower() == "after" else "미입고 해소 전"
    worksheet["B2"] = scenario_label
    header_end_row = header_row
    required_header_columns = {sku_column, order_qty_column, inbound_column}
    for merged_range in worksheet.merged_cells.ranges:
        if not (merged_range.min_row <= header_row <= merged_range.max_row):
            continue
        if any(merged_range.min_col <= column <= merged_range.max_col for column in required_header_columns):
            header_end_row = max(header_end_row, merged_range.max_row)
    data_start_row = header_end_row + 1
    sku_letter = get_column_letter(sku_column)
    order_qty_letter = get_column_letter(order_qty_column)
    inbound_letter = get_column_letter(inbound_column)
    before_formula = (
        f'COUNTIFS({order_qty_letter}{data_start_row}:{order_qty_letter}100000,">0",'
        f'{sku_letter}{data_start_row}:{sku_letter}100000,"<>",'
        f'{sku_letter}{data_start_row}:{sku_letter}100000,"<>0")'
    )
    after_formula = (
        f'SUMPRODUCT(--({order_qty_letter}{data_start_row}:{order_qty_letter}100000>'
        f'{inbound_letter}{data_start_row}:{inbound_letter}100000))'
    )
    worksheet.cell(summary_label_row, summary_label_column + 1).value = (
        f'=IF($B$2="미입고 해소 후",{after_formula},{before_formula})'
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_unified_esm_order_review_sheet(
    internal_df: pd.DataFrame,
    order_review_df: pd.DataFrame,
    arrival_calendar_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    base = internal_df.copy()
    unified = pd.DataFrame(index=base.index)
    unified["상품코드"] = _series_from_candidates(base, ["상품코드"])
    unified["상품명"] = _series_from_candidates(base, ["상품명"])
    unified["브랜드"] = _series_from_candidates(base, ["브랜드"])
    unified["기준 3개월 판매수량"] = pd.to_numeric(
        _series_from_candidates(base, ["기준 3개월 판매수량", "기준 3개월 판매수량(재고 PA+CA)"], 0),
        errors="coerce",
    ).fillna(0)
    unified["월평균 판매수량"] = pd.to_numeric(
        _series_from_candidates(base, ["월평균 판매수량", "월평균 판매수량(기준/3)", "월평균"], 0),
        errors="coerce",
    ).fillna(0)
    unified["일평균 판매수량"] = pd.to_numeric(
        _series_from_candidates(base, ["일평균 판매수량", "일평균 판매수량(기준/90)"], 0),
        errors="coerce",
    ).fillna(0)
    unified["유럽 현재 재고"] = pd.to_numeric(_series_from_candidates(base, ["유럽 현재 재고", "EU 현지 재고"], 0), errors="coerce").fillna(0)
    unified["유럽 가용재고"] = pd.to_numeric(_series_from_candidates(base, ["유럽 가용재고"], 0), errors="coerce").fillna(0)
    unified["본사 EU창고"] = pd.to_numeric(_series_from_candidates(base, ["본사 EU창고", "본사 EU창고 가용수량"], 0), errors="coerce").fillna(0)
    unified["EU 현지 커버일수"] = pd.to_numeric(_series_from_candidates(base, ["EU 현지 커버일수"], 0), errors="coerce").fillna(0)
    unified["운송중 수량"] = pd.to_numeric(_series_from_candidates(base, ["운송중 수량", "운송중"], 0), errors="coerce").fillna(0)
    unified["입고 전 예상 결품량"] = pd.to_numeric(_series_from_candidates(base, ["입고 전 예상 결품량"], 0), errors="coerce").fillna(0)
    unified["미입고 수량"] = pd.to_numeric(_series_from_candidates(base, ["미입고 수량", "미입고"], 0), errors="coerce").fillna(0)
    unified["유럽+운송 수량"] = pd.to_numeric(_series_from_candidates(base, ["유럽+운송 수량", "유럽+운송"], 0), errors="coerce").fillna(0)
    unified["운송 포함 보유개월"] = pd.to_numeric(_series_from_candidates(base, ["운송 포함 보유개월", "운송포함 보유개월"], 0), errors="coerce").fillna(0)
    unified["예상 입고일"] = _lookup_order_values(base, order_review_df, ["첫 입고 예정일", "최초 ETA"])
    unified["안전재고 필요 수량"] = pd.to_numeric(_series_from_candidates(base, ["안전재고 필요 수량", "안전재고", "안전재고수량"], 0), errors="coerce").fillna(0)
    unified["부족 수량"] = pd.to_numeric(_series_from_candidates(base, ["부족 수량", "부족수량"], 0), errors="coerce").fillna(0)
    unified["발주 필요 수량"] = pd.to_numeric(_series_from_candidates(base, ["발주 필요 수량", "발주필요수량"], 0), errors="coerce").fillna(0)
    unified["긴급 보충 필요 수량"] = pd.to_numeric(_series_from_candidates(base, ["긴급 보충 필요 수량"], 0), errors="coerce").fillna(0)
    unified["입고 전 결품 위험 여부"] = _series_from_candidates(base, ["입고 전 결품 위험 여부"], "N")
    unified["권장 긴급 액션"] = _series_from_candidates(base, ["권장 긴급 액션"])
    unified["발주필요금액(KRW)"] = pd.to_numeric(
        _lookup_order_values(base, order_review_df, ["발주필요금액_KRW", "발주필요금액"]),
        errors="coerce",
    ).fillna(0)
    unified["EU 입고단가(EUR)"] = pd.to_numeric(
        _lookup_order_values(base, order_review_df, ["EU 입고단가(EUR)", "EU 입고단가"]),
        errors="coerce",
    ).fillna(0)
    unified["예상 소진일"] = _first_non_blank(
        _series_from_candidates(base, ["예상 소진일", "고갈 예정일"]),
        _lookup_order_values(base, order_review_df, ["예상 소진일", "고갈 예상일"]),
    )
    unified["쇼티지 예상일수"] = pd.to_numeric(_series_from_candidates(base, ["쇼티지 예상일수", "쇼티지 예상 일수"], 0), errors="coerce")
    unified["우선 액션"] = _series_from_candidates(base, ["우선 액션"])
    unified["판단메모"] = _lookup_order_values(base, order_review_df, ["판단메모"])
    unified["판단 사유"] = _lookup_order_values(base, order_review_df, ["판단 사유"])
    unified["추천 운송수단"] = _series_from_candidates(base, ["추천 운송수단", "운송 검토안"])
    unified["추천 사유"] = _first_non_blank(
        _series_from_candidates(base, ["추천 사유"]),
        _lookup_order_values(base, order_review_df, ["운송수단 추천 사유", "추천 사유"]),
    )
    unified["권장 운송안"] = _first_non_blank(
        _series_from_candidates(base, ["권장 운송안"]),
        _lookup_order_values(base, order_review_df, ["권장 대응 / 운송 검토안", "운송 검토안"]),
    )
    unified["권장 수량 요약"] = _series_from_candidates(base, ["권장 수량 요약"])
    unified["항공 검토수량"] = pd.to_numeric(_series_from_candidates(base, ["항공 검토수량"], 0), errors="coerce").fillna(0)
    unified["철송 검토수량"] = pd.to_numeric(_series_from_candidates(base, ["철송 검토수량"], 0), errors="coerce").fillna(0)
    unified["해운 검토수량"] = pd.to_numeric(_series_from_candidates(base, ["해운 검토수량"], 0), errors="coerce").fillna(0)
    unified["담당자 발주량"] = _series_from_candidates(base, ["담당자 발주량"])
    unified["발주 차이"] = _series_from_candidates(base, ["발주 차이"])
    unified["발주 신호등"] = _series_from_candidates(base, ["발주 신호등"], "미입력")
    unified["발주 메모"] = _series_from_candidates(base, ["발주 메모"])
    unified["검토일"] = _series_from_candidates(base, ["검토일"])
    unified["담당자"] = _series_from_candidates(base, ["담당자"])

    action_priority = {
        "ETA 지연 위험": 0,
        "발주 필요": 1,
        "본사 EU창고 이동 검토": 2,
        "본사 일부 이동 + 발주 필요": 2,
        "운송중 도착 대기": 3,
        "확인 필요": 4,
        "운송 도착 후 검토": 4,
    }

    def priority_value(value: object) -> int:
        text = str(display_text(value) or "")
        if "ETA 지연 위험" in text:
            return 0
        if text == "발주 필요":
            return 1
        if "본사 EU창고" in text or "본사 일부 이동" in text:
            return 2
        if "운송중 도착 대기" in text:
            return 3
        if "FOC/단가 확인 필요" in text or "발주금액 확인 필요" in text or "ETA 일정 확인" in text or "확인 필요" in text or "입고예정 확인필요" in text:
            return 4
        return 9

    unified["_정렬_우선액션"] = unified["우선 액션"].map(priority_value)
    unified["_정렬_쇼티지"] = pd.to_numeric(unified["쇼티지 예상일수"], errors="coerce").fillna(0)
    unified["_정렬_발주금액"] = pd.to_numeric(unified["발주필요금액(KRW)"], errors="coerce").fillna(0)
    unified["_정렬_발주수량"] = pd.to_numeric(unified["발주 필요 수량"], errors="coerce").fillna(0)
    unified["_정렬_커버일수"] = pd.to_numeric(unified["EU 현지 커버일수"], errors="coerce").fillna(999999)
    unified = unified.sort_values(
        ["_정렬_발주금액", "_정렬_발주수량", "_정렬_우선액션", "_정렬_쇼티지", "_정렬_커버일수"],
        ascending=[False, False, True, False, True],
        kind="mergesort",
    )
    unified, _ = order_review_report_mod.apply_low_unit_price_review_label(
        unified,
        price_columns=["EU 입고단가(EUR)", "EU 입고단가"],
        qty_columns=["발주 필요 수량", "발주필요수량"],
    )
    base_date_value = None
    parsed_review_date = kpi_mod.parse_date_series(unified["검토일"]).dropna() if "검토일" in unified.columns else pd.Series(dtype="datetime64[ns]")
    if not parsed_review_date.empty:
        base_date_value = parsed_review_date.iloc[0].date()
    unified = order_review_mod.apply_order_review_reference_flags(unified, base_date=base_date_value)

    sales_3m = pd.to_numeric(unified["기준 3개월 판매수량"], errors="coerce").fillna(0)
    monthly_sales = pd.to_numeric(unified["월평균 판매수량"], errors="coerce").fillna(0)
    eu_stock = pd.to_numeric(unified["유럽 가용재고"], errors="coerce").fillna(0)
    shipping_stock = pd.to_numeric(unified["운송중 수량"], errors="coerce").fillna(0)
    combined_stock = eu_stock + shipping_stock
    safety_stock = pd.to_numeric(unified["안전재고 필요 수량"], errors="coerce").fillna(0)
    open_po_qty = pd.to_numeric(unified["미입고 수량"], errors="coerce").fillna(0)
    hq_eu_stock = pd.to_numeric(unified["본사 EU창고"], errors="coerce").fillna(0)
    order_required_qty = pd.to_numeric(unified["발주 필요 수량"], errors="coerce").fillna(0)
    target_months = pd.to_numeric(_series_from_candidates(base, ["목표 운영 개월 수"], 3.0), errors="coerce").fillna(3.0)
    target_stock = monthly_sales * target_months
    pipeline_qty = shipping_stock
    available_projection = eu_stock - safety_stock + pipeline_qty
    pipeline_display = pipeline_qty.round(0).astype(int).astype(object).where(pipeline_qty.ne(0), "-")
    replenishment_qty = (safety_stock - combined_stock).clip(lower=0)

    def cover_months(qty: pd.Series) -> pd.Series:
        display_monthly = monthly_sales.round(0)
        return np.where(display_monthly.gt(0), qty / display_monthly.replace(0, np.nan), 0)

    risk_status = (
        unified["입고 전 결품 위험 여부"].astype(str).eq("Y")
        | unified["즉시OOS여부"].astype(str).eq("Y")
        | unified["우선 액션"].astype(str).str.contains("OOS|품절|결품|ETA 지연", na=False)
    )
    reference_df = pd.DataFrame(
        {
            "No": range(1, len(unified) + 1),
            "SKU": unified["상품코드"],
            "상품명": unified["상품명"],
            "브랜드": unified["브랜드"],
            "3개월 판매량": sales_3m.round(0).astype(int),
            "월평균 판매량": monthly_sales.round(0).astype(int),
            "안전재고": safety_stock.round(0).astype(int),
            "1. 유럽 재고": eu_stock.round(0).astype(int),
            "2. 운송 재고": shipping_stock.round(0).astype(int),
            "3. 유럽+운송": combined_stock.round(0).astype(int),
            "1. 유럽 재고(M)": pd.Series(cover_months(eu_stock), index=unified.index).round(1),
            "2. 운송중(M)": pd.Series(cover_months(shipping_stock), index=unified.index).round(1),
            "3. 합산(M)": pd.Series(cover_months(combined_stock), index=unified.index).round(1),
            "상태": np.where(risk_status, "OOS", "-"),
            "유럽 재고": eu_stock.round(0).astype(int),
            "(-) 출고 물량\n(미래 안전재고 기간\n동안 소진될 예상량)": safety_stock.round(0).astype(int),
            "(+) 입고물량\n(파이프라인 합계)": pipeline_display,
            "가용 재고": available_projection.round(0).astype(int),
            "미입고 현황": open_po_qty.round(0).astype(int),
            "EU 창고 재고": hq_eu_stock.round(0).astype(int),
            "발주 필요수량": order_required_qty.round(0).astype(int),
            "필요 보충량": replenishment_qty.round(0).astype(int),
        }
    )

    if arrival_calendar_df is not None and not arrival_calendar_df.empty and {"SKU", "도착일", "수량"}.issubset(arrival_calendar_df.columns):
        arrival_work = arrival_calendar_df[["SKU", "도착일", "수량"]].copy()
        arrival_work["SKU"] = arrival_work["SKU"].map(_clean_merge_key)
        arrival_work["도착일"] = kpi_mod.parse_date_series(arrival_work["도착일"]).dt.date
        arrival_work["수량"] = pd.to_numeric(arrival_work["수량"], errors="coerce").fillna(0)
        arrival_work = arrival_work[arrival_work["SKU"].ne("") & arrival_work["도착일"].notna() & arrival_work["수량"].gt(0)]
        if not arrival_work.empty:
            pivot = arrival_work.pivot_table(index="SKU", columns="도착일", values="수량", aggfunc="sum", fill_value=0)
            pivot = pivot.reindex(columns=sorted(pivot.columns))
            sku_keys = unified["상품코드"].map(_clean_merge_key)
            for arrival_date in pivot.columns:
                label = f"{arrival_date.year}-{arrival_date.month}-{arrival_date.day}"
                mapped_qty = sku_keys.map(pivot[arrival_date]).fillna(0).round(0).astype(int)
                if mapped_qty.sum() > 0:
                    reference_df[label] = mapped_qty.to_numpy()

    fixed_columns = [col for col in UNIFIED_ESM_ORDER_REVIEW_COLUMNS if col in reference_df.columns]
    dynamic_columns = [col for col in reference_df.columns if col not in fixed_columns]
    return reference_df[fixed_columns + dynamic_columns].reset_index(drop=True)


def validate_unified_esm_order_review_sheet(
    old_internal_df: pd.DataFrame,
    unified_df: pd.DataFrame,
    old_order_df: pd.DataFrame,
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []

    def add(level: str, item: str, detail: str, value: object = "") -> None:
        rows.append({"구분": level, "항목": item, "내용": detail, "건수/값": value})

    def norm_series(df: pd.DataFrame, candidates: list[str], default: object = "") -> pd.Series:
        return _series_from_candidates(df, candidates, default)

    internal_sku = norm_series(old_internal_df, ["상품코드"]).map(_clean_merge_key)
    unified_sku = norm_series(unified_df, ["상품코드"]).map(_clean_merge_key)

    if len(old_internal_df) == len(unified_df):
        add("통과", "통합 시트 행 수", "기존 ESM 내부 검토와 ESM 발주 검토 행 수가 같습니다.", len(unified_df))
    else:
        add("경고", "통합 시트 행 수", "기존 ESM 내부 검토와 ESM 발주 검토 행 수가 다릅니다.", f"{len(old_internal_df):,} -> {len(unified_df):,}")

    old_sku_set = set(internal_sku[internal_sku.ne("")])
    new_sku_set = set(unified_sku[unified_sku.ne("")])
    missing_skus = sorted(old_sku_set - new_sku_set)
    extra_skus = sorted(new_sku_set - old_sku_set)
    add(
        "통과" if not missing_skus else "경고",
        "상품코드 기준 누락 SKU",
        "기존 ESM 내부 검토 데이터에는 있으나 ESM 발주 검토에는 없는 SKU 수입니다.",
        len(missing_skus),
    )
    if missing_skus:
        add("경고", "누락 SKU 목록", "상위 20개 누락 SKU입니다.", ", ".join(missing_skus[:20]))
    add(
        "통과" if not extra_skus else "경고",
        "상품코드 기준 추가 SKU",
        "ESM 발주 검토에만 존재하는 SKU 수입니다. merge로 행이 늘었는지 확인하는 보조 지표입니다.",
        len(extra_skus),
    )

    old_counts = internal_sku[internal_sku.ne("")].value_counts()
    new_counts = unified_sku[unified_sku.ne("")].value_counts()
    expanded_skus = sorted(sku for sku, count in new_counts.items() if count > old_counts.get(sku, 0))
    add(
        "통과" if not expanded_skus else "경고",
        "merge 행 증가 SKU",
        "통합 후 기존보다 행 수가 늘어난 SKU 수입니다.",
        len(expanded_skus),
    )
    if expanded_skus:
        add("경고", "merge 행 증가 SKU 목록", "상위 20개 행 증가 SKU입니다.", ", ".join(expanded_skus[:20]))

    blank_sku_count = int(unified_sku.eq("").sum())
    if blank_sku_count:
        add("경고", "상품코드 공백", "ESM 발주 검토에 상품코드가 비어 있는 행이 있습니다.", blank_sku_count)
    else:
        add("통과", "상품코드 공백", "ESM 발주 검토의 상품코드 공백 행이 없습니다.", 0)

    internal_dup_count = int(internal_sku[internal_sku.ne("")].duplicated(keep=False).sum())
    if internal_dup_count:
        add("경고", "ESM 내부 검토 상품코드 중복", "통합 시트는 기존 행을 삭제하지 않고 보존했으며, 중복 SKU 값 검증은 첫 매칭 기준으로 기록합니다.", internal_dup_count)
    else:
        add("통과", "ESM 내부 검토 상품코드 중복", "ESM 내부 검토 상품코드 중복이 없습니다.", 0)

    compare_cols = [
        ("상품코드", ["상품코드"], ["상품코드"]),
        ("브랜드", ["브랜드"], ["브랜드"]),
        ("유럽 가용재고", ["유럽 가용재고"], ["유럽 가용재고"]),
        ("운송중 수량", ["운송중 수량"], ["운송중 수량", "운송중"]),
        ("입고 전 예상 결품량", ["입고 전 예상 결품량"], ["입고 전 예상 결품량"]),
        ("미입고 수량", ["미입고 수량"], ["미입고 수량", "미입고"]),
        ("발주 필요 수량", ["발주 필요 수량"], ["발주 필요 수량", "발주필요수량"]),
        ("긴급 보충 필요 수량", ["긴급 보충 필요 수량"], ["긴급 보충 필요 수량"]),
        ("우선 액션", ["우선 액션"], ["우선 액션"]),
    ]
    internal_key = internal_sku
    unified_by_sku = unified_df.assign(_검증키=_series_from_candidates(unified_df, ["상품코드"]).map(_clean_merge_key)).drop_duplicates("_검증키").set_index("_검증키")
    for label, internal_candidates, unified_candidates in compare_cols:
        internal_values = _series_from_candidates(old_internal_df, internal_candidates)
        unified_col = order_review_mod.first_existing_col(unified_df, unified_candidates)
        if unified_col is None:
            add("경고", f"{label} 값 검증", "통합 시트에서 검증 대상 컬럼을 찾지 못했습니다.", "")
            continue
        matched_values = internal_key.map(unified_by_sku[unified_col]) if not unified_by_sku.empty else pd.Series(index=old_internal_df.index, dtype=object)
        left_num = pd.to_numeric(internal_values, errors="coerce")
        right_num = pd.to_numeric(matched_values, errors="coerce")
        numeric_mask = left_num.notna() | right_num.notna()
        numeric_diff = (left_num.fillna(0) - right_num.fillna(0)).abs() > 0.000001
        left_text = internal_values.fillna("").astype(str).str.strip()
        right_text = matched_values.fillna("").astype(str).str.strip()
        text_diff = left_text != right_text
        diff_count = int((numeric_diff & numeric_mask).sum() + (text_diff & ~numeric_mask).sum())
        add("통과" if diff_count == 0 else "경고", f"{label} 값 검증", "기존 ESM 내부 검토와 통합 시트의 상품코드 기준 값 비교입니다.", diff_count)

    order_sku = _series_from_candidates(old_order_df, ["상품코드"]).map(_clean_merge_key)
    order_dup_count = int(order_sku[order_sku.ne("")].duplicated(keep=False).sum())
    if order_dup_count:
        add("경고", "발주 검토 상품코드 중복", "중복 상품코드는 임의 병합하지 않고 고유하게 매칭 가능한 값만 통합했습니다.", order_dup_count)
    else:
        add("통과", "발주 검토 상품코드 중복", "발주 검토 상품코드 중복이 없습니다.", 0)

    order_amount = _series_from_candidates(old_order_df, ["발주필요금액_KRW", "발주필요금액"], np.nan)
    unique_amount_lookup, _ = _build_unique_lookup_df(
        pd.DataFrame({"상품코드": order_sku, "발주필요금액_KRW": order_amount}),
        ["상품코드"],
    )
    if unique_amount_lookup.empty:
        add("경고", "발주필요금액_KRW 매칭", "발주 검토에서 발주필요금액_KRW 매칭 기준을 찾지 못했습니다.", 0)
    else:
        unified_keys = unified_sku
        matched_count = int(unified_keys.isin(set(unique_amount_lookup["상품코드"])).sum())
        add("통과", "발주필요금액_KRW 매칭", "발주 검토 금액 컬럼을 상품코드 기준으로 통합 시트에 매칭했습니다.", matched_count)
        unified_amount = pd.to_numeric(_series_from_candidates(unified_df, ["발주필요금액(KRW)", "발주필요금액_KRW"], np.nan), errors="coerce").fillna(0)
        expected_amount = pd.to_numeric(unified_keys.map(unique_amount_lookup.set_index("상품코드")["발주필요금액_KRW"]), errors="coerce").fillna(0)
        amount_diff = (unified_amount - expected_amount).abs() > 0.000001
        comparable = unified_keys.isin(set(unique_amount_lookup["상품코드"]))
        diff_count = int((amount_diff & comparable).sum())
        add(
            "통과" if diff_count == 0 else "경고",
            "발주필요금액_KRW 값 검증",
            "상품코드 기준으로 기존 발주 검토 데이터의 발주필요금액_KRW와 통합 시트 값을 비교했습니다.",
            diff_count,
        )

    return pd.DataFrame(rows, columns=["구분", "항목", "내용", "건수/값"])


def format_unified_esm_order_review_sheet(
    ws,
    settings: dict | None = None,
    dashboard_summary: list[tuple[str, object]] | None = None,
) -> None:
    from copy import copy
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    top_rows = 8
    header_row = top_rows + 1
    group_row = top_rows
    data_start_row = header_row + 1

    if str(ws["A1"].value or "") in {"코드", "No"}:
        ws.insert_rows(1, top_rows)

    dark_fill = PatternFill("solid", fgColor="4F6228")
    light_fill = PatternFill("solid", fgColor="E2F0D9")
    param_value_fill = PatternFill("solid", fgColor="FFF2CC")
    summary_value_fill = PatternFill("solid", fgColor="EBF1DE")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    order_header_fill = PatternFill("solid", fgColor="F4B183")
    alert_header_fill = PatternFill("solid", fgColor="FFFF00")
    risk_fill = PatternFill("solid", fgColor="FFC7CE")
    order_needed_fill = PatternFill("solid", fgColor="FCE4D6")
    white_font = Font(color="FFFFFF", bold=True)
    green_font = Font(color="4F6228", bold=True)
    header_font = Font(name="맑은 고딕", color="4F6228", bold=True, size=9)
    group_header_font = Font(name="맑은 고딕", color="375623", bold=True, size=8)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    raw_max_column = ws.max_column
    headers = {str(ws.cell(header_row, col).value or ""): col for col in range(1, raw_max_column + 1)}
    eligibility_col = headers.get(ORDER_REVIEW_ELIGIBILITY_HELPER)
    data_max_column = eligibility_col - 1 if eligibility_col == raw_max_column else raw_max_column
    if eligibility_col:
        ws.column_dimensions[get_column_letter(eligibility_col)].hidden = True
    def numeric_sum(header: str) -> float:
        col_idx = headers.get(header)
        if not col_idx:
            return 0.0
        values = [ws.cell(row_idx, col_idx).value for row_idx in range(data_start_row, ws.max_row + 1)]
        return float(pd.to_numeric(pd.Series(values), errors="coerce").fillna(0).sum())

    sku_col = headers.get("SKU")
    if sku_col:
        sku_values = [
            str(ws.cell(row_idx, sku_col).value or "").strip()
            for row_idx in range(data_start_row, ws.max_row + 1)
        ]
        sku_count = len({value for value in sku_values if value})
    else:
        sku_count = max(ws.max_row - header_row, 0)

    summary_start_col = 5
    summary_value_col = summary_start_col + 1
    def col_letter(header: str) -> str | None:
        col_idx = headers.get(header)
        return get_column_letter(col_idx) if col_idx else None

    sku_letter = col_letter("SKU")
    eu_letter = col_letter("1. 현지 재고")
    combined_letter = col_letter("3. 현지+운송")
    order_qty_letter = (
        col_letter("발주필요수량")
        or col_letter("발주 필요수량")
        or col_letter("최종 발주 필요수량")
    )
    open_po_letter = col_letter("미입고 현황")
    scenario_label = "미입고 해소 후" if str((settings or {}).get("inbound_scenario", "before")).strip().lower() == "after" else "미입고 해소 전"
    before_order_needed_formula = (
        f'COUNTIFS({order_qty_letter}{data_start_row}:{order_qty_letter}100000,">0",'
        f'{sku_letter}{data_start_row}:{sku_letter}100000,"<>",'
        f'{sku_letter}{data_start_row}:{sku_letter}100000,"<>0")'
        if order_qty_letter and sku_letter
        else "0"
    )
    after_order_needed_formula = (
        f'SUMPRODUCT(--({order_qty_letter}{data_start_row}:{order_qty_letter}100000>{open_po_letter}{data_start_row}:{open_po_letter}100000))'
        if order_qty_letter and open_po_letter
        else "0"
    )
    ws["A2"] = "미입고 반영 기준"
    ws["B2"] = scenario_label
    scenario_validation = DataValidation(
        type="list",
        formula1='"미입고 해소 전,미입고 해소 후"',
        allow_blank=False,
    )
    scenario_validation.errorTitle = "미입고 반영 기준"
    scenario_validation.error = "목록에서 미입고 해소 전 또는 미입고 해소 후를 선택해 주세요."
    scenario_validation.promptTitle = "미입고 반영 기준"
    scenario_validation.prompt = "미입고 해소 전/후를 선택하면 요약 수치가 해당 기준으로 계산됩니다."
    ws.add_data_validation(scenario_validation)
    scenario_validation.add(ws["B2"])
    ws["A2"].fill = light_fill
    ws["A2"].font = Font(name="맑은 고딕", color="375623", bold=True, size=9)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].border = border
    ws["B2"].fill = param_value_fill
    ws["B2"].font = Font(name="맑은 고딕", color="7F6000", bold=True, size=10)
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B2"].border = border
    summary = [
        ("SKU", f'=COUNTIFS({sku_letter}{data_start_row}:{sku_letter}1048576,"<>",{sku_letter}{data_start_row}:{sku_letter}1048576,"<>0")' if sku_letter else sku_count),
        ("현지 재고 수량", f"=SUM({eu_letter}{data_start_row}:{eu_letter}100000)" if eu_letter else numeric_sum("1. 현지 재고")),
        ("현지+운송 수량", f"=SUM({combined_letter}{data_start_row}:{combined_letter}100000)" if combined_letter else numeric_sum("3. 현지+운송")),
        (
            "OOS SKU",
            f'=COUNTIFS({eu_letter}{data_start_row}:{eu_letter}100000,"<=0",'
            f'{sku_letter}{data_start_row}:{sku_letter}100000,"<>",'
            f'{sku_letter}{data_start_row}:{sku_letter}100000,"<>0")'
            if eu_letter and sku_letter
            else 0,
        ),
        (
            "발주필요 SKU",
            f'=IF($B$2="미입고 해소 후",{after_order_needed_formula},{before_order_needed_formula})',
        ),
    ]
    for offset, (label, value) in enumerate(summary, start=1):
        label_cell = ws.cell(offset, summary_start_col, label)
        value_cell = ws.cell(offset, summary_value_col, value)
        label_cell.fill = dark_fill
        label_cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=9)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.fill = summary_value_fill
        value_cell.font = Font(name="맑은 고딕", color="4F6228", bold=True, size=11)
        value_cell.number_format = "#,##0"
        value_cell.alignment = Alignment(horizontal="right", vertical="center")
        label_cell.border = value_cell.border = border
        if len(label) >= 18:
            ws.row_dimensions[offset].height = max(ws.row_dimensions[offset].height or 15, 30)

    parameter_start = 8
    settings = settings or {}
    lead_times = (
        settings.get("lead_times")
        if "lead_times" in settings
        else DEFAULT_LEAD_TIME_DAYS
    ) or {}
    base_date = settings.get("base_date") or korea_today()
    parameter_labels = [
        "안전재고(M)",
        "Air\nL/T(일)",
        "해운\nL/T(일)",
        "철송\nL/T(일)",
        "트럭킹\nL/T(일)",
        "기준일",
    ]
    parameter_values = [
        float(settings.get("safety_months", 3)),
        int(lead_times["항공"]) if "항공" in lead_times else "-",
        int(lead_times["해운"]) if "해운" in lead_times else "-",
        int(lead_times["철송"]) if "철송" in lead_times else "-",
        int(lead_times["트럭"]) if "트럭" in lead_times else "-",
        base_date,
    ]
    for idx, label in enumerate(parameter_labels, start=parameter_start):
        cell = ws.cell(3, idx, label)
        cell.fill = light_fill
        cell.font = Font(name="맑은 고딕", color="375623", bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for idx, value in enumerate(parameter_values, start=parameter_start):
        cell = ws.cell(4, idx, value)
        cell.fill = param_value_fill
        cell.font = Font(name="맑은 고딕", color="7F6000", bold=True, size=12)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        if isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
    order_total_col = (
        headers.get("발주필요수량")
        or headers.get("발주 필요수량")
        or headers.get("최종 발주 필요수량")
        or 17
    )
    order_total_label = ws.cell(5, order_total_col, "총 발주 필요량")
    order_total_value = ws.cell(6, order_total_col, f"=SUM({order_qty_letter}{data_start_row}:{order_qty_letter}100000)" if order_qty_letter else 0)
    order_total_label.fill = dark_fill
    order_total_label.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=9)
    order_total_label.alignment = Alignment(horizontal="center", vertical="center")
    order_total_value.fill = orange_fill
    order_total_value.font = Font(name="맑은 고딕", color="833C00", bold=True, size=9)
    order_total_value.alignment = Alignment(horizontal="center", vertical="center")
    order_total_value.number_format = "#,##0"
    order_total_label.border = order_total_value.border = border

    fixed_col_count = len(UNIFIED_ESM_ORDER_REVIEW_COLUMNS)
    first_eta_col = fixed_col_count + 1
    if data_max_column >= first_eta_col:
        eta_total_label = ws.cell(5, first_eta_col, "ETA 날짜별 총 도착 수량")
        if data_max_column > first_eta_col:
            ws.merge_cells(start_row=5, start_column=first_eta_col, end_row=5, end_column=data_max_column)
        eta_total_label.fill = dark_fill
        eta_total_label.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=9)
        eta_total_label.alignment = Alignment(horizontal="center", vertical="center")
        eta_total_label.border = border

        for col_idx in range(first_eta_col, data_max_column + 1):
            col_letter = get_column_letter(col_idx)
            eta_total_value = ws.cell(6, col_idx, f"=SUM({col_letter}{data_start_row}:{col_letter}100000)")
            eta_total_value.fill = light_fill
            eta_total_value.font = Font(name="맑은 고딕", color="375623", bold=True, size=9)
            eta_total_value.alignment = Alignment(horizontal="center", vertical="center")
            eta_total_value.number_format = "#,##0"
            eta_total_value.border = border

    group_specs = [
        (headers.get("1. 현지 재고"), headers.get("3. 현지+운송"), "재고 수량"),
        (headers.get("1. 현지 재고(M)"), headers.get("3. 합산(M)"), "재고 보유개월수"),
    ]
    grouped_header_cols = set()
    for start_col, end_col, title in group_specs:
        if not start_col or not end_col:
            continue
        grouped_header_cols.update(range(start_col, end_col + 1))
        ws.merge_cells(start_row=group_row, start_column=start_col, end_row=group_row, end_column=end_col)
        cell = ws.cell(group_row, start_col, title)
        cell.fill = light_fill
        cell.font = group_header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    simulation_end_col = headers.get("3. 합산(M)", 13)
    for col_idx in range(1, data_max_column + 1):
        group_cell = ws.cell(group_row, col_idx)
        header_cell = ws.cell(header_row, col_idx)
        if col_idx not in grouped_header_cols:
            group_cell.fill = light_fill if col_idx <= simulation_end_col else orange_fill
            group_cell.border = border
        header_cell.fill = orange_fill if col_idx > simulation_end_col else light_fill
        header_cell.font = Font(name="맑은 고딕", color="375623" if col_idx <= simulation_end_col else "833C00", bold=True, size=9)
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cell.border = border

    display_header_labels = {
        "1. 현지 재고": "1. 현지\n재고\n(가용재고)",
        "2. 운송 재고": "2. 운송\n재고",
        "3. 현지+운송": "3. 현지\n+운송",
        "1. 현지 재고(M)": "1. 현지\n재고(M)",
        "2. 운송중(M)": "2. 운송중\n(M)",
        "3. 합산(M)": "3. 합산\n(M)",
        "알림": "⚠ 알람",
        "미입고 현황": "미입고 현황",
        "현지 창고 재고": "현지 창고 재고",
        "발주필요수량": "발주\n필요수량",
    }
    for source_header, display_label in display_header_labels.items():
        col_idx = headers.get(source_header)
        if col_idx:
            ws.cell(header_row, col_idx).value = display_label

    for col_idx in range(fixed_col_count + 1, data_max_column + 1):
        total_cell = ws.cell(group_row, col_idx)
        col_letter = get_column_letter(col_idx)
        total_cell.value = f"=SUM({col_letter}{data_start_row}:{col_letter}{ws.max_row})"
        total_cell.fill = light_fill
        total_cell.font = header_font
        total_cell.number_format = "#,##0"
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.border = border
        header_cell = ws.cell(header_row, col_idx)
        header_cell.fill = light_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cell.border = border

    widths = {
        "코드": 15.14,
        "No": 5,
        "SKU": 63.57,
        "브랜드": 13.71,
        "3개월 판매량": 12.14,
        "월평균 판매량": 14.29,
        "안전재고": 9,
        "1. 현지 재고": 10.29,
        "1. 현지\n재고\n(가용재고)": 10.29,
        "2. 운송 재고": 13,
        "2. 운송\n재고": 13,
        "3. 현지+운송": 11,
        "3. 현지\n+운송": 11,
        "1. 현지 재고(M)": 8,
        "1. 현지\n재고(M)": 8,
        "2. 운송중(M)": 9,
        "2. 운송중\n(M)": 9,
        "3. 합산(M)": 8,
        "3. 합산\n(M)": 8,
        "알림": 11,
        "⚠ 알람": 11,
        "미입고 현황": 13,
        "현지 창고 재고": 17.57,
        "발주필요수량": 14.71,
        "발주\n필요수량": 14.71,
    }
    for col_idx in range(1, data_max_column + 1):
        header = str(ws.cell(header_row, col_idx).value or "")
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = widths.get(header, 12 if col_idx > fixed_col_count else 13)
    if order_total_col:
        order_total_letter = get_column_letter(order_total_col)
        ws.column_dimensions[order_total_letter].width = max(ws.column_dimensions[order_total_letter].width or 0, 14.71)

    decimal_headers = {"1. 현지 재고(M)", "2. 운송중(M)", "3. 합산(M)", "1. 현지\n재고(M)", "2. 운송중\n(M)", "3. 합산\n(M)"}
    text_headers = {"코드", "SKU", "브랜드", "알림", "⚠ 알람"}
    data_end_row = ws.max_row
    data_row_count = max(data_end_row - data_start_row + 1, 0)
    fast_data_style = data_row_count > 300
    alert_col = headers.get("알림")
    open_po_col = headers.get("미입고 현황")
    hq_stock_col = headers.get("현지 창고 재고")
    order_required_col = (
        headers.get("발주필요수량")
        or headers.get("발주 필요수량")
        or headers.get("최종 발주 필요수량")
    )
    for col_idx, fill, font_color in [
        (alert_col, alert_header_fill, "7F6000"),
        (open_po_col, orange_fill, "833C00"),
        (hq_stock_col, orange_fill, "833C00"),
        (order_required_col, order_header_fill, "833C00"),
    ]:
        if not col_idx:
            continue
        for row_idx in (group_row, header_row):
            cell = ws.cell(row_idx, col_idx)
            cell.fill = fill
            cell.font = Font(name="맑은 고딕", color=font_color, bold=True, size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    if fast_data_style and alert_col:
        status_letter = get_column_letter(alert_col)
        ws.conditional_formatting.add(
            f"A{data_start_row}:{get_column_letter(data_max_column)}{data_end_row}",
            FormulaRule(formula=[f'ISNUMBER(SEARCH("OOS",${status_letter}{data_start_row}))'], fill=risk_fill),
        )
        ws.conditional_formatting.add(
            f"A{data_start_row}:{get_column_letter(data_max_column)}{data_end_row}",
            FormulaRule(formula=[f'ISNUMBER(SEARCH("발주필요",${status_letter}{data_start_row}))'], fill=order_needed_fill),
        )
    def fill_for_status(status_text: str):
        compact_status = status_text.replace(" ", "")
        if "OOS" in compact_status:
            return risk_fill
        if "발주필요" in compact_status:
            return order_needed_fill
        return None

    header_values = {col_idx: str(ws.cell(header_row, col_idx).value or "") for col_idx in range(1, data_max_column + 1)}
    sku_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_alignment = Alignment(horizontal="right", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")

    def apply_stock_eta_number_format(cell, header: str) -> None:
        if header in text_headers:
            cell.number_format = "@"
            return
        cell.number_format = "#,##0.##" if header in decimal_headers else "#,##0"

    monthly_sales_col = headers.get("월평균 판매량")
    combined_months_col = headers.get("3. 합산(M)")
    if fast_data_style:
        for row_idx in range(data_start_row, data_end_row + 1):
            status_text = str(ws.cell(row_idx, alert_col).value or "") if alert_col else ""
            row_fill = fill_for_status(status_text)
            for col_idx in range(1, data_max_column + 1):
                cell = ws.cell(row_idx, col_idx)
                cell.border = border
                apply_stock_eta_number_format(cell, header_values[col_idx])
                if row_fill:
                    cell.fill = row_fill
    else:
        for row_idx in range(data_start_row, data_end_row + 1):
            status_text = str(ws.cell(row_idx, alert_col).value or "") if alert_col else ""
            row_fill = fill_for_status(status_text)
            for col_idx in range(1, data_max_column + 1):
                cell = ws.cell(row_idx, col_idx)
                header = header_values[col_idx]
                cell.border = border
                cell.alignment = sku_alignment if header == "SKU" else center_alignment if col_idx in {alert_col, order_required_col} else right_alignment
                if row_fill:
                    cell.fill = row_fill
                if header in text_headers:
                    apply_stock_eta_number_format(cell, header)
                    if header != "SKU":
                        cell.alignment = center_alignment
                elif header in decimal_headers:
                    apply_stock_eta_number_format(cell, header)
                else:
                    apply_stock_eta_number_format(cell, header)

    if alert_col and eu_letter and order_qty_letter:
        for row_idx in range(data_start_row, data_end_row + 1):
            ws.cell(row_idx, alert_col).value = (
                f'=IF({eu_letter}{row_idx}<=0,"{ALERT_URGENT_LABEL}",'
                f'IF({order_qty_letter}{row_idx}>0,"{ALERT_LOCAL_SHORT_LABEL}","-"))'
            )

    safety_col = headers.get("안전재고")
    combined_col = headers.get("3. 현지+운송")
    safety_param_letter = get_column_letter(parameter_start)
    if monthly_sales_col and safety_col:
        monthly_letter = get_column_letter(monthly_sales_col)
        safety_letter = get_column_letter(safety_col)
        for row_idx in range(data_start_row, data_end_row + 1):
            ws.cell(row_idx, safety_col).value = f"=ROUND({monthly_letter}{row_idx}*${safety_param_letter}$4,0)"
    if order_required_col and monthly_sales_col and combined_col and safety_col:
        monthly_letter = get_column_letter(monthly_sales_col)
        combined_data_letter = get_column_letter(combined_col)
        safety_letter = get_column_letter(safety_col)
        eligibility_letter = get_column_letter(eligibility_col) if eligibility_col else None
        for row_idx in range(data_start_row, data_end_row + 1):
            cell = ws.cell(row_idx, order_required_col)
            shortage_formula = (
                f"IF({monthly_letter}{row_idx}=0,0,"
                f"MAX(0,{safety_letter}{row_idx}-{combined_data_letter}{row_idx}))"
            )
            cell.value = (
                f"=IF(${eligibility_letter}{row_idx}=0,0,{shortage_formula})"
                if eligibility_letter
                else f"={shortage_formula}"
            )
            cell.number_format = "#,##0"

    ws.row_dimensions[group_row].height = 13.5
    ws.row_dimensions[header_row].height = 40.5
    compact_heights = {1: 6, 2: 19.5, 3: 27, 4: 19.5, 5: 19.5, 6: 19.5, 7: 10.5}
    for row_idx, height in compact_heights.items():
        ws.row_dimensions[row_idx].height = height
    for col_letter, width in {
        "A": 15.14,
        "B": 5,
        "C": 63.57,
        "D": 13.71,
        "E": 12.14,
        "F": 14.29,
        "G": 9,
    }.items():
        ws.column_dimensions[col_letter].width = max(ws.column_dimensions[col_letter].width or 0, width)
    for idx, label in enumerate(parameter_labels, start=parameter_start):
        width = 12 if label == "기준일" else 7.5
        if label.startswith("안전재고"):
            width = 8.5
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)

    for col_idx in range(1, data_max_column + 1):
        if col_idx in grouped_header_cols:
            continue
        group_cell = ws.cell(group_row, col_idx)
        header_cell = ws.cell(header_row, col_idx)
        header_value = header_cell.value
        if header_value in (None, ""):
            continue
        group_cell.value = header_value
        group_cell.fill = copy(header_cell.fill)
        group_cell.font = copy(header_cell.font)
        group_cell.number_format = header_cell.number_format
        group_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        group_cell.border = border
        ws.merge_cells(start_row=group_row, start_column=col_idx, end_row=header_row, end_column=col_idx)

    if order_required_col:
        ws.cell(group_row, order_required_col).comment = Comment(
            "발주필요수량\n= MAX(0, 안전재고 목표수량 - (현지재고(가용재고) + 운송재고))\n\n"
            "발주제외/확인필요 SKU는 안전재고 개월 수를 변경해도 0으로 유지됩니다.\n"
            "발주필요수량 자체에는 미입고 현황과 현지 창고 재고를 차감하지 않습니다.\n"
            "상단의 '발주필요 SKU'는 상단의 '미입고 반영 기준' 선택값에 따라 집계합니다.",
            "USER",
        )

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(data_max_column)}{data_end_row}"
    ws.freeze_panes = f"D{data_start_row}"
    for row_idx in range(1, data_end_row + 1):
        for col_idx in range(1, data_max_column + 1):
            ws.cell(row_idx, col_idx).protection = Protection(locked=False)
    if alert_col:
        for row_idx in range(data_start_row, data_end_row + 1):
            ws.cell(row_idx, alert_col).protection = Protection(locked=True, hidden=True)
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.sort = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
