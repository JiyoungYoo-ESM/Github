from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TARGET_SHEETS = ["요약", "발주 검토", "ESM 내부 검토", "ETA 타임라인_검토"]
HEADER_KEYS = {
    "상품코드",
    "상품명",
    "브랜드",
    "발주필요수량",
    "발주 필요 수량",
    "우선 액션",
    "SKU",
    "월평균",
    "운송중",
    "권장 운송안",
    "항공 검토수량",
    "도착일",
    "운송수단",
}
TEXT_HEADERS = {
    "상품코드",
    "SKU",
    "바코드",
    "브랜드",
    "상품명",
    "우선 액션",
    "운송 검토안",
    "권장 대응 / 운송 검토안",
    "추천 운송수단",
    "권장 운송안",
    "권장 수량 요약",
    "판단 사유",
    "운송수단 추천 사유",
    "추천 사유",
    "판단메모",
    "발주 신호등",
    "발주 메모",
    "담당자",
    "제품상태",
    "구분",
}
DATE_HEADERS = {
    "예상 소진일",
    "첫 입고 예정일",
    "검토일",
    "최초 ETA",
    "ETA",
    "도착일",
    "출고일",
    "고갈 예정일",
    "고갈 예상일",
    "예상 도착 가능일",
}


def non_empty_bounds(ws) -> tuple[int, int]:
    max_row = 1
    max_col = 1
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.value not in (None, ""):
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    return max_row, max_col


def row_values(ws, row_idx: int, max_col: int) -> list[str]:
    values = []
    for col_idx in range(1, max_col + 1):
        value = ws.cell(row_idx, col_idx).value
        values.append(str(value).strip() if value is not None else "")
    return values


def find_header_rows(ws, summary: bool = False) -> list[int]:
    max_row, max_col = non_empty_bounds(ws)
    found: list[int] = []
    for row_idx in range(1, max_row + 1):
        values = row_values(ws, row_idx, max_col)
        keys = sum(1 for value in values if value in HEADER_KEYS)
        has_action_table = "우선 액션" in values and ("상품코드" in values or "SKU" in values)
        has_arrival_table = "도착일" in values and "운송수단" in values
        has_kpi_table = "KPI" in values and "값" in values
        if keys >= 2 or has_action_table or has_arrival_table or (summary and has_kpi_table):
            found.append(row_idx)
    if summary:
        return found
    return found[:1]


def header_map(ws, header_row: int, max_col: int) -> dict[str, int]:
    return {
        str(ws.cell(header_row, col_idx).value or "").strip(): col_idx
        for col_idx in range(1, max_col + 1)
        if str(ws.cell(header_row, col_idx).value or "").strip()
    }


def is_amount_header(header: str) -> bool:
    return "_KRW" in header or "_EUR" in header or "금액" in header


def is_month_or_cover_header(header: str) -> bool:
    return "개월" in header or "커버일수" in header or "커버" in header


def is_qty_header(header: str) -> bool:
    keywords = ["수량", "판매수량", "가용재고", "운송중", "재고", "수요", "파이프라인", "합계"]
    return any(keyword in header for keyword in keywords)


def is_date_header(header: str) -> bool:
    if header in DATE_HEADERS:
        return True
    if header.startswith("ETA "):
        return False
    return "일" in header and not any(token in header for token in ["공백일수", "커버일수", "경과/잔여일"])


def style_header(ws, header_rows: Iterable[int], max_col: int) -> None:
    header_fill = PatternFill("solid", fgColor="2F5496")
    for header_row in header_rows:
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(header_row, col_idx)
            if isinstance(cell, MergedCell) or cell.value in (None, ""):
                continue
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def apply_table_formats(ws, header_row: int) -> None:
    max_row, max_col = non_empty_bounds(ws)
    headers_by_col = {col: str(ws.cell(header_row, col).value or "").strip() for col in range(1, max_col + 1)}
    thin_border = Border(
        left=Side(style="thin", color="D9DEE8"),
        right=Side(style="thin", color="D9DEE8"),
        top=Side(style="thin", color="D9DEE8"),
        bottom=Side(style="thin", color="D9DEE8"),
    )

    for row in ws.iter_rows(min_row=header_row, max_row=max_row, min_col=1, max_col=max_col):
        if all(cell.value in (None, "") for cell in row):
            continue
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            header = headers_by_col.get(cell.column, "")
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=header in {"상품명", "판단메모", "판단 사유", "추천 사유", "권장 수량 요약"})
            if cell.row == header_row:
                continue
            if header in TEXT_HEADERS:
                cell.number_format = "@"
            elif is_date_header(header):
                cell.number_format = "yyyy-mm-dd"
            elif is_amount_header(header):
                cell.number_format = "#,##0"
            elif is_month_or_cover_header(header):
                cell.number_format = "0.0"
            elif is_qty_header(header) or header.startswith("ETA "):
                cell.number_format = "#,##0"


def apply_column_widths(ws, header_row: int) -> None:
    max_row, max_col = non_empty_bounds(ws)
    for col_idx in range(1, max_col + 1):
        header = str(ws.cell(header_row, col_idx).value or "").strip()
        letter = get_column_letter(col_idx)
        if header in {"상품명", "주요 상품", "판단메모", "판단 사유", "추천 사유", "권장 수량 요약"}:
            width = 34
        elif header in {"상품코드", "SKU", "바코드"}:
            width = 16
        elif header == "브랜드":
            width = 16
        elif is_amount_header(header) or is_qty_header(header) or is_month_or_cover_header(header) or header.startswith("ETA "):
            width = 14
        elif is_date_header(header):
            width = 14
        else:
            max_len = 0
            for row_idx in range(header_row, min(max_row, header_row + 50) + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            width = min(max(max_len + 2, 10), 24)
        ws.column_dimensions[letter].width = width


def apply_action_highlights(ws, header_row: int) -> None:
    max_row, max_col = non_empty_bounds(ws)
    headers = header_map(ws, header_row, max_col)
    action_col = headers.get("우선 액션")
    required_qty_col = headers.get("발주필요수량")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    red_fill = PatternFill("solid", fgColor="FCE4E4")

    for row_idx in range(header_row + 1, max_row + 1):
        if all(ws.cell(row_idx, col_idx).value in (None, "") for col_idx in range(1, max_col + 1)):
            continue
        if action_col:
            action_text = str(ws.cell(row_idx, action_col).value or "")
            if "입고 전 품절 위험 SKU" in action_text or "입고 전 품절 위험" in action_text:
                for col_idx in range(1, max_col + 1):
                    ws.cell(row_idx, col_idx).fill = red_fill
            elif "추가 발주 검토" in action_text or "신규 발주 검토" in action_text:
                for col_idx in range(1, max_col + 1):
                    ws.cell(row_idx, col_idx).fill = yellow_fill
        if required_qty_col:
            value = ws.cell(row_idx, required_qty_col).value
            try:
                numeric_value = float(str(value).replace(",", ""))
            except (TypeError, ValueError):
                numeric_value = 0
            if numeric_value > 0:
                for col_idx in range(1, max_col + 1):
                    cell = ws.cell(row_idx, col_idx)
                    if isinstance(cell, MergedCell):
                        continue
                    font = copy(cell.font)
                    font.bold = True
                    cell.font = font


def format_summary_sheet(ws) -> list[int]:
    max_row, max_col = non_empty_bounds(ws)
    if ws["A1"].value:
        ws["A1"].font = Font(size=18, bold=True, color="1F2937")
    for row_idx in range(1, min(max_row, 25) + 1):
        label = str(ws.cell(row_idx, 4).value or "")
        value = ws.cell(row_idx, 5)
        if label and value.value not in (None, ""):
            value.font = Font(bold=True)
    header_rows = find_header_rows(ws, summary=True)
    style_header(ws, header_rows, max_col)
    for header_row in header_rows:
        apply_table_formats(ws, header_row)
        apply_column_widths(ws, header_row)
    if header_rows:
        ws.freeze_panes = ws.cell(min(header_rows) + 1, 1).coordinate
    return header_rows


def format_table_sheet(ws) -> list[int]:
    max_row, max_col = non_empty_bounds(ws)
    header_rows = find_header_rows(ws)
    if not header_rows:
        return []
    header_row = header_rows[0]
    style_header(ws, [header_row], max_col)
    apply_table_formats(ws, header_row)
    apply_column_widths(ws, header_row)
    apply_action_highlights(ws, header_row)
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
    return [header_row]


def workbook_signature(wb) -> dict[str, object]:
    return {
        "sheetnames": list(wb.sheetnames),
        "dimensions": {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets},
        "headers": {
            ws.title: [tuple(row_values(ws, row_idx, non_empty_bounds(ws)[1])) for row_idx in find_header_rows(ws, summary=ws.title == "요약")]
            for ws in wb.worksheets
            if ws.title in TARGET_SHEETS
        },
    }


def format_workbook(input_path: Path, output_path: Path) -> dict[str, object]:
    wb = load_workbook(input_path)
    before = workbook_signature(wb)
    formatted_headers: dict[str, list[int]] = {}

    for sheet_name in TARGET_SHEETS:
        if sheet_name not in wb.sheetnames:
            formatted_headers[sheet_name] = []
            continue
        ws = wb[sheet_name]
        if sheet_name == "요약":
            formatted_headers[sheet_name] = format_summary_sheet(ws)
        else:
            formatted_headers[sheet_name] = format_table_sheet(ws)

    after = workbook_signature(wb)
    if before["sheetnames"] != after["sheetnames"]:
        raise RuntimeError("Sheet order or names changed unexpectedly.")
    if before["dimensions"] != after["dimensions"]:
        raise RuntimeError("Sheet dimensions changed unexpectedly.")
    if before["headers"] != after["headers"]:
        raise RuntimeError("Detected header/column order change unexpectedly.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "input": str(input_path),
        "output": str(output_path),
        "sheetnames_equal": before["sheetnames"] == after["sheetnames"],
        "dimensions_equal": before["dimensions"] == after["dimensions"],
        "headers_equal": before["headers"] == after["headers"],
        "formatted_headers": formatted_headers,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Apply openpyxl-only readability formatting to an order review workbook.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    result = format_workbook(args.input, args.output)
    print(result)


if __name__ == "__main__":
    main()
