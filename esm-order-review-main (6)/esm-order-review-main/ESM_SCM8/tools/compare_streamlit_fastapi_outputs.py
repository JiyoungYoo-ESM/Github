from __future__ import annotations

import argparse
import math
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


REPORT_PATH = Path("outputs") / "compare_streamlit_fastapi_report.xlsx"
NUMERIC_TOLERANCE = 1e-6
MAX_CONSOLE_ROWS = 25

IMPORTANT_COLUMNS: dict[str, list[str]] = {
    "상품코드": ["상품코드", "상품 코드", "품목코드", "itemcode", "productcode"],
    "SKU": ["SKU", "sku"],
    "추가 발주 필요 수량": ["추가 발주 필요 수량", "추가발주필요수량"],
    "최종 발주 필요수량": ["최종 발주 필요수량", "최종발주필요수량", "발주\n필요량"],
    "발주필요수량": ["발주필요수량", "발주 필요 수량", "미입고 미포함\n발주 필요수량"],
    "발주필요금액": ["발주필요금액", "발주 필요 금액", "추가 발주 필요 금액", "발주금액"],
    "가용 재고": ["가용 재고", "가용재고", "가용수량", "가용 재고 수량", "유럽 현재 재고"],
    "가용재고 부족분": ["가용재고 부족분", "가용 재고 부족분", "가용재고부족분"],
    "미입고 현황": ["미입고 현황", "미입고현황", "미입고수량", "미입고 수량"],
    "상태": ["상태"],
    "우선 액션": ["우선 액션", "우선액션", "최종 액션"],
    "발주검토": ["발주검토", "발주검토여부", "발주 검토", "구분"],
}

KEY_COLUMN_PRIORITY = ["상품코드", "SKU"]
TEXT_LIKE_COLUMNS = {
    "상품코드",
    "sku",
    "바코드",
    "상품명",
    "제품명",
    "브랜드",
    "상태",
    "우선액션",
    "최종액션",
    "발주검토",
    "발주검토여부",
}


@dataclass
class SheetTable:
    name: str
    hidden: bool
    header_row: int
    used_rows: int
    used_cols: int
    columns: list[str]
    df: pd.DataFrame


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def comparable_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if pd.isna(value):
        return ""
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return pd.to_datetime(value).strftime("%Y-%m-%d")
    return str(value).strip()


def make_unique_columns(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    columns: list[str] = []
    for idx, value in enumerate(values, start=1):
        base = comparable_text(value) or f"Column_{idx}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        columns.append(base if count == 1 else f"{base}__{count}")
    return columns


def trim_matrix(rows: list[list[Any]]) -> tuple[list[list[Any]], int, int]:
    last_row = 0
    last_col = 0
    for row_idx, row in enumerate(rows, start=1):
        row_last_col = 0
        for col_idx, value in enumerate(row, start=1):
            if not is_empty(value):
                row_last_col = col_idx
        if row_last_col:
            last_row = row_idx
            last_col = max(last_col, row_last_col)
    trimmed = [
        list(row[:last_col]) + [None] * max(0, last_col - len(row))
        for row in rows[:last_row]
    ]
    return trimmed, last_row, last_col


def header_score(row: list[Any]) -> tuple[int, int, int, int]:
    aliases = {compact_text(alias) for items in IMPORTANT_COLUMNS.values() for alias in items}
    non_empty = sum(not is_empty(value) for value in row)
    text_count = sum(isinstance(value, str) and value.strip() != "" for value in row)
    important_hits = sum(compact_text(value) in aliases for value in row if not is_empty(value))
    sku_hits = sum(compact_text(value) in {"상품코드", "sku"} for value in row if not is_empty(value))
    return important_hits, sku_hits, text_count, non_empty


def detect_header_row(matrix: list[list[Any]]) -> int:
    if not matrix:
        return 0
    best_idx = 0
    best_score = (-1, -1, -1, -1)
    for idx, row in enumerate(matrix[: min(len(matrix), 120)]):
        score = header_score(row)
        if score > best_score:
            best_idx = idx
            best_score = score
    return best_idx


def workbook_sheet_names(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_sheet(path: Path, sheet_name: str) -> SheetTable:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        matrix, used_rows, used_cols = trim_matrix(rows)
        if not matrix:
            return SheetTable(sheet_name, ws.sheet_state != "visible", 0, 0, 0, [], pd.DataFrame())

        header_idx = detect_header_row(matrix)
        columns = make_unique_columns(matrix[header_idx][:used_cols])
        data_rows: list[list[Any]] = []
        for row in matrix[header_idx + 1 :]:
            values = list(row[:used_cols]) + [None] * max(0, used_cols - len(row))
            if any(not is_empty(value) for value in values):
                data_rows.append(values)

        return SheetTable(
            name=sheet_name,
            hidden=ws.sheet_state != "visible",
            header_row=header_idx + 1,
            used_rows=used_rows,
            used_cols=used_cols,
            columns=columns,
            df=pd.DataFrame(data_rows, columns=columns),
        )
    finally:
        wb.close()


def find_column(columns: list[str], canonical: str) -> str | None:
    aliases = IMPORTANT_COLUMNS.get(canonical, [canonical])
    alias_set = {compact_text(alias) for alias in aliases}
    for column in columns:
        if compact_text(column) in alias_set:
            return column
    return None


def find_key_columns(streamlit: SheetTable, fastapi: SheetTable) -> tuple[str, str, str] | None:
    for canonical in KEY_COLUMN_PRIORITY:
        streamlit_col = find_column(streamlit.columns, canonical)
        fastapi_col = find_column(fastapi.columns, canonical)
        if streamlit_col and fastapi_col:
            return canonical, streamlit_col, fastapi_col
    return None


def to_numeric_series(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    text = series.map(comparable_text)
    text = text.str.replace(",", "", regex=False)
    text = text.str.replace("EUR", "", regex=False)
    text = text.str.replace("KRW", "", regex=False)
    text = text.str.replace("원", "", regex=False)
    text = text.str.replace("개", "", regex=False)
    text = text.str.replace("%", "", regex=False)
    text = text.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
    text = text.replace({"": None, "-": None, "nan": None, "None": None})
    return pd.to_numeric(text, errors="coerce")


def is_numeric_like(streamlit_series: pd.Series, fastapi_series: pd.Series, column_name: str) -> bool:
    if compact_text(column_name) in {compact_text(name) for name in TEXT_LIKE_COLUMNS}:
        return False
    combined = pd.concat([streamlit_series, fastapi_series], ignore_index=True)
    non_empty_count = int(combined.map(lambda value: not is_empty(value)).sum())
    if non_empty_count == 0:
        return False
    numeric_count = int(to_numeric_series(combined).notna().sum())
    return numeric_count >= max(1, int(non_empty_count * 0.6))


def numeric_sum(series: pd.Series) -> float:
    return float(to_numeric_series(series).fillna(0).sum())


def nearly_equal(left: float, right: float) -> bool:
    return abs(left - right) <= NUMERIC_TOLERANCE


def values_equal(streamlit_value: Any, fastapi_value: Any, numeric: bool) -> bool:
    if numeric:
        return nearly_equal(float(streamlit_value or 0), float(fastapi_value or 0))
    return comparable_text(streamlit_value) == comparable_text(fastapi_value)


def aggregate_by_sku(df: pd.DataFrame, key_col: str, value_col: str, numeric: bool) -> pd.Series:
    working = df[[key_col, value_col]].copy()
    working[key_col] = working[key_col].map(comparable_text)
    working = working[working[key_col] != ""]
    if numeric:
        working[value_col] = to_numeric_series(working[value_col]).fillna(0)
        return working.groupby(key_col, dropna=False)[value_col].sum()

    def join_unique(values: pd.Series) -> str:
        cleaned = sorted({comparable_text(value) for value in values if comparable_text(value)})
        return " | ".join(cleaned)

    return working.groupby(key_col, dropna=False)[value_col].agg(join_unique)


def is_stock_eta_sheet(sheet_name: str) -> bool:
    name = compact_text(sheet_name)
    return ("재고" in name or "stock" in name) and "eta" in name


def is_date_column_name(column_name: str) -> bool:
    text = comparable_text(column_name)
    if not text:
        return False
    if re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", text):
        return True
    parsed = pd.to_datetime(text, errors="coerce")
    return not pd.isna(parsed)


def compare_numeric_sums(streamlit: SheetTable, fastapi: SheetTable) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for column in sorted(set(streamlit.columns) & set(fastapi.columns)):
        if not is_numeric_like(streamlit.df[column], fastapi.df[column], column):
            continue
        streamlit_sum = numeric_sum(streamlit.df[column])
        fastapi_sum = numeric_sum(fastapi.df[column])
        diff = fastapi_sum - streamlit_sum
        rows.append(
            {
                "sheet": streamlit.name,
                "column": column,
                "streamlit_sum": streamlit_sum,
                "fastapi_sum": fastapi_sum,
                "diff": diff,
                "status": "PASS" if nearly_equal(streamlit_sum, fastapi_sum) else "FAIL",
            }
        )
    return rows


def compare_sku_values(streamlit: SheetTable, fastapi: SheetTable) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    key_info = find_key_columns(streamlit, fastapi)
    if not key_info:
        return rows

    key_name, streamlit_key, fastapi_key = key_info
    streamlit_skus = set(streamlit.df[streamlit_key].map(comparable_text)) - {""}
    fastapi_skus = set(fastapi.df[fastapi_key].map(comparable_text)) - {""}

    for sku in sorted(streamlit_skus - fastapi_skus):
        rows.append(
            {
                "sheet": streamlit.name,
                "key_column": key_name,
                "sku": sku,
                "column": "<SKU>",
                "streamlit_value": "present",
                "fastapi_value": "missing",
                "diff": "",
                "status": "FAIL",
            }
        )
    for sku in sorted(fastapi_skus - streamlit_skus):
        rows.append(
            {
                "sheet": streamlit.name,
                "key_column": key_name,
                "sku": sku,
                "column": "<SKU>",
                "streamlit_value": "missing",
                "fastapi_value": "present",
                "diff": "",
                "status": "FAIL",
            }
        )

    common_skus = streamlit_skus & fastapi_skus
    for canonical in IMPORTANT_COLUMNS:
        if canonical in KEY_COLUMN_PRIORITY:
            continue
        streamlit_col = find_column(streamlit.columns, canonical)
        fastapi_col = find_column(fastapi.columns, canonical)
        if not streamlit_col or not fastapi_col:
            continue
        numeric = is_numeric_like(streamlit.df[streamlit_col], fastapi.df[fastapi_col], canonical)
        streamlit_agg = aggregate_by_sku(streamlit.df, streamlit_key, streamlit_col, numeric)
        fastapi_agg = aggregate_by_sku(fastapi.df, fastapi_key, fastapi_col, numeric)
        for sku in sorted(common_skus):
            streamlit_value = streamlit_agg.get(sku, 0 if numeric else "")
            fastapi_value = fastapi_agg.get(sku, 0 if numeric else "")
            if values_equal(streamlit_value, fastapi_value, numeric):
                continue
            diff = float(fastapi_value or 0) - float(streamlit_value or 0) if numeric else ""
            rows.append(
                {
                    "sheet": streamlit.name,
                    "key_column": key_name,
                    "sku": sku,
                    "column": canonical,
                    "streamlit_value": streamlit_value,
                    "fastapi_value": fastapi_value,
                    "diff": diff,
                    "status": "FAIL",
                }
            )
    return rows


def compare_stock_eta_date_sums(streamlit: SheetTable, fastapi: SheetTable) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not is_stock_eta_sheet(streamlit.name):
        return rows
    common_columns = sorted(set(streamlit.columns) & set(fastapi.columns), key=streamlit.columns.index)
    for column in common_columns:
        if not is_date_column_name(column):
            continue
        streamlit_sum = numeric_sum(streamlit.df[column])
        fastapi_sum = numeric_sum(fastapi.df[column])
        diff = fastapi_sum - streamlit_sum
        rows.append(
            {
                "sheet": streamlit.name,
                "date_column": column,
                "streamlit_qty_sum": streamlit_sum,
                "fastapi_qty_sum": fastapi_sum,
                "diff": diff,
                "status": "PASS" if nearly_equal(streamlit_sum, fastapi_sum) else "FAIL",
            }
        )
    return rows


def compare_workbooks(streamlit_path: Path, fastapi_path: Path) -> dict[str, pd.DataFrame]:
    streamlit_sheets = workbook_sheet_names(streamlit_path)
    fastapi_sheets = workbook_sheet_names(fastapi_path)
    common_sheets = [sheet for sheet in streamlit_sheets if sheet in fastapi_sheets]

    sheet_rows: list[dict[str, Any]] = []
    column_rows: list[dict[str, Any]] = []
    numeric_rows: list[dict[str, Any]] = []
    sku_rows: list[dict[str, Any]] = []
    stock_eta_rows: list[dict[str, Any]] = []

    sheet_list_equal = streamlit_sheets == fastapi_sheets
    for sheet in streamlit_sheets:
        if sheet not in fastapi_sheets:
            sheet_rows.append({"sheet": sheet, "status": "FAIL", "reason": "missing in FastAPI output"})
    for sheet in fastapi_sheets:
        if sheet not in streamlit_sheets:
            sheet_rows.append({"sheet": sheet, "status": "FAIL", "reason": "added in FastAPI output"})

    for sheet in common_sheets:
        streamlit_table = read_sheet(streamlit_path, sheet)
        fastapi_table = read_sheet(fastapi_path, sheet)
        row_status = "PASS" if len(streamlit_table.df) == len(fastapi_table.df) else "FAIL"
        col_count_status = "PASS" if len(streamlit_table.columns) == len(fastapi_table.columns) else "FAIL"
        col_list_status = "PASS" if streamlit_table.columns == fastapi_table.columns else "FAIL"
        hidden_status = "PASS" if streamlit_table.hidden == fastapi_table.hidden else "FAIL"
        sheet_rows.append(
            {
                "sheet": sheet,
                "status": "PASS"
                if row_status == col_count_status == col_list_status == hidden_status == "PASS"
                else "FAIL",
                "reason": "",
                "streamlit_hidden": streamlit_table.hidden,
                "fastapi_hidden": fastapi_table.hidden,
                "hidden_status": hidden_status,
                "streamlit_header_row": streamlit_table.header_row,
                "fastapi_header_row": fastapi_table.header_row,
                "streamlit_used_rows": streamlit_table.used_rows,
                "fastapi_used_rows": fastapi_table.used_rows,
                "streamlit_data_rows": len(streamlit_table.df),
                "fastapi_data_rows": len(fastapi_table.df),
                "row_count_status": row_status,
                "streamlit_column_count": len(streamlit_table.columns),
                "fastapi_column_count": len(fastapi_table.columns),
                "column_count_status": col_count_status,
                "column_list_status": col_list_status,
            }
        )
        column_rows.append(
            {
                "sheet": sheet,
                "status": col_list_status,
                "streamlit_columns": "\n".join(streamlit_table.columns),
                "fastapi_columns": "\n".join(fastapi_table.columns),
                "removed_columns": "\n".join(
                    [col for col in streamlit_table.columns if col not in fastapi_table.columns]
                ),
                "added_columns": "\n".join(
                    [col for col in fastapi_table.columns if col not in streamlit_table.columns]
                ),
            }
        )
        numeric_rows.extend(compare_numeric_sums(streamlit_table, fastapi_table))
        sku_rows.extend(compare_sku_values(streamlit_table, fastapi_table))
        stock_eta_rows.extend(compare_stock_eta_date_sums(streamlit_table, fastapi_table))

    summary_rows = [
        {"item": "streamlit_file", "value": str(streamlit_path)},
        {"item": "fastapi_file", "value": str(fastapi_path)},
        {"item": "sheet_list_equal", "value": sheet_list_equal},
        {"item": "streamlit_sheet_count", "value": len(streamlit_sheets)},
        {"item": "fastapi_sheet_count", "value": len(fastapi_sheets)},
        {"item": "missing_sheet_count", "value": len(set(streamlit_sheets) - set(fastapi_sheets))},
        {"item": "added_sheet_count", "value": len(set(fastapi_sheets) - set(streamlit_sheets))},
        {"item": "sheet_structure_fail_count", "value": sum(row.get("status") == "FAIL" for row in sheet_rows)},
        {"item": "column_list_fail_count", "value": sum(row.get("status") == "FAIL" for row in column_rows)},
        {"item": "numeric_sum_fail_count", "value": sum(row.get("status") == "FAIL" for row in numeric_rows)},
        {"item": "sku_value_diff_count", "value": len(sku_rows)},
        {
            "item": "stock_eta_date_sum_fail_count",
            "value": sum(row.get("status") == "FAIL" for row in stock_eta_rows),
        },
    ]
    fail = (
        not sheet_list_equal
        or any(row.get("status") == "FAIL" for row in sheet_rows)
        or any(row.get("status") == "FAIL" for row in column_rows)
        or any(row.get("status") == "FAIL" for row in numeric_rows)
        or bool(sku_rows)
        or any(row.get("status") == "FAIL" for row in stock_eta_rows)
    )
    summary_rows.insert(0, {"item": "overall_status", "value": "FAIL" if fail else "PASS"})

    return {
        "Summary": pd.DataFrame(summary_rows),
        "Sheet Structure": pd.DataFrame(sheet_rows),
        "Columns": pd.DataFrame(column_rows),
        "Numeric Sums": pd.DataFrame(numeric_rows),
        "SKU Diffs": pd.DataFrame(sku_rows),
        "Stock ETA Date Sums": pd.DataFrame(stock_eta_rows),
    }


def write_report(report: dict[str, pd.DataFrame], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in report.items():
            df.to_excel(writer, sheet_name=sheet_name[:31], index=False)


def failing_rows(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    if "status" not in df.columns:
        return df
    return df[df["status"].eq("FAIL")]


def print_report(report: dict[str, pd.DataFrame], report_path: Path | None) -> None:
    summary = report["Summary"]
    status = summary.loc[summary["item"].eq("overall_status"), "value"].iloc[0]
    print(f"OVERALL: {status}")
    for _, row in summary.iloc[1:].iterrows():
        print(f"- {row['item']}: {row['value']}")

    sections = [
        ("Sheet structure differences", report["Sheet Structure"]),
        ("Column differences", report["Columns"]),
        ("Numeric sum differences", report["Numeric Sums"]),
        ("SKU-level differences", report["SKU Diffs"]),
        ("Stock ETA date sum differences", report["Stock ETA Date Sums"]),
    ]
    for title, df in sections:
        failed = failing_rows(df)
        if failed.empty:
            print(f"\n{title}: PASS")
            continue
        print(f"\n{title}: {len(failed)} issue(s)")
        print(failed.head(MAX_CONSOLE_ROWS).to_string(index=False))
        if len(failed) > MAX_CONSOLE_ROWS:
            print(f"... {len(failed) - MAX_CONSOLE_ROWS} more row(s) in the Excel report")

    if report_path is not None:
        print(f"\nReport saved: {report_path}")


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare Streamlit and FastAPI SCM Excel outputs."
    )
    parser.add_argument("streamlit_output", type=Path, help="Excel output generated by Streamlit")
    parser.add_argument("fastapi_output", type=Path, help="Excel output generated by FastAPI")
    parser.add_argument(
        "--report",
        type=Path,
        default=REPORT_PATH,
        help=f"Excel comparison report path. Default: {REPORT_PATH}",
    )
    parser.add_argument(
        "--no-report",
        action="store_true",
        help="Print console output only and do not write the Excel report.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    streamlit_path = args.streamlit_output.resolve()
    fastapi_path = args.fastapi_output.resolve()
    if not streamlit_path.exists():
        print(f"ERROR: Streamlit output file does not exist: {streamlit_path}", file=sys.stderr)
        return 2
    if not fastapi_path.exists():
        print(f"ERROR: FastAPI output file does not exist: {fastapi_path}", file=sys.stderr)
        return 2

    report = compare_workbooks(streamlit_path, fastapi_path)
    report_path: Path | None = None
    if not args.no_report:
        report_path = args.report.resolve()
        try:
            write_report(report, report_path)
        except Exception as exc:  # noqa: BLE001
            print(f"WARNING: could not save report: {exc}", file=sys.stderr)
            report_path = None

    print_report(report, report_path)
    status = report["Summary"].loc[report["Summary"]["item"].eq("overall_status"), "value"].iloc[0]
    return 0 if status == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
