from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


IMPORTANT_COLUMNS: dict[str, list[str]] = {
    "상품코드": ["상품코드", "상품 코드", "품목코드", "itemcode"],
    "SKU": ["SKU", "sku"],
    "추가 발주 필요 수량": ["추가 발주 필요 수량", "추가발주필요수량"],
    "발주필요수량": ["발주필요수량", "발주 필요 수량", "최종 발주 필요수량"],
    "발주필요금액": ["발주필요금액", "발주 필요 금액", "추가 발주 필요 금액"],
    "가용재고 부족분": ["가용재고 부족분", "가용 재고 부족분"],
    "운송중 수량": ["운송중 수량", "운송중수량", "운송 재고"],
    "미입고 수량": ["미입고 수량", "미입고수량", "미입고 현황"],
    "상태": ["상태"],
    "우선 액션": ["우선 액션", "우선액션"],
    "발주검토": ["발주검토", "발주검토여부"],
}

KEY_COLUMN_PRIORITY = ["상품코드", "SKU"]
TEXT_LIKE_COLUMN_NAMES = {
    "상품코드",
    "sku",
    "바코드",
    "상품명",
    "제품명",
    "브랜드",
    "상태",
    "우선액션",
    "발주검토",
    "발주검토여부",
}
NUMERIC_TOLERANCE = 1e-6
MAX_CONSOLE_ROWS = 30


@dataclass
class SheetTable:
    name: str
    header_row: int
    used_rows: int
    used_cols: int
    columns: list[str]
    df: pd.DataFrame


def compact_text(value: Any) -> str:
    return "".join(str(value or "").split()).lower()


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def comparable_text(value: Any) -> str:
    if value is None:
        return ""
    if pd.isna(value):
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            pass
    return str(value).strip()


def make_unique_columns(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    columns: list[str] = []
    for idx, value in enumerate(values, start=1):
        name = comparable_text(value) or f"Column_{idx}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        columns.append(name if count == 1 else f"{name}__{count}")
    return columns


def trim_matrix(rows: list[list[Any]]) -> tuple[list[list[Any]], int, int]:
    last_row = 0
    last_col = 0
    for row_idx, row in enumerate(rows, start=1):
        row_last_col = 0
        for col_idx, value in enumerate(row, start=1):
            if not is_empty(value):
                row_last_col = col_idx
        if row_last_col:
            last_row = row_idx
            last_col = max(last_col, row_last_col)
    trimmed = [list(row[:last_col]) + [None] * max(0, last_col - len(row)) for row in rows[:last_row]]
    return trimmed, last_row, last_col


def header_score(row: list[Any]) -> tuple[int, int, int]:
    compact_specials = {compact_text(alias) for aliases in IMPORTANT_COLUMNS.values() for alias in aliases}
    non_empty = sum(not is_empty(value) for value in row)
    text_count = sum(isinstance(value, str) and value.strip() != "" for value in row)
    special_hits = sum(compact_text(value) in compact_specials for value in row if not is_empty(value))
    return special_hits, text_count, non_empty


def detect_header_row(matrix: list[list[Any]]) -> int:
    if not matrix:
        return 0
    search_rows = matrix[: min(len(matrix), 100)]
    best_idx = 0
    best_score = (-1, -1, -1)
    for idx, row in enumerate(search_rows):
        score = header_score(row)
        if score > best_score:
            best_idx = idx
            best_score = score
    return best_idx


def read_sheet(workbook_path: Path, sheet_name: str) -> SheetTable:
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
    matrix, used_rows, used_cols = trim_matrix(raw_rows)
    if not matrix:
        wb.close()
        return SheetTable(sheet_name, 0, 0, 0, [], pd.DataFrame())

    header_idx = detect_header_row(matrix)
    columns = make_unique_columns(matrix[header_idx][:used_cols])
    data_rows = []
    for row in matrix[header_idx + 1 :]:
        values = list(row[:used_cols]) + [None] * max(0, used_cols - len(row))
        if any(not is_empty(value) for value in values):
            data_rows.append(values)
    wb.close()
    return SheetTable(
        name=sheet_name,
        header_row=header_idx + 1,
        used_rows=used_rows,
        used_cols=used_cols,
        columns=columns,
        df=pd.DataFrame(data_rows, columns=columns),
    )


def workbook_sheet_names(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def find_column(columns: list[str], canonical: str) -> str | None:
    aliases = IMPORTANT_COLUMNS.get(canonical, [canonical])
    alias_set = {compact_text(alias) for alias in aliases}
    for column in columns:
        if compact_text(column) in alias_set:
            return column
    return None


def find_key_columns(before: SheetTable, after: SheetTable) -> tuple[str, str, str] | None:
    for canonical in KEY_COLUMN_PRIORITY:
        before_col = find_column(before.columns, canonical)
        after_col = find_column(after.columns, canonical)
        if before_col and after_col:
            return canonical, before_col, after_col
    return None


def to_numeric_series(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    text = series.map(comparable_text)
    text = text.str.replace(",", "", regex=False)
    text = text.str.replace("₩", "", regex=False)
    text = text.str.replace("€", "", regex=False)
    text = text.str.replace("원", "", regex=False)
    text = text.str.replace("개", "", regex=False)
    text = text.str.replace("%", "", regex=False)
    text = text.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
    text = text.replace({"": None, "-": None, "nan": None, "None": None})
    return pd.to_numeric(text, errors="coerce")


def is_numeric_like(before_series: pd.Series, after_series: pd.Series, column_name: str) -> bool:
    if compact_text(column_name) in {compact_text(name) for name in TEXT_LIKE_COLUMN_NAMES}:
        return False
    combined = pd.concat([before_series, after_series], ignore_index=True)
    non_empty_count = combined.map(lambda value: not is_empty(value)).sum()
    if non_empty_count == 0:
        return False
    numeric_count = to_numeric_series(combined).notna().sum()
    return numeric_count >= max(1, int(non_empty_count * 0.6))


def numeric_sum(value: pd.Series) -> float:
    return float(to_numeric_series(value).fillna(0).sum())


def values_equal(before_value: Any, after_value: Any, numeric: bool) -> bool:
    if numeric:
        before_num = float(before_value or 0)
        after_num = float(after_value or 0)
        return abs(before_num - after_num) <= NUMERIC_TOLERANCE
    return comparable_text(before_value) == comparable_text(after_value)


def aggregate_by_sku(df: pd.DataFrame, key_col: str, value_col: str, numeric: bool) -> pd.Series:
    working = df[[key_col, value_col]].copy()
    working[key_col] = working[key_col].map(comparable_text)
    working = working[working[key_col] != ""]
    if numeric:
        working[value_col] = to_numeric_series(working[value_col]).fillna(0)
        return working.groupby(key_col, dropna=False)[value_col].sum()

    def join_unique(values: pd.Series) -> str:
        unique_values = sorted({comparable_text(value) for value in values if comparable_text(value) != ""})
        return " | ".join(unique_values)

    return working.groupby(key_col, dropna=False)[value_col].agg(join_unique)


def compare_sheet_data(before: SheetTable, after: SheetTable) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    numeric_rows: list[dict[str, Any]] = []
    sku_rows: list[dict[str, Any]] = []

    for column in sorted(set(before.columns) & set(after.columns)):
        if not is_numeric_like(before.df[column], after.df[column], column):
            continue
        before_sum = numeric_sum(before.df[column])
        after_sum = numeric_sum(after.df[column])
        diff = after_sum - before_sum
        numeric_rows.append(
            {
                "sheet": before.name,
                "column": column,
                "before_sum": before_sum,
                "after_sum": after_sum,
                "diff": diff,
                "status": "PASS" if abs(diff) <= NUMERIC_TOLERANCE else "FAIL",
            }
        )

    key_info = find_key_columns(before, after)
    if not key_info:
        return numeric_rows, sku_rows

    key_name, before_key, after_key = key_info
    before_skus = set(before.df[before_key].map(comparable_text)) - {""}
    after_skus = set(after.df[after_key].map(comparable_text)) - {""}
    for sku in sorted(before_skus - after_skus):
        sku_rows.append(
            {
                "sheet": before.name,
                "key_column": key_name,
                "sku": sku,
                "column": "<SKU>",
                "before_value": "present",
                "after_value": "missing",
                "diff": "",
                "status": "FAIL",
            }
        )
    for sku in sorted(after_skus - before_skus):
        sku_rows.append(
            {
                "sheet": before.name,
                "key_column": key_name,
                "sku": sku,
                "column": "<SKU>",
                "before_value": "missing",
                "after_value": "present",
                "diff": "",
                "status": "FAIL",
            }
        )

    common_skus = before_skus & after_skus
    for canonical in IMPORTANT_COLUMNS:
        if canonical in KEY_COLUMN_PRIORITY:
            continue
        before_col = find_column(before.columns, canonical)
        after_col = find_column(after.columns, canonical)
        if not before_col or not after_col:
            continue
        numeric = is_numeric_like(before.df[before_col], after.df[after_col], canonical)
        before_agg = aggregate_by_sku(before.df, before_key, before_col, numeric)
        after_agg = aggregate_by_sku(after.df, after_key, after_col, numeric)
        for sku in sorted(common_skus):
            before_value = before_agg.get(sku, 0 if numeric else "")
            after_value = after_agg.get(sku, 0 if numeric else "")
            if values_equal(before_value, after_value, numeric):
                continue
            diff = float(after_value or 0) - float(before_value or 0) if numeric else ""
            sku_rows.append(
                {
                    "sheet": before.name,
                    "key_column": key_name,
                    "sku": sku,
                    "column": canonical,
                    "before_value": before_value,
                    "after_value": after_value,
                    "diff": diff,
                    "status": "FAIL",
                }
            )
    return numeric_rows, sku_rows


def compare_workbooks(before_path: Path, after_path: Path) -> dict[str, pd.DataFrame]:
    before_sheets = workbook_sheet_names(before_path)
    after_sheets = workbook_sheet_names(after_path)
    common_sheets = [sheet for sheet in before_sheets if sheet in after_sheets]

    sheet_rows: list[dict[str, Any]] = []
    column_rows: list[dict[str, Any]] = []
    numeric_rows: list[dict[str, Any]] = []
    sku_rows: list[dict[str, Any]] = []

    for sheet in before_sheets:
        if sheet not in after_sheets:
            sheet_rows.append({"sheet": sheet, "status": "FAIL", "reason": "missing in after"})
    for sheet in after_sheets:
        if sheet not in before_sheets:
            sheet_rows.append({"sheet": sheet, "status": "FAIL", "reason": "added in after"})

    for sheet in common_sheets:
        before = read_sheet(before_path, sheet)
        after = read_sheet(after_path, sheet)
        row_status = "PASS" if len(before.df) == len(after.df) else "FAIL"
        col_status = "PASS" if len(before.columns) == len(after.columns) else "FAIL"
        columns_status = "PASS" if before.columns == after.columns else "FAIL"
        sheet_rows.append(
            {
                "sheet": sheet,
                "status": "PASS" if row_status == col_status == columns_status == "PASS" else "FAIL",
                "reason": "",
                "before_header_row": before.header_row,
                "after_header_row": after.header_row,
                "before_used_rows": before.used_rows,
                "after_used_rows": after.used_rows,
                "before_row_count": len(before.df),
                "after_row_count": len(after.df),
                "row_count_status": row_status,
                "before_column_count": len(before.columns),
                "after_column_count": len(after.columns),
                "column_count_status": col_status,
                "column_list_status": columns_status,
            }
        )
        if before.columns != after.columns:
            column_rows.append(
                {
                    "sheet": sheet,
                    "status": "FAIL",
                    "before_columns": "\n".join(before.columns),
                    "after_columns": "\n".join(after.columns),
                    "removed_columns": "\n".join([col for col in before.columns if col not in after.columns]),
                    "added_columns": "\n".join([col for col in after.columns if col not in before.columns]),
                }
            )
        else:
            column_rows.append(
                {
                    "sheet": sheet,
                    "status": "PASS",
                    "before_columns": "\n".join(before.columns),
                    "after_columns": "\n".join(after.columns),
                    "removed_columns": "",
                    "added_columns": "",
                }
            )
        sheet_numeric_rows, sheet_sku_rows = compare_sheet_data(before, after)
        numeric_rows.extend(sheet_numeric_rows)
        sku_rows.extend(sheet_sku_rows)

    summary_rows = [
        {"item": "before_file", "value": str(before_path)},
        {"item": "after_file", "value": str(after_path)},
        {"item": "before_sheet_count", "value": len(before_sheets)},
        {"item": "after_sheet_count", "value": len(after_sheets)},
        {"item": "missing_sheet_count", "value": len(set(before_sheets) - set(after_sheets))},
        {"item": "added_sheet_count", "value": len(set(after_sheets) - set(before_sheets))},
        {"item": "sheet_fail_count", "value": sum(row.get("status") == "FAIL" for row in sheet_rows)},
        {"item": "column_fail_count", "value": sum(row.get("status") == "FAIL" for row in column_rows)},
        {"item": "numeric_sum_fail_count", "value": sum(row.get("status") == "FAIL" for row in numeric_rows)},
        {"item": "sku_value_diff_count", "value": len(sku_rows)},
    ]
    failed = any(int(row["value"]) > 0 for row in summary_rows[4:])
    summary_rows.insert(0, {"item": "overall_status", "value": "FAIL" if failed else "PASS"})

    return {
        "Summary": pd.DataFrame(summary_rows),
        "Sheet Compare": pd.DataFrame(sheet_rows),
        "Column Compare": pd.DataFrame(column_rows),
        "Numeric Sums": pd.DataFrame(numeric_rows),
        "SKU Diffs": pd.DataFrame(sku_rows),
    }


def write_report(report: dict[str, pd.DataFrame], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in report.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)


def print_console_summary(report: dict[str, pd.DataFrame], output_path: Path | None) -> None:
    summary = report["Summary"]
    status = summary.loc[summary["item"].eq("overall_status"), "value"].iloc[0]
    print(f"OVERALL: {status}")
    for _, row in summary.iloc[1:].iterrows():
        print(f"- {row['item']}: {row['value']}")

    for title, df in [
        ("Sheet/row/column differences", report["Sheet Compare"]),
        ("Column list differences", report["Column Compare"]),
        ("Numeric sum differences", report["Numeric Sums"]),
        ("SKU value differences", report["SKU Diffs"]),
    ]:
        if df.empty:
            print(f"\n{title}: PASS")
            continue
        failed_df = df[df.get("status", "FAIL").eq("FAIL")] if "status" in df.columns else df
        if failed_df.empty:
            print(f"\n{title}: PASS")
            continue
        print(f"\n{title}: {len(failed_df)} issue(s)")
        print(failed_df.head(MAX_CONSOLE_ROWS).to_string(index=False))
        if len(failed_df) > MAX_CONSOLE_ROWS:
            print(f"... {len(failed_df) - MAX_CONSOLE_ROWS} more row(s) in the Excel report")

    if output_path is not None:
        print(f"\nReport saved: {output_path}")


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare pre/post refactor SCM Excel outputs without changing app logic."
    )
    parser.add_argument("before", type=Path, help="Excel file generated before refactoring")
    parser.add_argument("after", type=Path, help="Excel file generated after refactoring")
    parser.add_argument(
        "--report",
        type=Path,
        default=Path("outputs") / "compare_report.xlsx",
        help="Report path. Default: outputs/compare_report.xlsx",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    before_path = args.before.resolve()
    after_path = args.after.resolve()
    if not before_path.exists():
        print(f"ERROR: before file does not exist: {before_path}", file=sys.stderr)
        return 2
    if not after_path.exists():
        print(f"ERROR: after file does not exist: {after_path}", file=sys.stderr)
        return 2

    report = compare_workbooks(before_path, after_path)
    report_path: Path | None = args.report.resolve() if args.report else None
    if report_path is not None:
        try:
            write_report(report, report_path)
        except Exception as exc:  # noqa: BLE001
            print(f"WARNING: could not save report: {exc}", file=sys.stderr)
            report_path = None
    print_console_summary(report, report_path)
    status = report["Summary"].loc[report["Summary"]["item"].eq("overall_status"), "value"].iloc[0]
    return 0 if status == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
