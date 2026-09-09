from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
import json
import math
import os
import re

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple


DEFAULT_OLLAMA_BASE_URL = "http://localhost:11434"
DEFAULT_OLLAMA_MODEL = "gemma4"
OLLAMA_TIMEOUT_SECONDS = 90


@dataclass(frozen=True)
class AiContext:
    prompt_context: str
    summary: dict[str, Any]
    file_mapping: dict[str, str]
    settings: dict[str, Any]
    context_summary: dict[str, int]


def clean_column_name(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def jsonable_value(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        if pd.isna(value):
            return None
        return value.isoformat()
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if pd.isna(value):
        return None
    return value


def read_excel_sheet(path: Path, sheet_name: str, header: int) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet_name, header=header)
    df = df.dropna(how="all").copy()
    df.columns = [str(column).strip() for column in df.columns]
    return df


def resolve_sheet_name(path: Path, preferred_names: list[str], fallback_index: int) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for name in preferred_names:
            if name in wb.sheetnames:
                return name
        if fallback_index < len(wb.sheetnames):
            return wb.sheetnames[fallback_index]
        raise ValueError(f"필수 시트를 찾을 수 없습니다. 확인한 시트명: {', '.join(preferred_names)}")
    finally:
        wb.close()


def read_cell(path: Path, sheet_name: str, coordinate: str) -> object:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        row_idx, col_idx = coordinate_to_tuple(coordinate)
        ws = wb[sheet_name]
        for row in ws.iter_rows(
            min_row=row_idx,
            max_row=row_idx,
            min_col=col_idx,
            max_col=col_idx,
            values_only=True,
        ):
            return row[0] if row else None
        return None
    finally:
        wb.close()


def read_adjacent_value(path: Path, sheet_name: str, label: str, max_rows: int = 20) -> object:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), values_only=True):
            for index, value in enumerate(row[:-1]):
                if str(value or "").strip() == label:
                    return row[index + 1]
        return None
    finally:
        wb.close()


def find_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    normalized_candidates = {clean_column_name(candidate) for candidate in candidates}
    for column in df.columns:
        if clean_column_name(column) in normalized_candidates:
            return str(column)
    return None


def numeric_sum(df: pd.DataFrame, candidates: list[str]) -> float | None:
    column = find_column(df, candidates)
    if not column:
        return None
    values = pd.to_numeric(df[column], errors="coerce").fillna(0)
    return float(values.sum())


def selected_columns(df: pd.DataFrame, candidates: list[str]) -> list[str]:
    selected: list[str] = []
    normalized_seen: set[str] = set()
    for candidate in candidates:
        column = find_column(df, [candidate])
        if column and column not in selected:
            selected.append(column)
            normalized_seen.add(clean_column_name(column))
    for column in df.columns:
        normalized = clean_column_name(column)
        if normalized not in normalized_seen and len(selected) < 12:
            selected.append(str(column))
            normalized_seen.add(normalized)
    return selected


def dataframe_to_ai_context(
    df: pd.DataFrame,
    *,
    title: str,
    max_rows: int = 50,
    columns: list[str] | None = None,
) -> str:
    if df.empty:
        return f"## {title}\n데이터 없음\n"

    view = df.copy()
    if columns:
        existing = [column for column in columns if column in view.columns]
        if existing:
            view = view[existing]

    view = view.head(max_rows).copy()
    for column in view.columns:
        view[column] = view[column].map(jsonable_value)

    lines = [f"## {title}", f"- 전체 행 수: {len(df)}", f"- 아래 표는 상위 {len(view)}행만 포함합니다."]
    lines.append("| " + " | ".join(str(column) for column in view.columns) + " |")
    lines.append("| " + " | ".join("---" for _ in view.columns) + " |")
    for row in view.to_dict(orient="records"):
        values = [str(row.get(column, "") if row.get(column, "") is not None else "") for column in view.columns]
        lines.append("| " + " | ".join(value.replace("\n", " ")[:180] for value in values) + " |")
    return "\n".join(lines) + "\n"


def numeric_summary_text(df: pd.DataFrame, title: str, candidates_by_label: dict[str, list[str]]) -> str:
    lines = [f"## {title} 주요 숫자 합계"]
    found = False
    for label, candidates in candidates_by_label.items():
        value = numeric_sum(df, candidates)
        if value is None:
            continue
        found = True
        lines.append(f"- {label}: {value:,.2f}")
    if not found:
        lines.append("- 계산 가능한 주요 숫자 합계 없음")
    return "\n".join(lines) + "\n"


def read_file_mapping(path: Path, sheet_name: str) -> dict[str, str]:
    try:
        df = read_excel_sheet(path, sheet_name, header=0)
    except Exception:
        return {}

    key_col = find_column(df, ["파일키"])
    filename_col = find_column(df, ["파일명"])
    if not key_col or not filename_col:
        return {}

    mapping: dict[str, str] = {}
    for _, row in df.iterrows():
        key = str(row.get(key_col, "")).strip()
        filename = str(row.get(filename_col, "")).strip()
        if key and filename and filename != "미업로드":
            mapping[key] = filename
    return mapping


def build_context_from_output_excel(output_path: Path) -> AiContext:
    report_source_sheet = resolve_sheet_name(output_path, ["_보고서원본"], 7)
    check_required_sheet = resolve_sheet_name(output_path, ["확인필요"], 3)
    stock_eta_sheet = resolve_sheet_name(output_path, ["재고 ETA"], 1)
    arrival_sheet = resolve_sheet_name(output_path, ["입고 예정 목록"], 2)
    input_files_sheet = resolve_sheet_name(output_path, ["_입력파일현황"], 5)

    order_df = read_excel_sheet(output_path, report_source_sheet, header=0)
    check_required_df = read_excel_sheet(output_path, check_required_sheet, header=8)
    stock_eta_df = read_excel_sheet(output_path, stock_eta_sheet, header=7)
    arrival_df = read_excel_sheet(output_path, arrival_sheet, header=19)

    file_mapping = read_file_mapping(output_path, input_files_sheet)
    eur_krw_rate = read_cell(output_path, arrival_sheet, "J1")
    stock_eta_review_required = read_adjacent_value(output_path, stock_eta_sheet, "검토필요 SKU")
    stock_eta_order_required = read_adjacent_value(output_path, stock_eta_sheet, "발주필요 SKU")

    category_col = find_column(order_df, ["구분"])
    order_required_sku = int(order_df[category_col].astype(str).eq("필요").sum()) if category_col else 0
    order_amount_eur = numeric_sum(order_df, ["미입고 미포함\n발주금액(EUR)", "미입고 미포함 발주금액(EUR)"])
    order_qty = numeric_sum(order_df, ["미입고 미포함\n발주 필요수량", "미입고 미포함 발주 필요수량"])
    stock_eta_qty = numeric_sum(stock_eta_df, ["발주필요수량", "최종 발주 필요수량"])
    arrival_qty = numeric_sum(arrival_df, ["수량"])
    arrival_amount_krw = numeric_sum(arrival_df, ["도착예정금액(KRW)"])

    summary = {
        "total_sku": int(len(order_df)),
        "review_required_sku": int(stock_eta_review_required) if isinstance(stock_eta_review_required, (int, float)) else None,
        "order_required_sku": int(stock_eta_order_required) if isinstance(stock_eta_order_required, (int, float)) else order_required_sku,
        "check_required_rows": int(len(check_required_df)),
        "stock_eta_rows": int(len(stock_eta_df)),
        "order_amount_eur": order_amount_eur,
        "order_qty": order_qty,
        "stock_eta_required_qty": stock_eta_qty,
        "arrival_qty": arrival_qty,
        "arrival_amount_krw": arrival_amount_krw,
    }
    settings = {
        "eur_krw_rate": float(eur_krw_rate) if isinstance(eur_krw_rate, (int, float)) else eur_krw_rate,
        "rate_source": "output_excel",
    }
    context_summary = {
        "order_review_rows": int(min(len(order_df), 50)),
        "check_required_rows": int(min(len(check_required_df), 50)),
        "stock_eta_rows": int(min(len(stock_eta_df), 30)),
        "arrival_rows": int(min(len(arrival_df), 30)),
    }

    order_columns = selected_columns(
        order_df,
        [
            "구분",
            "SKU",
            "상품명",
            "브랜드",
            "월평균",
            "안전재고",
            "1. 유럽\n재고",
            "2. 운송\n재고",
            "미입고 현황",
            "미입고 미포함\n발주 필요수량",
            "미입고 미포함\n발주금액(EUR)",
            "상태",
        ],
    )
    check_columns = selected_columns(
        check_required_df,
        ["확인 구분", "상품코드", "상품명", "브랜드", "수량", "예상 입고일", "운송수단", "확인 내용"],
    )
    stock_eta_columns = selected_columns(
        stock_eta_df,
        ["No", "SKU", "브랜드", "3개월 판매량", "월평균 판매량", "재고 ETA 상태", "발주필요수량"],
    )
    arrival_columns = selected_columns(
        arrival_df,
        ["도착일", "운송수단", "브랜드", "SKU", "상품명", "수량", "도착예정금액(KRW)", "출고일", "운송 L/T"],
    )

    context_parts = [
        "# 분석 결과 요약",
        json.dumps({"summary": summary, "settings": settings, "file_mapping": file_mapping}, ensure_ascii=False, indent=2),
        numeric_summary_text(
            order_df,
            "발주 검토",
            {
                "미입고 미포함 발주 필요수량": ["미입고 미포함\n발주 필요수량", "미입고 미포함 발주 필요수량"],
                "미입고 미포함 발주금액(EUR)": ["미입고 미포함\n발주금액(EUR)", "미입고 미포함 발주금액(EUR)"],
                "미입고 포함 발주 필요수량": ["미입고 포함\n발주 필요수량", "미입고 포함 발주 필요수량"],
                "미입고 포함 발주금액(EUR)": ["미입고 포함\n발주금액(EUR)", "미입고 포함 발주금액(EUR)"],
            },
        ),
        numeric_summary_text(
            arrival_df,
            "입고 예정",
            {
                "도착 예정 수량": ["수량"],
                "도착예정금액(KRW)": ["도착예정금액(KRW)"],
            },
        ),
        dataframe_to_ai_context(order_df, title="발주 검토 상위 행", max_rows=50, columns=order_columns),
        dataframe_to_ai_context(check_required_df, title="확인필요 상위 행", max_rows=50, columns=check_columns),
        dataframe_to_ai_context(stock_eta_df, title="재고 ETA 상위 행", max_rows=30, columns=stock_eta_columns),
        dataframe_to_ai_context(arrival_df, title="입고 예정 목록 상위 행", max_rows=30, columns=arrival_columns),
    ]

    return AiContext(
        prompt_context="\n".join(context_parts),
        summary=summary,
        file_mapping=file_mapping,
        settings=settings,
        context_summary=context_summary,
    )


def build_ai_prompt(question: str, context: AiContext) -> str:
    return f"""너는 SCM/물류/재고 분석을 돕는 AI 어시스턴트다.

반드시 지킬 원칙:
- 제공된 데이터 범위 안에서만 답변한다.
- 데이터에 없는 내용은 추측하지 않는다.
- 확인 불가한 내용은 "현재 데이터만으로는 확인이 어렵습니다"라고 답한다.
- 답변은 팀장님 보고에 쓸 수 있도록 간결하게 작성한다.
- 재고, 발주, 운송중, 미입고, ETA, 품절 위험 관점으로 해석한다.
- 숫자가 있으면 요약해서 설명한다.
- 원본 엑셀 전체가 아니라 제한된 분석 요약과 핵심 행만 제공되었다는 점을 고려한다.

[분석 context]
{context.prompt_context}

[사용자 질문]
{question}

[답변]
"""


def ollama_settings() -> tuple[str, str]:
    base_url = os.environ.get("OLLAMA_BASE_URL", DEFAULT_OLLAMA_BASE_URL).rstrip("/")
    model = os.environ.get("OLLAMA_MODEL", DEFAULT_OLLAMA_MODEL).strip() or DEFAULT_OLLAMA_MODEL
    return base_url, model


def call_ollama(prompt: str) -> tuple[str, str]:
    base_url, model = ollama_settings()
    url = f"{base_url}/api/generate"
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
    }

    try:
        response = requests.post(url, json=payload, timeout=OLLAMA_TIMEOUT_SECONDS)
        response.raise_for_status()
    except requests.exceptions.ConnectionError as exc:
        raise RuntimeError(f"Ollama가 실행 중인지 확인해주세요. 예: ollama run {model}") from exc
    except requests.exceptions.Timeout as exc:
        raise RuntimeError("Ollama 응답 시간이 초과되었습니다. 모델 상태를 확인해주세요.") from exc
    except requests.exceptions.RequestException as exc:
        raise RuntimeError(f"Ollama 호출 실패: {exc}") from exc

    data = response.json()
    answer = str(data.get("response") or "").strip()
    if not answer:
        raise RuntimeError("Ollama 응답이 비어 있습니다.")
    return answer, model
