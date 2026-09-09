"""데이터셋 XLSX 입출력.

report_template_export 분해(IMPROVEMENT_PLAN.md 2단계). 마스킹된 블록을 검증용/교환용
XLSX 데이터셋으로 쓰고 다시 읽는다. openpyxl은 함수 내부에서 지연 import한다."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from backend.services.reports.formatting import (
    excel_safe_html_snapshot,
    parse_json_object,
    safe_sheet_name,
    stringify_report_value,
)
from backend.services.reports.models import (
    AUDIENCE_LABELS,
    MaskingSummary,
    ReportBlockPayload,
    SECTION_ORDER,
)
from backend.services.reports.report_chrome import audience_badge

def write_report_dataset_xlsx(
    output_path: Path,
    *,
    audience: str,
    blocks: list[ReportBlockPayload],
    title: str | None,
    masking_summary: MaskingSummary | None = None,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is required to generate report dataset XLSX files. Run pip install -r requirements.txt.",
        ) from exc

    workbook = Workbook()
    meta_sheet = workbook.active
    meta_sheet.title = "report_meta"
    meta_sheet.append(["key", "value"])
    meta_sheet.append(["title", title or "Report Cart"])
    meta_sheet.append(["audience", audience])
    meta_sheet.append(["audience_label", AUDIENCE_LABELS.get(audience, audience)])
    meta_sheet.append(["generated_at", datetime.now().isoformat(timespec="seconds")])
    meta_sheet.append(["block_count", len(blocks)])
    if masking_summary:
        meta_sheet.append(["audience_badge", audience_badge(audience)])
        meta_sheet.append(["anonymized_competitor_count", masking_summary.anonymized_competitor_count])
        meta_sheet.append(["removed_columns", ", ".join(masking_summary.removed_columns)])

    blocks_sheet = workbook.create_sheet("report_blocks")
    blocks_sheet.append(["order", "block_id", "block_type", "block_kind", "section", "size", "title", "subtitle", "meta", "params_json", "snapshot_json", "html_snapshot"])
    for index, block in enumerate(blocks, start=1):
        blocks_sheet.append(
            [
                index,
                block.id,
                block.type or "summary",
                block.kind or "ranking",
                block.section or "summary",
                block.size or "full",
                block.title,
                block.subtitle,
                block.meta,
                json.dumps(block.params or {}, ensure_ascii=False),
                json.dumps(block.snapshot or {}, ensure_ascii=False),
                excel_safe_html_snapshot(block.html_snapshot),
            ]
        )

    params_sheet = workbook.create_sheet("block_params")
    params_sheet.append(["order", "block_id", "key", "value"])
    for index, block in enumerate(blocks, start=1):
        for key, value in (block.params or {}).items():
            params_sheet.append([index, block.id, key, stringify_report_value(value)])

    snapshots_sheet = workbook.create_sheet("block_snapshots")
    snapshots_sheet.append(["order", "block_id", "snapshot_json"])
    for index, block in enumerate(blocks, start=1):
        if block.snapshot:
            snapshots_sheet.append([index, block.id, json.dumps(block.snapshot, ensure_ascii=False)])

    if masking_summary:
        masking_sheet = workbook.create_sheet("masking_summary")
        masking_sheet.append(["key", "value"])
        masking_sheet.append(["audience", audience])
        masking_sheet.append(["badge", audience_badge(audience)])
        masking_sheet.append(["removed_columns", ", ".join(masking_summary.removed_columns) or "-"])
        masking_sheet.append(["anonymized_competitor_count", masking_summary.anonymized_competitor_count])
        masking_sheet.append(["client_identity_policy", masking_summary.client_identity_policy])
        for original, alias in masking_summary.competitor_aliases.items():
            masking_sheet.append([f"alias:{alias}", original if audience == "internal" else alias])

    for section in SECTION_ORDER:
        section_blocks = [block for block in blocks if (block.section or "summary") == section]
        if not section_blocks:
            continue
        sheet = workbook.create_sheet(safe_sheet_name(section))
        param_keys = sorted({key for block in section_blocks for key in (block.params or {}).keys()})
        sheet.append(["order", "block_id", "block_kind", "title", "subtitle", "meta", *[f"param_{key}" for key in param_keys]])
        for index, block in enumerate(blocks, start=1):
            if (block.section or "summary") != section:
                continue
            params = block.params or {}
            sheet.append(
                [
                    index,
                    block.id,
                    block.kind or "ranking",
                    block.title,
                    block.subtitle,
                    block.meta,
                    *[stringify_report_value(params.get(key, "")) for key in param_keys],
                ]
            )

    style_workbook_headers(workbook, header_color="E90035", Font=Font, PatternFill=PatternFill, get_column_letter=get_column_letter)
    workbook.save(output_path)


def read_report_dataset_xlsx(dataset_path: Path) -> list[ReportBlockPayload]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is required to read report dataset XLSX files. Run pip install -r requirements.txt.",
        ) from exc

    workbook = load_workbook(dataset_path, data_only=True, read_only=True)
    if "report_blocks" not in workbook.sheetnames:
        raise HTTPException(status_code=500, detail="report_dataset.xlsx is missing the report_blocks sheet.")

    sheet = workbook["report_blocks"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(value or "") for value in rows[0]]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        record = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
        records.append(record)

    records.sort(key=lambda item: int(item.get("order") or 0))
    return [
        ReportBlockPayload(
            id=str(record.get("block_id") or ""),
            type=str(record.get("block_type") or "summary"),
            kind=str(record.get("block_kind") or ""),
            section=str(record.get("section") or ""),
            size=str(record.get("size") or "full"),
            title=str(record.get("title") or ""),
            subtitle=str(record.get("subtitle") or ""),
            meta=str(record.get("meta") or ""),
            params=parse_json_object(record.get("params_json")),
            snapshot=parse_json_object(record.get("snapshot_json")) or None,
            html_snapshot=str(record.get("html_snapshot") or "") or None,
        )
        for record in records
    ]


def style_workbook_headers(workbook, *, header_color: str, Font, PatternFill, get_column_letter) -> None:
    header_fill = PatternFill("solid", fgColor=header_color)
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            width = min(42, max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width




__all__ = [
    "write_report_dataset_xlsx",
    "read_report_dataset_xlsx",
    "style_workbook_headers",
]
