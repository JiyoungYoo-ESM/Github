"""Region and brand sections for the branded report PPTX template."""

from __future__ import annotations

import re

from backend.services.reports.models import *  # noqa: F401,F403
from backend.services.reports.formatting import *  # noqa: F401,F403
from backend.services.reports.report_data import *  # noqa: F401,F403
from backend.services.reports.compute_cross import *  # noqa: F401,F403
from backend.services.reports.compute_snapshot import *  # noqa: F401,F403
from backend.services.reports.compute_series import *  # noqa: F401,F403
from backend.services.reports.report_chrome import *  # noqa: F401,F403
from backend.services.reports.pptx_prims import *  # noqa: F401,F403
from backend.services.reports.pptx_template_common import *  # noqa: F401,F403


_MAX_PPTX_RANKING_TABLE_ROWS = 60

def region_page_descriptors(blocks: list[ReportBlockPayload]) -> tuple[list[dict], list[dict[str, object]]]:
    """region 섹션 블록을 국가 순위(ranking)와 권역 비중(share)으로 분리해 페이지 기술자를 만든다.

    두 종류를 한 순위표에 합치면 권역 합계(예: 유럽)가 개별 국가와 섞여 이중집계되므로,
    국가 순위와 권역 비중을 각각 별도 표/슬라이드로 렌더링한다.
    insight 패널도 해당 페이지의 표 데이터 기준으로 다시 채운다.
    """
    share_blocks = [block for block in blocks if str(block.type or "").lower() == "region_share" or (block.kind or "") == "share"]
    country_blocks = [block for block in blocks if block not in share_blocks]
    country_rows = region_section_rows(country_blocks)
    share_rows = [
        row
        for row in section_rows(share_blocks)
        if first_value(row, ("권역", "region"), "") not in {None, "", "-"}
    ]

    descriptors: list[dict] = []
    for title, rows in (("01 국가별 판매 순위 및 성장", country_rows), ("권역별 매출 비중", share_rows)):
        if not rows:
            continue
        chunks = chunk_rows(rows, REGION_ROWS_PER_SLIDE)
        for page_index, chunk in enumerate(chunks):
            suffix = f" ({page_index + 1}/{len(chunks)})" if len(chunks) > 1 else ""
            descriptors.append(
                {
                    "rows": chunk,
                    "insight_rows": chunk,
                    "title": title + suffix,
                    "rank_start": page_index * REGION_ROWS_PER_SLIDE + 1,
                }
            )
    if not descriptors:
        descriptors = [{"rows": [], "insight_rows": [], "title": "01 국가·권역별 판매 순위 및 성장", "rank_start": 1}]
    insight_rows = list(descriptors[0].get("insight_rows") or [])
    return descriptors, insight_rows

def region_section_rows(blocks: list[ReportBlockPayload]) -> list[dict[str, object]]:
    """국가 순위표에 실제 국가 단위 행만 모으고 성장 카드는 동일 국가에 결합한다.

    국가 상세·상위 SKU·시즌성 카드의 하위 행을 순위표에 그대로 합치면 국가명이
    없는 브랜드/SKU/월 행이 ``-``로 표시된다. 상세 카드는 국가 요약 한 행만
    사용하고, SKU·월 행은 국가 순위 데이터로 취급하지 않는다.
    """
    ranked: dict[str, dict[str, object]] = {}
    growth: dict[str, dict[str, object]] = {}

    def country_label(row: dict[str, object]) -> str:
        return stringify_report_value(first_value(row, ("국가", "country"), "")).strip()

    def upsert(row: dict[str, object]) -> None:
        label = country_label(row)
        if not label or label == "-":
            return
        key = label.casefold()
        candidate = dict(row)
        current = ranked.get(key)
        if current is None:
            ranked[key] = candidate
            return
        if row_sales_value(candidate) > row_sales_value(current):
            merged, secondary = candidate, current
        else:
            merged, secondary = current, candidate
        for field, value in secondary.items():
            existing = merged.get(field)
            if (existing is None or existing == "" or existing == "-") and not (value is None or value == "" or value == "-"):
                merged[field] = value
        ranked[key] = merged

    for block in blocks:
        rows = snapshot_preview_rows(block)
        block_type = str(block.type or "").strip().lower()
        params = block.params or {}

        if block_type == "country_growth":
            basis = str(params.get("comparison_basis") or "").lower()
            growth_key = "YoY" if basis == "yoy" else "MoM" if basis == "mom" else "성장률"
            for row in rows:
                label = country_label(row)
                if not label:
                    continue
                growth[label.casefold()] = {
                    **row,
                    growth_key: first_value(row, ("성장률", "YoY", "MoM", "성장률MoM"), ""),
                    "_growth_basis": basis,
                    "_growth_target_month": params.get("target_month") or "",
                    "_growth_comparison_month": params.get("comparison_month") or "",
                    "_growth_period_label": params.get("period_label") or "",
                }
            continue

        if block_type == "country_rank":
            for row in rows:
                candidate = dict(row)
                candidate["_yoy_target_month"] = params.get("yoy_target_month") or ""
                candidate["_yoy_comparison_month"] = params.get("yoy_comparison_month") or ""
                candidate["_mom_target_month"] = params.get("mom_target_month") or ""
                candidate["_mom_comparison_month"] = params.get("mom_comparison_month") or ""
                upsert(candidate)
            continue

        if block_type == "country_detail":
            label = stringify_report_value(params.get("country") or "").strip()
            summary = next((row for row in rows if str(row.get("구분") or "") == "국가 요약"), {})
            if label:
                detail_row = dict(summary)
                detail_row["국가"] = label
                detail_row["권역"] = params.get("region") or detail_row.get("권역") or ""
                if row_sales_value(detail_row) <= 0 and params.get("amount") not in {None, "", "-"}:
                    detail_row["매출액"] = params.get("amount")
                if first_value(detail_row, ("점유율",), "") in {None, "", "-"} and params.get("share") not in {None, "", "-"}:
                    detail_row["점유율"] = format_percent_value(float(params["share"]))
                upsert(detail_row)
            continue

        if block_type == "country_sku_season":
            continue

        for row in rows:
            if country_label(row) and (row_sales_value(row) > 0 or block_type in {"", "summary"}):
                upsert(row)

    for key, growth_row in growth.items():
        current = ranked.get(key)
        if current is None:
            continue
        for field in (
            "YoY",
            "MoM",
            "성장률MoM",
            "성장률",
            "_growth_basis",
            "_growth_target_month",
            "_growth_comparison_month",
            "_growth_period_label",
        ):
            value = growth_row.get(field)
            if value not in {None, "", "-"}:
                current[field] = value

    if not ranked:
        for row in growth.values():
            upsert(row)

    return top_rows_by_sales(list(ranked.values()), len(ranked))

def fill_template_region(prs, *, audience: str, blocks: list[ReportBlockPayload]) -> dict | None:
    intro = prs.slides[2]
    data = prs.slides[3]
    descriptors, insight_rows = region_page_descriptors(blocks)
    eur_krw_rate = blocks_exchange_rate(blocks)
    set_shape_text(intro, 117, "국가별 판매 순위")
    set_shape_text(intro, 118, "국가·권역별 매출 순위와 핵심 시장 신호를 정리합니다.")
    first = descriptors[0]
    set_shape_text(data, 127, first["title"])
    set_template_badge(data, 130, audience)
    replace_region_table(data, first["rows"], audience, eur_krw_rate=eur_krw_rate, rank_start=first["rank_start"])
    fill_region_insights(data, insight_rows, eur_krw_rate=eur_krw_rate)
    if len(descriptors) <= 1:
        return None

    def render_page(slide, page_index: int) -> None:
        desc = descriptors[page_index]
        set_shape_text(slide, 127, desc["title"])
        replace_region_table(slide, desc["rows"], audience, eur_krw_rate=eur_krw_rate, rank_start=desc["rank_start"])
        fill_region_insights(slide, list(desc.get("insight_rows") or desc["rows"]), eur_krw_rate=eur_krw_rate)

    return {"base_slide": data, "page_count": len(descriptors), "render_page": render_page}


def replace_region_table(slide, rows: list[dict[str, object]], audience: str, *, eur_krw_rate: float | None = None, rank_start: int = 1) -> None:
    try:
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError:
        return

    clear_named_dynamic_shapes(slide, "Silicon2 Region Dynamic")
    for shape_id in (*range(132, 138), *range(142, 167), 168, 169, 170):
        delete_shape_by_id(slide, shape_id)

    include_krw = audience == "internal"
    growth_key, growth_label = select_growth_metric(rows, REGION_GROWTH_METRICS)
    include_growth = has_growth_metric(rows, growth_key)
    growth_header = growth_metric_header(growth_label)
    headers = ["순위", "국가/권역", "매출 규모", "원화 환산", "시장 점유율"] if include_krw else ["순위", "국가/권역", "매출 규모", "시장 점유율"]
    if include_growth:
        headers.append(growth_header)
    table_rows = build_region_table_rows(rows, include_krw=include_krw, eur_krw_rate=eur_krw_rate, rank_start=rank_start, growth_key=growth_key, include_growth=include_growth)
    if not table_rows:
        table_rows = [["-", "표시할 데이터 없음", "-", "-", "-"]] if include_krw else [["-", "표시할 데이터 없음", "-", "-"]]
        if include_growth:
            table_rows[0].append("-")

    row_count = len(table_rows) + 1
    left = Inches(0.78)
    top = Inches(2.28)
    width = Inches(5.95)
    height = Inches(0.44 + 0.46 * len(table_rows))
    table_shape = slide.shapes.add_table(row_count, len(headers), left, top, width, height)
    table_shape.name = "Silicon2 Region Dynamic Table"
    table = table_shape.table
    if include_krw:
        column_widths = [0.48, 1.20, 1.00, 1.08, 1.05, 1.14] if include_growth else [0.55, 1.55, 1.15, 1.35, 1.35]
    else:
        column_widths = [0.58, 1.52, 1.28, 1.32, 1.25] if include_growth else [0.62, 1.88, 1.55, 1.90]
    for index, column_width in enumerate(column_widths):
        table.columns[index].width = Inches(column_width)
    table.rows[0].height = Inches(0.56 if "\n" in growth_header else 0.44)
    for row_index in range(1, row_count):
        table.rows[row_index].height = Inches(0.46)

    for column_index, header in enumerate(headers):
        apply_region_table_cell(
            table.cell(0, column_index),
            header,
            RGBColor(191, 24, 28),
            RGBColor(255, 255, 255),
            bold=True,
            align=PP_ALIGN.CENTER,
            font_size=7.8 if include_krw else 8.5,
            MSO_ANCHOR=MSO_ANCHOR,
            Pt=Pt,
        )

    for row_index, values in enumerate(table_rows, start=1):
        fill_color = RGBColor(247, 249, 252) if row_index % 2 else RGBColor(255, 255, 255)
        for column_index, value in enumerate(values):
            center_columns = {0, 4} if include_krw else {0, 3}
            if include_growth:
                center_columns.add(len(headers) - 1)
            right_columns = {2, 3} if include_krw else {2}
            align = PP_ALIGN.CENTER if column_index in center_columns else PP_ALIGN.LEFT
            if column_index in right_columns:
                align = PP_ALIGN.RIGHT
            apply_region_table_cell(
                table.cell(row_index, column_index),
                value,
                fill_color,
                RGBColor(38, 49, 68),
                bold=column_index in {0, 1},
                align=align,
                font_size=7.4 if include_krw else 8.2,
                MSO_ANCHOR=MSO_ANCHOR,
                Pt=Pt,
            )


def build_region_table_rows(
    rows: list[dict[str, object]],
    *,
    include_krw: bool = False,
    eur_krw_rate: float | None = None,
    rank_start: int = 1,
    growth_key: str | None = None,
    include_growth: bool | None = None,
) -> list[list[str]]:
    table_rows: list[list[str]] = []
    all_rows = rows
    rows = rows[:_MAX_PPTX_RANKING_TABLE_ROWS]
    growth_key = growth_key or select_growth_metric(rows, REGION_GROWTH_METRICS)[0]
    include_growth = has_growth_metric(rows, growth_key) if include_growth is None else include_growth
    for offset, row in enumerate(rows):
        index = rank_start + offset
        name = first_value(row, ("국가", "권역", "region", "country"), "-")
        values = [
            stringify_report_value(index),
            stringify_report_value(name),
            display_amount(row),
        ]
        if include_krw:
            values.append(display_krw_amount(row, eur_krw_rate=eur_krw_rate) or "-")
        # 점유율 분모는 표시 상한으로 자르기 전의 전체 rows 기준이어야 실제 시장 점유율과 일치한다.
        values.append(region_share_display(row, all_rows))
        if include_growth:
            values.append(growth_metric_value(row, growth_key))
        table_rows.append(values)
    return table_rows

def select_growth_metric(rows: list[dict[str, object]], metrics: tuple[tuple[str, str], ...]) -> tuple[str, str]:
    for key, label in metrics:
        if any(growth_metric_value(row, key) != "-" for row in rows):
            return key, label
    return metrics[0][0], ""


def has_growth_metric(rows: list[dict[str, object]], key: str) -> bool:
    return any(growth_metric_value(row, key) != "-" for row in rows)


def growth_metric_label_for_key(key: str, metrics: tuple[tuple[str, str], ...]) -> str:
    for metric_key, label in metrics:
        if metric_key == key:
            return label
    return ""


def growth_metric_value(row: dict[str, object], key: str) -> str:
    return clean_metric_display(first_value(row, (key,), "-"))


def growth_metric_header(label: str) -> str:
    return f"성장률({label})" if label else "성장률"


def growth_metric_sentence_label(label: str) -> str:
    return f"{label} 성장률" if label else "성장률"


def region_share_display(row: dict[str, object], rows: list[dict[str, object]], *, total_sales: float | None = None) -> str:
    share = clean_metric_display(first_value(row, ("점유율", "시장 점유율", "percent", "share"), ""))
    if share != "-":
        return share
    if total_sales is None:
        total_sales = sum(max(row_sales_value(item), 0) for item in rows)
    sales = row_sales_value(row)
    if total_sales and sales > 0:
        return format_percent(sales / total_sales)
    return "-"


def apply_region_table_cell(
    cell,
    text: object,
    fill_color,
    font_color,
    *,
    bold: bool,
    align,
    font_size: float,
    MSO_ANCHOR,
    Pt,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill_color
    text_frame = cell.text_frame
    text_frame.clear()
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = Pt(5)
    text_frame.margin_right = Pt(5)
    text_frame.margin_top = Pt(2)
    text_frame.margin_bottom = Pt(2)
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    # 긴 SKU 상품명이 잘리지 않도록 넉넉히 자르고 최대 두 줄로 표시한다.
    display_text = clamp_text(stringify_report_value(text), 72)
    run.text = display_text
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(font_size if len(display_text) <= 48 else max(6.2, font_size - 1.4))
    run.font.bold = bold
    run.font.color.rgb = font_color


def fill_region_insights(slide, rows: list[dict[str, object]], *, eur_krw_rate: float | None = None) -> None:
    set_shape_text(slide, 138, "국가·권역 시장 핵심 요약")
    insights = build_region_insights(rows, eur_krw_rate=eur_krw_rate)
    for shape_id, insight in zip((139, 140, 141), [*insights, "", ""]):
        set_shape_text(slide, shape_id, insight)


def format_report_month(value: object) -> str:
    text = stringify_report_value(value).strip()
    match = re.fullmatch(r"(\d{4})[-./](\d{1,2})(?:[-./]\d{1,2})?", text)
    if not match:
        return text
    return f"{match.group(1)}년 {int(match.group(2))}월"


def growth_period_note(row: dict[str, object], growth_key: str, growth_label: str) -> str:
    metric = f"{growth_key} {growth_label}".lower()
    basis = "yoy" if "yoy" in metric else "mom" if "mom" in metric else str(row.get("_growth_basis") or "").lower()
    target = row.get(f"_{basis}_target_month") if basis else ""
    comparison = row.get(f"_{basis}_comparison_month") if basis else ""
    if not target and str(row.get("_growth_basis") or "").lower() == basis:
        target = row.get("_growth_target_month")
        comparison = row.get("_growth_comparison_month")
    target_label = format_report_month(target)
    comparison_label = format_report_month(comparison)
    if target_label and comparison_label:
        return f"기준월 {target_label} vs 비교월 {comparison_label}"
    return ""


def build_region_insights(rows: list[dict[str, object]], *, eur_krw_rate: float | None = None) -> list[str]:
    ranked_rows = top_rows_by_sales(rows, len(rows))
    if not ranked_rows:
        return ["표시할 국가·권역 데이터가 없습니다."]

    rate = eur_krw_rate or infer_eur_krw_rate_from_rows(ranked_rows)
    total_sales = sum(max(row_sales_value(row), 0) for row in ranked_rows)
    growth_key, growth_label = select_growth_metric(ranked_rows, REGION_GROWTH_METRICS)
    top = ranked_rows[0]
    top_label = stringify_report_value(first_value(top, ("국가", "권역", "region", "country"), "상위 시장"))
    top_rank = stringify_report_value(first_value(top, ("순위", "rank"), "")).strip()
    top_amount = amount_with_krw(top, eur_krw_rate=rate)
    top_share = region_share_display(top, ranked_rows, total_sales=total_sales)
    top_growth = growth_metric_value(top, growth_key)
    top_sentence = f"{top_label} 매출 비중:\n{top_amount}"
    if top_share != "-":
        top_sentence += f", 점유율 {top_share}"
    if top_rank and top_rank != "1":
        top_sentence += f"입니다. 전체 순위 {top_rank}위로, 표시된 범위 안에서 매출 규모가 가장 큽니다."
    else:
        top_sentence += "로 웹 화면 기준 상위 시장입니다."
    if top_growth != "-":
        period_note = growth_period_note(top, growth_key, growth_label)
        growth_name = growth_metric_sentence_label(growth_label)
        if period_note:
            top_sentence += f"\n{growth_name}: {top_growth} ({period_note})"
        else:
            top_sentence += f"\n{growth_name}: {top_growth}"

    insights = [top_sentence]
    if len(ranked_rows) >= 2:
        second = ranked_rows[1]
        second_label = stringify_report_value(first_value(second, ("국가", "권역", "region", "country"), "2위 시장"))
        gap = row_sales_value(top) - row_sales_value(second)
        if gap > 0:
            insights.append(f"1·2위 매출 차이:\n{second_label}보다 {format_amount_with_krw_value(gap, eur_krw_rate=rate)} 높아 상위 시장 집중도가 나타납니다.")
        else:
            insights.append(f"상위 시장 격차:\n{top_label}와 {second_label}의 매출 규모가 유사합니다.")
    if len(ranked_rows) >= 3 and total_sales > 0:
        top_three_rows = ranked_rows[:3]
        top_three_total = sum(max(row_sales_value(row), 0) for row in top_three_rows)
        web_share_total = sum(percent_numeric_value(region_share_display(row, ranked_rows)) for row in top_three_rows)
        is_country = any(first_value(row, ("국가", "country"), "") for row in top_three_rows)
        entity = "국가" if is_country else "시장"
        three_entities = "3개국" if is_country else "3개 시장"
        ranks = [int(numeric_value(first_value(row, ("순위", "rank"), 0))) for row in top_three_rows]
        if ranks == [1, 2, 3]:
            summary_label = f"매출 상위 {three_entities} 합계"
        elif all(rank > 0 for rank in ranks):
            summary_label = f"전체 {min(ranks)}~{max(ranks)}위 {three_entities} 합계"
        else:
            summary_label = f"이 페이지 상위 {three_entities} 합계"
        if web_share_total > 0:
            insights.append(f"{summary_label}:\n{format_amount_with_krw_value(top_three_total, eur_krw_rate=rate)}이며, 분석 대상 전체 {entity} 매출의 {format_percent_value(web_share_total)}를 차지합니다.")
        else:
            top_three_share = format_percent(top_three_total / total_sales)
            insights.append(f"{summary_label}:\n{format_amount_with_krw_value(top_three_total, eur_krw_rate=rate)}이며, 이 페이지에 포함된 {entity} 매출의 {top_three_share}입니다.")
    return insights[:3]

def fill_template_brand(prs, *, audience: str, blocks: list[ReportBlockPayload]) -> dict | None:
    intro = prs.slides[4]
    data = prs.slides[5]
    rows = brand_section_rows(blocks)
    eur_krw_rate = blocks_exchange_rate(blocks) or infer_eur_krw_rate_from_rows(rows)
    pages = chunk_rows(rows, BRAND_ROWS_PER_SLIDE)
    page_count = len(pages)
    brand_growth_key, brand_growth_label = select_growth_metric(rows, BRAND_GROWTH_METRICS)
    include_brand_growth = has_growth_metric(rows, brand_growth_key)
    brand_growth_header = brand_growth_metric_header(brand_growth_label, brand_growth_period_label(rows))
    if audience == "internal":
        set_shape_text(intro, 176, "브랜드별 판매 순위")
        set_shape_text(intro, 177, "브랜드별 매출, 원화 환산과 성장률을 함께 비교합니다.")
        base_title = "02 브랜드별 판매 순위 및 원화 환산"
        table_header = "원화 환산"
    else:
        set_shape_text(intro, 176, "브랜드 경쟁 순위")
        set_shape_text(intro, 177, "자사 브랜드는 실명 유지, 경쟁사는 보고서별 별칭으로 익명화해 매출과 성장 지표만 비교합니다.")
        base_title = "02 브랜드별 판매 순위 및 공유지표 비교"
        table_header = "공유지표"
    set_section_data_title(data, 186, base_title, page_index=0, page_count=page_count)
    set_template_badge(data, 189, audience)
    replace_brand_table(
        data,
        pages[0],
        audience,
        table_header,
        eur_krw_rate=eur_krw_rate,
        rank_start=1,
        growth_key=brand_growth_key,
        include_growth=include_brand_growth,
        growth_header=brand_growth_header,
    )
    top = top_rows_by_sales(rows, 1)[0] if rows else {}
    brand = stringify_report_value(first_value(top, ("브랜드명", "브랜드", "brand", "display_brand"), "선택 브랜드"))
    amount = amount_with_krw(top, eur_krw_rate=eur_krw_rate) if audience == "internal" else display_amount(top)
    growth = growth_metric_value(top, brand_growth_key) if top else "-"
    growth_note = f" ({growth_metric_sentence_label(brand_growth_label)} {growth})" if growth != "-" else ""
    if audience == "internal":
        summary = f"{brand}은(는) 선택한 브랜드 비교군에서 가장 큰 매출 규모({amount})입니다. 매출 규모, 원화 환산과 성장 흐름을 함께 확인할 수 있습니다{growth_note}."
    else:
        summary = f"{brand}은(는) 선택한 브랜드 비교군에서 가장 큰 매출 규모({amount})입니다. 외부용 보고서에서는 경쟁사 실명과 원가·마진 지표가 데이터 레이어에서 제거됩니다{growth_note}."
    replace_brand_chart(data, rows, audience, summary, eur_krw_rate=eur_krw_rate)
    def render_page(slide, page_index: int) -> None:
        set_section_data_title(slide, 186, base_title, page_index=page_index, page_count=page_count)
        replace_brand_table(
            slide,
            pages[page_index],
            audience,
            table_header,
            eur_krw_rate=eur_krw_rate,
            rank_start=page_index * BRAND_ROWS_PER_SLIDE + 1,
            growth_key=brand_growth_key,
            include_growth=include_brand_growth,
            growth_header=brand_growth_header,
        )

    return {"base_slide": data, "page_count": page_count, "render_page": render_page}


def replace_brand_table(
    slide,
    rows: list[dict[str, object]],
    audience: str,
    table_header: str,
    *,
    eur_krw_rate: float | None = None,
    rank_start: int = 1,
    growth_key: str | None = None,
    include_growth: bool | None = None,
    growth_header: str | None = None,
) -> None:
    try:
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError:
        return

    clear_named_dynamic_shapes(slide, "Silicon2 Brand Dynamic Table")
    for shape_id in (191, *range(195, 229)):
        delete_shape_by_id(slide, shape_id)
    include_krw = audience == "internal"
    if growth_key is None:
        growth_key, growth_label = select_growth_metric(rows, BRAND_GROWTH_METRICS)
    else:
        growth_label = growth_metric_label_for_key(growth_key, BRAND_GROWTH_METRICS)
    include_growth = has_growth_metric(rows, growth_key) if include_growth is None else include_growth
    growth_header = growth_header or brand_growth_metric_header(growth_label, brand_growth_period_label(rows))
    headers = ["순위", "브랜드", "매출 규모", "원화 환산"] if include_krw else ["순위", "브랜드", "매출 규모", table_header]
    if include_growth:
        headers.append(growth_header)
    table_rows = build_brand_table_rows(rows, audience, eur_krw_rate=eur_krw_rate, rank_start=rank_start, growth_key=growth_key, include_growth=include_growth)
    if not table_rows:
        table_rows = [["-", "표시할 데이터 없음", "-", "-"]]
        if include_growth:
            table_rows[0].append("-")

    row_count = len(table_rows) + 1
    table_shape = slide.shapes.add_table(row_count, len(headers), Inches(0.78), Inches(2.38), Inches(5.95), Inches(0.44 + 0.46 * len(table_rows)))
    table_shape.name = "Silicon2 Brand Dynamic Table"
    table = table_shape.table
    if include_krw:
        column_widths = [0.58, 1.55, 1.28, 1.30, 1.24] if include_growth else [0.62, 1.88, 1.55, 1.90]
    else:
        column_widths = [0.58, 1.55, 1.28, 1.3, 1.24] if include_growth else [0.62, 1.88, 1.55, 1.90]
    for index, column_width in enumerate(column_widths):
        table.columns[index].width = Inches(column_width)
    table.rows[0].height = Inches(0.44)
    for row_index in range(1, row_count):
        table.rows[row_index].height = Inches(0.46)

    for column_index, header in enumerate(headers):
        apply_region_table_cell(
            table.cell(0, column_index),
            header,
            RGBColor(191, 24, 28),
            RGBColor(255, 255, 255),
            bold=True,
            align=PP_ALIGN.CENTER,
            font_size=7.4 if "\n" in str(header) else 8.5,
            MSO_ANCHOR=MSO_ANCHOR,
            Pt=Pt,
        )
    for row_index, values in enumerate(table_rows, start=1):
        fill_color = RGBColor(247, 249, 252) if row_index % 2 else RGBColor(255, 255, 255)
        for column_index, value in enumerate(values):
            center_columns = {0} if include_krw else {0, 3}
            if include_growth:
                center_columns.add(len(headers) - 1)
            right_columns = {2, 3} if include_krw else {2}
            align = PP_ALIGN.CENTER if column_index in center_columns else PP_ALIGN.LEFT
            if column_index in right_columns:
                align = PP_ALIGN.RIGHT
            apply_region_table_cell(
                table.cell(row_index, column_index),
                value,
                fill_color,
                RGBColor(38, 49, 68),
                bold=column_index in {0, 1},
                align=align,
                font_size=8.2,
                MSO_ANCHOR=MSO_ANCHOR,
                Pt=Pt,
            )


def build_brand_table_rows(
    rows: list[dict[str, object]],
    audience: str,
    *,
    eur_krw_rate: float | None = None,
    rank_start: int = 1,
    growth_key: str | None = None,
    include_growth: bool | None = None,
) -> list[list[str]]:
    table_rows: list[list[str]] = []
    rows = rows[:_MAX_PPTX_RANKING_TABLE_ROWS]
    include_krw = audience == "internal"
    growth_key = growth_key or select_growth_metric(rows, BRAND_GROWTH_METRICS)[0]
    include_growth = has_growth_metric(rows, growth_key) if include_growth is None else include_growth
    for offset, row in enumerate(rows):
        index = rank_start + offset
        values = [
            stringify_report_value(index),
            stringify_report_value(first_value(row, ("브랜드명", "브랜드", "brand", "display_brand"), "-")),
            display_amount(row),
        ]
        if include_krw:
            values.append(display_krw_amount(row, eur_krw_rate=eur_krw_rate) or "-")
        if audience == "internal":
            if include_growth:
                values.append(growth_metric_value(row, growth_key))
        else:
            values.append("비공개")
            if include_growth:
                values.append(growth_metric_value(row, growth_key))
        table_rows.append(values)
    return table_rows


def brand_growth_metric_header(label: str, period_label: str = "") -> str:
    base = "매출 MoM" if label == "MoM" else "매출 YoY" if label == "YoY" else "매출 성장률"
    if period_label and period_label != "비교 기간 없음":
        return f"{base}\n{period_label}"
    return base


def brand_growth_period_label(rows: list[dict[str, object]]) -> str:
    for row in rows:
        value = stringify_report_value(first_value(row, ("비교기간", "period_label", "comparison_period"), "")).strip()
        if value and value != "-":
            return value
    return ""


def replace_brand_chart(slide, rows: list[dict[str, object]], audience: str, summary: str, *, eur_krw_rate: float | None = None) -> None:
    try:
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
    except ImportError:
        return

    for shape_id in (185, 192, 193, 194, 200):
        delete_shape_by_id(slide, shape_id)

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.12), Inches(2.28), Inches(4.95), Inches(3.16))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(248, 250, 252)
    panel.line.color.rgb = RGBColor(226, 232, 240)
    panel.line.width = Pt(0.75)
    try:
        panel.adjustments[0] = 0.06
        panel.shadow.inherit = False
    except Exception:
        pass

    title = "브랜드별 매출 및 원화 환산 비교" if audience == "internal" else "브랜드별 매출 및 공유지표 비교"
    add_brand_chart_text(slide, Inches(7.38), Inches(2.60), Inches(4.18), Inches(0.25), title, 9.5, RGBColor(15, 23, 42), bold=True, Pt=Pt, PP_ALIGN=PP_ALIGN)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.38), Inches(2.90), Inches(4.18), Inches(0.01))
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(191, 24, 28)
    divider.line.fill.background()

    chart_rows = top_rows_by_sales(rows, 4)
    if not chart_rows:
        add_brand_chart_text(slide, Inches(7.38), Inches(3.32), Inches(4.18), Inches(0.28), "표시할 브랜드 데이터가 없습니다.", 9, RGBColor(100, 116, 139), Pt=Pt, PP_ALIGN=PP_ALIGN)
    else:
        max_sales = max((row_sales_value(row) for row in chart_rows), default=0) or 1
        for index, row in enumerate(chart_rows):
            brand = clamp_text(stringify_report_value(first_value(row, ("브랜드명", "브랜드", "brand", "display_brand"), f"브랜드 {index + 1}")), 16)
            amount = display_amount(row)
            krw_amount = display_krw_amount(row, eur_krw_rate=eur_krw_rate) if audience == "internal" else ""
            sales = max(row_sales_value(row), 0)
            ratio = max(0.08, min(sales / max_sales, 1.0))
            y = Inches(3.12 + index * 0.40)
            is_self = is_self_brand(brand, row)
            bar_color = RGBColor(191, 24, 28) if index == 0 or is_self else RGBColor(235, 55, 92)
            if index >= 2 and not is_self:
                bar_color = RGBColor(251, 151, 166)

            add_brand_chart_text(slide, Inches(7.38), y + Inches(0.05), Inches(1.04), Inches(0.18), brand, 7.5, RGBColor(38, 49, 68), bold=index == 0 or is_self, Pt=Pt, PP_ALIGN=PP_ALIGN)
            bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.42), y, Inches(3.16), Inches(0.25))
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(238, 242, 247)
            bg.line.fill.background()

            bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.42), y, Inches(3.16 * ratio), Inches(0.25))
            bar.fill.solid()
            bar.fill.fore_color.rgb = bar_color
            bar.line.fill.background()
            value_color = RGBColor(255, 255, 255) if ratio >= 0.72 else RGBColor(38, 49, 68)
            add_brand_chart_text(slide, Inches(10.74), y + Inches(0.045), Inches(0.78), Inches(0.15), amount, 7, value_color, bold=True, align=PP_ALIGN.RIGHT, Pt=Pt, PP_ALIGN=PP_ALIGN)
            if krw_amount and krw_amount != "-":
                add_brand_chart_text(slide, Inches(8.42), y + Inches(0.27), Inches(3.16), Inches(0.11), krw_amount, 5.8, RGBColor(100, 116, 139), bold=True, align=PP_ALIGN.RIGHT, Pt=Pt, PP_ALIGN=PP_ALIGN)

    add_brand_chart_text(slide, Inches(7.38), Inches(4.68), Inches(4.16), Inches(0.54), clamp_text(summary, 130), 8.2, RGBColor(51, 65, 85), bold=True, Pt=Pt, PP_ALIGN=PP_ALIGN)

__all__ = [
    "region_page_descriptors",
    "region_section_rows",
    "fill_template_region",
    "replace_region_table",
    "build_region_table_rows",
    "select_growth_metric",
    "has_growth_metric",
    "growth_metric_label_for_key",
    "growth_metric_value",
    "growth_metric_header",
    "growth_metric_sentence_label",
    "region_share_display",
    "apply_region_table_cell",
    "fill_region_insights",
    "format_report_month",
    "growth_period_note",
    "build_region_insights",
    "fill_template_brand",
    "replace_brand_table",
    "build_brand_table_rows",
    "brand_growth_metric_header",
    "brand_growth_period_label",
    "replace_brand_chart",
]

