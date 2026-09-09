"""Create general-account XLSX copies without monetary source data."""

from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook

from backend.auth.amount_permissions import is_amount_field


_TABLE_HEADER_TOKENS = (
    "sku",
    "product",
    "상품",
    "브랜드",
    "brand",
    "country",
    "국가",
    "quantity",
    "qty",
    "수량",
    "점유율",
    "비중",
    "순위",
    "rank",
)


def _looks_like_table_header(row) -> bool:
    values = [cell.value for cell in row if cell.value not in (None, "")]
    normalized = [str(value).replace(" ", "").lower() for value in values if isinstance(value, str)]
    return (
        len(values) >= 3
        or any(any(token in value for token in _TABLE_HEADER_TOKENS) for value in normalized)
    )


def _amount_columns_and_summary_values(worksheet) -> tuple[set[int], set[tuple[int, int]]]:
    amount_columns: set[int] = set()
    summary_values: set[tuple[int, int]] = set()
    scan_limit = min(worksheet.max_row, 30)
    for row in worksheet.iter_rows(min_row=1, max_row=scan_limit):
        table_header = _looks_like_table_header(row)
        for cell in row:
            if not isinstance(cell.value, str) or not is_amount_field(cell.value):
                continue
            if table_header:
                amount_columns.add(cell.column)
                continue
            # A short amount label in a summary area is normally followed by
            # its value. Clear both even when the value is in another column.
            summary_values.add((cell.row, cell.column))
            if cell.column < worksheet.max_column:
                summary_values.add((cell.row, cell.column + 1))
    return amount_columns, summary_values


def create_amount_safe_workbook(source_path: Path) -> Path:
    workbook = load_workbook(source_path, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            amount_columns, summary_values = _amount_columns_and_summary_values(worksheet)
            for row, column in summary_values:
                worksheet.cell(row=row, column=column).value = None
            for column in sorted(amount_columns, reverse=True):
                worksheet.delete_cols(column)

        temporary = tempfile.NamedTemporaryFile(
            prefix="esm_no_amount_",
            suffix=".xlsx",
            delete=False,
        )
        temporary.close()
        target = Path(temporary.name)
        workbook.save(target)
        return target
    finally:
        workbook.close()


__all__ = ["create_amount_safe_workbook"]
