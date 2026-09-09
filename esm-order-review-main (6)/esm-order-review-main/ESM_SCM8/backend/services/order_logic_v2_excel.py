"""Excel export for the entity-aware order-logic v2 result payload.

The approved order calculation remains value-only and comes from the core
service.  Review-only scheduling fields and the logistics outlook are written
as formulas so users can audit the v4.1 reference logic without duplicating or
changing the approved order quantity formula.
"""

from __future__ import annotations

import json
import os
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Any


ORDER_HEADER_ROW = 7
ORDER_DATA_START_ROW = ORDER_HEADER_ROW + 1
LOGISTICS_HEADER_ROW = 6
LOGISTICS_DATA_START_ROW = LOGISTICS_HEADER_ROW + 1
CHECK_REQUIRED_HEADER_ROW = 5
CHECK_REQUIRED_DATA_START_ROW = CHECK_REQUIRED_HEADER_ROW + 1
CALC_SHEET_NAME = "_계산기준"
SHIPPING_SOURCE_SHEET_NAME = "_운송원천"
SHIPPING_SOURCE_CASE_KEY_COLUMN = 8
TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "order_logic_v2_template.xlsx"
_COMPACT_TEMPLATE_CACHE: dict[tuple[str, int], bytes] = {}
_COMPACT_TEMPLATE_CACHE_LOCK = Lock()

# Visible columns come from the supplied reference workbook. Keep these
# mappings separate from the calculation payload keys so presentation changes
# do not change the approved order logic.
ORDER_TEMPLATE_COLUMNS: dict[str, int] = {
    "sku_code": 1,
    "product_name": 2,
    "brand": 3,
    "data_status": 4,
    "grade": 5,
    "sales_13w_qty": 6,
    "monthly_sales": 7,
    "demand_avg": 8,
    "safety_stock": 9,
    "reorder_point": 10,
    "target_stock": 11,
    "incoming_qty": 12,
    "next_eta": 13,
    "eu_available_qty": 14,
    "transit_qty": 15,
    "local_available_qty": 16,
    "inventory_position": 17,
    "moi": 18,
    "depletion_weeks": 19,
    "suggested_qty": 20,
    "upper_suggested_qty": 21,
    "transport_recommendation": 22,
    "order_signal": 23,
    "order_slack_weeks": 24,
    "early_warning_date": 25,
    "conservative_slack_weeks": 26,
    "expected_order_date": 27,
    "demand_sigma": 28,
    "z_applied": 29,
}

# CMS 미입고현황의 상태 구분을 기존 V2 Excel의 미입고 옆에 표시한다.
# ``inbound_completed_qty``는 조회·검증용으로만 표시한다. 이미 재고에 반영된
# 수량이므로 재고 위치(IP)나 제안수량 산식에는 다시 더하지 않는다.
INBOUND_STATUS_COLUMNS: tuple[tuple[str, str, float], ...] = (
    ("open_qty", "① 미입고 수량", 18),
    ("pnfm_qty", "② PNFM확정 수량", 19),
    ("inbound_progress_qty", "③ 입고진행중 수량", 19),
    ("inbound_completed_qty", "④ 입고완료 수량", 18),
)

LOGISTICS_TEMPLATE_COLUMNS: dict[str, int] = {
    "sku_code": 1,
    "product_name": 2,
    "barcode": 3,
    "sales_13w_qty": 4,
    "monthly_sales": 5,
    "demand_avg": 6,
    "local_available_qty": 7,
    "transit_qty": 8,
    "moi": 9,
    "safety_stock": 10,
    "delayed_qty": 11,
    "week_start": 12,
    "week_end": 24,
    "after_eta_qty": 25,
    "reorder_point": 26,
    "sku_case_key": 27,
}


@dataclass(frozen=True)
class ExportColumn:
    key: str
    label: str
    width: float
    number_format: str | None = None
    alignment: str = "left"


# The template inserts inbound-status columns before the ETA column. Excel
# column dimensions do not move with ``insert_cols`` in openpyxl, so these two
# long headers need an explicit final width after all structural changes.
ORDER_COLUMN_WIDTH_OVERRIDES: dict[str, float] = {
    "next_eta": 24,
    "early_warning_date": 28,
}


# Keep the existing V2 review order. The engine/API retain the conservative
# quantity for compatibility, but the user-facing export shows only the
# incoming-trusted proposal quantity.
BASE_COLUMNS: tuple[ExportColumn, ...] = (
    ExportColumn("sku_code", "상품코드", 18),
    ExportColumn("product_name", "상품명", 34),
    ExportColumn("brand", "브랜드", 16),
    ExportColumn("data_status", "상태", 16, alignment="center"),
    ExportColumn("grade", "등급", 10, alignment="center"),
    ExportColumn("demand_avg", "주평균", 13, "#,##0.0", "right"),
    ExportColumn("cv", "CV\n(σ÷주평균)", 13, "0.00", "right"),
    ExportColumn("safety_stock", "안전재고\n(상하한 적용)", 15, "#,##0", "right"),
    ExportColumn("reorder_point", "발주점\n(R/S 미사용)", 14, "#,##0", "right"),
    ExportColumn("target_stock", "목표재고", 14, "#,##0", "right"),
    ExportColumn("incoming_qty", "미입고", 13, "#,##0", "right"),
    ExportColumn(
        "eu_available_qty",
        "본사 창고\n가용재고",
        16,
        "#,##0",
        "right",
    ),
    ExportColumn("transit_qty", "운송중", 13, "#,##0", "right"),
    ExportColumn(
        "next_eta",
        "다음 입고예정\n(운송중 최소 ETA)",
        ORDER_COLUMN_WIDTH_OVERRIDES["next_eta"],
        "yyyy-mm-dd",
        "center",
    ),
    ExportColumn("local_available_qty", "현지 가용재고", 15, "#,##0", "right"),
    ExportColumn("inventory_position", "보유전체", 14, "#,##0", "right"),
    ExportColumn("suggested_qty", "제안수량\n(미입고 인정)", 15, "#,##0", "right"),
    ExportColumn("depletion_weeks", "고갈주수\n(현지+운송)", 14, "0.0", "right"),
    ExportColumn("order_signal", "신호등", 16, alignment="center"),
    ExportColumn("memo", "메모", 24),
    ExportColumn(
        "order_slack_weeks",
        "발주까지 여유\n(R/S 정의 미결)",
        15,
        "0.0",
        "right",
    ),
    ExportColumn(
        "expected_order_date",
        "예상 발주일\n(R/S 정의 미결)",
        16,
        "yyyy-mm-dd",
        "center",
    ),
    ExportColumn(
        "conservative_slack_weeks",
        "보수적 발주 여유\n(R/S 정의 미결)",
        17,
        "0.0",
        "right",
    ),
    ExportColumn(
        "early_warning_date",
        "조기경보일\n(R/S 정의 미결)",
        ORDER_COLUMN_WIDTH_OVERRIDES["early_warning_date"],
        "yyyy-mm-dd",
        "center",
    ),
)

AMOUNT_COLUMNS: tuple[ExportColumn, ...] = (
    ExportColumn("unit_price_eur", "단가(EUR)", 14, '€#,##0.00', "right"),
    ExportColumn("suggested_amount_eur", "제안금액(EUR)", 16, '€#,##0.00', "right"),
    ExportColumn("unit_price_krw", "단가(KRW)", 16, '₩#,##0.00', "right"),
    ExportColumn("suggested_amount_krw", "제안금액(KRW)", 18, '₩#,##0', "right"),
)

# CMS 원화 단가/금액 열. 본사는 원천 통화 자체가 KRW이므로 동일 값이
# 중복되지 않게 현지통화 열만 KRW 라벨로 바꾸어 표시한다.
KRW_CONVERSION_COLUMN_KEYS = frozenset(
    {"unit_price_krw", "suggested_amount_krw"}
)


def _entity_currency_code(entity_code: object) -> str:
    code = str(entity_code or "PL").strip().upper()
    if code == "HQ":
        return "KRW"
    if code == "USA":
        return "USD"
    return "EUR"


def _is_day_unit_result(result: dict[str, Any]) -> bool:
    """HQ/OPO reports day-grain demand; PL and USA report week-grain demand."""

    return str(result.get("period_unit") or "").lower() == "day"


def _demand_period_label(result: dict[str, Any]) -> str:
    if _is_day_unit_result(result):
        count = result.get("demand_period_count") or 91
        return f"판매량\n({count}일)"
    count = result.get("demand_period_count") or 13
    return f"판매량\n({count}주)"


def _check_required_sales_label(result: dict[str, Any]) -> str:
    """Return the plain demand-period label used by the review sheet."""

    if _is_day_unit_result(result):
        count = result.get("demand_period_count") or 91
        return f"{count}일 판매수량"
    count = result.get("demand_period_count") or 13
    return f"{count}주 판매수량"


def _demand_average_label(result: dict[str, Any]) -> str:
    return "일평균" if _is_day_unit_result(result) else "주평균"


def _logistics_demand_average_label(result: dict[str, Any]) -> str:
    return f"{_demand_average_label(result)}\n판매량"


def _amount_columns(currency_code: str) -> tuple[ExportColumn, ...]:
    code = str(currency_code).strip().upper()
    if code == "KRW":
        # 본사 원천 금액은 이미 원화다. `EUR` 라벨만 치환하면 원화 환산 열이
        # 같은 이름으로 중복되고, 그 열 수식이 존재하지 않는 환율 셀을 참조해
        # 항상 빈 값이 된다. 그래서 환산 열 자체를 내보내지 않는다.
        # 각 열의 소수 정밀도는 원천 그대로 두고 통화기호만 바꾼다. 단가는
        # OPO 수량가중 평균원가라 소수를 버리면 검산이 어긋난다.
        return tuple(
            ExportColumn(
                column.key,
                column.label.replace("EUR", "KRW"),
                column.width,
                column.number_format.replace("€", "₩")
                if column.number_format
                else None,
                column.alignment,
            )
            for column in AMOUNT_COLUMNS
            if column.key not in KRW_CONVERSION_COLUMN_KEYS
        )
    if code == "USD":
        currency, symbol = "USD", "$"
    else:
        currency, symbol = "EUR", "€"
    return tuple(
        ExportColumn(
            column.key,
            column.label.replace("EUR", currency),
            column.width,
            column.number_format.replace("€", symbol) if column.number_format else None,
            column.alignment,
        )
        for column in AMOUNT_COLUMNS
    )


def generate_order_logic_v2_excel(
    result: dict[str, Any],
    include_amounts: bool,
    entity_code: str = "PL",
) -> bytes:
    """Return an XLSX snapshot for an order-logic v2 result payload.

    ``result`` contains a ``rows`` list whose dictionaries follow the v2 row
    contract. Missing optional values are exported as blank cells. The
    supplied workbook template owns the visible layout; amount columns remain
    optional and are appended only when the caller has amount permission.
    """

    if not isinstance(result, dict):
        raise TypeError("result must be a dictionary")
    rows = result.get("rows", [])
    if rows is None:
        rows = []
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise TypeError("result['rows'] must be a list of dictionaries")

    return _generate_order_logic_v2_from_template(
        result=result,
        rows=rows,
        include_amounts=include_amounts,
        entity_code=entity_code,
    )


def _order_logic_v2_template_path() -> Path:
    configured = str(os.environ.get("ORDER_LOGIC_V2_TEMPLATE_PATH") or "").strip()
    path = Path(configured) if configured else TEMPLATE_PATH
    if not path.is_file():
        raise RuntimeError(f"Order-logic v2 Excel template is missing: {path}")
    return path


def _compact_template_bytes(template_path: Path, load_workbook: Any) -> bytes:
    """Cache a small styled template while preserving the reference layout.

    The checked-in workbook reserves thousands of blank styled rows. Loading
    and serializing those rows for every export makes the download scale with
    the template size instead of the selected result. Keep two alternating
    data rows as style sources and let the normal population code extend them
    only when the current result needs more rows.
    """

    stat = template_path.stat()
    cache_key = (str(template_path.resolve()), stat.st_mtime_ns)
    cached = _COMPACT_TEMPLATE_CACHE.get(cache_key)
    if cached is not None:
        return cached

    with _COMPACT_TEMPLATE_CACHE_LOCK:
        cached = _COMPACT_TEMPLATE_CACHE.get(cache_key)
        if cached is not None:
            return cached

        workbook = load_workbook(
            BytesIO(template_path.read_bytes()),
            data_only=False,
        )
        try:
            compact_data_rows = (
                ("발주제안", ORDER_DATA_START_ROW),
                ("물류전망", LOGISTICS_DATA_START_ROW),
                ("확인필요", CHECK_REQUIRED_DATA_START_ROW),
                (SHIPPING_SOURCE_SHEET_NAME, 2),
            )
            for sheet_name, first_data_row in compact_data_rows:
                if sheet_name in workbook.sheetnames:
                    _trim_template_rows(
                        workbook[sheet_name],
                        first_data_row=first_data_row,
                        data_row_count=2,
                    )
            output = BytesIO()
            workbook.save(output)
            cached = output.getvalue()
        finally:
            workbook.close()

        _COMPACT_TEMPLATE_CACHE[cache_key] = cached
        return cached


def _copy_cell_style(source: Any, target: Any) -> None:
    # A shallow StyleArray copy is required because a few output columns change
    # their number format after the row style is applied. Alignment and
    # protection are already part of that array, so copying them again is
    # redundant and expensive for HQ exports with tens of thousands of rows.
    target._style = copy(source._style)


def _reset_order_group_header_row(
    sheet: Any,
    *,
    visible_end: int,
) -> None:
    """Remove stale template group formatting from the spacer row above headers."""
    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.min_row <= 6 <= merged_range.max_row:
            sheet.unmerge_cells(str(merged_range))

    # Row 5 is an unstyled spacer in the template. Copying it clears both the
    # old gray fills and their borders while keeping the row in place above the
    # actual header row (7).
    for column_index in range(1, max(visible_end, sheet.max_column) + 1):
        cell = sheet.cell(6, column_index)
        cell.value = None
        _copy_cell_style(sheet.cell(5, column_index), cell)

def _copy_template_row_style(
    sheet: Any,
    *,
    source_row: int,
    target_row: int,
    max_column: int,
) -> None:
    for column_index in range(1, max_column + 1):
        _copy_cell_style(
            sheet.cell(source_row, column_index),
            sheet.cell(target_row, column_index),
        )
    source_height = sheet.row_dimensions[source_row].height
    if source_height is not None:
        sheet.row_dimensions[target_row].height = source_height


def _clear_template_data(sheet: Any, *, first_data_row: int) -> None:
    for row in sheet.iter_rows(min_row=first_data_row, max_row=sheet.max_row):
        for cell in row:
            cell.value = None


def _trim_template_rows(
    sheet: Any,
    *,
    first_data_row: int,
    data_row_count: int,
) -> None:
    keep_count = max(1, data_row_count)
    last_kept_row = first_data_row + keep_count - 1
    if sheet.max_row > last_kept_row:
        sheet.delete_rows(last_kept_row + 1, sheet.max_row - last_kept_row)


def _template_order_formula(
    key: str,
    *,
    row: int,
    columns: dict[str, int],
) -> str | None:
    from openpyxl.utils import get_column_letter, quote_sheetname

    def ref(column_key: str) -> str:
        return f"{get_column_letter(columns[column_key])}{row}"

    calc_ref = quote_sheetname(CALC_SHEET_NAME)
    demand_avg = ref("demand_avg")
    reorder_point = ref("reorder_point")
    inventory_position = ref("inventory_position")
    depletion_weeks = ref("depletion_weeks")
    demand_sigma = ref("demand_sigma")
    z_applied = ref("z_applied")
    blank_guard = (
        f'OR({demand_avg}="",{demand_avg}<=0,'
        f'{reorder_point}="",{inventory_position}="")'
    )

    if key == "monthly_sales":
        return (
            f'=IF({ref("sku_code")}="","",'
            f"'물류전망'!E{LOGISTICS_DATA_START_ROW + row - ORDER_DATA_START_ROW})"
        )
    if key == "moi":
        available_key = (
            "local_available_qty"
            if "local_available_qty" in columns
            else "eu_available_qty"
        )
        return (
            f'=IF(OR({ref("monthly_sales")}="",{ref("monthly_sales")}=0,'
            f'{inventory_position}=""),"–",'
            f"ROUND(({ref(available_key)}+{ref('transit_qty')})/"
            f"{ref('monthly_sales')},1))"
        )
    if key == "transport_recommendation":
        return (
            f'=IF(OR({depletion_weeks}="",{depletion_weeks}="–"),"–",'
            f'IF({calc_ref}!$B$2="CASH",'
            f'IF({depletion_weeks}>{calc_ref}!$B$8,"철송",'
            f'IF({depletion_weeks}>{calc_ref}!$B$6,"항공","🚨항공긴급")),'
            f'IF({depletion_weeks}>{calc_ref}!$B$10,"해운(여유)",'
            f'IF({depletion_weeks}>{calc_ref}!$B$8,"철송",'
            f'IF({depletion_weeks}>{calc_ref}!$B$6,"항공","🚨항공긴급")))))'
        )
    if key == "order_slack_weeks":
        return (
            f'=IF({blank_guard},"–",MAX(0,ROUND('
            f"({inventory_position}-{reorder_point})/{demand_avg},1)))"
        )
    if key == "expected_order_date":
        return (
            f'=IF({blank_guard},"–",{calc_ref}!$B$1+MAX(0,'
            f"({inventory_position}-{reorder_point})/{demand_avg})*7)"
        )
    if key == "conservative_slack_weeks":
        return (
            f'=IF(OR({blank_guard},{demand_sigma}="",{z_applied}=""),"–",'
            f"MAX(0,ROUND(((SQRT(({z_applied}*{demand_sigma})^2+"
            f"4*{demand_avg}*MAX(0,{inventory_position}-{reorder_point}))-"
            f"{z_applied}*{demand_sigma})/(2*{demand_avg}))^2,1)))"
        )
    if key == "early_warning_date":
        return (
            f'=IF(OR({blank_guard},{demand_sigma}="",{z_applied}=""),"–",'
            f"{calc_ref}!$B$1+((SQRT(({z_applied}*{demand_sigma})^2+"
            f"4*{demand_avg}*MAX(0,{inventory_position}-{reorder_point}))-"
            f"{z_applied}*{demand_sigma})/(2*{demand_avg}))^2*7)"
        )
    return None


def _populate_template_order_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
    result: dict[str, Any],
    include_amounts: bool,
    entity_code: str,
    style: Any,
) -> None:
    from openpyxl.utils import get_column_letter

    warehouse_names = {
        "PL": "EU",
        "USA": "미주",
        "UK": "영국",
        "ME": "중동",
        "MX": "멕시코",
        "MY": "말레이시아",
        "VN": "베트남",
        "HQ": "본사",
    }
    warehouse_name = warehouse_names.get(
        str(entity_code or "PL").strip().upper(),
        str(entity_code or "PL").strip().upper(),
    )
    currency_code = _entity_currency_code(entity_code)
    columns = dict(ORDER_TEMPLATE_COLUMNS)
    normalized_entity_code = str(
        entity_code or result.get("entity_code") or ""
    ).strip().upper()
    is_hq = normalized_entity_code == "HQ"

    # The template is intentionally kept as the UI source of truth. Insert
    # the new visible columns at export time so the bundled binary template is
    # not rewritten and its existing styles remain the starting point.
    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.min_row <= 4 and merged_range.max_row <= 4:
            sheet.unmerge_cells(str(merged_range))
    inbound_insert_at = columns["incoming_qty"] + 1
    inbound_extra_count = len(INBOUND_STATUS_COLUMNS) - 1
    if inbound_extra_count:
        sheet.insert_cols(inbound_insert_at, inbound_extra_count)
        for key, column_index in tuple(columns.items()):
            if column_index >= inbound_insert_at:
                columns[key] = column_index + inbound_extra_count
    incoming_source_column = columns["incoming_qty"]
    for offset, (key, label, width) in enumerate(INBOUND_STATUS_COLUMNS):
        column_index = incoming_source_column + offset
        columns[key] = column_index
        if offset:
            for row_index in range(1, sheet.max_row + 1):
                _copy_cell_style(
                    sheet.cell(ORDER_HEADER_ROW, incoming_source_column),
                    sheet.cell(ORDER_HEADER_ROW, column_index),
                )
                _copy_cell_style(
                    sheet.cell(row_index, incoming_source_column),
                    sheet.cell(row_index, column_index),
                )
        sheet.cell(ORDER_HEADER_ROW, column_index).value = label
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    # The bundled template still contains the historical conservative quantity
    # column. Remove it from the exported workbook while keeping the field in
    # the result contract for API/backward compatibility.
    upper_column = columns.pop("upper_suggested_qty")
    sheet.delete_cols(upper_column, 1)
    for key, column_index in tuple(columns.items()):
        if column_index > upper_column:
            columns[key] = column_index - 1

    # The transport recommendation is not part of the user-facing V2 export.
    transport_column = columns.pop("transport_recommendation")
    sheet.delete_cols(transport_column, 1)
    for key, column_index in tuple(columns.items()):
        if column_index > transport_column:
            columns[key] = column_index - 1

    # HQ의 `local_available_qty`는 현지 SKO 재고가 아니라 OPO 본사창고
    # 가용재고다. 공통 결과 계약은 유지하되, HQ Excel에서는 이 값을 본사
    # 가용재고 컬럼에 표시하고 혼동을 주는 현지 컬럼은 노출하지 않는다.
    if is_hq:
        local_column = columns.pop("local_available_qty")
        sheet.delete_cols(local_column, 1)
        for key, column_index in tuple(columns.items()):
            if column_index > local_column:
                columns[key] = column_index - 1

    amount_columns = _amount_columns(currency_code) if include_amounts else ()
    amount_start = columns["demand_sigma"]
    if amount_columns:
        sheet.insert_cols(amount_start, amount_columns.__len__())
        visible_end = amount_start + len(amount_columns) - 1
        for range_name in ("A1:AA1", "A2:AA2", "A3:AA3", "A4:AA4"):
            if range_name in {str(item) for item in sheet.merged_cells.ranges}:
                sheet.unmerge_cells(range_name)
        for row_index in range(1, 5):
            sheet.merge_cells(
                start_row=row_index,
                start_column=1,
                end_row=row_index,
                end_column=visible_end,
            )
        for offset, column in enumerate(amount_columns):
            column_index = amount_start + offset
            columns[column.key] = column_index
            header = sheet.cell(ORDER_HEADER_ROW, column_index)
            _copy_cell_style(sheet.cell(ORDER_HEADER_ROW, 20), header)
            header.value = column.label
            header.alignment = copy(sheet.cell(ORDER_HEADER_ROW, 20).alignment)
            sheet.column_dimensions[get_column_letter(column_index)].width = column.width
            for row_index in range(ORDER_DATA_START_ROW, sheet.max_row + 1):
                _copy_cell_style(sheet.cell(row_index, 20), sheet.cell(row_index, column_index))
        columns["demand_sigma"] = visible_end + 1
        columns["z_applied"] = visible_end + 2

    # The reference template reserves thousands of styled rows. Keep only the
    # rows needed for this export before clearing/populating them; otherwise
    # every download needlessly serializes the unused template tail.
    _trim_template_rows(
        sheet,
        first_data_row=ORDER_DATA_START_ROW,
        data_row_count=len(rows),
    )
    visible_end = max(
        column_index
        for key, column_index in columns.items()
        if key not in {"demand_sigma", "z_applied"}
    )
    if not amount_columns:
        for row_index in range(1, 5):
            sheet.merge_cells(
                start_row=row_index,
                start_column=1,
                end_row=row_index,
                end_column=visible_end,
            )
    _reset_order_group_header_row(
        sheet,
        visible_end=visible_end,
    )

    # Re-apply widths by logical column key after inserting/deleting columns.
    # openpyxl keeps the old physical column dimensions in place, which makes
    # the long ETA and early-warning headers inherit unrelated narrow widths.
    for key, width in ORDER_COLUMN_WIDTH_OVERRIDES.items():
        column_index = columns.get(key)
        if column_index:
            sheet.column_dimensions[get_column_letter(column_index)].width = width

    # Derive the visible demand period from the result contract so PL/USA show
    # 13 weeks and HQ shows 91 days.
    relabels = [("sales_13w_qty", _demand_period_label(result))]
    if _is_day_unit_result(result):
        # The periodic-review model has no reorder point, so the trigger wording
        # is removed instead of labelling an always-empty column.
        relabels.extend(
            [
                ("demand_avg", _demand_average_label(result)),
                ("reorder_point", "발주점\n(R/S 미사용)"),
            ]
        )
    for key, label in relabels:
        column_index = columns.get(key)
        if column_index:
            sheet.cell(ORDER_HEADER_ROW, column_index).value = label

    _clear_template_data(sheet, first_data_row=ORDER_DATA_START_ROW)
    helper_end = max(columns["demand_sigma"], columns["z_applied"])
    for index, payload_row in enumerate(rows):
        excel_row = ORDER_DATA_START_ROW + index
        source_row = ORDER_DATA_START_ROW if index % 2 == 0 else ORDER_DATA_START_ROW + 1
        _copy_template_row_style(
            sheet,
            source_row=source_row,
            target_row=excel_row,
            max_column=helper_end,
        )
        logistics_row = LOGISTICS_DATA_START_ROW + index
        static_values = {
            "sku_code": payload_row.get("sku_code"),
            "product_name": payload_row.get("product_name"),
            "brand": payload_row.get("brand"),
            "data_status": payload_row.get("data_status"),
            "grade": payload_row.get("grade"),
            "demand_avg": payload_row.get("demand_avg"),
            "safety_stock": payload_row.get("safety_stock"),
            "reorder_point": payload_row.get("reorder_point"),
            "target_stock": payload_row.get("target_stock"),
            "open_qty": payload_row.get("open_qty", payload_row.get("incoming_qty")),
            "pnfm_qty": payload_row.get("pnfm_qty"),
            "inbound_progress_qty": payload_row.get("inbound_progress_qty"),
            "inbound_completed_qty": payload_row.get("inbound_completed_qty"),
            "next_eta": payload_row.get("next_eta"),
            "eu_available_qty": payload_row.get("eu_available_qty"),
            "transit_qty": payload_row.get("transit_qty"),
            "local_available_qty": payload_row.get("local_available_qty"),
            "inventory_position": payload_row.get("inventory_position"),
            "depletion_weeks": payload_row.get("depletion_weeks"),
            "suggested_qty": payload_row.get("suggested_qty"),
            "order_signal": payload_row.get("order_signal"),
            "demand_sigma": payload_row.get("demand_sigma"),
            "z_applied": payload_row.get("z_applied"),
        }
        inbound_status_source_present = payload_row.get(
            "inbound_status_source_present"
        )
        if inbound_status_source_present is False:
            for key in (
                "open_qty",
                "pnfm_qty",
                "inbound_progress_qty",
                "inbound_completed_qty",
            ):
                static_values[key] = "-"
        if is_hq:
            static_values["eu_available_qty"] = payload_row.get("local_available_qty")
            static_values.pop("local_available_qty", None)
        for key, value in static_values.items():
            sheet.cell(excel_row, columns[key]).value = _row_value(key, value)
        sheet.cell(excel_row, columns["sales_13w_qty"]).value = (
            f"='물류전망'!D{logistics_row}"
        )
        sheet.cell(excel_row, columns["monthly_sales"]).value = (
            f"='물류전망'!E{logistics_row}"
        )
        for key in (
            "moi",
            "order_slack_weeks",
            "early_warning_date",
            "conservative_slack_weeks",
            "expected_order_date",
        ):
            sheet.cell(excel_row, columns[key]).value = _template_order_formula(
                key,
                row=excel_row,
                columns=columns,
            )
        for column in amount_columns:
            column_index = columns[column.key]
            sheet.cell(excel_row, column_index).value = _row_value(
                column.key,
                payload_row.get(column.key),
            )
            sheet.cell(excel_row, column_index).number_format = column.number_format or "General"
        _style_status_cells(
            sheet,
            excel_row,
            columns,
            style,
        )

    sheet.cell(
        ORDER_HEADER_ROW,
        columns["eu_available_qty"],
    ).value = (
        "본사 창고\n가용재고"
    )
    if "local_available_qty" in columns:
        sheet.cell(
            ORDER_HEADER_ROW,
            columns["local_available_qty"],
        ).value = (
            "현지 가용재고"
        )
    sheet["A1"] = "발주분석 V2 BETA"
    sheet["A2"] = f"적용 시나리오 : {_scenario_label(result)}"
    sheet["A3"] = (
        f"{'일별' if _is_day_unit_result(result) else '주간'} "
        "수요·변동성·리드타임·서비스 수준을 반영한 통계 기반 제안입니다. "
        "실제 발주 확정 전 데이터 기준시점, ETA 및 단가를 확인하세요."
    )
    sheet["A4"] = _amount_notice(include_amounts, result=result)
    sheet.auto_filter.ref = (
        f"A{ORDER_HEADER_ROW}:{get_column_letter(visible_end)}"
        f"{max(ORDER_HEADER_ROW, ORDER_HEADER_ROW + len(rows))}"
    )
    sheet.freeze_panes = f"A{ORDER_DATA_START_ROW}"
    sheet.print_title_rows = f"1:{ORDER_HEADER_ROW}"
    sheet.column_dimensions[get_column_letter(columns["demand_sigma"])].hidden = True
    sheet.column_dimensions[get_column_letter(columns["z_applied"])].hidden = True
    _trim_template_rows(
        sheet,
        first_data_row=ORDER_DATA_START_ROW,
        data_row_count=len(rows),
    )
    # Re-apply the signal style after the template rows have been trimmed.
    # The template alternates row fills, so leaving this to the copied row style
    # can make the same signal appear yellow on one row and red on the next.
    for index in range(len(rows)):
        _style_status_cells(
            sheet,
            ORDER_DATA_START_ROW + index,
            columns,
            style,
        )


def _populate_template_shipping_source(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
) -> int:
    _clear_template_data(sheet, first_data_row=2)
    output_row = 2
    for row_index, payload_row in enumerate(rows, start=1):
        sku_code = str(payload_row.get("sku_code") or "").strip()
        if not sku_code:
            continue
        details = payload_row.get("shipping_eta_details")
        if not isinstance(details, list) or not details:
            schedule = payload_row.get("shipping_schedule")
            details = [
                {
                    **item,
                    "eta_status": "원천 ETA",
                    "ship_date": None,
                    "transport_label": None,
                    "lead_time_days": None,
                }
                for item in schedule
                if isinstance(item, dict)
            ] if isinstance(schedule, list) else []
        for item in details:
            if not isinstance(item, dict):
                continue
            quantity = _positive_number(item.get("qty"))
            if quantity is None:
                continue
            try:
                eta = _parsed_date(item.get("eta"))
            except ValueError:
                eta = None
            try:
                ship_date = _parsed_date(item.get("ship_date"))
            except ValueError:
                ship_date = None
            if output_row > sheet.max_row:
                _copy_template_row_style(sheet, source_row=2, target_row=output_row, max_column=8)
            values = (
                sku_code,
                eta,
                quantity,
                item.get("eta_status"),
                ship_date,
                item.get("transport_label"),
                item.get("lead_time_days"),
                _sku_case_key(row_index),
            )
            for column_index, value in enumerate(values, start=1):
                sheet.cell(output_row, column_index).value = _row_value("", value)
            sheet.cell(output_row, 2).number_format = "yyyy-mm-dd"
            sheet.cell(output_row, 3).number_format = "#,##0"
            sheet.cell(output_row, 5).number_format = "yyyy-mm-dd"
            sheet.cell(output_row, 7).number_format = "0.0"
            output_row += 1
    _trim_template_rows(
        sheet,
        first_data_row=2,
        data_row_count=output_row - 2,
    )
    return max(2, output_row - 1)


def _populate_template_logistics_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
    result: dict[str, Any],
    shipping_source_last_row: int,
) -> None:
    from openpyxl.utils import get_column_letter, quote_sheetname

    # Match the sheet size to the current result before writing the repeated
    # ETA formulas and styles. The bundled template contains a large reserve
    # area for historical exports.
    _trim_template_rows(
        sheet,
        first_data_row=LOGISTICS_DATA_START_ROW,
        data_row_count=len(rows),
    )
    _clear_template_data(sheet, first_data_row=LOGISTICS_DATA_START_ROW)
    as_of = _parsed_date(result.get("as_of"), fallback=date.today())
    current_week_start = as_of - timedelta(days=as_of.weekday())
    week_starts = tuple(
        current_week_start + timedelta(weeks=index) for index in range(14)
    )
    sheet["A1"] = "물류전망 · 운송중 도착 시뮬레이션"
    sheet["A2"] = (
        "주차별 숫자는 해당 주의 ETA 입고예정수량입니다. 셀 배경색은 "
        "현지가용재고+누적도착량−예상수요 기준의 재고 상태를 표시합니다."
    )
    eta_actual_total = sum(_positive_number(row.get("eta_actual_qty")) or 0 for row in rows)
    eta_estimated_total = sum(_positive_number(row.get("eta_estimated_qty")) or 0 for row in rows)
    eta_missing_total = sum(_positive_number(row.get("eta_missing_qty")) or 0 for row in rows)
    sheet["A3"] = (
        "순수 R/S 전환으로 기존 발주점 기준 색상은 적용하지 않습니다. "
        "빨강 배경=재고 소진 · 회색 배경=판매없음. "
        f"ETA 근거: 원천 ETA {eta_actual_total:,.0f}개 · "
        f"출고일+운송수단별 평균 L/T 추정 {eta_estimated_total:,.0f}개 · "
        f"ETA 미확인 {eta_missing_total:,.0f}개. "
        "원천 ETA는 CMS에 입력된 예정일이며 도착 확정을 의미하지 않습니다. "
        "추정 ETA에는 리드타임 표준편차(σL)를 더하지 않습니다."
    )
    sheet.cell(LOGISTICS_HEADER_ROW, 11).value = (
        f"지연\n({current_week_start:%m/%d} 이전 ETA)"
    )
    sheet.cell(LOGISTICS_HEADER_ROW, 4).value = _demand_period_label(result)
    sheet.cell(
        LOGISTICS_HEADER_ROW, 6
    ).value = _logistics_demand_average_label(result)
    for offset, week_start in enumerate(week_starts[:13]):
        sheet.cell(LOGISTICS_HEADER_ROW, 12 + offset).value = (
            f"W{week_start.isocalendar().week} {week_start:%m/%d}"
        )
    sheet.cell(LOGISTICS_HEADER_ROW, 25).value = (
        f"W{week_starts[13].isocalendar().week}+\n{week_starts[13]:%m/%d}~ ETA"
    )

    source_ref = quote_sheetname(SHIPPING_SOURCE_SHEET_NAME)
    calc_ref = quote_sheetname(CALC_SHEET_NAME)
    source_case_key_range = (
        f"{source_ref}!$H$2:$H${shipping_source_last_row}"
    )
    source_eta_range = f"{source_ref}!$B$2:$B${shipping_source_last_row}"
    source_qty_range = f"{source_ref}!$C$2:$C${shipping_source_last_row}"
    for index, payload_row in enumerate(rows):
        excel_row = LOGISTICS_DATA_START_ROW + index
        source_row = LOGISTICS_DATA_START_ROW if index % 2 == 0 else LOGISTICS_DATA_START_ROW + 1
        _copy_template_row_style(sheet, source_row=source_row, target_row=excel_row, max_column=27)
        case_key = _sku_case_key(index + 1)
        values = {
            1: payload_row.get("sku_code"),
            2: payload_row.get("product_name"),
            3: payload_row.get("barcode"),
            4: payload_row.get("sales_13w_qty"),
            6: payload_row.get("demand_avg"),
            7: payload_row.get("local_available_qty"),
            8: payload_row.get("transit_qty"),
            10: payload_row.get("safety_stock"),
        }
        for column_index, value in values.items():
            sheet.cell(excel_row, column_index).value = _row_value("", value)
        sheet.cell(excel_row, 5).value = (
            f'=IF($A{excel_row}="","",ROUND(D{excel_row}/3,0))'
        )
        sheet.cell(excel_row, 9).value = (
            f'=IF(OR($E{excel_row}="",$E{excel_row}=0),"–",'
            f'ROUND(($G{excel_row}+$H{excel_row})/$E{excel_row},1))'
        )
        sheet.cell(excel_row, 11).value = (
            f'=IF($A{excel_row}="","",SUMIFS({source_qty_range},'
            f"{source_case_key_range},$AA{excel_row},{source_eta_range},"
            f'"<>",{source_eta_range},"<"&{calc_ref}!$B$3))'
        )
        for week_index in range(13):
            column_index = 12 + week_index
            start_offset = week_index * 7
            end_offset = (week_index + 1) * 7
            sheet.cell(excel_row, column_index).value = (
                f'=IF($A{excel_row}="","",SUMIFS({source_qty_range},'
                f"{source_case_key_range},$AA{excel_row},{source_eta_range},"
                f'">="&{calc_ref}!$B$3+{start_offset},{source_eta_range},'
                f'"<"&{calc_ref}!$B$3+{end_offset}))'
            )
        sheet.cell(excel_row, 25).value = (
            f'=IF($A{excel_row}="","",SUMIFS({source_qty_range},'
            f"{source_case_key_range},$AA{excel_row},{source_eta_range},"
            f'">="&{calc_ref}!$B$3+91))'
        )
        sheet.cell(excel_row, 26).value = _row_value(
            "", payload_row.get("reorder_point")
        )
        sheet.cell(excel_row, 27).value = case_key
        for column_index in range(4, 26):
            sheet.cell(excel_row, column_index).number_format = (
                "#,##0.0" if column_index in {6, 9} else "#,##0;-#,##0;;"
            )
    sheet.column_dimensions["Z"].hidden = True
    sheet.column_dimensions["AA"].hidden = True
    last_data_row = max(LOGISTICS_HEADER_ROW, LOGISTICS_HEADER_ROW + len(rows))
    sheet.auto_filter.ref = f"A{LOGISTICS_HEADER_ROW}:Y{last_data_row}"
    sheet.freeze_panes = f"A{LOGISTICS_DATA_START_ROW}"
    sheet.print_title_rows = f"1:{LOGISTICS_HEADER_ROW}"
    _trim_template_rows(
        sheet,
        first_data_row=LOGISTICS_DATA_START_ROW,
        data_row_count=len(rows),
    )
    # The template colors ETA quantities by projected stock state. Keep those
    # background colors, but use one readable body-text color for every state.
    for conditional_range in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[conditional_range]:
            if rule.dxf and rule.dxf.font:
                font = copy(rule.dxf.font)
                font.color = "FF20232A"
                rule.dxf.font = font


def _populate_template_check_required_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
    result: dict[str, Any],
    entity_code: str,
    warehouse_name: str,
    style: Any,
) -> None:
    check_rows = [row for row in rows if row.get("check_required") is True]
    title, description = _check_required_sheet_messages(
        entity_code=entity_code,
        warehouse_name=warehouse_name,
    )
    _trim_template_rows(
        sheet,
        first_data_row=CHECK_REQUIRED_DATA_START_ROW,
        data_row_count=len(check_rows),
    )
    _clear_template_data(sheet, first_data_row=CHECK_REQUIRED_DATA_START_ROW)
    sheet["A1"] = title
    sheet["A2"] = (
        f"{description} "
        "단가와 등급을 산출할 수 없어 발주제안에서 제외했습니다. "
        "재고 마스터 등록 여부를 확인한 뒤 다시 분석하세요."
    )
    sheet["A3"] = (
        f"확인필요 SKU {len(check_rows):,}개. 수량은 판매·운송중·미입고 원천에서 집계한 값이며 "
        "재고 마스터가 없어 재고 상태는 판정하지 않았습니다."
    )
    sheet.cell(
        CHECK_REQUIRED_HEADER_ROW, 4
    ).value = _check_required_sales_label(result)
    sheet.cell(CHECK_REQUIRED_HEADER_ROW, 6).value = "미입고/입고예정 수량"
    for index, payload_row in enumerate(check_rows):
        excel_row = CHECK_REQUIRED_DATA_START_ROW + index
        source_row = CHECK_REQUIRED_DATA_START_ROW if index % 2 == 0 else CHECK_REQUIRED_DATA_START_ROW + 1
        _copy_template_row_style(sheet, source_row=source_row, target_row=excel_row, max_column=7)
        values = (
            payload_row.get("sku_code"),
            payload_row.get("product_name"),
            payload_row.get("brand"),
            _positive_number(payload_row.get("sales_13w_qty")) or 0,
            _positive_number(payload_row.get("transit_qty")) or 0,
            _positive_number(payload_row.get("incoming_qty")) or 0,
            payload_row.get("check_required_reason") or "",
        )
        for column_index, value in enumerate(values, start=1):
            sheet.cell(excel_row, column_index).value = _row_value("", value)
        for column_index in range(4, 7):
            sheet.cell(excel_row, column_index).number_format = "#,##0"
    last_data_row = max(CHECK_REQUIRED_HEADER_ROW, CHECK_REQUIRED_HEADER_ROW + len(check_rows))
    sheet.auto_filter.ref = f"A{CHECK_REQUIRED_HEADER_ROW}:G{last_data_row}"
    sheet.freeze_panes = f"A{CHECK_REQUIRED_DATA_START_ROW}"
    sheet.print_title_rows = f"1:{CHECK_REQUIRED_HEADER_ROW}"
    _trim_template_rows(
        sheet,
        first_data_row=CHECK_REQUIRED_DATA_START_ROW,
        data_row_count=len(check_rows),
    )


def _generate_order_logic_v2_from_template(
    *,
    result: dict[str, Any],
    rows: list[dict[str, Any]],
    include_amounts: bool,
    entity_code: str,
) -> bytes:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:  # pragma: no cover - enforced by requirements
        raise RuntimeError(
            "openpyxl is required to generate order-logic v2 XLSX files"
        ) from exc

    template_path = _order_logic_v2_template_path()
    workbook = load_workbook(
        BytesIO(_compact_template_bytes(template_path, load_workbook)),
        data_only=False,
    )
    try:
        order_sheet = workbook["발주제안"]
        logistics_sheet = workbook["물류전망"]
        check_required_sheet = workbook["확인필요"]
        calc_sheet = workbook[CALC_SHEET_NAME]
        shipping_source_sheet = workbook[SHIPPING_SOURCE_SHEET_NAME]
        style = _WorkbookStyle(
            Font=Font,
            PatternFill=PatternFill,
            Alignment=Alignment,
            Border=Border,
            Side=Side,
        )
        warehouse_names = {
            "PL": "EU",
            "USA": "미주",
            "UK": "영국",
            "ME": "중동",
            "MX": "멕시코",
            "MY": "말레이시아",
            "VN": "베트남",
            "HQ": "본사",
        }
        warehouse_name = warehouse_names.get(
            str(entity_code or "PL").strip().upper(),
            str(entity_code or "PL").strip().upper(),
        )
        _set_document_properties(workbook, result)
        _populate_template_order_sheet(
            order_sheet,
            rows=rows,
            result=result,
            include_amounts=include_amounts,
            entity_code=entity_code,
            style=style,
        )
        _write_calc_basis_sheet(calc_sheet, result=result)
        shipping_source_last_row = _populate_template_shipping_source(
            shipping_source_sheet,
            rows=rows,
        )
        _populate_template_logistics_sheet(
            logistics_sheet,
            rows=rows,
            result=result,
            shipping_source_last_row=shipping_source_last_row,
        )
        _populate_template_check_required_sheet(
            check_required_sheet,
            rows=rows,
            result=result,
            entity_code=entity_code,
            warehouse_name=warehouse_name,
            style=style,
        )
        calc_sheet.sheet_state = "hidden"
        shipping_source_sheet.sheet_state = "hidden"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


@dataclass(frozen=True)
class _WorkbookStyle:
    Font: Any
    PatternFill: Any
    Alignment: Any
    Border: Any
    Side: Any

    @property
    def red_fill(self) -> Any:
        return self.PatternFill("solid", fgColor="E90035")

    @property
    def dark_fill(self) -> Any:
        return self.PatternFill("solid", fgColor="111318")

    @property
    def light_red_fill(self) -> Any:
        return self.PatternFill("solid", fgColor="FFF0F3")

    @property
    def note_fill(self) -> Any:
        return self.PatternFill("solid", fgColor="FFF8E7")

    @property
    def gray_fill(self) -> Any:
        return self.PatternFill("solid", fgColor="F6F7F9")

    @property
    def white_font(self) -> Any:
        return self.Font(name="Malgun Gothic", color="FFFFFF", bold=True)

    @property
    def body_font(self) -> Any:
        return self.Font(name="Malgun Gothic", color="20232A", size=10)

    @property
    def thin_bottom_border(self) -> Any:
        side = self.Side(style="thin", color="E5E7EB")
        return self.Border(bottom=side)


def _set_document_properties(workbook: Any, result: dict[str, Any]) -> None:
    from openpyxl.packaging.core import DocumentProperties

    logic_version = str(result.get("logic_version") or "v2")
    entity_code = str(result.get("entity_code") or "PL").strip().upper()
    workbook.properties = DocumentProperties(
        creator="Silicon2 SCM",
        title=f"발주분석 V2 {logic_version} BETA",
        subject=f"{entity_code} 발주 제안 검토용",
        description="통계 기반 발주 제안 결과. 실제 확정 전 담당자 검토 필요.",
        keywords=f"Silicon2, SCM, {entity_code}, order logic, BETA",
    )


def _scenario_label(result: dict[str, Any]) -> str:
    mode = str(
        result.get("applied_mode")
        or result.get("policy_mode")
        or ""
    ).strip().upper()
    if mode == "SHORTAGE":
        return "쇼티지 방어"
    if mode == "CASH":
        return "현금흐름 우선"
    return mode or "확인 필요"


def _write_order_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
    columns: tuple[ExportColumn, ...],
    include_amounts: bool,
    result: dict[str, Any],
    style: _WorkbookStyle,
    get_column_letter: Any,
) -> None:
    last_column_letter = get_column_letter(len(columns))
    index_by_key = {
        column.key: index for index, column in enumerate(columns, start=1)
    }
    helper_columns = {
        "demand_sigma": len(columns) + 1,
        "z_applied": len(columns) + 2,
    }
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    sheet["A1"] = "발주분석 V2 BETA"
    sheet["A1"].fill = style.red_fill
    sheet["A1"].font = style.Font(
        name="Malgun Gothic", color="FFFFFF", bold=True, size=16
    )
    sheet["A1"].alignment = style.Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    sheet["A2"] = (
        f"적용 시나리오: {_scenario_label(result)} · "
        "주간 수요·변동성·리드타임·서비스 수준을 반영한 통계 기반 제안입니다. "
        "실제 발주 확정 전 데이터 기준시점, ETA 및 단가를 확인하세요."
    )
    sheet["A2"].fill = style.light_red_fill
    sheet["A2"].font = style.Font(name="Malgun Gothic", color="7A1028", size=10)
    sheet["A2"].alignment = style.Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 28

    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(columns))
    sheet["A3"] = _amount_notice(include_amounts, result=result)
    sheet["A3"].fill = style.note_fill
    sheet["A3"].font = style.Font(name="Malgun Gothic", color="7A4B00", size=10)
    sheet["A3"].alignment = style.Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 26

    for column_index, column in enumerate(columns, start=1):
        cell = sheet.cell(ORDER_HEADER_ROW, column_index, column.label)
        cell.fill = style.dark_fill
        cell.font = style.white_font
        cell.alignment = style.Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = column.width
    sheet.row_dimensions[ORDER_HEADER_ROW].height = 38

    for excel_row, payload_row in enumerate(rows, start=ORDER_DATA_START_ROW):
        for column_index, column in enumerate(columns, start=1):
            formula = _order_formula(
                column.key,
                excel_row=excel_row,
                index_by_key=index_by_key,
                helper_columns=helper_columns,
                get_column_letter=get_column_letter,
            )
            value = (
                formula
                if formula is not None
                else _row_value(column.key, payload_row.get(column.key))
            )
            cell = sheet.cell(excel_row, column_index, value)
            cell.font = style.body_font
            cell.border = style.thin_bottom_border
            cell.alignment = style.Alignment(
                horizontal=column.alignment,
                vertical="center",
                wrap_text=column.key in {"product_name", "memo", "warnings"},
            )
            if column.number_format and value not in (None, ""):
                cell.number_format = column.number_format
            if excel_row % 2 == 1:
                cell.fill = style.gray_fill

        for helper_key, helper_column in helper_columns.items():
            helper_cell = sheet.cell(
                excel_row,
                helper_column,
                _row_value(helper_key, payload_row.get(helper_key)),
            )
            helper_cell.number_format = "0.0000"

        _style_status_cells(sheet, excel_row, columns, style)
        sheet.row_dimensions[excel_row].height = 24

    for helper_column in helper_columns.values():
        sheet.column_dimensions[get_column_letter(helper_column)].hidden = True

    last_data_row = max(ORDER_HEADER_ROW, ORDER_HEADER_ROW + len(rows))
    sheet.auto_filter.ref = (
        f"A{ORDER_HEADER_ROW}:{last_column_letter}{last_data_row}"
    )
    sheet.freeze_panes = f"A{ORDER_DATA_START_ROW}"
    sheet.print_title_rows = f"1:{ORDER_HEADER_ROW}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryBelow = True


def _order_formula(
    key: str,
    *,
    excel_row: int,
    index_by_key: dict[str, int],
    helper_columns: dict[str, int],
    get_column_letter: Any,
) -> str | None:
    def ref(column_key: str) -> str:
        return f"{get_column_letter(index_by_key[column_key])}{excel_row}"

    demand_avg = ref("demand_avg")
    reorder_point = ref("reorder_point")
    inventory_position = ref("inventory_position")
    depletion_weeks = ref("depletion_weeks")
    demand_sigma = (
        f"{get_column_letter(helper_columns['demand_sigma'])}{excel_row}"
    )
    z_applied = f"{get_column_letter(helper_columns['z_applied'])}{excel_row}"
    blank_guard = (
        f'OR({demand_avg}="",{demand_avg}<=0,'
        f'{reorder_point}="",{inventory_position}="")'
    )

    if key == "transport_recommendation":
        return (
            f'=IF(OR({depletion_weeks}="",{depletion_weeks}="–"),"–",'
            f"IF('{CALC_SHEET_NAME}'!$B$2=\"CASH\","
            f'IF({depletion_weeks}>\'{CALC_SHEET_NAME}\'!$B$8,"철송",'
            f'IF({depletion_weeks}>\'{CALC_SHEET_NAME}\'!$B$6,"항공",'
            f'"🚨항공긴급")),'
            f'IF({depletion_weeks}>\'{CALC_SHEET_NAME}\'!$B$10,"해운(여유)",'
            f'IF({depletion_weeks}>\'{CALC_SHEET_NAME}\'!$B$8,"철송",'
            f'IF({depletion_weeks}>\'{CALC_SHEET_NAME}\'!$B$6,"항공",'
            f'"🚨항공긴급")))))'
        )
    if key == "order_slack_weeks":
        return (
            f'=IF({blank_guard},"–",MAX(0,ROUND('
            f"({inventory_position}-{reorder_point})/{demand_avg},1)))"
        )
    if key == "expected_order_date":
        return (
            f'=IF({blank_guard},"–",\'{CALC_SHEET_NAME}\'!$B$1+MAX(0,'
            f"({inventory_position}-{reorder_point})/{demand_avg})*7)"
        )
    if key == "conservative_slack_weeks":
        return (
            f'=IF(OR({blank_guard},{demand_sigma}="",{z_applied}=""),"–",'
            f"MAX(0,ROUND(((SQRT(({z_applied}*{demand_sigma})^2+"
            f"4*{demand_avg}*MAX(0,{inventory_position}-{reorder_point}))-"
            f"{z_applied}*{demand_sigma})/(2*{demand_avg}))^2,1)))"
        )
    if key == "early_warning_date":
        return (
            f'=IF(OR({blank_guard},{demand_sigma}="",{z_applied}=""),"–",'
            f"'{CALC_SHEET_NAME}'!$B$1+((SQRT(("
            f"{z_applied}*{demand_sigma})^2+4*{demand_avg}*MAX(0,"
            f"{inventory_position}-{reorder_point}))-{z_applied}*"
            f"{demand_sigma})/(2*{demand_avg}))^2*7)"
        )
    return None


def _parsed_date(value: object, *, fallback: date | None = None) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
        except ValueError:
            pass
    if fallback is not None:
        return fallback
    raise ValueError("엑셀 계산 기준일을 날짜로 인식할 수 없습니다.")


def _setting_number(
    settings: object,
    key: str,
    *,
    default: float,
) -> float:
    if not isinstance(settings, dict):
        return default
    value = settings.get(key)
    if value is None or isinstance(value, bool):
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return number if number > 0 else default


def _write_calc_basis_sheet(sheet: Any, *, result: dict[str, Any]) -> None:
    as_of = _parsed_date(result.get("as_of"), fallback=date.today())
    current_week_start = as_of - timedelta(days=as_of.weekday())
    settings = result.get("settings")
    currency_code = str(result.get("currency_code") or "EUR").upper()
    if currency_code not in {"EUR", "USD", "KRW"}:
        currency_code = "EUR"
    if str(result.get("period_unit") or "").lower() == "day":
        # HQ/OPO reviews in day units with one measured domestic lead time and
        # no currency conversion, so the shipping-mode ladder does not apply.
        values = (
            ("분석 기준일", as_of),
            ("적용 시나리오", str(result.get("policy_mode") or "").upper()),
            ("수요 기준기간", f"{result.get('period_start')} ~ {result.get('period_end')}"),
            ("", ""),
            ("국내조달 L/T(일)", _setting_number(settings, "lt_days", default=0.0)),
            ("L/T 표준편차(일)", _setting_number(settings, "sigma_l_days", default=0.0)),
            ("정기 검토주기 R(일)", _setting_number(settings, "review_days", default=28.0)),
            ("보호기간 P(일)", "=B5+B7"),
            ("안전재고 하한(일분)", _setting_number(settings, "ss_floor_days", default=14.0)),
            ("안전재고 상한(일분)", _setting_number(settings, "ss_cap_days", default=35.0)),
            ("", ""),
            ("발주금액 통화", currency_code),
            ("환율 적용", "원화 원천값이라 환율을 적용하지 않음"),
        )
        for row_index, (label, value) in enumerate(values, start=1):
            sheet.cell(row_index, 1, label)
            sheet.cell(row_index, 2, value)
        sheet["B1"].number_format = "yyyy-mm-dd"
        for row_index in (5, 6, 7, 8, 9, 10):
            sheet.cell(row_index, 2).number_format = "0.00"
        return
    values = (
        ("분석 기준일", as_of),
        ("적용 시나리오", str(result.get("policy_mode") or "").upper()),
        ("이번 주 월요일", current_week_start),
        ("", ""),
        (
            "항공 L/T(일)",
            _setting_number(settings, "lt_air_days", default=16.4),
        ),
        ("항공 보호기간(주)", "=(B5+7)/7"),
        (
            "철송 L/T(일)",
            _setting_number(settings, "lt_rail_days", default=36.6),
        ),
        ("철송 보호기간(주)", "=(B7+7)/7"),
        (
            "해운 L/T(일)",
            _setting_number(settings, "lt_sea_days", default=72.9),
        ),
        ("해운 보호기간(주)", "=(B9+7)/7"),
        ("", ""),
        ("원화 단가 필드", "CMS stock/local unit_cost_krw"),
        ("환율 적용 기준", "재고 조회 기준일 고시환율(API 반영값)"),
        ("원화 계산", "제안수량 × unit_price_krw"),
        ("환율 출처", str(result.get("exchange_rate_source") or "CMS_STOCK_DAILY_RATE")),
    )
    for row_index, (label, value) in enumerate(values, start=1):
        sheet.cell(row_index, 1, label)
        sheet.cell(row_index, 2, value)
    sheet["B1"].number_format = "yyyy-mm-dd"
    sheet["B3"].number_format = "yyyy-mm-dd"
    for row_index in (5, 6, 7, 8, 9, 10):
        sheet.cell(row_index, 2).number_format = "0.00"


def _positive_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _sku_case_key(row_index: int) -> str:
    return f"SKU-{row_index:06d}"


def _check_required_sheet_messages(
    *,
    entity_code: str,
    warehouse_name: str,
) -> tuple[str, str]:
    """Return entity-specific review copy without changing the SKU population."""
    normalized_entity_code = str(entity_code or "PL").strip().upper()
    if normalized_entity_code == "USA":
        return (
            "USA 최초 입고 대상 SKU · 재고 마스터 미등록",
            "USA 현지·본사 미주창고 재고 마스터에 없는 SKU입니다.",
        )
    if normalized_entity_code == "HQ":
        return (
            "본사 오포창고 최초 입고 대상 SKU · 재고 마스터 미등록",
            "본사 오포창고 재고 마스터에 없는 SKU입니다.",
        )
    return (
        "SKO 최초 입고 대상 SKU · 재고 마스터 미등록",
        f"현지·본사 {warehouse_name}창고 재고 마스터에 없는 SKU입니다.",
    )


def _write_check_required_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
    style: "_WorkbookStyle",
    entity_code: str,
    warehouse_name: str,
    get_column_letter: Any,
) -> None:
    """Write SKUs that are absent from both stock masters for manual review.

    Mirrors the legacy analysis' "확인필요" sheet: a SKU that appears only in a
    non-stock source (sales/shipping/open-PO) cannot be costed or graded, so it
    is listed here instead of the order suggestion.
    """

    columns = (
        ("상품코드", 22),
        ("상품명", 34),
        ("브랜드", 18),
        ("13주 판매수량", 14),
        ("운송중 수량", 14),
        ("미입고 수량", 14),
        ("확인필요 사유", 60),
    )
    column_count = len(columns)
    last_column_letter = get_column_letter(column_count)

    check_rows = [row for row in rows if row.get("check_required") is True]
    title, description = _check_required_sheet_messages(
        entity_code=entity_code,
        warehouse_name=warehouse_name,
    )

    sheet.sheet_view.showGridLines = False

    sheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=column_count
    )
    sheet["A1"] = title
    sheet["A1"].fill = style.red_fill
    sheet["A1"].font = style.Font(
        name="Malgun Gothic", color="FFFFFF", bold=True, size=16
    )
    sheet["A1"].alignment = style.Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(
        start_row=2, start_column=1, end_row=2, end_column=column_count
    )
    sheet["A2"] = (
        f"{description} "
        "단가와 등급을 산출할 수 없어 발주제안에서 제외했습니다. "
        "재고 마스터 등록 여부를 확인한 뒤 다시 분석하세요."
    )
    sheet["A2"].fill = style.light_red_fill
    sheet["A2"].font = style.Font(name="Malgun Gothic", color="7A1028", size=10)
    sheet["A2"].alignment = style.Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 28

    sheet.merge_cells(
        start_row=3, start_column=1, end_row=3, end_column=column_count
    )
    sheet["A3"] = (
        f"확인필요 SKU {len(check_rows):,}개. "
        "수량은 판매·운송중·미입고 원천에서 집계한 값이며 "
        "재고 마스터가 없어 재고 상태는 판정하지 않았습니다."
    )
    sheet["A3"].fill = style.note_fill
    sheet["A3"].font = style.Font(name="Malgun Gothic", color="7A4B00", size=10)
    sheet["A3"].alignment = style.Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 26

    header_row = CHECK_REQUIRED_HEADER_ROW
    for column_index, (label, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=column_index, value=label)
        cell.font = style.white_font
        cell.fill = style.dark_fill
        cell.alignment = style.Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.row_dimensions[header_row].height = 38

    excel_row = CHECK_REQUIRED_DATA_START_ROW
    for row in check_rows:
        values = (
            str(row.get("sku_code") or ""),
            str(row.get("product_name") or ""),
            str(row.get("brand") or ""),
            _positive_number(row.get("sales_13w_qty")) or 0,
            _positive_number(row.get("transit_qty")) or 0,
            _positive_number(row.get("incoming_qty")) or 0,
            str(row.get("check_required_reason") or ""),
        )
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=column_index, value=value)
            cell.font = style.body_font
            cell.border = style.thin_bottom_border
            cell.alignment = style.Alignment(
                horizontal="left" if column_index <= 3 else "right",
                vertical="center",
                wrap_text=column_index in {2, column_count},
            )
            if column_index == column_count:
                cell.alignment = style.Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )
            if 4 <= column_index <= 6:
                cell.number_format = "#,##0"
            if excel_row % 2 == 1:
                cell.fill = style.gray_fill
        sheet.row_dimensions[excel_row].height = 24
        excel_row += 1

    if excel_row == CHECK_REQUIRED_DATA_START_ROW:
        note = sheet.cell(
            row=excel_row,
            column=1,
            value="확인필요 대상 SKU가 없습니다.",
        )
        note.font = style.body_font
        note.alignment = style.Alignment(horizontal="left", vertical="center")
        sheet.merge_cells(
            start_row=excel_row,
            start_column=1,
            end_row=excel_row,
            end_column=column_count,
        )
        sheet.row_dimensions[excel_row].height = 24

    last_data_row = max(header_row, header_row + len(check_rows))
    sheet.auto_filter.ref = f"A{header_row}:{last_column_letter}{last_data_row}"
    sheet.freeze_panes = f"A{CHECK_REQUIRED_DATA_START_ROW}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _write_shipping_source_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
) -> int:
    sheet.append(
        (
            "상품코드",
            "ETA",
            "수량",
            "ETA 구분",
            "출고일",
            "운송수단",
            "적용 L/T(일)",
            "SKU 대소문자 구분키",
        )
    )
    for row_index, row in enumerate(rows, start=1):
        sku_code = str(row.get("sku_code") or "").strip()
        sku_case_key = _sku_case_key(row_index)
        details = row.get("shipping_eta_details")
        if not sku_code:
            continue
        if not isinstance(details, list):
            details = []
        if not details:
            schedule = row.get("shipping_schedule")
            if isinstance(schedule, list):
                details = [
                    {
                        **item,
                        "eta_status": "원천 ETA",
                        "ship_date": None,
                        "transport_label": None,
                        "lead_time_days": None,
                    }
                    for item in schedule
                    if isinstance(item, dict)
                ]
        for item in details:
            if not isinstance(item, dict):
                continue
            quantity = _positive_number(item.get("qty"))
            if quantity is None:
                continue
            try:
                eta = _parsed_date(item.get("eta"))
            except ValueError:
                eta = None
            try:
                ship_date = _parsed_date(item.get("ship_date"))
            except ValueError:
                ship_date = None
            sheet.append(
                (
                    sku_code,
                    eta,
                    quantity,
                    item.get("eta_status"),
                    ship_date,
                    item.get("transport_label"),
                    item.get("lead_time_days"),
                    sku_case_key,
                )
            )
            if eta is not None:
                sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd"
            sheet.cell(sheet.max_row, 3).number_format = "#,##0"
            if ship_date is not None:
                sheet.cell(sheet.max_row, 5).number_format = "yyyy-mm-dd"
            sheet.cell(sheet.max_row, 7).number_format = "0.0"
    return max(2, sheet.max_row)


def _write_logistics_sheet(
    sheet: Any,
    *,
    rows: list[dict[str, Any]],
    style: _WorkbookStyle,
    get_column_letter: Any,
    FormulaRule: Any,
    shipping_source_last_row: int,
    result: dict[str, Any],
) -> None:
    as_of = _parsed_date(result.get("as_of"), fallback=date.today())
    current_week_start = as_of - timedelta(days=as_of.weekday())
    week_starts = tuple(
        current_week_start + timedelta(weeks=index) for index in range(14)
    )
    eta_actual_total = sum(
        _positive_number(row.get("eta_actual_qty")) or 0.0
        for row in rows
    )
    eta_estimated_total = sum(
        _positive_number(row.get("eta_estimated_qty")) or 0.0
        for row in rows
    )
    eta_missing_total = sum(
        _positive_number(row.get("eta_missing_qty")) or 0.0
        for row in rows
    )
    visible_headers = [
        "상품코드",
        "상품명",
        "바코드",
        "판매량\n(13주)",
        "월평균\n판매량",
        "주평균",
        "현재고\n(현지가용)",
        "운송중\n합계",
        "안전재고\n(상하한 적용)",
        f"지연\n({current_week_start:%m/%d} 이전 ETA)",
    ]
    visible_headers.extend(
        f"W{week_start.isocalendar().week} {week_start:%m/%d}"
        for week_start in week_starts[:13]
    )
    visible_headers.extend(
        (
            f"W{week_starts[13].isocalendar().week}+\n"
            f"{week_starts[13]:%m/%d}~ ETA",
            "원천 ETA\n물량",
            "출고일 추정\nETA 물량",
        )
    )
    # 도착일을 특정할 수 없는 물량이 실제로 존재할 때만 컬럼을 노출한다.
    # 상시 0으로 비어 있는 컬럼이 표 오른쪽 끝을 차지하는 것을 막고,
    # 결측이 발생한 회차에서만 눈에 띄게 한다.
    show_eta_missing_column = eta_missing_total > 0
    if show_eta_missing_column:
        visible_headers.append("ETA 미확인\n물량")
    visible_column_count = len(visible_headers)
    last_visible_column = get_column_letter(visible_column_count)
    reorder_helper_column = visible_column_count + 1
    reorder_helper_letter = get_column_letter(reorder_helper_column)
    case_key_helper_column = visible_column_count + 2
    case_key_helper_letter = get_column_letter(case_key_helper_column)

    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=visible_column_count,
    )
    sheet["A1"] = "물류전망 · 운송중 도착 시뮬레이션"
    sheet["A1"].fill = style.red_fill
    sheet["A1"].font = style.Font(
        name="Malgun Gothic", color="FFFFFF", bold=True, size=16
    )
    sheet["A1"].alignment = style.Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=visible_column_count,
    )
    sheet["A2"] = (
        "주차별 숫자는 해당 주의 ETA 입고예정수량입니다. "
        "셀 배경색은 현지가용재고+누적도착량−예상수요 기준의 재고 상태를 표시합니다."
    )
    sheet["A2"].fill = style.light_red_fill
    sheet["A2"].font = style.Font(name="Malgun Gothic", color="7A1028", size=10)
    sheet["A2"].alignment = style.Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 28

    sheet.merge_cells(
        start_row=3,
        start_column=1,
        end_row=3,
        end_column=visible_column_count,
    )
    eta_notice = (
        f"ETA 근거: 원천 ETA {eta_actual_total:,.0f}개 · "
        f"출고일+운송수단별 평균 L/T 추정 {eta_estimated_total:,.0f}개 · "
        f"ETA 미확인 {eta_missing_total:,.0f}개. "
        "원천 ETA는 CMS에 입력된 예정일이며 도착 확정을 의미하지 않습니다. "
        "추정 ETA에는 리드타임 표준편차(σL)를 더하지 않습니다."
    )
    missing_notice = (
        "ETA 미확인 물량은 주차별 도착량에 포함하지 않습니다. "
        if show_eta_missing_column
        else ""
    )
    sheet["A3"] = (
        "순수 R/S 전환으로 기존 발주점 기준 색상은 적용하지 않습니다. "
        "빨강 배경=재고 소진 · 회색 배경=판매없음. "
        f"{missing_notice}"
        f"{eta_notice}"
    )
    sheet["A3"].fill = style.note_fill
    sheet["A3"].font = style.Font(name="Malgun Gothic", color="7A4B00", size=10)
    sheet["A3"].alignment = style.Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 40

    widths = [18, 34, 18, 13, 13, 13, 15, 14, 16, 16]
    widths.extend([13] * 13)
    widths.extend((17, 14, 16))
    if show_eta_missing_column:
        widths.append(15)
    for column_index, (header, width) in enumerate(
        zip(visible_headers, widths, strict=True),
        start=1,
    ):
        cell = sheet.cell(LOGISTICS_HEADER_ROW, column_index, header)
        cell.fill = style.dark_fill
        cell.font = style.white_font
        cell.alignment = style.Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.row_dimensions[LOGISTICS_HEADER_ROW].height = 42

    source_case_key_range = (
        f"'{SHIPPING_SOURCE_SHEET_NAME}'!"
        f"${get_column_letter(SHIPPING_SOURCE_CASE_KEY_COLUMN)}$2:"
        f"${get_column_letter(SHIPPING_SOURCE_CASE_KEY_COLUMN)}"
        f"${shipping_source_last_row}"
    )
    source_eta_range = (
        f"'{SHIPPING_SOURCE_SHEET_NAME}'!$B$2:$B${shipping_source_last_row}"
    )
    source_qty_range = (
        f"'{SHIPPING_SOURCE_SHEET_NAME}'!$C$2:$C${shipping_source_last_row}"
    )

    for row_index, payload_row in enumerate(rows, start=1):
        excel_row = LOGISTICS_DATA_START_ROW + row_index - 1
        sku_case_key = _sku_case_key(row_index)
        static_values = (
            payload_row.get("sku_code"),
            payload_row.get("product_name"),
            payload_row.get("barcode"),
            payload_row.get("sales_13w_qty"),
            None,
            payload_row.get("demand_avg"),
            payload_row.get("local_available_qty"),
            payload_row.get("transit_qty"),
            payload_row.get("safety_stock"),
        )
        for column_index, value in enumerate(static_values, start=1):
            cell = sheet.cell(
                excel_row,
                column_index,
                _row_value("", value),
            )
            cell.font = style.body_font
            cell.border = style.thin_bottom_border
            cell.alignment = style.Alignment(
                horizontal="left" if column_index <= 3 else "right",
                vertical="center",
                wrap_text=column_index == 2,
            )
            if column_index >= 4 and value not in (None, ""):
                cell.number_format = "#,##0.0" if column_index == 6 else "#,##0"
            if excel_row % 2 == 1:
                cell.fill = style.gray_fill

        sheet.cell(excel_row, 5, f'=IF($A{excel_row}="","",ROUND(D{excel_row}/3,0))')
        sheet.cell(excel_row, 5).number_format = "#,##0"

        delay_formula = (
            f'=IF($A{excel_row}="","",SUMIFS({source_qty_range},'
            f"{source_case_key_range},${case_key_helper_letter}{excel_row},"
            f'{source_eta_range},"<>",{source_eta_range},'
            f'"<"&\'{CALC_SHEET_NAME}\'!$B$3))'
        )
        sheet.cell(excel_row, 10, delay_formula)

        for week_index in range(13):
            column_index = 11 + week_index
            start_offset = week_index * 7
            end_offset = (week_index + 1) * 7
            formula = (
                f'=IF($A{excel_row}="","",SUMIFS({source_qty_range},'
                f"{source_case_key_range},${case_key_helper_letter}{excel_row},"
                f"{source_eta_range},"
                f'">="&\'{CALC_SHEET_NAME}\'!$B$3+{start_offset},'
                f"{source_eta_range},"
                f'"<"&\'{CALC_SHEET_NAME}\'!$B$3+{end_offset}))'
            )
            sheet.cell(excel_row, column_index, formula)

        after_column = 24
        sheet.cell(
            excel_row,
            after_column,
            (
                f'=IF($A{excel_row}="","",SUMIFS({source_qty_range},'
                f"{source_case_key_range},${case_key_helper_letter}{excel_row},"
                f"{source_eta_range},"
                f'">="&\'{CALC_SHEET_NAME}\'!$B$3+91))'
            ),
        )
        sheet.cell(
            excel_row,
            after_column + 1,
            _row_value("", payload_row.get("eta_actual_qty")),
        )
        sheet.cell(
            excel_row,
            after_column + 2,
            _row_value("", payload_row.get("eta_estimated_qty")),
        )
        if show_eta_missing_column:
            sheet.cell(
                excel_row,
                after_column + 3,
                _row_value("", payload_row.get("eta_missing_qty")),
            )
        sheet.cell(
            excel_row,
            reorder_helper_column,
            _row_value("", payload_row.get("reorder_point")),
        )
        sheet.cell(excel_row, case_key_helper_column, sku_case_key)
        for column_index in range(5, visible_column_count + 1):
            cell = sheet.cell(excel_row, column_index)
            cell.font = style.body_font
            cell.border = style.thin_bottom_border
            cell.alignment = style.Alignment(horizontal="right", vertical="center")
            if column_index >= 10:
                cell.number_format = "#,##0;-#,##0;;"
            if excel_row % 2 == 1 and cell.fill.fill_type is None:
                cell.fill = style.gray_fill
        sheet.row_dimensions[excel_row].height = 24

    last_data_row = max(
        LOGISTICS_HEADER_ROW,
        LOGISTICS_HEADER_ROW + len(rows),
    )
    sheet.auto_filter.ref = (
        f"A{LOGISTICS_HEADER_ROW}:{last_visible_column}{last_data_row}"
    )
    sheet.freeze_panes = f"A{LOGISTICS_DATA_START_ROW}"
    sheet.print_title_rows = f"1:{LOGISTICS_HEADER_ROW}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.column_dimensions[reorder_helper_letter].hidden = True
    sheet.column_dimensions[case_key_helper_letter].hidden = True

    if rows:
        weekly_range = (
            f"K{LOGISTICS_DATA_START_ROW}:W{last_data_row}"
        )
        projected = (
            "$G6+SUM($J6:K6)-$F6*(COLUMN()-COLUMN($J6))"
        )
        rules = (
            (
                'OR($F6="",$F6=0)',
                style.PatternFill("solid", fgColor="EEF0F3"),
                style.body_font,
            ),
            (
                f"AND($F6>0,{projected}<=0)",
                style.PatternFill("solid", fgColor="FFE4E8"),
                style.Font(
                    name="Malgun Gothic", color="20232A", bold=True, size=10
                ),
            ),
            (
                f"AND($F6>0,{projected}>0,{projected}<$I6)",
                style.PatternFill("solid", fgColor="FFE7CC"),
                style.Font(
                    name="Malgun Gothic", color="20232A", bold=True, size=10
                ),
            ),
            (
                f"AND($F6>0,{projected}>=$I6,{projected}<${reorder_helper_letter}6)",
                style.PatternFill("solid", fgColor="FFF3CD"),
                style.body_font,
            ),
            (
                f"AND($F6>0,{projected}>=${reorder_helper_letter}6)",
                style.PatternFill("solid", fgColor="E8F7EE"),
                style.body_font,
            ),
        )
        for formula, fill, font in rules:
            sheet.conditional_formatting.add(
                weekly_range,
                FormulaRule(
                    formula=[formula],
                    fill=fill,
                    font=font,
                    stopIfTrue=True,
                ),
            )
        sheet.conditional_formatting.add(
            f"J{LOGISTICS_DATA_START_ROW}:J{last_data_row}",
            FormulaRule(
                formula=[f"J{LOGISTICS_DATA_START_ROW}>0"],
                fill=style.PatternFill("solid", fgColor="FFE4E8"),
                font=style.Font(
                    name="Malgun Gothic", color="20232A", bold=True, size=10
                ),
            ),
        )
        sheet.conditional_formatting.add(
            f"AA{LOGISTICS_DATA_START_ROW}:AA{last_data_row}",
            FormulaRule(
                formula=[f"AA{LOGISTICS_DATA_START_ROW}>0"],
                fill=style.PatternFill("solid", fgColor="FFF3CD"),
                font=style.Font(
                    name="Malgun Gothic", color="20232A", bold=True, size=10
                ),
            ),
        )


def _style_status_cells(
    sheet: Any,
    excel_row: int,
    columns: dict[str, int] | tuple[ExportColumn, ...],
    style: _WorkbookStyle,
) -> None:
    if isinstance(columns, dict):
        index_by_key = columns
    else:
        index_by_key = {
            column.key: index for index, column in enumerate(columns, start=1)
        }

    signal_cell = sheet.cell(excel_row, index_by_key["order_signal"])
    signal = str(signal_cell.value or "")
    if "확인" in signal:
        signal_cell.fill = style.PatternFill("solid", fgColor="FFFFF3CD")
        signal_cell.font = style.Font(
            name="Malgun Gothic", color="FF7A4B00", bold=True, size=10
        )
    elif "발주" in signal:
        signal_cell.fill = style.PatternFill("solid", fgColor="FFFFE4E8")
        signal_cell.font = style.Font(
            name="Malgun Gothic", color="FFB4002D", bold=True, size=10
        )
    elif "충분" in signal:
        signal_cell.fill = style.PatternFill("solid", fgColor="FFE8F7EE")
        signal_cell.font = style.Font(
            name="Malgun Gothic", color="FF087A3D", bold=True, size=10
        )

    data_status_cell = sheet.cell(excel_row, index_by_key["data_status"])
    data_status = str(data_status_cell.value or "")
    if data_status.startswith("⚠") or "확인" in data_status:
        data_status_cell.fill = style.PatternFill("solid", fgColor="FFFFF3CD")
        data_status_cell.font = style.Font(
            name="Malgun Gothic", color="FF7A4B00", bold=True, size=10
        )
    elif data_status == "판매없음":
        data_status_cell.fill = style.gray_fill
        data_status_cell.font = style.Font(
            name="Malgun Gothic", color="6B7280", bold=True, size=10
        )
    elif data_status in {"정상", "충분"}:
        data_status_cell.fill = style.PatternFill("solid", fgColor="FFE8F7EE")
        data_status_cell.font = style.Font(
            name="Malgun Gothic", color="FF087A3D", bold=True, size=10
        )

def _amount_notice(include_amounts: bool, *, result: dict[str, Any]) -> str:
    if include_amounts:
        result_currency = str(result.get("currency_code") or "").strip().upper()
        if result_currency == "KRW":
            # 본사 발주금액은 원화 원천값이므로 환산할 환율이 없다. 이를 조회
            # 실패로 안내하면 정상 금액을 신뢰할 수 없는 값처럼 보이게 한다.
            return (
                "※ 단가와 제안금액은 발주 검토용 참고값입니다. "
                "본사 발주금액은 원화 원천값이라 환율을 적용하지 않습니다."
            )
        return (
            "※ 단가와 제안금액은 발주 검토용 참고값입니다. "
            "원화 금액은 CMS 재고 API의 조회 기준일 원화 단가(unit_cost_krw)를 "
            "제안수량에 곱한 값입니다."
        )
    return (
        "※ 금액 조회 권한이 없어 단가·금액 열을 제외했습니다. "
        "참고금액은 권한이 있는 화면에서 확인하세요."
    )


def _row_value(key: str, value: Any) -> Any:
    if key == "warnings":
        return _excel_safe_text(_warning_text(value))
    if key == "grade":
        return {
            "MAJOR": "주력",
            "MINOR": "일반",
        }.get(str(value), value)
    if key == "order_signal":
        return {
            "-": "–",
            "발주": "🔴발주",
            "확인후발주": "⚠확인후발주",
            "충분": "🟢충분",
        }.get(str(value), value)
    if key == "next_eta" and isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
        except ValueError:
            return _excel_safe_text(value)
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    if isinstance(value, str):
        return _excel_safe_text(value)
    return value


def _warning_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (list, tuple)):
        return "\n".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return str(value)


def _excel_safe_text(value: str) -> str:
    # A single hyphen is the user-facing placeholder for unavailable values,
    # not a formula. Keep it as-is so Excel displays `-` instead of `'-`.
    if value == "-":
        return value
    if value.startswith(("=", "+", "-", "@")):
        # Prevent exported user-controlled text from becoming an Excel formula.
        return f"'{value}"
    return value


__all__ = [
    "AMOUNT_COLUMNS",
    "BASE_COLUMNS",
    "ORDER_HEADER_ROW",
    "generate_order_logic_v2_excel",
]
