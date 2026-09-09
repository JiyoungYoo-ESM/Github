from __future__ import annotations

import argparse
import json
import math
import tempfile
import zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


TOLERANCE = 1e-6


def as_number(value: object, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return number if math.isfinite(number) else default


def close_enough(left: float, right: float, tolerance: float = TOLERANCE) -> bool:
    return abs(left - right) <= max(tolerance, tolerance * max(abs(left), abs(right)))


def find_latest(downloads: Path, pattern: str) -> Path:
    candidates = list(downloads.glob(pattern))
    if not candidates:
        raise FileNotFoundError(f"No file matched {pattern!r} in {downloads}")
    return max(candidates, key=lambda path: (path.stat().st_mtime, path.name))


@contextmanager
def normalized_workbook_copy(path: Path):
    """Yield a temporary standards-compliant copy without changing the source workbook."""
    with tempfile.TemporaryDirectory(prefix="order-analysis-backtest-") as temp_dir:
        target = Path(temp_dir) / path.name
        invalid_style = False
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                    if b'errorStyle="error"' in payload:
                        invalid_style = True
                        payload = payload.replace(b'errorStyle="error"', b'errorStyle="stop"')
                destination.writestr(info, payload)
        yield target, invalid_style


def audit_order_export(path: Path) -> dict[str, object]:
    with normalized_workbook_copy(path) as (normalized_path, invalid_style):
        formula_book = load_workbook(normalized_path, read_only=False, data_only=False)
        value_book = load_workbook(normalized_path, read_only=False, data_only=True)
        try:
            formula_sheet = formula_book.worksheets[0]
            value_sheet = value_book.worksheets[0]
            value_calc = value_book.worksheets[1]

            start_row = 6
            end_row = formula_sheet.max_row
            records: list[dict[str, object]] = []
            formula_pattern_failures = 0
            scenario_stock_failures = 0
            scenario_qty_failures = 0
            scenario_amount_failures = 0
            moi_failures = 0
            visible_cache_failures = 0
            negative_value_count = 0

            formula_rows = formula_sheet.iter_rows(min_row=start_row, max_row=end_row, max_col=13, values_only=True)
            value_rows = value_sheet.iter_rows(min_row=start_row, max_row=end_row, max_col=13, values_only=True)
            calc_rows = value_calc.iter_rows(min_row=start_row, max_row=end_row, max_col=14, values_only=True)
            for row_number, (formula_values, visible_values, calc_values) in enumerate(
                zip(formula_rows, value_rows, calc_rows, strict=True), start=start_row
            ):
                product_code = str(visible_values[0] or "").strip()
                if not product_code:
                    continue
                (
                    before_stock,
                    after_stock,
                    before_moi,
                    after_moi,
                    before_qty,
                    after_qty,
                    before_amount,
                    after_amount,
                    before_open_po,
                    after_open_po,
                    before_reason,
                    after_reason,
                    before_status,
                    after_status,
                ) = calc_values

                sales_3m = as_number(visible_values[5])
                expected_after_stock = as_number(before_stock) + as_number(before_open_po)
                expected_after_qty = max(as_number(before_qty) - as_number(before_open_po), 0)
                unit_amount = as_number(before_amount) / as_number(before_qty) if as_number(before_qty) > 0 else 0
                expected_after_amount = expected_after_qty * unit_amount

                if not close_enough(as_number(after_stock), expected_after_stock):
                    scenario_stock_failures += 1
                if not close_enough(as_number(after_qty), expected_after_qty):
                    scenario_qty_failures += 1
                if not close_enough(as_number(after_amount), expected_after_amount):
                    scenario_amount_failures += 1
                before_demand = as_number(before_stock) / as_number(before_moi) if as_number(before_moi) > 0 else 0
                after_demand = as_number(after_stock) / as_number(after_moi) if as_number(after_moi) > 0 else 0
                invalid_zero_moi = (as_number(before_stock) == 0 and as_number(before_moi) != 0) or (
                    as_number(after_stock) == 0 and as_number(after_moi) != 0
                )
                inconsistent_demand = before_demand > 0 and after_demand > 0 and not close_enough(before_demand, after_demand)
                if invalid_zero_moi or inconsistent_demand:
                    moi_failures += 1

                numeric_values = [
                    before_stock,
                    after_stock,
                    before_moi,
                    after_moi,
                    before_qty,
                    after_qty,
                    before_amount,
                    after_amount,
                    before_open_po,
                    after_open_po,
                    visible_values[5],
                    visible_values[10],
                ]
                negative_value_count += sum(as_number(value) < 0 for value in numeric_values)

                expected_visible = {
                    5: before_stock,
                    8: before_qty,
                    9: before_amount,
                    10: before_open_po,
                    12: before_reason,
                    13: before_status,
                }
                for column, expected in expected_visible.items():
                    actual = visible_values[column - 1]
                    if isinstance(expected, (int, float)):
                        if not close_enough(as_number(actual), as_number(expected)):
                            visible_cache_failures += 1
                    elif str(actual or "") != str(expected or ""):
                        visible_cache_failures += 1

                for column in (5, 8, 9, 10, 12, 13):
                    formula = formula_values[column - 1]
                    if not isinstance(formula, str) or not formula.startswith("=IF($B$2="):
                        formula_pattern_failures += 1
                moi_formula = formula_values[6]
                if not isinstance(moi_formula, str) or "$B$2=" not in moi_formula or "_계산값" not in moi_formula:
                    formula_pattern_failures += 1

                records.append(
                    {
                        "product_code": product_code,
                        "before_qty": as_number(before_qty),
                        "after_qty": as_number(after_qty),
                        "before_amount": as_number(before_amount),
                        "after_amount": as_number(after_amount),
                        "before_status": str(before_status or ""),
                        "after_status": str(after_status or ""),
                    }
                )

            product_codes = [record["product_code"] for record in records]
            duplicates = sum(count - 1 for count in Counter(product_codes).values() if count > 1)
            before_qty_sum = sum(float(record["before_qty"]) for record in records)
            after_qty_sum = sum(float(record["after_qty"]) for record in records)
            before_amount_sum = sum(float(record["before_amount"]) for record in records)
            after_amount_sum = sum(float(record["after_amount"]) for record in records)
            target_count = int(as_number(value_sheet["J2"].value))
            summary_qty = as_number(value_sheet["B3"].value)
            summary_amount = as_number(value_sheet["D3"].value)

            return {
                "file": path.name,
                "invalid_data_validation_style": invalid_style,
                "generated_at": str(value_sheet["F3"].value or ""),
                "row_count": len(records),
                "target_count": target_count,
                "duplicate_product_code_rows": duplicates,
                "missing_product_code_rows": (end_row - start_row + 1) - len(records),
                "negative_numeric_values": negative_value_count,
                "formula_pattern_failures": formula_pattern_failures,
                "scenario_stock_failures": scenario_stock_failures,
                "scenario_qty_failures": scenario_qty_failures,
                "scenario_amount_failures": scenario_amount_failures,
                "moi_failures": moi_failures,
                "visible_cache_failures": visible_cache_failures,
                "before_qty_sum": before_qty_sum,
                "after_qty_sum": after_qty_sum,
                "before_amount_sum_krw": before_amount_sum,
                "after_amount_sum_krw": after_amount_sum,
                "summary_qty": summary_qty,
                "summary_amount_krw": summary_amount,
                "summary_qty_matches": close_enough(summary_qty, before_qty_sum),
                "summary_amount_matches": close_enough(summary_amount, before_amount_sum),
                "target_count_matches": target_count == len(records),
            }
        finally:
            formula_book.close()
            value_book.close()


def month_day_to_date(value: object, year: int) -> datetime | None:
    text = str(value or "").strip()
    if not text or text == "-":
        return None
    for pattern in ("%m/%d", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, pattern)
            return parsed.replace(year=year) if pattern == "%m/%d" else parsed
        except ValueError:
            continue
    return None


def audit_stock_gap_export(path: Path) -> dict[str, object]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        header_row = 4 if sheet.cell(4, 1).value else 5
        headers = [str(sheet.cell(header_row, column).value or "").strip() for column in range(1, sheet.max_column + 1)]
        index = {header: position + 1 for position, header in enumerate(headers)}
        product_code_column = index.get("상품코드")
        stock_column = index.get("현재고")
        speed_column = index.get("판매 속도")
        stockout_column = index.get("소진 예정")
        eta_column = index.get("최초 ETA")
        gap_column = index.get("공백") or index.get("공백 일수")
        shortage_column = index.get("부족 수량")
        status_column = index.get("상태")
        if not all((stock_column, speed_column, stockout_column, eta_column, gap_column, shortage_column, status_column)):
            raise ValueError(f"Unexpected stock-gap columns: {headers}")

        generated_text = str(sheet["B2"].value or "")
        year = 2026
        for token in generated_text.replace(".", " ").split():
            if token.isdigit() and len(token) == 4:
                year = int(token)
                break

        records: list[dict[str, object]] = []
        gap_date_failures = 0
        gap_date_ambiguous_rows = 0
        shortage_reconciliation_exceptions = 0
        shortage_exception_displayed_total = 0
        shortage_exception_formula_total = 0.0
        negative_value_count = 0
        missing_status_count = 0
        for values in sheet.iter_rows(
            min_row=header_row + 1,
            max_row=sheet.max_row,
            max_col=sheet.max_column,
            values_only=True,
        ):
            status = str(values[status_column - 1] or "").strip()
            if not status:
                continue
            product_code = str(values[product_code_column - 1] or "").strip() if product_code_column else ""
            stock = as_number(values[stock_column - 1])
            speed = as_number(values[speed_column - 1])
            gap_days = int(round(as_number(values[gap_column - 1])))
            shortage = int(round(as_number(values[shortage_column - 1])))
            stockout = month_day_to_date(values[stockout_column - 1], year)
            eta = month_day_to_date(values[eta_column - 1], year)

            if stockout and eta:
                same_year_gap = max((eta.date() - stockout.date()).days, 0)
                if gap_days == 0 and same_year_gap > 0:
                    gap_date_ambiguous_rows += 1
                if eta < stockout and gap_days > 0:
                    eta = eta.replace(year=eta.year + 1)
                expected_gap = max((eta.date() - stockout.date()).days, 0)
                if gap_days > 0 and gap_days != expected_gap:
                    gap_date_failures += 1
            if gap_days > 0 and speed > 0 and shortage > 0:
                lower = max(speed - 0.05, 0) * gap_days
                upper = (speed + 0.05) * gap_days
                if shortage < math.floor(lower) or shortage > math.ceil(upper):
                    shortage_reconciliation_exceptions += 1
                    shortage_exception_displayed_total += shortage
                    shortage_exception_formula_total += speed * gap_days
            negative_value_count += sum(value < 0 for value in (stock, speed, gap_days, shortage))
            missing_status_count += int(not status)
            records.append(
                {
                    "product_code": product_code,
                    "status": status,
                    "gap_days": gap_days,
                    "shortage": shortage,
                }
            )

        status_counts = Counter(record["status"] for record in records)
        risk_statuses = {"재고공백"}
        risk_rows = [record for record in records if record["status"] in risk_statuses]
        waiting_rows = [record for record in records if record["status"] == "ETA 없음"]
        duplicate_product_codes = 0
        if product_code_column:
            codes = [record["product_code"] for record in records if record["product_code"]]
            duplicate_product_codes = sum(count - 1 for count in Counter(codes).values() if count > 1)

        summary_risk = int(as_number(sheet["B3"].value))
        summary_long = int(as_number(sheet["D3"].value))
        summary_shortage = int(round(as_number(sheet["F3"].value)))
        summary_waiting = int(as_number(sheet["I3"].value))
        computed_long = sum(int(record["gap_days"]) >= 7 for record in risk_rows)
        computed_shortage = sum(int(record["shortage"]) for record in risk_rows)
        target_count = int(as_number(sheet["I2"].value))

        return {
            "file": path.name,
            "generated_at": generated_text,
            "row_count": len(records),
            "target_count": target_count,
            "target_count_matches": target_count == len(records),
            "status_counts": dict(status_counts),
            "duplicate_product_code_rows": duplicate_product_codes,
            "missing_product_code_rows": sum(not record["product_code"] for record in records) if product_code_column else None,
            "missing_status_rows": missing_status_count,
            "negative_numeric_values": negative_value_count,
            "gap_date_failures": gap_date_failures,
            "gap_date_ambiguous_rows": gap_date_ambiguous_rows,
            "shortage_reconciliation_exceptions": shortage_reconciliation_exceptions,
            "shortage_exception_displayed_total": shortage_exception_displayed_total,
            "shortage_exception_visible_formula_total": round(shortage_exception_formula_total),
            "shortage_exception_delta": round(shortage_exception_displayed_total - shortage_exception_formula_total),
            "summary_risk": summary_risk,
            "computed_risk": len(risk_rows),
            "summary_long_gap": summary_long,
            "computed_long_gap": computed_long,
            "summary_shortage": summary_shortage,
            "computed_shortage": computed_shortage,
            "summary_waiting_eta": summary_waiting,
            "computed_waiting_eta": len(waiting_rows),
            "summary_matches": {
                "risk": summary_risk == len(risk_rows),
                "long_gap": summary_long == computed_long,
                "shortage": summary_shortage == computed_shortage,
                "waiting_eta": summary_waiting == len(waiting_rows),
            },
        }
    finally:
        book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Backtest exported order-analysis workbooks without modifying them.")
    parser.add_argument("--downloads", type=Path, default=Path.home() / "Downloads")
    parser.add_argument("--order", type=Path)
    parser.add_argument("--stock-gap", type=Path)
    parser.add_argument("--output", type=Path, default=Path(__file__).with_name("results.json"))
    args = parser.parse_args()

    order_path = args.order or find_latest(args.downloads, "ESM_order_review_filtered_all_*.xlsx")
    stock_gap_path = args.stock_gap or find_latest(args.downloads, "ESM_stock_gap_filtered_all_*.xlsx")
    result = {
        "as_of": datetime.now().astimezone().isoformat(),
        "order_export": audit_order_export(order_path),
        "stock_gap_export": audit_stock_gap_export(stock_gap_path),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
